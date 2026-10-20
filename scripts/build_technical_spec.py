"""Genera la especificación técnica de PEA-i UPC en formato DOCX."""

from __future__ import annotations

from pathlib import Path
import runpy
import tempfile
from typing import Any

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from PIL import Image, ImageDraw, ImageFont


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "docs/especificacion_tecnica.docx"
FONT = "/usr/share/fonts/TTF/DejaVuSans.ttf"
BOLD_FONT = "/usr/share/fonts/TTF/DejaVuSans-Bold.ttf"
NAVY = (25, 52, 78)
PALE = (232, 239, 246)
GRAY = (90, 103, 115)
# Corte de los indicadores publicados: captura del 19 de octubre de 2026.
INDICATOR_CUTOFF = 2026
CLASS_TERMS = {"TOP": "Nc.Top", "A": "Nc.A", "B": "Nc.B", "ASC": "AP", "DPC": "DPC",
               "FRH_A": "Fr.A", "FRH_B": "Fr.B"}
CLASS_NAMES = {
    "TOP": "Productos tipo TOP (nuevo conocimiento de la más alta calidad)",
    "A": "Productos tipo A", "B": "Productos tipo B",
    "ASC": "Apropiación social del conocimiento",
    "DPC": "Divulgación pública de la ciencia",
    "FRH_A": "Formación de recurso humano tipo A (doctorado)",
    "FRH_B": "Formación de recurso humano tipo B",
}


def font(size: int, bold: bool = False) -> ImageFont.FreeTypeFont:
    return ImageFont.truetype(BOLD_FONT if bold else FONT, size)


def box(draw: ImageDraw.ImageDraw, xy: tuple[int, int, int, int],
        lines: list[str], fill: tuple[int, int, int] = PALE) -> None:
    draw.rounded_rectangle(xy, radius=17, fill=fill, outline=NAVY, width=3)
    left, top, right, bottom = xy
    widths = [draw.textbbox((0, 0), line, font=font(28, i == 0))[2] for i, line in enumerate(lines)]
    total_height = len(lines) * 40
    y = top + (bottom - top - total_height) // 2
    for i, (line, width) in enumerate(zip(lines, widths)):
        draw.text((left + (right - left - width) // 2, y), line,
                  font=font(28, i == 0), fill=NAVY)
        y += 40


def arrow(draw: ImageDraw.ImageDraw, start: tuple[int, int],
          end: tuple[int, int]) -> None:
    draw.line((start, end), fill=NAVY, width=5)
    x1, y1 = start
    x2, y2 = end
    if abs(x2 - x1) >= abs(y2 - y1):
        sign = 1 if x2 > x1 else -1
        draw.polygon([(x2, y2), (x2 - 18 * sign, y2 - 10),
                      (x2 - 18 * sign, y2 + 10)], fill=NAVY)
    else:
        sign = 1 if y2 > y1 else -1
        draw.polygon([(x2, y2), (x2 - 10, y2 - 18 * sign),
                      (x2 + 10, y2 - 18 * sign)], fill=NAVY)


def architecture_image(path: Path) -> None:
    canvas = Image.new("RGB", (1600, 545), "white")
    draw = ImageDraw.Draw(canvas)
    box(draw, (25, 60, 350, 215), ["Usuario", "Tkinter"])
    box(draw, (475, 60, 800, 215), ["Mensajes JSON", "stdin / stdout"])
    box(draw, (925, 60, 1250, 215), ["Motor C++", "Repository"])
    box(draw, (1350, 60, 1570, 215), ["Archivo", "único CSV"])
    arrow(draw, (355, 135), (468, 135))
    arrow(draw, (805, 135), (918, 135))
    arrow(draw, (1255, 135), (1343, 135))
    box(draw, (25, 315, 550, 475), ["Fuentes", "URL / PDF / CSV / Office"])
    box(draw, (690, 315, 1215, 475), ["Importador Python", "vista previa y validación"])
    arrow(draw, (555, 395), (683, 395))
    arrow(draw, (950, 308), (950, 225))
    draw.text((1230, 380), "Motor Python autónomo", font=font(24), fill=GRAY)
    draw.text((1230, 418), "misma interfaz", font=font(24), fill=GRAY)
    canvas.save(path)


def model_image(path: Path) -> None:
    canvas = Image.new("RGB", (1600, 730), "white")
    draw = ImageDraw.Draw(canvas)
    box(draw, (575, 260, 1025, 460), ["PRODUCTO", "hecho contado una vez"], (220, 235, 245))
    box(draw, (15, 270, 395, 445), ["GRUPO", "grupo_producto"])
    box(draw, (1205, 270, 1585, 445), ["INVESTIGADOR", "autoría"])
    box(draw, (575, 15, 1025, 175), ["TIEMPO", "año / fecha"])
    box(draw, (390, 555, 790, 715), ["TIPO", "familia / tipología"])
    box(draw, (815, 555, 1240, 715), ["ESTADO", "validación automática"])
    arrow(draw, (400, 357), (565, 357))
    arrow(draw, (1195, 357), (1035, 357))
    arrow(draw, (800, 180), (800, 250))
    arrow(draw, (635, 545), (700, 470))
    arrow(draw, (1015, 545), (925, 470))
    canvas.save(path)


def use_case_image(path: Path) -> None:
    labels = [
        "Cargar o iniciar sin datos",
        "Gestionar registros y relaciones",
        "Asignar categoría a un producto",
        "Importar ficha GrupLAC",
        "Importar ficha CvLAC",
        "Filtrar y consultar estadísticas",
        "Calcular indicador de grupo",
        "Calcular indicador de investigador",
        "Revisar cola y deshacer",
        "Guardar o recuperar el archivo",
    ]
    canvas = Image.new("RGB", (1600, 16 + len(labels) * 92), "white")
    draw = ImageDraw.Draw(canvas)
    middle = canvas.height // 2
    draw.ellipse((120, middle - 180, 200, middle - 100), outline=NAVY, width=5)
    draw.line((160, middle - 100, 160, middle + 85), fill=NAVY, width=5)
    draw.line((75, middle - 45, 245, middle - 45), fill=NAVY, width=5)
    draw.line((160, middle + 85, 85, middle + 195), fill=NAVY, width=5)
    draw.line((160, middle + 85, 235, middle + 195), fill=NAVY, width=5)
    draw.text((60, middle + 215), "Usuario", font=font(30, True), fill=NAVY)
    for index, line in enumerate(labels):
        top = 8 + index * 92
        xy = (475, top, 1540, top + 76)
        draw.ellipse(xy, fill=PALE, outline=NAVY, width=3)
        width = draw.textbbox((0, 0), line, font=font(29, True))[2]
        draw.text((475 + (1065 - width) // 2, top + 20),
                  line, font=font(29, True), fill=NAVY)
        arrow(draw, (280, middle - 20), (465, top + 38))
    canvas.save(path)


def dashed(draw: ImageDraw.ImageDraw, start: tuple[int, int], end: tuple[int, int]) -> None:
    (x1, y), (x2, _) = start, end
    step = 1 if x2 > x1 else -1
    for x in range(x1, x2, 26 * step):
        draw.line((x, y, x + 14 * step if abs(x2 - x) > 14 else x2, y), fill=NAVY, width=4)


def sequence_image(path: Path, title: str, lifelines: list[str],
                   messages: list[tuple[int, int, str] | tuple[int, int, str, bool]]) -> None:
    """Diagrama de secuencia UML: líneas de vida y mensajes numerados; True marca una respuesta."""
    step, first = 64, 190
    height = first + len(messages) * step + 30
    canvas = Image.new("RGB", (1600, height), "white")
    draw = ImageDraw.Draw(canvas)
    draw.text((30, 16), title, font=font(30, True), fill=NAVY)
    span = 1600 // len(lifelines)
    xs = [span // 2 + index * span for index in range(len(lifelines))]
    for x, name in zip(xs, lifelines):
        box(draw, (x - span // 2 + 12, 70, x + span // 2 - 12, 140), [name])
        for y in range(140, height - 20, 24):
            draw.line((x, y, x, min(y + 12, height - 20)), fill=GRAY, width=3)
    for index, message in enumerate(messages, start=1):
        source, target, label = message[:3]
        response = len(message) > 3 and bool(message[3])
        y = first + (index - 1) * step + 32
        start, end = (xs[source], y), (xs[target], y)
        if response:
            dashed(draw, start, end)
            sign = 1 if end[0] > start[0] else -1
            draw.line((end, (end[0] - 18 * sign, y - 10)), fill=NAVY, width=4)
            draw.line((end, (end[0] - 18 * sign, y + 10)), fill=NAVY, width=4)
        else:
            arrow(draw, start, end)
        text = f"{index}. {label}"
        left = min(start[0], end[0]) + 12
        bounds = draw.textbbox((left, y - 34), text, font=font(21))
        draw.rectangle((bounds[0] - 4, bounds[1] - 2, bounds[2] + 4, bounds[3] + 2), fill="white")
        draw.text((left, y - 34), text, font=font(21), fill=NAVY)
    canvas.save(path)


def shade(cell, color: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), color)
    tc_pr.append(shd)


def borders(cell) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    edges = OxmlElement("w:tcBorders")
    for edge in ("top", "left", "bottom", "right"):
        item = OxmlElement("w:" + edge)
        item.set(qn("w:val"), "single")
        item.set(qn("w:sz"), "4")
        item.set(qn("w:color"), "D9D9D9")
        edges.append(item)
    tc_pr.append(edges)


def margins(cell) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    cell_mar = OxmlElement("w:tcMar")
    for edge, value in (("top", "90"), ("left", "100"), ("bottom", "90"), ("right", "100")):
        node = OxmlElement("w:" + edge)
        node.set(qn("w:w"), value)
        node.set(qn("w:type"), "dxa")
        cell_mar.append(node)
    tc_pr.append(cell_mar)


def add_table(doc: Document, headers: list[str], rows: list[list[str]],
              widths: list[float]) -> None:
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    for row_index, values in enumerate([headers] + rows):
        row = table.rows[0] if row_index == 0 else table.add_row()
        tr_pr = row._tr.get_or_add_trPr()
        tr_pr.append(OxmlElement("w:cantSplit"))
        for col_index, value in enumerate(values):
            cell = row.cells[col_index]
            cell.width = Cm(widths[col_index])
            cell.text = str(value)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            borders(cell)
            margins(cell)
            if row_index == 0:
                shade(cell, "1C3955")
            elif row_index % 2 == 0:
                shade(cell, "F3F7FA")
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_after = Pt(0)
                paragraph.paragraph_format.line_spacing = 1.05
                if col_index == 0:
                    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                for run in paragraph.runs:
                    run.font.name = "Liberation Sans"
                    run.font.size = Pt(8.4 if len(headers) > 2 else 8.7)
                    if row_index == 0:
                        run.font.bold = True
                        run.font.color.rgb = RGBColor(255, 255, 255)
    header = table.rows[0]._tr
    tr_pr = header.get_or_add_trPr()
    repeat = OxmlElement("w:tblHeader")
    repeat.set(qn("w:val"), "true")
    tr_pr.append(repeat)
    doc.add_paragraph()


def paragraph(doc: Document, text: str, style: str | None = None) -> None:
    doc.add_paragraph(text, style=style)


def bullet(doc: Document, text: str) -> None:
    doc.add_paragraph(text, style="List Bullet")


def heading(doc: Document, text: str, level: int = 1) -> None:
    doc.add_heading(text, level=level)


def figure(doc: Document, path: Path, caption: str) -> None:
    image = doc.add_paragraph()
    image.alignment = WD_ALIGN_PARAGRAPH.CENTER
    image.add_run().add_picture(str(path), width=Cm(14.5))
    label = doc.add_paragraph(caption)
    label.alignment = WD_ALIGN_PARAGRAPH.CENTER
    label.style = "Caption"


def setup(doc: Document) -> None:
    section = doc.sections[0]
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(1.8)
    section.bottom_margin = Cm(1.8)
    section.left_margin = Cm(2.1)
    section.right_margin = Cm(2.1)
    normal = doc.styles["Normal"]
    normal.font.name = "Liberation Sans"
    normal.font.size = Pt(9.2)
    normal.font.color.rgb = RGBColor(25, 25, 25)
    normal.paragraph_format.space_after = Pt(5)
    normal.paragraph_format.line_spacing = 1.12
    for name, size, above, below in (
        ("Title", 20, 0, 13), ("Heading 1", 13, 12, 6),
        ("Heading 2", 10.7, 11, 5),
    ):
        style = doc.styles[name]
        style.font.name = "Liberation Sans"
        style.font.size = Pt(size)
        style.font.bold = name != "Title"
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.paragraph_format.space_before = Pt(above)
        style.paragraph_format.space_after = Pt(below)
        style.paragraph_format.keep_with_next = True
        ppr = style.element.get_or_add_pPr()
        for border in ppr.findall(qn("w:pBdr")):
            ppr.remove(border)
    doc.styles["Title"].font.bold = True
    subtitle = doc.styles["Subtitle"]
    subtitle.font.name = "Liberation Sans"
    subtitle.font.size = Pt(9)
    subtitle.font.italic = False
    subtitle.font.color.rgb = RGBColor(0, 0, 0)
    subtitle.paragraph_format.space_after = Pt(9)
    doc.styles["Caption"].font.name = "Liberation Sans"
    doc.styles["Caption"].font.size = Pt(8.5)
    doc.styles["Caption"].font.color.rgb = RGBColor(70, 70, 70)
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    footer.style.font.name = "Liberation Sans"
    footer.style.font.size = Pt(8)
    footer.style.font.color.rgb = RGBColor(90, 90, 90)
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    footer._p.append(field)


def number(value: int) -> str:
    return f"{value:,}".replace(",", ".")


def decimal(value: float, digits: int = 6) -> str:
    return f"{value:.{digits}f}".replace(".", ",")


def engine_facts() -> dict[str, Any]:
    """Lee el catálogo, las estimaciones y las cifras de la captura desde el motor Python."""
    pea = runpy.run_path(str(ROOT / "src/python/Taller2_REMR.py"))
    repo = pea["load_repository"](ROOT / "data/pea_upc.csv")
    products = repo.rows("productos")
    years = [int(row["anio"]) for row in products if row["anio"]]
    statistics = repo.statistics(limit=0)
    indicators = []
    for kind, key, label in (("grupos", "G-00000000002099", "GISICO (GrupLAC)"),
                             ("investigadores", "I-0000494917", "Adith Bismarck Pérez Orozco (CvLAC)")):
        side = "left" if kind == "grupos" else "right"
        relation = "grupos_productos" if kind == "grupos" else "autorias"
        linked = sum(1 for row in repo.related(relation, key, side) if row["activo"] == "1")
        indicators.append((label, linked, repo.indicator(kind, key, end=INDICATOR_CUTOFF)))
    return {
        "pea": pea,
        "counts": {kind: len(repo.rows(kind)) for kind in pea["ALL_FIELDS"]},
        "active_memberships": sum(1 for row in repo.rows("membresias") if row["activo"] == "1"),
        "people_with_category": sum(1 for row in repo.rows("investigadores") if row["categoria"]),
        "years": (min(years), max(years)),
        "validation": statistics["por_validacion"],
        "rules": statistics["por_regla"],
        "categorized": sum(1 for row in products if row["categoria"]),
        "estimated": sum(1 for row in products if row["categoria_origen"] == "estimada"),
        "manual": sum(1 for row in products if row["categoria_origen"] == "manual"),
        "indicators": indicators,
    }


def build() -> None:
    facts = engine_facts()
    pea = facts["pea"]
    catalog = pea["MODEL_2024"]
    coefficients = pea["MODEL_COEFFICIENTS"]
    doc = Document()
    doc.core_properties.title = "Especificación técnica del Programa Estadístico de Análisis de Investigación"
    doc.core_properties.subject = "Modelo de datos, estructuras, algoritmos y contratos de PEA-i"
    doc.core_properties.author = ""
    doc.core_properties.last_modified_by = ""
    setup(doc)
    doc.add_paragraph("UNIVERSIDAD POPULAR DEL CESAR  ·  INGENIERÍA DE SISTEMAS", style="Subtitle")
    doc.add_paragraph("Especificación técnica del Programa Estadístico de Análisis de Investigación",
                      style="Title")
    paragraph(doc, "PEA-i administra grupos de investigación, investigadores, productos y planes de la "
                   "Universidad Popular del Cesar. El modelo conserva cada entidad una sola vez y representa "
                   "sus relaciones mediante multilistas. Esta especificación define las entradas y salidas, "
                   "estructuras, reglas de integridad, consultas estadísticas, persistencia e interoperabilidad "
                   "de las soluciones C++ y Python, y el cálculo del indicador de producción científica de "
                   "grupos e investigadores según el modelo de medición Minciencias 2024 (apartado 9).")

    heading(doc, "1 Objetivo y alcance")
    paragraph(doc, "Las operaciones de dominio son alta, consulta, edición, desactivación, reactivación y baja "
                   "de entidades y relaciones. C++ ofrece menús de consola y un protocolo local de solicitudes; "
                   "Python ofrece una interfaz Tkinter y conserva un motor autónomo con el mismo esquema de datos. "
                   "Tkinter usa el motor C++ en modo API: toma un binario compilado vigente, en Windows el "
                   "ejecutable incluido si sus huellas coinciden, o lo compila con CMake. Si no lo consigue, "
                   "ejecuta el motor Python y lo indica en la barra de estado.")
    paragraph(doc, "Las salidas comprenden fichas y relaciones consultables, revisión pendiente, historial de "
                   "deshacer, distribuciones descriptivas, el indicador de producción de un grupo o de un "
                   "investigador y un archivo único de datos. La adquisición "
                   "externa usa fichas públicas GrupLAC o CvLAC, CSV, Excel, Word y PDF con texto extraíble. "
                   "El conjunto incluido reúne fichas públicas de 66 grupos UPC y una ficha CvLAC; no es un "
                   "censo institucional.")

    heading(doc, "2 Arquitectura y casos de uso")
    paragraph(doc, "La interfaz Tkinter envía una solicitud JSON por línea a la entrada estándar de "
                   "pea_cpp --api y recibe una respuesta JSON por la salida estándar. No intervienen puertos "
                   "ni servicios de red en la conexión entre procesos. Repository mantiene listas, "
                   "multilistas, pila y cola en memoria y coordina el guardado del archivo. Python obtiene y "
                   "previsualiza las fuentes web; los registros confirmados se incorporan mediante "
                   "operaciones del motor activo.")
    with tempfile.TemporaryDirectory(prefix="pea-doc-") as folder:
        temp = Path(folder)
        architecture = temp / "arquitectura.png"
        model = temp / "modelo.png"
        use_case = temp / "casos.png"
        architecture_image(architecture)
        model_image(model)
        use_case_image(use_case)
        figure(doc, architecture, "Figura 1. Procesos y flujo de datos de la aplicación.")
        paragraph(doc, "La ventana Tkinter organiza el trabajo en una barra lateral con tres grupos: Análisis "
                       "(Dashboard y Gráficos), Registros (Grupos, Investigadores, Productos, Planes y "
                       "Relaciones: integrantes, autorías y grupos-productos) y Seguimiento (Cola de revisión). "
                       "Admite tema claro u oscuro y Ctrl+Z para deshacer.")
        heading(doc, "Protocolo entre procesos", 2)
        paragraph(doc, "Cada solicitud es un objeto JSON en una línea con la acción y sus parámetros como "
                       "cadenas; la respuesta es {\"ok\":true,\"result\":…} o {\"ok\":false,\"error\":…}. Al conectarse, la interfaz "
                       "envía ping y exige la respuesta {\"backend\":\"cpp\",\"protocol\":\"4\"}; con otra "
                       "versión no usa el proceso. pea_cpp acepta --data ARCHIVO para abrir un archivo al iniciar.")
        add_table(doc, ["Grupo", "Acciones del protocolo 4"], [
            ["Sesión y archivo", "ping, state, reset, load, save, shutdown"],
            ["Consulta", "rows, page, get, related, related_page, summary, field_usage, product_issues, statistics"],
            ["Modelo 2024", "catalog (169 subtipos con clase, ventana y pesos), indicator (tipo, ID, años inicial y final, validación)"],
            ["Cambios", "new_id, create, update, delete, toggle, preview_csv, import_csv"],
            ["Cola de revisión", "queue_rows, queue_page, queue_front, enqueue_review, enqueue_rejected, process_review, discard_review"],
            ["Historial", "undo, clear_history"],
        ], [3.6, 13.2])
        paragraph(doc, "Los casos de uso cubren inicio vacío o desde el archivo de datos, mantenimiento de entidades y "
                       "vínculos, asignación de la categoría del modelo a un producto, incorporación de fichas "
                       "GrupLAC y CvLAC, consulta estadística, cálculo del indicador de un grupo o de un "
                       "investigador, revisión, deshacer y guardado.")
        figure(doc, use_case, "Figura 2. Casos de uso principales del usuario.")

        heading(doc, "3 Modelo de información")
        paragraph(doc, "El producto es la unidad que se cuenta en las estadísticas. Cada grupo, investigador, "
                       "producto o plan tiene un ID estable. Las multilistas vinculan entidades sin duplicarlas. "
                       "Membresías unen grupo e investigador; autorías unen investigador y producto; "
                       "grupos_productos unen grupo y producto. Un plan pertenece a un grupo.")
        figure(doc, model, "Figura 3. Dimensiones del hipercubo lógico alrededor del producto.")

    heading(doc, "Variables de entrada y salida", 2)
    add_table(doc,
              ["Registro", "Entrada principal", "Salida y relación"],
              [
                  ["Grupo", "ID, nombre, código GrupLAC, unidad, categoría, líneas, estado y fuente.",
                   "Ficha y productos vinculados; integrantes y planes."],
                  ["Investigador", "ID, nombre, CvLAC, afiliación, categoría, contacto público, estado y fuente.",
                   "Ficha, grupos y productos de autoría."],
                  ["Producto", "ID, título, año o fecha, familia, tipología, categoría del modelo 2024 y su origen "
                               "(estimada o manual), observación, DOI, URL y fuente; la validación se calcula.",
                   "Ficha, grupos, autores, clase, peso relativo y valor en puntos de su categoría; "
                   "contribución a estadísticas e indicadores."],
                  ["Plan", "ID, grupo, plazo, objetivo, indicador, meta y actividad.",
                   "Plan consultable dentro del grupo."],
              ], [2.5, 7.3, 7.0])
    paragraph(doc, "Las salidas estadísticas son total de productos únicos, conteos por año, tipología, "
                   "categoría, validación y reglas incumplidas, más listas paginadas. Los campos vacíos no se "
                   "completan con valores supuestos. Las variables del indicador de producción se describen en "
                   "el apartado 9 y las de persistencia en el diccionario del apartado 13.")

    heading(doc, "4 Estructuras de datos y complejidad")
    add_table(doc,
              ["Estructura", "Uso concreto", "Costos relevantes"],
              [
                  ["Lista doble", "Entidades: grupos, investigadores, productos y planes. Nodos anterior/siguiente; índices auxiliares por ID.",
                   "Inserción al final O(1); recorrido O(n); búsqueda por ID O(1) promedio con índice."],
                  ["Multilista", "Membresías, autorías y vínculos grupo-producto recorridos desde cualquiera de sus extremos.",
                   "Inserción O(1) promedio con índice; recorrer vecinos O(k)."],
                  ["Pila", "Deshacer las últimas 30 operaciones. C++ guarda las operaciones inversas de cada cambio; Python guarda instantáneas y también aplica las inversas que C++ dejó en el historial.",
                   "Apilar y desapilar O(1); el tamaño del registro de deshacer depende del cambio."],
                  ["Cola", "Revisiones pendientes de productos en orden de llegada; se encolan uno a uno o todos los rechazados en una sola acción.",
                   "Encolar, consultar frente y desencolar O(1); enviar los rechazados O(P)."],
              ], [2.8, 8.0, 6.0])
    paragraph(doc, "Los nodos son la fuente de verdad. Los diccionarios y mapas sirven de índices auxiliares. "
                   "La consulta estadística recorre productos activos y vínculos pertinentes; no existe una "
                   "matriz densa con todas las combinaciones posibles. Python materializa temporalmente los "
                   "productos de la vista; C++ recorre la lista principal y limita la página devuelta.")

    heading(doc, "5 Operaciones y reglas de integridad")
    bullet(doc, "Alta, búsqueda, consulta, edición, desactivación, reactivación y baja de entidades y relaciones; los IDs y los extremos de una relación no se editan.")
    bullet(doc, "No se elimina en cascada un grupo, investigador o producto con relaciones pendientes; primero se resuelven esos vínculos o se desactiva el registro.")
    bullet(doc, "Los códigos GrupLAC, CvLAC y DOI no vacíos son únicos. El DOI se compara sin distinguir mayúsculas ni el prefijo doi.org. Las fechas siguen AAAA-MM-DD; el año del producto debe concordar con su fecha.")
    bullet(doc, "La validación de un producto es automática: queda validado si tiene año, tipología, una autoría activa de un investigador activo o un vínculo activo con un grupo activo, y una URL verificable o un DOI válido. Basta el autor o el grupo porque GrupLAC publica eventos, proyectos, consultorías y redes sin autores, y CvLAC productos que ningún grupo declara. Un DOI registrado con formato inválido también lo rechaza. Se recalcula al cambiar el producto, sus vínculos o el estado de un grupo o investigador enlazado, y no equivale a la clasificación oficial de Minciencias.")
    bullet(doc, "La categoría de un producto es uno de los 169 subtipos del modelo Minciencias 2024 (por ejemplo ART_A1 o TP_B) o queda vacía; un código desconocido se rechaza. Su origen es «estimada» cuando la asigna el importador y «manual» cuando la elige una persona; cambiar la categoría a mano la marca como manual y vaciarla borra el origen. La categoría de grupos e investigadores sigue siendo su clasificación publicada.")
    bullet(doc, "La cola conserva trabajos pendientes y no admite dos veces el mismo producto; la pila permite deshacer. Ambas se guardan junto con las entidades y relaciones.")

    heading(doc, "6 Persistencia e interoperabilidad")
    paragraph(doc, "Todos los datos se guardan en un solo archivo CSV UTF-8, data/pea_upc.csv, dividido en "
                   "secciones: manifest, grupos, investigadores, productos, planes, membresías, autorías, "
                   "vínculos grupo-producto, cola de revisión e historial, siempre en ese orden. Cada sección "
                   "empieza con una fila de una sola celda #nombre, sigue su encabezado y luego sus filas. "
                   "Ambos motores leen el archivo de arriba abajo y crean cada registro al leerlo: las "
                   "relaciones aparecen después de las entidades que enlazan. El manifiesto declara la "
                   "versión 5, que añade categoria y categoria_origen a los productos. Un archivo de la "
                   "versión 4 se abre con sus productos sin categoría y se guarda como versión 5; su historial "
                   "de deshacer sigue siendo aplicable. Las carpetas de las versiones 1 a 3, con un CSV por "
                   "tabla, se abren por la carpeta o su manifest.csv y se guardan como archivo único; la "
                   "categoría libre de un producto de las versiones 1 y 2 se conserva en la observación. "
                   "El formato neutraliza los "
                   "valores que una hoja de cálculo interpretaría como fórmulas y recupera su valor original "
                   "al cargar. Se valida la integridad referencial. El guardado escribe un temporal, conserva "
                   "en .backup el archivo previo y lo reemplaza de una sola vez.")
    paragraph(doc, "Las dos aplicaciones pueden abrir por turnos el mismo archivo. No existe bloqueo de "
                   "escritura multiusuario; abrir simultáneamente un archivo editable puede producir "
                   "conflictos. La interfaz abre data/pea_upc.csv al iniciar y guarda en él; cada "
                   "importación se guarda de inmediato, de modo que los datos crecen con cada fuente. Si el archivo "
                   "no puede leerse, la interfaz ofrece abrir la copia .backup.")

    heading(doc, "7 Adquisición e importación")
    paragraph(doc, "Ambos programas importan CSV del esquema PEA-i con vista previa y conteo de filas "
                   "aceptables o rechazadas. Python también admite Excel XLSX y Word DOCX mediante "
                   "dependencias opcionales, y consulta URL públicas HTTP/HTTPS de HTML, texto, "
                   "CSV o PDF de texto y muestra su contenido antes de crear registros. La extracción "
                   "estructurada se limita a fichas GrupLAC/CvLAC y metadatos explícitos de artículos. "
                   "Las páginas de estructura desconocida se limitan a vista previa y captura supervisada. "
                   "Volver a importar una ficha GrupLAC agrega lo nuevo y actualiza los campos que la fuente "
                   "publica (clasificación y fecha de creación del grupo, período de los integrantes, tipología "
                   "y categoría estimada de productos), sin borrar registros ni reemplazar las observaciones de "
                   "revisión ni una categoría elegida manualmente.")
    paragraph(doc, "De GrupLAC se leen la producción bibliográfica y técnica, la apropiación social, la "
                   "divulgación pública (eventos, redes, contenidos, consultorías, informes), la producción en "
                   "arte, los trabajos dirigidos, los proyectos y los programas y cursos de formación. En las "
                   "secciones nuevas solo se incorporan las entradas que tienen una categoría del modelo. De "
                   "CvLAC se leen artículos, libros, capítulos, documentos de trabajo, otros artículos, software, "
                   "innovaciones, empresas de base tecnológica y trabajos dirigidos; cada producto queda con la "
                   "autoría del investigador y, si ya existe por su DOI o por título y año, se fusiona con el "
                   "registro importado desde GrupLAC. La fecha de creación del grupo se toma del año y mes de "
                   "formación de la ficha.")
    paragraph(doc, "La importación de PDF de texto requiere pdftotext de Poppler. Los escaneos sin texto "
                   "seleccionable requieren OCR externo. C++ dispone de importación CSV autónoma.")

    heading(doc, "Controles de seguridad de las fuentes", 2)
    paragraph(doc, "La descarga limita la respuesta a 8 MB, permite solo HTTP y HTTPS en los puertos 80 y "
                   "443, rechaza destinos locales o privados y verifica cada redirección. La conexión usa "
                   "una dirección pública previamente validada, sin proxy heredado del entorno. El tiempo "
                   "de espera por operación es 15 segundos; no constituye un límite total de duración.")
    paragraph(doc, "Los archivos locales PDF, XLSX y DOCX tienen un límite de 32 MB. Los documentos Office "
                   "se inspeccionan antes de extraerlos, con un máximo de 10.000 entradas y 128 MB "
                   "de tamaño declarado sin comprimir. El conversor PDF tiene 25 segundos para extraer "
                   "hasta cinco millones de caracteres y límites adicionales de CPU y memoria en POSIX. "
                   "Cada tabla admite hasta 500.000 filas y 100.000 caracteres por campo.")
    paragraph(doc, "El ejecutable Windows incluido se usa cuando su SHA-256 y la huella de las fuentes "
                   "coinciden. Esas huellas detectan cambios accidentales, pero no autentican al autor. "
                   "Los parámetros privados de URL se omiten al guardar la procedencia; se retienen solo "
                   "identificadores públicos necesarios para las fichas GrupLAC y CvLAC.")

    heading(doc, "8 Hipercubo lógico y estadísticas")
    paragraph(doc, "El hipercubo se implementa como consulta multidimensional calculada bajo demanda: "
                   "cada producto es un hecho único; grupo e investigador son dimensiones "
                   "relacionales, y año, tipología y validación son atributos de análisis. "
                   "La vista selecciona todos los productos o un grupo, investigador o producto. Después "
                   "se aplican rango de años y estado de validación; el resultado se agrupa por año, "
                   "tipología, validación y reglas incumplidas.")
    paragraph(doc, "Total(vista, filtros) = número de IDs de producto activos distintos que satisfacen "
                   "los vínculos activos y todos los filtros. Un producto unido a varios grupos o autores "
                   "se cuenta una sola vez dentro de la vista. Si se fija un rango de años, los productos sin "
                   "año quedan fuera; sin rango, cuentan en el total y aparecen como Sin dato en la "
                   "distribución por año.")
    paragraph(doc, "No se almacena un cubo OLAP materializado. Cada solicitud selecciona una dimensión "
                   "relacional; la consulta no admite una intersección simultánea de grupo e investigador. "
                   "La consola C++ muestra los conteos en tablas y cifras, con ventanas de todos los años, "
                   "últimos 2, últimos 5 o un rango.")
    paragraph(doc, "En Tkinter, el Dashboard reúne los filtros (vista, ventana de observación y validación), "
                   "cuatro indicadores (productos, validados, rechazados y sin año), el panel del indicador de "
                   "producción cuando la vista es un grupo o un investigador (apartado 9), un histograma por "
                   "año, la distribución por tipología y la tabla paginada de productos de la vista. La página "
                   "Gráficos amplía cinco distribuciones: por año en histograma, por tipología, por categoría "
                   "del modelo, por validación y por reglas incumplidas.")

    heading(doc, "9 Indicador de producción científica")
    paragraph(doc, "El indicador sigue el Anexo 1 de la Convocatoria Nacional de Reconocimiento y Medición "
                   "de Grupos 2024 de Minciencias (M601PR04G01, versión 02, apartados 2.2.9, 2.2.10, 3.3 a 3.10). "
                   "Para un grupo se calcula con los productos vinculados al grupo; para un investigador, con la "
                   "misma fórmula y los productos de su autoría:")
    formula = doc.add_paragraph()
    formula.alignment = WD_ALIGN_PARAGRAPH.CENTER
    formula.add_run("Ind = 3,7·Nc.Top + 2,3·Nc.A + 0,4·Nc.B + 1,5·AP + 0,5·DPC + 1·Fr.A + 0,2·Fr.B "
                    "+ 0,1·Cohe + 0,3·Coop").bold = True
    paragraph(doc, "Cada término es un índice entre 0 y 1, de modo que el indicador vale como máximo 10. "
                   "Ambos motores implementan el mismo cálculo y una prueba de integración compara sus "
                   "resultados.")

    heading(doc, "Variables de entrada y salida del modelo", 2)
    add_table(doc, ["Variable", "Origen o cálculo", "Uso"], [
        ["Entrada: categoría del producto", "Uno de los 169 subtipos del catálogo; estimada por el importador o elegida a mano.",
         "Define la clase, la ventana, el peso relativo y el peso global."],
        ["Entrada: año del producto", "Campo anio.", "Decide si el producto cae en la ventana del subtipo."],
        ["Entrada: vínculos activos", "grupos_productos (grupo) o autorías (investigador).", "Seleccionan los productos de la entidad."],
        ["Entrada: membresías activas", "Integrantes activos de grupos activos.", "Alimentan la cohesión y la cooperación."],
        ["Entrada: fecha de creación del grupo", "Año y mes de formación en GrupLAC.", "Limita el período de observación."],
        ["Entrada: ventana y validación", "Año de corte (por defecto, el actual), año inicial opcional y filtro validado/rechazado.",
         "Recortan el período y los productos contados."],
        ["Salida: λ por subtipo", "ln(productos del subtipo / período + 1).", "Conteo escalado del subtipo."],
        ["Salida: valor por clase", "Σ λ × peso global de los subtipos de la clase.", "Producción de la clase."],
        ["Salida: Cohe y Coop", "Coautorías internas y grupos relacionados por producto.", "Colaboración."],
        ["Salida: índices y aportes", "Valor / máximo de la población; aporte = coeficiente × índice.", "Términos de la fórmula."],
        ["Salida: indicador", "Suma de los nueve aportes.", "Resultado entre 0 y 10."],
        ["Salida: conteos", "Productos contados, estimados, sin categoría y tamaño de la población.", "Contexto del resultado."],
    ], [4.0, 7.2, 5.6])

    heading(doc, "Cálculo", 2)
    bullet(doc, "Peso global del subtipo = peso de la más alta calidad de su familia (Tabla 6) × peso relativo "
                "(apartado 2.2.9). Ejemplos del Anexo: ART_A2 = 100 × 5,5 = 550; CAP_LIB_B = 60 × 8 = 480; "
                "PA1 = 500 × 10 = 5.000. La interfaz llama «valor en puntos» al peso global.")
    bullet(doc, "Período del subtipo = mínimo entre su ventana (5, 7 o 10 años, Tabla 5), la edad del grupo al "
                "corte (corte − año de creación + 1) y, si hay año inicial, corte − inicio + 1. Cuenta un "
                "producto activo, vinculado por una relación activa, con categoría y con año dentro de "
                "[corte − período + 1, corte]. Un grupo creado después del corte no cuenta productos.")
    bullet(doc, "λ = ln(n / período + 1); valor del subtipo = λ × peso global; valor de la clase = suma de sus subtipos.")
    bullet(doc, "Cohesión (sobre los P productos contados de clases TOP, A y B) = máx(0, Autores / P − 1), "
                "donde Autores suma, por producto, los autores activos que son integrantes activos del grupo. "
                "Para un investigador cuentan él mismo y los integrantes de sus grupos activos.")
    bullet(doc, "Cooperación = máx(0, Σ grupos relacionados por producto / P − 1); los grupos relacionados son "
                "los grupos activos de los autores del producto, más el propio grupo cuando la entidad es un grupo.")
    bullet(doc, "Índice = valor / máximo del mismo término entre todas las entidades activas del mismo tipo, "
                "calculadas con los mismos filtros (0 si el máximo es 0). Aporte = coeficiente × índice.")
    add_table(doc, ["Término", "Clase", "Coeficiente", "Subtipos", "Ejemplos (peso global)"], [
        [CLASS_TERMS[name], CLASS_NAMES[name], decimal(coefficients[name], 1),
         str(sum(1 for item in catalog.values() if item.clase == name)),
         ", ".join(f"{item.codigo} ({number(round(item.peso_global))})"
                   for item in [item for item in catalog.values() if item.clase == name][:3])]
        for name in CLASS_TERMS
    ] + [["Cohe", "Índice de cohesión", "0,1", "—", "Coautoría entre integrantes"],
         ["Coop", "Índice de cooperación", "0,3", "—", "Coautoría entre grupos"]],
              [1.6, 5.4, 1.8, 1.6, 6.4])

    heading(doc, "Decisiones de interpretación", 2)
    bullet(doc, "Población de referencia: Minciencias normaliza contra los grupos de la misma gran área OCDE "
                "del país. PEA-i normaliza contra las entidades activas cargadas (los grupos UPC o los "
                "investigadores del archivo), por lo que el resultado es relativo a esa población y no una "
                "clasificación oficial. Con una sola entidad cargada, cada término no nulo vale 1.")
    bullet(doc, "Categorías estimadas: GrupLAC y CvLAC no publican la categoría de cada producto, que depende "
                "de índices externos (JCR, SJR, Publindex). El importador asigna la categoría más baja de la "
                "familia que corresponde a la sección y la tipología, la marca como estimada y la reemplaza "
                "en cada reimportación; una categoría elegida a mano nunca se sobrescribe.")
    bullet(doc, "La validación automática no excluye productos del indicador salvo que se filtre por "
                "validado o rechazado. Los productos sin año o sin categoría no cuentan; los sin categoría "
                "de los últimos diez años se informan aparte.")
    estimates = pea["CATEGORY_ESTIMATES"]
    order = list(CLASS_TERMS)
    estimate_rows = sorted(
        ([section if section != "*" else "Cualquiera", kind if kind != "*" else "Cualquiera", code,
          CLASS_TERMS[catalog[code].clase]] for (section, kind), code in estimates.items()),
        key=lambda row: (order.index(catalog[row[2]].clase), row[0], row[1]))
    paragraph(doc, "Tabla de estimación usada por el importador (sección GrupLAC, tipología declarada y "
                   "categoría asignada); las entradas de CvLAC se traducen a la sección GrupLAC equivalente:")
    add_table(doc, ["Sección GrupLAC", "Tipología", "Categoría", "Término"], estimate_rows, [6.6, 5.6, 2.5, 2.1])

    heading(doc, "Presentación y resultados", 2)
    paragraph(doc, "En Tkinter, al elegir la vista Grupo o Investigador en el Dashboard aparece el panel "
                   "«Ind. x de 10» con la población de referencia, los productos contados, estimados y sin "
                   "categoría, una barra por término (coeficiente y aporte) y el botón «Detalle por subtipo» "
                   "con cantidad, período, λ, peso global y valor de cada subtipo. Los filtros de ventana y "
                   "validación del Dashboard se aplican al indicador. El formulario de producto permite elegir "
                   "la categoría y muestra su clase, peso relativo, valor en puntos y ventana. En la consola "
                   "C++, la opción 12 calcula el indicador de un grupo o de un investigador y lista la tabla "
                   "de pesos del modelo.")
    add_table(doc, ["Entidad", "Productos vinculados", "Contados", "Estimados", "Sin categoría", "Población", "Indicador"], [
        [label, number(linked), number(result["productos"]), number(result["estimadas"]),
         number(result["sin_categoria"]), number(result["poblacion"]), decimal(result["indicador"])]
        for label, linked, result in facts["indicators"]
    ], [4.4, 2.2, 1.7, 1.8, 1.9, 2.0, 2.8])
    paragraph(doc, f"Cifras calculadas sobre data/pea_upc.csv con corte {INDICATOR_CUTOFF}, sin año inicial "
                   "ni filtro de validación. Los productos vinculados fuera de las ventanas de observación no "
                   "se cuentan.")
    with tempfile.TemporaryDirectory(prefix="pea-doc-") as folder:
        temp = Path(folder)
        group_sequence = temp / "secuencia_grupo.png"
        person_sequence = temp / "secuencia_investigador.png"
        sequence_image(group_sequence, "Indicador de un grupo desde su URL GrupLAC",
                       ["Usuario", "Tkinter", "Importador", "GrupLAC", "Motor C++"], [
                           (0, 1, "Importar → URL pública (ficha GrupLAC)"),
                           (1, 2, "fetch_web_page(url)"),
                           (2, 3, "GET visualizagr.jsp?nro=…"),
                           (3, 2, "HTML de la ficha", True),
                           (2, 1, "Vista previa: grupo, integrantes, productos con categoría estimada", True),
                           (0, 1, "Importar datos detectados"),
                           (1, 4, "create / update: grupo, investigadores, productos, vínculos"),
                           (1, 4, "save(data/pea_upc.csv)"),
                           (0, 1, "Dashboard: vista Grupo, ventana y validación"),
                           (1, 4, "indicator(grupos, ID, inicio, corte, estado)"),
                           (4, 1, "JSON: clases, Cohe, Coop, subtipos, indicador", True),
                           (1, 0, "Panel «Ind. x de 10» con barras por término", True),
                       ])
        sequence_image(person_sequence, "Indicador de un investigador desde su URL CvLAC",
                       ["Usuario", "Tkinter", "Importador", "CvLAC", "Motor C++"], [
                           (0, 1, "Importar → URL pública (ficha CvLAC)"),
                           (1, 2, "fetch_web_page(url)"),
                           (2, 3, "GET generarCurriculoCv.do?cod_rh=…"),
                           (3, 2, "HTML de la hoja de vida", True),
                           (2, 1, "Vista previa: investigador y productos con categoría estimada", True),
                           (0, 1, "Importar datos detectados"),
                           (1, 4, "create / update: investigador, productos, autorías"),
                           (1, 4, "save(data/pea_upc.csv)"),
                           (0, 1, "Dashboard: vista Investigador, ventana y validación"),
                           (1, 4, "indicator(investigadores, ID, inicio, corte, estado)"),
                           (4, 1, "JSON: clases, Cohe, Coop, subtipos, indicador", True),
                           (1, 0, "Panel «Ind. x de 10» con barras por término", True),
                       ])
        figure(doc, group_sequence, "Figura 4. Diagrama de secuencia: indicador de un grupo desde GrupLAC.")
        figure(doc, person_sequence, "Figura 5. Diagrama de secuencia: indicador de un investigador desde CvLAC.")

    heading(doc, "10 Historias de usuario y aceptación")
    add_table(doc,
              ["Historia", "Necesidad", "Criterio de aceptación"],
              [
                  ["HU01", "Como gestor quiero iniciar vacío o abrir datos guardados.",
                   "La aplicación permite elegir y conserva cambios al cerrar y reabrir."],
                  ["HU02", "Como gestor quiero relacionar grupos, personas y productos.",
                   "Los vínculos se consultan desde ambos extremos sin duplicar entidades."],
                  ["HU03", "Como revisor quiero procesar productos en orden y deshacer errores.",
                   "La cola atiende FIFO; la pila restaura la acción reversible más reciente."],
                  ["HU04", "Como analista quiero filtrar los resultados.",
                   "Las vistas por grupo, investigador y producto respetan años y validación."],
                  ["HU05", "Como gestor quiero importar una fuente pública o CSV.",
                   "La vista previa identifica la fuente, muestra campos y registra aceptación o rechazo."],
                  ["HU06", "Como analista quiero un dashboard con indicadores y gráficos.",
                   "Para la vista y los filtros elegidos muestra totales, histograma por año y barras por tipología, categoría, validación y reglas."],
                  ["HU07", "Como analista quiero el indicador de producción del grupo cuya URL GrupLAC proporciono.",
                   "Tras importar la ficha, la vista Grupo muestra el indicador, sus nueve términos y el detalle por subtipo; C++ y Python dan el mismo valor."],
                  ["HU08", "Como analista quiero el indicador del investigador cuya URL CvLAC proporciono.",
                   "Tras importar la hoja de vida, la vista Investigador calcula el indicador con la misma fórmula sobre sus productos."],
                  ["HU09", "Como gestor quiero corregir la categoría de un producto.",
                   "Elijo uno de los 169 subtipos; el producto queda con origen manual y una reimportación no lo cambia."],
              ], [1.8, 7.2, 7.8])

    heading(doc, "11 Correspondencia entre requisitos y diseño")
    rows = [
        ("R01", "Dos soluciones", "Consola C++17; interfaz Tkinter y motor Python autónomo."),
        ("R02", "Grupos", "Registro con ID, código GrupLAC, categoría, líneas, fuente y estado."),
        ("R03", "Investigadores", "Registro con ID, CvLAC, afiliación, categoría, fuente y estado."),
        ("R04", "Productos", "Registro con fecha, tipología, categoría del modelo, peso relativo y valor en puntos, DOI y URL; validación automática."),
        ("R05", "Integrantes y planes", "Membresía grupo-investigador y plan con grupo_id."),
        ("R06", "Entradas y salidas", "Campos del apartado 3; variables del modelo del apartado 9; CSV y distribuciones del apartado 8."),
        ("R07", "Fuentes externas", "CSV en ambos motores; URL pública (GrupLAC y CvLAC), PDF, XLSX y DOCX en Python."),
        ("R08", "Listas", "Lista doble enlazada por tipo de entidad."),
        ("R09", "Multilistas", "Vínculos indexados y recorribles desde ambos extremos."),
        ("R10", "Pilas", "Deshacer LIFO de hasta 30 acciones, persistido en la sección historial."),
        ("R11", "Colas", "Revisión FIFO persistida en la sección cola_validacion; envío de todos los rechazados en una acción."),
        ("R12", "CRUD y persistencia", "Operaciones de dominio y archivo único versión 5, con lectura de la versión 4 y migración de las carpetas 1 a 3."),
        ("R13", "Categoría y validación", "Categoría del modelo 2024 en productos (estimada o manual) y clasificación en grupos e investigadores; validación de productos por reglas, recalculada en cada cambio."),
        ("R14", "Ventanas temporales", "Dos años, cinco años o intervalo personalizado."),
        ("R15", "Vistas estadísticas", "Total, grupo, investigador o producto sin duplicar IDs."),
        ("R16", "Presentación C++", "Menús, tablas y resúmenes numéricos en consola, incluido el indicador (opción 12)."),
        ("R17", "Presentación Python", "Dashboard con indicadores, panel del indicador de producción, histograma por año y barras por tipología, categoría, validación y reglas incumplidas; tabla paginada."),
        ("R18", "Inicio con o sin datos", "Carga inicial de data/pea_upc.csv, inicio vacío y apertura de otro archivo."),
        ("R19", "Archivos fuente", "Taller2_REMR.cpp y Taller2_REMR.py."),
        ("R20", "Especificación", "Modelo, estructuras, algoritmos, contratos, diagrama de casos de uso y diagramas de secuencia UML."),
        ("R21", "Indicador del grupo", "Fórmula Minciencias 2024 sobre los productos del grupo de la URL GrupLAC (apartado 9)."),
        ("R22", "Indicador del investigador", "Misma fórmula sobre los productos del investigador de la URL CvLAC (apartado 9)."),
    ]
    add_table(doc, ["ID", "Requisito", "Diseño correspondiente"], [list(row) for row in rows],
              [1.25, 4.5, 11.05])

    heading(doc, "12 Restricciones del sistema y alcance de los datos")
    counts = facts["counts"]
    rules = facts["rules"]
    paragraph(doc, "El archivo data/pea_upc.csv reúne fichas públicas GrupLAC de 66 grupos UPC y la ficha "
                   f"CvLAC del investigador del enunciado. Contiene {number(counts['investigadores'])} "
                   f"investigadores, {number(counts['membresias'])} membresías "
                   f"({number(facts['active_memberships'])} vigentes), {number(counts['planes'])} planes, "
                   f"{number(counts['productos'])} productos únicos con año entre {facts['years'][0]} y "
                   f"{facts['years'][1]}, {number(counts['grupos_productos'])} vínculos grupo-producto y "
                   f"{number(counts['autorias'])} autorías. Hay {number(facts['people_with_category'])} "
                   "investigadores con categoría CvLAC capturada. "
                   f"{number(facts['categorized'])} productos tienen categoría del modelo "
                   f"({number(facts['estimated'])} estimadas y {number(facts['manual'])} manuales). "
                   f"La validación automática deja {number(facts['validation'].get('validado', 0))} productos "
                   f"validados y {number(facts['validation'].get('rechazado', 0))} rechazados"
                   + ("; las reglas incumplidas son: " + ", ".join(
                       f"{pea['RULE_LABELS'][code].lower()} {number(count)}" for code, count in rules.items())
                      if rules else "") + ". "
                   "Tres DOI publicados con defectos en GrupLAC se corrigieron al importar: uno partido por un "
                   "espacio se reparó y dos truncados no se guardan como DOI. Las fichas públicas pueden "
                   "cambiar; este conjunto no constituye un censo institucional actualizado.")
    paragraph(doc, "La extracción de URL depende del contenido público entregado por el servidor. Los PDF "
                   "escaneados requieren OCR externo; las páginas que generan sus datos exclusivamente con "
                   "JavaScript o exigen autenticación no suministran registros estructurados por esta vía. "
                   "La persistencia no implementa bloqueo de escritura entre procesos: un archivo editable "
                   "debe abrirse en una instancia a la vez.")
    paragraph(doc, "El motor Python autónomo guarda instantáneas en el historial; el motor C++ conserva "
                   "registros compactos de cambios. La consulta estadística construye agregados bajo demanda "
                   "y no reserva una celda para cada combinación de dimensiones. El tamaño de la memoria "
                   "depende del número de entidades, vínculos y longitud de los campos almacenados.")

    heading(doc, "13 Diccionario de persistencia")
    add_table(doc, ["Sección (en orden)", "Campos guardados"],
              [["#" + name, ", ".join(fields)] for name, fields in pea["DATA_SECTIONS"].items()],
              [4.0, 12.8])
    paragraph(doc, f"El manifiesto declara la versión {pea['SCHEMA_VERSION']}. En los productos, categoria "
                   "es un código del catálogo del modelo o queda vacía, y categoria_origen vale estimada, "
                   "manual o queda vacía.")

    heading(doc, "14 Ejecución y referencias")
    paragraph(doc, "En Windows de 64 bits, con Python 3.10 o posterior y Tkinter, se abre "
                   "Iniciar PEA-i.pyw mediante doble clic; el repositorio incluye pea_cpp.exe. "
                   "Para la consola se ejecuta bin/windows/pea_cpp.exe. En Linux se ejecuta "
                   "python3 src/python/Taller2_REMR.py; CMake compila C++17 si no existe el binario. "
                   "La interfaz abre otro archivo con --data ARCHIVO y fuerza el motor Python con "
                   "--python-backend. La consola acepta pea_cpp --data ARCHIVO; sin él pregunta si iniciar "
                   "vacía o abrir un archivo. Las pruebas automáticas están en tests/.")
    bullet(doc, "Enunciado del taller: docs/ENUNCIADO_TALLER_2.md, Universidad Popular del Cesar, 2026.")
    bullet(doc, "Enunciado del Parcial 2 de Estructura de Datos, Universidad Popular del Cesar, octubre de 2026.")
    bullet(doc, "Minciencias, Anexo 1. Convocatoria nacional para el reconocimiento y medición de grupos de "
                "investigación, desarrollo tecnológico o de innovación y para el reconocimiento de investigadores "
                "del SNCTI 2024 (M601PR04G01, versión 02).")
    bullet(doc, "Esquema y reglas: apartados 5, 6 y 13; protocolo entre procesos: apartado 2; indicador: apartado 9.")
    bullet(doc, "Procedencia de datos: data/README.md y fichas públicas GrupLAC y CvLAC enlazadas allí.")
    bullet(doc, "Código fuente: src/cpp/Taller2_REMR.cpp y src/python/Taller2_REMR.py.")
    bullet(doc, "Videos de presentación: docs/video_taller.mp4 (Taller 2) y docs/video_parcial.mp4 (Parcial 2).")

    heading(doc, "15 Algoritmos principales")
    paragraph(doc, "Consulta estadística: la multilista selecciona los IDs de producto asociados a la "
                   "vista. Se recorren los productos activos, se aplican los filtros y se incrementan "
                   "los conteos. La lista visible se pagina sin alterar el total. Con P productos, R "
                   "vínculos de la vista y D valores distintos de distribución, Python emplea O(P + R) "
                   "tiempo promedio y O(P + R + D) memoria temporal. C++ usa conjuntos y mapas ordenados: "
                   "el tiempo es O(P + R log(R + 1) + P log(R + 1) + P log(D + 1)) y la memoria adicional "
                   "es O(R + L + D), donde L es el tamaño de página solicitado.")
    paragraph(doc, "Guardado: se escriben todas las secciones en un archivo temporal; se conserva el "
                   "archivo anterior en .backup; y se sustituye el archivo de destino de una sola vez. "
                   "Carga: se recorre el archivo en orden; cada marca debe ser la siguiente sección "
                   "esperada y cada encabezado debe coincidir con el esquema; cada fila se valida y se "
                   "inserta al leerla, de modo que los vínculos encuentran sus extremos ya cargados.")
    paragraph(doc, "Revisión: encolar añade al final si el producto no espera ya en la cola; enviar los "
                   "rechazados encola, en una sola acción, cada producto activo rechazado con sus reglas "
                   "incumplidas como motivo. Procesar lee el frente, exige una nota, la guarda en la "
                   "observación del producto, apila el registro de deshacer y retira el frente; la "
                   "validación no cambia porque es automática. Descartar retira el frente sin nota. En C++ "
                   "el registro de procesar contiene dos inversas: restaurar la observación y devolver el "
                   "trabajo al frente. Deshacer aplica el registro más reciente.")
    paragraph(doc, "Indicador: se construye una vez el índice auxiliar investigador → grupos activos a partir "
                   "de las membresías. Para la entidad pedida y para cada entidad activa del mismo tipo se "
                   "recorren sus vínculos activos en la multilista; cada producto contado se clasifica por su "
                   "subtipo y, si es de nuevo conocimiento (TOP, A o B), se recorren sus autorías para la "
                   "cohesión y la cooperación. Luego se toman los máximos por término y se normaliza. Con M "
                   "membresías, V vínculos de la población y A autorías de sus productos de nuevo "
                   "conocimiento, el tiempo es O(M + V + A) promedio en Python y O((M + V + A) log n) en C++, "
                   "y la memoria adicional es O(M + S), donde S es el número de subtipos con productos.")

    heading(doc, "16 Contratos de operación")
    add_table(doc, ["Operación", "Entrada y precondición", "Salida y postcondición"], [
        ["Crear producto", "ID y título no vacíos; ID y DOI sin colisión; año y fecha concordantes.",
         "Un nodo activo queda indexado por ID; no se crean relaciones implícitas."],
        ["Vincular grupo y producto", "IDs de ambos extremos existentes; par no duplicado.",
         "La relación activa se consulta desde grupo y desde producto."],
        ["Cerrar revisión", "Cola no vacía y nota de revisión no vacía.",
         "La nota queda en la observación, el producto sale del frente y la acción entra en la pila; la validación se recalcula sola."],
        ["Enviar rechazados", "Ninguna; los productos ya encolados se omiten.",
         "Cada producto activo rechazado queda al final de la cola con sus reglas como motivo; una sola entrada de deshacer; devuelve cuántos entraron."],
        ["Deshacer", "Historial no vacío.",
         "Se restaura el estado previo a la última acción, incluida la cola, y el historial disminuye en uno."],
        ["Consultar estadísticas", "Vista válida; ID seleccionado cuando corresponde; rango inicial no mayor que el final.",
         "Total único y distribuciones calculadas tras filtrar; paginar no modifica el total."],
        ["Calcular indicador", "Tipo grupos o investigadores; ID existente y activo; año inicial no mayor que el corte; validación vacía, validado o rechazado.",
         "Siete clases, cohesión y cooperación con valor, máximo, índice y aporte; subtipos contados; indicador entre 0 y 10. No modifica datos."],
        ["Asignar categoría", "Código del catálogo de 169 subtipos o vacío.",
         "El producto queda con origen manual (o sin origen si se vacía) y la reimportación conserva la elección."],
        ["Guardar archivo", "Datos con integridad referencial y ruta de destino escribible.",
         "Se escribe el archivo único del esquema 5 y se conserva el guardado anterior en .backup."],
    ], [3.15, 6.8, 6.85])

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    print(f"Documento generado: {OUTPUT}")


if __name__ == "__main__":
    build()
