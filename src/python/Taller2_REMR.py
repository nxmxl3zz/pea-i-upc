"""PEA-i UPC: interfaz Tkinter conectada al backend C++.

Ejecutar: python src/python/Taller2_REMR.py [--data ARCHIVO]
Sin --data se abre data/pea_upc.csv, el archivo único con todos los datos.
Con --python-backend se usa la implementación Python independiente de estructuras enlazadas.
La lógica de datos puede probarse sin un servidor gráfico.
"""

from __future__ import annotations

import argparse
import csv
from contextlib import contextmanager, nullcontext
from datetime import date, datetime
import hashlib
from html.parser import HTMLParser
import http.client
import importlib.util
import io
import ipaddress
import json
import math
import os
from pathlib import Path
import queue
import re
import shutil
import site
import socket
import subprocess
import sys
import tempfile
import threading
import time
from typing import Any, Callable, Iterator, NamedTuple
import unicodedata
from urllib.parse import parse_qs, urlencode, urlparse
from urllib.request import (HTTPRedirectHandler, HTTPHandler, HTTPSHandler,
                            ProxyHandler, Request, build_opener)
import uuid
import zipfile


# Versión 5: categorías del modelo 2024 y su origen en los productos.
# Se leen archivos v4 y v5 y carpetas v1 a v3; siempre se guarda un archivo único v5.
SCHEMA_VERSION = "5"
FILE_VERSIONS = ("4", "5")
FOLDER_VERSIONS = ("1", "2", "3")
PRODUCT_FIELDS_V4 = ("id", "titulo", "anio", "fecha", "familia", "tipologia", "validacion", "observacion", "doi", "url", "fuente", "activo")
ENTITY_FIELDS = {
    "grupos": ("id", "nombre", "codigo_gruplac", "fecha_creacion", "unidad", "responsable", "categoria", "descripcion", "objetivos", "mision", "vision", "lineas", "url", "fuente", "activo"),
    "investigadores": ("id", "nombre", "codigo_cvlac", "afiliacion", "categoria", "contacto", "url", "fuente", "activo"),
    "productos": ("id", "titulo", "anio", "fecha", "familia", "tipologia", "categoria", "categoria_origen", "validacion", "observacion", "doi", "url", "fuente", "activo"),
    "planes": ("id", "grupo_id", "nombre", "inicio", "fin", "objetivo", "indicador", "meta", "actividad", "activo"),
}
RELATION_FIELDS = {
    "membresias": ("grupo_id", "investigador_id", "rol", "inicio", "fin", "activo"),
    "autorias": ("producto_id", "investigador_id", "orden", "rol", "activo"),
    "grupos_productos": ("grupo_id", "producto_id", "origen", "activo"),
}
RELATION_ENDS = {
    "membresias": ("grupo_id", "investigador_id", "grupos", "investigadores"),
    "autorias": ("producto_id", "investigador_id", "productos", "investigadores"),
    "grupos_productos": ("grupo_id", "producto_id", "grupos", "productos"),
}
ALL_FIELDS = {**ENTITY_FIELDS, **RELATION_FIELDS}
MANIFEST_FIELDS = ("version", "guardado")
QUEUE_FIELDS = ("id", "producto_id", "motivo", "creado")
HISTORY_FIELDS = ("orden", "snapshot")
# Secciones del archivo único en el orden de lectura: cada relación aparece después de
# las entidades que enlaza, así que puede crearse en cuanto se lee.
DATA_SECTIONS = {"manifest": MANIFEST_FIELDS, **ALL_FIELDS,
                 "cola_validacion": QUEUE_FIELDS, "historial": HISTORY_FIELDS}
PAGE_SIZE = 100
REQUIRED = {
    "grupos": ("id", "nombre"), "investigadores": ("id", "nombre"),
    "productos": ("id", "titulo"), "planes": ("id", "grupo_id", "nombre"),
}
LABELS = {
    "id": "ID", "nombre": "Nombre", "codigo_gruplac": "Código GrupLAC",
    "fecha_creacion": "Fecha de creación", "unidad": "Unidad académica", "responsable": "Responsable",
    "categoria": "Categoría", "descripcion": "Descripción", "objetivos": "Objetivos",
    "mision": "Misión", "vision": "Visión", "lineas": "Líneas", "url": "URL",
    "fuente": "Fuente", "activo": "Activo", "codigo_cvlac": "Código CvLAC",
    "afiliacion": "Afiliación", "contacto": "Contacto público", "titulo": "Título",
    "anio": "Año", "fecha": "Fecha", "familia": "Familia", "tipologia": "Tipología",
    "validacion": "Validación", "observacion": "Observación", "doi": "DOI",
    "grupo_id": "Grupo (ID)", "investigador_id": "Investigador (ID)",
    "producto_id": "Producto (ID)", "inicio": "Inicio", "fin": "Fin",
    "objetivo": "Objetivo", "indicador": "Indicador", "meta": "Meta",
    "actividad": "Actividad", "rol": "Rol", "orden": "Orden",
    "origen": "Origen de asociación", "categoria_origen": "Origen de la categoría",
}
# Validación automática de productos. El orden de los códigos es parte del contrato
# con el backend C++: ambos motores devuelven las mismas reglas en el mismo orden.
# Un producto se atribuye con un autor o con un grupo activos: GrupLAC publica eventos y
# proyectos sin autores, y CvLAC productos que ningún grupo declara.
VALIDATION_RULES = {
    "anio": "Sin año de publicación",
    "tipologia": "Sin tipología",
    "vinculo": "Sin autor ni grupo activos",
    "doi": "DOI con formato inválido",
    "fuente": "Sin fuente verificable (URL o DOI)",
}
# Códigos anteriores que pueden seguir en motivos de cola guardados.
LEGACY_RULE_LABELS = {"autor": "Sin autor activo registrado", "grupo": "Sin grupo activo asociado"}
RULE_LABELS = {**VALIDATION_RULES, **LEGACY_RULE_LABELS}
VALIDATION_STATES = ("validado", "rechazado")
# Motivo de cola generado por la validación: "reglas:" seguido de códigos separados por comas.
AUTO_REASON_PREFIX = "reglas:"
LONG_FIELDS = {"descripcion", "objetivos", "mision", "vision", "lineas", "observacion", "objetivo", "indicador", "meta", "actividad", "fuente"}
MAX_FIELD_LEN = 100_000
MAX_HISTORY_BYTES = 64 * 1024 * 1024
csv.field_size_limit(MAX_HISTORY_BYTES)
MAX_CSV_ROWS = 500_000
MAX_PDF_TEXT = 5_000_000
MAX_PDF_MEMORY = 512 * 1024 * 1024
PDF_TIMEOUT = 25
MAX_LOCAL_FILE = 32 * 1024 * 1024
MAX_ARCHIVE_UNPACKED = 128 * 1024 * 1024
DEFAULT_DATA_FILE = Path(__file__).resolve().parents[2] / "data" / "pea_upc.csv"


MODEL_CLASSES: tuple[str, ...] = ("TOP", "A", "B", "ASC", "DPC", "FRH_A", "FRH_B")
MODEL_COEFFICIENTS: dict[str, float] = {
    "TOP": 3.7, "A": 2.3, "B": 0.4, "ASC": 1.5, "DPC": 0.5,
    "FRH_A": 1.0, "FRH_B": 0.2, "COHE": 0.1, "COOP": 0.3,
}
NEW_KNOWLEDGE_CLASSES = ("TOP", "A", "B")
CATEGORY_ORIGINS = ("estimada", "manual")


class Subtype(NamedTuple):
    codigo: str
    clase: str
    familia: str
    descripcion: str
    ventana: int
    peso_alta: float
    peso_relativo: float

    @property
    def peso_global(self) -> float:
        return self.peso_alta * self.peso_relativo


# Minciencias, Anexo 1 de la convocatoria 2024, §§2.2.9, 3.3, 3.6 y 3.7.
# Datos de referencia constantes; no dependen de archivos externos en ejecución.
MODEL_2024: dict[str, Subtype] = {subtype.codigo: subtype for subtype in (
    Subtype('ART_OPEN_A1', 'TOP', 'ART_R', 'Artículo de investigación A1 acceso abierto', 7, 100.0, 10.0),
    Subtype('ART_A1', 'TOP', 'ART_R', 'Artículo de investigación A1', 7, 100.0, 9.5),
    Subtype('ART_OPEN_A2', 'TOP', 'ART_R', 'Artículo de investigación A2 acceso abierto', 7, 100.0, 6.0),
    Subtype('ART_A2', 'TOP', 'ART_R', 'Artículo de investigación A2', 7, 100.0, 5.5),
    Subtype('ART_OPEN_B', 'A', 'ART_R', 'Artículo de investigación B acceso abierto', 7, 100.0, 3.5),
    Subtype('ART_B', 'A', 'ART_R', 'Artículo de investigación B', 7, 100.0, 3.0),
    Subtype('ART_OPEN_C', 'A', 'ART_R', 'Artículo de investigación C acceso abierto', 7, 100.0, 2.0),
    Subtype('ART_C', 'A', 'ART_R', 'Artículo de investigación C', 7, 100.0, 1.5),
    Subtype('ART_OPEN_D', 'B', 'ART_D', 'Artículo de investigación D acceso abierto', 5, 5.0, 10.0),
    Subtype('ART_D', 'B', 'ART_D', 'Artículo de investigación D', 5, 5.0, 9.0),
    Subtype('N_A1', 'B', 'N', 'Nota científica A1', 5, 4.0, 10.0),
    Subtype('N_A2', 'B', 'N', 'Nota científica A2', 5, 4.0, 8.5),
    Subtype('N_B', 'B', 'N', 'Nota científica B', 5, 4.0, 7.0),
    Subtype('N_C', 'B', 'N', 'Nota científica C', 5, 4.0, 5.5),
    Subtype('N_D', 'B', 'N', 'Nota científica D', 5, 4.0, 2.0),
    Subtype('LIB_A1', 'TOP', 'LIB', 'Libro resultado de investigación A1', 7, 300.0, 10.0),
    Subtype('LIB_A', 'TOP', 'LIB', 'Libro resultado de investigación A', 7, 300.0, 9.0),
    Subtype('LIB_B', 'A', 'LIB', 'Libro resultado de investigación B', 7, 300.0, 8.0),
    Subtype('LIB_C', 'B', 'LIB_C', 'Libro resultado de investigación C', 7, 15.0, 10.0),
    Subtype('CAP_LIB_A1', 'TOP', 'CAP_LIB', 'Capítulo en libro resultado de investigación A1', 5, 60.0, 10.0),
    Subtype('CAP_LIB_A', 'TOP', 'CAP_LIB', 'Capítulo en libro resultado de investigación A', 5, 60.0, 9.0),
    Subtype('CAP_LIB_B', 'A', 'CAP_LIB', 'Capítulo en libro resultado de investigación B', 5, 60.0, 8.0),
    Subtype('CAP_LIB_C', 'B', 'CAP_LIB_C', 'Capítulo en libro resultado de investigación C', 5, 3.0, 10.0),
    Subtype('LIB_FOR1', 'B', 'LIB_FOR', 'Libro de formación Q1', 5, 60.0, 10.0),
    Subtype('PA1', 'TOP', 'PAT', 'Patente de invención A1', 10, 500.0, 10.0),
    Subtype('PA2', 'TOP', 'PAT', 'Patente de invención A2', 10, 500.0, 7.0),
    Subtype('PA3', 'A', 'PAT', 'Patente de invención A3', 10, 500.0, 6.0),
    Subtype('PA4', 'A', 'PAT', 'Patente de invención A4', 10, 500.0, 5.5),
    Subtype('PB1', 'B', 'PAT', 'Patente de invención B1', 10, 500.0, 5.0),
    Subtype('PB2', 'B', 'PAT', 'Patente de invención B2', 10, 500.0, 3.5),
    Subtype('PB3', 'B', 'PAT', 'Patente de invención B3', 10, 500.0, 3.0),
    Subtype('PB4', 'B', 'PAT', 'Patente de invención B4', 10, 500.0, 2.6),
    Subtype('PB5', 'B', 'PAT', 'Patente de invención B5', 10, 500.0, 2.5),
    Subtype('PC', 'B', 'PAT', 'Patente de invención C', 10, 500.0, 1.8),
    Subtype('MA1', 'TOP', 'MOD', 'Patente de modelo de utilidad A1', 10, 500.0, 6.0),
    Subtype('MA2', 'TOP', 'MOD', 'Patente de modelo de utilidad A2', 10, 500.0, 4.2),
    Subtype('MA3', 'A', 'MOD', 'Patente de modelo de utilidad A3', 10, 500.0, 3.6),
    Subtype('MA4', 'A', 'MOD', 'Patente de modelo de utilidad A4', 10, 500.0, 3.33),
    Subtype('MB1', 'B', 'MOD', 'Patente de modelo de utilidad B1', 10, 500.0, 3.0),
    Subtype('MB2', 'B', 'MOD', 'Patente de modelo de utilidad B2', 10, 500.0, 2.1),
    Subtype('MB3', 'B', 'MOD', 'Patente de modelo de utilidad B3', 10, 500.0, 1.8),
    Subtype('MB4', 'B', 'MOD', 'Patente de modelo de utilidad B4', 10, 500.0, 1.7),
    Subtype('MB5', 'B', 'MOD', 'Patente de modelo de utilidad B5', 10, 500.0, 1.5),
    Subtype('MC', 'B', 'MOD', 'Patente de modelo de utilidad C', 10, 500.0, 1.1),
    Subtype('VV_A1', 'TOP', 'VV', 'Variedad vegetal A1', 10, 300.0, 10.0),
    Subtype('VV_A2', 'TOP', 'VV', 'Variedad vegetal A2', 10, 300.0, 8.0),
    Subtype('VV_A3', 'A', 'VV', 'Variedad vegetal A3', 10, 300.0, 5.0),
    Subtype('VV_A4', 'A', 'VV', 'Variedad vegetal A4', 10, 300.0, 2.5),
    Subtype('VV_B1', 'B', 'VV', 'Variedad vegetal B1', 10, 300.0, 5.0),
    Subtype('VV_B2', 'B', 'VV', 'Variedad vegetal B2', 10, 300.0, 4.0),
    Subtype('VV_B3', 'B', 'VV', 'Variedad vegetal B3', 10, 300.0, 2.5),
    Subtype('VV_B4', 'B', 'VV', 'Variedad vegetal B4', 10, 300.0, 1.0),
    Subtype('VA_A', 'TOP', 'VA', 'Nueva raza animal o población mejorada A', 10, 300.0, 10.0),
    Subtype('VA_B', 'A', 'VA', 'Nueva raza animal o población mejorada B', 10, 300.0, 5.0),
    Subtype('AAD_A1', 'TOP', 'AAD', 'Obra o producto de creación (AAD) A1', 10, 100.0, 10.0),
    Subtype('AAD_A', 'TOP', 'AAD', 'Obra o producto de creación (AAD) A', 10, 100.0, 8.0),
    Subtype('AAD_B', 'A', 'AAD', 'Obra o producto de creación (AAD) B', 10, 100.0, 6.0),
    Subtype('AAD_C', 'B', 'AAD', 'Obra o producto de creación (AAD) C', 10, 100.0, 4.0),
    Subtype('DI_A', 'A', 'DI', 'Diseño industrial A', 5, 35.0, 10.0),
    Subtype('DI_B', 'B', 'DI', 'Diseño industrial B', 5, 35.0, 5.0),
    Subtype('ECI', 'A', 'ECI', 'Esquema de circuito integrado', 5, 35.0, 4.0),
    Subtype('SF', 'B', 'SF', 'Software', 5, 35.0, 8.0),
    Subtype('PP', 'A', 'PP', 'Planta piloto', 5, 35.0, 4.0),
    Subtype('PI', 'A', 'PI', 'Prototipo industrial', 5, 35.0, 4.0),
    Subtype('SD', 'A', 'SD', 'Signo distintivo', 5, 35.0, 4.0),
    Subtype('PN', 'B', 'PN', 'Producto nutracéutico', 5, 35.0, 6.0),
    Subtype('CC', 'B', 'CC', 'Colección científica', 5, 35.0, 10.0),
    Subtype('NRC_A', 'B', 'NRC', 'Nuevo registro científico A', 5, 35.0, 8.0),
    Subtype('NRC_B', 'B', 'NRC', 'Nuevo registro científico B', 5, 35.0, 5.0),
    Subtype('SE', 'A', 'SE', 'Secreto empresarial', 5, 100.0, 5.0),
    Subtype('EBT_A', 'A', 'EBT', 'Empresa de base tecnológica A', 5, 100.0, 10.0),
    Subtype('EBT_B', 'B', 'EBT', 'Empresa de base tecnológica B', 5, 100.0, 8.0),
    Subtype('ICC_A', 'A', 'ICC', 'Empresa creativa y cultural A', 5, 100.0, 10.0),
    Subtype('ICC_B', 'B', 'ICC', 'Empresa creativa y cultural B', 5, 100.0, 4.0),
    Subtype('IG_A1', 'A', 'IG', 'Innovación en gestión empresarial A1', 5, 100.0, 10.0),
    Subtype('IG_A2', 'A', 'IG', 'Innovación en gestión empresarial A2', 5, 100.0, 6.0),
    Subtype('IG_B1', 'B', 'IG', 'Innovación en gestión empresarial B1', 5, 100.0, 5.0),
    Subtype('IG_B2', 'B', 'IG', 'Innovación en gestión empresarial B2', 5, 100.0, 3.0),
    Subtype('IPP', 'B', 'IPP', 'Innovación en procedimientos y servicios', 5, 100.0, 5.0),
    Subtype('RNL_A', 'TOP', 'RNL', 'Regulación, norma, reglamento o legislación A', 5, 100.0, 10.0),
    Subtype('RNL_B', 'A', 'RNL', 'Regulación, norma, reglamento o legislación B', 5, 100.0, 8.0),
    Subtype('RNR', 'A', 'RNR', 'Normatividad del espectro radioeléctrico', 5, 100.0, 8.0),
    Subtype('RNT', 'B', 'RNT', 'Norma técnica', 5, 100.0, 7.0),
    Subtype('RNPC', 'B', 'RNPC', 'Guía de práctica clínica', 5, 100.0, 7.0),
    Subtype('GMCF', 'B', 'GMCF', 'Guía de manejo clínico forense', 5, 100.0, 7.0),
    Subtype('MADV', 'B', 'MADV', 'Manual o modelo de atención diferencial a víctimas', 5, 100.0, 7.0),
    Subtype('PAU', 'B', 'PAU', 'Protocolo de atención a usuarios/víctimas', 5, 100.0, 7.0),
    Subtype('PVE', 'B', 'PVE', 'Protocolo de vigilancia epidemiológica', 5, 100.0, 7.0),
    Subtype('AL', 'B', 'AL', 'Acuerdo de ley', 5, 100.0, 8.0),
    Subtype('RNPL', 'B', 'RNPL', 'Proyecto de ley', 5, 100.0, 6.0),
    Subtype('CT', 'B', 'CT', 'Concepto técnico', 5, 15.0, 10.0),
    Subtype('MR', 'B', 'MR', 'Registro de acuerdo de licencia (AAD)', 5, 14.0, 10.0),
    Subtype('FIS', 'ASC', 'FIS', 'Proceso ASC para asuntos de interés social', 5, 200.0, 10.0),
    Subtype('GPP_A', 'ASC', 'GPP', 'Proceso ASC para insumos de política pública A', 5, 200.0, 10.0),
    Subtype('GPP_B', 'ASC', 'GPP', 'Proceso ASC para insumos de política pública B', 5, 200.0, 7.0),
    Subtype('GPP_C', 'ASC', 'GPP', 'Proceso ASC para insumos de política pública C', 5, 200.0, 5.0),
    Subtype('FCP_A', 'ASC', 'FCP', 'Proceso ASC para cadenas productivas A', 5, 200.0, 10.0),
    Subtype('FCP_B', 'ASC', 'FCP', 'Proceso ASC para cadenas productivas B', 5, 200.0, 5.0),
    Subtype('FCP_C', 'ASC', 'FCP', 'Proceso ASC para cadenas productivas C', 5, 200.0, 3.0),
    Subtype('TCCG_A', 'ASC', 'TCCG', 'Proceso ASC Centro de Ciencia y grupo A', 5, 200.0, 10.0),
    Subtype('TCCG_B', 'ASC', 'TCCG', 'Proceso ASC Centro de Ciencia y grupo B', 5, 200.0, 7.0),
    Subtype('TCCG_C', 'ASC', 'TCCG', 'Proceso ASC Centro de Ciencia y grupo C', 5, 200.0, 4.0),
    Subtype('TCCG_D', 'ASC', 'TCCG', 'Proceso ASC Centro de Ciencia y grupo D', 5, 200.0, 1.0),
    Subtype('EC_A', 'DPC', 'EC', 'Evento científico con componente de apropiación A', 5, 100.0, 10.0),
    Subtype('EC_B', 'DPC', 'EC', 'Evento científico con componente de apropiación B', 5, 100.0, 6.0),
    Subtype('RC_A', 'DPC', 'RC', 'Participación en red de conocimiento A', 5, 100.0, 10.0),
    Subtype('RC_B', 'DPC', 'RC', 'Participación en red de conocimiento B', 5, 100.0, 6.0),
    Subtype('TC_A', 'DPC', 'TC', 'Taller de creación A', 5, 100.0, 10.0),
    Subtype('TC_B', 'DPC', 'TC', 'Taller de creación B', 5, 100.0, 8.0),
    Subtype('TC_C', 'DPC', 'TC', 'Taller de creación C', 5, 100.0, 6.0),
    Subtype('ECA_A', 'DPC', 'ECA', 'Evento artístico, de arquitectura o diseño A', 5, 100.0, 8.0),
    Subtype('ECA_B', 'DPC', 'ECA', 'Evento artístico, de arquitectura o diseño B', 5, 100.0, 6.0),
    Subtype('WP', 'DPC', 'WP', 'Documento de trabajo (working paper)', 5, 100.0, 10.0),
    Subtype('NSG', 'DPC', 'NSG', 'Nueva secuencia genética', 5, 100.0, 10.0),
    Subtype('ERL', 'DPC', 'ERL', 'Edición de revista o libro de divulgación', 5, 100.0, 6.0),
    Subtype('IFI', 'DPC', 'IFI', 'Informe final de investigación', 5, 100.0, 2.0),
    Subtype('INF', 'DPC', 'INF', 'Informe técnico', 5, 100.0, 5.0),
    Subtype('CON_CT', 'DPC', 'CON_CT', 'Consultoría científico-tecnológica', 5, 100.0, 7.5),
    Subtype('CON_AAD', 'DPC', 'CON_AAD', 'Consultoría en arte, arquitectura y diseño', 5, 100.0, 7.5),
    Subtype('PEE_A1', 'DPC', 'PEE', 'Publicación editorial no especializada A1', 5, 100.0, 10.0),
    Subtype('PEE_A2', 'DPC', 'PEE', 'Publicación editorial no especializada A2', 5, 100.0, 8.0),
    Subtype('PEE_B1', 'DPC', 'PEE', 'Publicación editorial no especializada B1', 5, 100.0, 7.0),
    Subtype('PEE_B2', 'DPC', 'PEE', 'Publicación editorial no especializada B2', 5, 100.0, 5.0),
    Subtype('PEE_C1', 'DPC', 'PEE', 'Publicación editorial no especializada C1', 5, 100.0, 5.0),
    Subtype('PEE_C2', 'DPC', 'PEE', 'Publicación editorial no especializada C2', 5, 100.0, 3.0),
    Subtype('PCD_A1', 'DPC', 'PCD', 'Producción de contenido digital A1', 5, 100.0, 10.0),
    Subtype('PCD_A2', 'DPC', 'PCD', 'Producción de contenido digital A2', 5, 100.0, 8.0),
    Subtype('PCD_B1', 'DPC', 'PCD', 'Producción de contenido digital B1', 5, 100.0, 7.0),
    Subtype('PCD_B2', 'DPC', 'PCD', 'Producción de contenido digital B2', 5, 100.0, 5.0),
    Subtype('PCD_C1', 'DPC', 'PCD', 'Producción de contenido digital C1', 5, 100.0, 5.0),
    Subtype('PCD_C2', 'DPC', 'PCD', 'Producción de contenido digital C2', 5, 100.0, 3.0),
    Subtype('TRM_A1', 'DPC', 'TRM', 'Estrategia y contenido transmedia A1', 5, 100.0, 10.0),
    Subtype('TRM_A2', 'DPC', 'TRM', 'Estrategia y contenido transmedia A2', 5, 100.0, 8.0),
    Subtype('TRM_B1', 'DPC', 'TRM', 'Estrategia y contenido transmedia B1', 5, 100.0, 7.0),
    Subtype('TRM_B2', 'DPC', 'TRM', 'Estrategia y contenido transmedia B2', 5, 100.0, 5.0),
    Subtype('TRM_C1', 'DPC', 'TRM', 'Estrategia y contenido transmedia C1', 5, 100.0, 5.0),
    Subtype('TRM_C2', 'DPC', 'TRM', 'Estrategia y contenido transmedia C2', 5, 100.0, 3.0),
    Subtype('DW_A1', 'DPC', 'DW', 'Desarrollo web A1', 5, 100.0, 10.0),
    Subtype('DW_A2', 'DPC', 'DW', 'Desarrollo web A2', 5, 100.0, 8.0),
    Subtype('DW_B1', 'DPC', 'DW', 'Desarrollo web B1', 5, 100.0, 7.0),
    Subtype('DW_B2', 'DPC', 'DW', 'Desarrollo web B2', 5, 100.0, 5.0),
    Subtype('DW_C1', 'DPC', 'DW', 'Desarrollo web C1', 5, 100.0, 5.0),
    Subtype('DW_C2', 'DPC', 'DW', 'Desarrollo web C2', 5, 100.0, 3.0),
    Subtype('LIB_FOR2', 'DPC', 'LIB_FOR_DPC', 'Libro de formación Q2', 5, 100.0, 10.0),
    Subtype('LIB_FOR3', 'DPC', 'LIB_FOR_DPC', 'Libro de formación Q3', 5, 100.0, 10.0),
    Subtype('BOL', 'DPC', 'BOL', 'Boletín divulgativo de resultado de investigación', 5, 100.0, 3.0),
    Subtype('LIB_DIV', 'DPC', 'LIB_DIV', 'Libro de divulgación o compilación de divulgación', 5, 100.0, 10.0),
    Subtype('GC', 'DPC', 'GC', 'Generación de contenidos', 5, 100.0, 4.0),
    Subtype('MAN_GUI', 'DPC', 'MAN_GUI', 'Manual o guía especializada', 5, 100.0, 2.0),
    Subtype('TD_A', 'FRH_A', 'TD', 'Dirección de tesis de doctorado A', 5, 160.0, 10.0),
    Subtype('TD_B', 'FRH_A', 'TD', 'Dirección de tesis de doctorado B', 5, 160.0, 5.0),
    Subtype('TM_A', 'FRH_B', 'TM', 'Dirección de trabajo de grado de maestría A', 5, 70.0, 10.0),
    Subtype('TM_B', 'FRH_B', 'TM', 'Dirección de trabajo de grado de maestría B', 5, 70.0, 5.0),
    Subtype('TP_A', 'FRH_B', 'TP', 'Dirección de trabajo de pregrado A', 5, 20.0, 10.0),
    Subtype('TP_B', 'FRH_B', 'TP', 'Dirección de trabajo de pregrado B', 5, 20.0, 5.0),
    Subtype('PID_A', 'FRH_B', 'PID', 'Proyecto de investigación y desarrollo A', 5, 50.0, 10.0),
    Subtype('PID_B', 'FRH_B', 'PID', 'Proyecto de investigación y desarrollo B', 5, 50.0, 6.0),
    Subtype('PID_C', 'FRH_B', 'PID', 'Proyecto de investigación y desarrollo C', 5, 50.0, 2.0),
    Subtype('PIC_A', 'FRH_B', 'PIC', 'Proyecto de investigación-creación A', 5, 50.0, 10.0),
    Subtype('PIC_B', 'FRH_B', 'PIC', 'Proyecto de investigación-creación B', 5, 50.0, 6.0),
    Subtype('PIC_C', 'FRH_B', 'PIC', 'Proyecto de investigación-creación C', 5, 50.0, 2.0),
    Subtype('PF_A', 'FRH_B', 'PF', 'Proyecto ID+I con formación A', 5, 50.0, 10.0),
    Subtype('PF_B', 'FRH_B', 'PF', 'Proyecto ID+I con formación B', 5, 50.0, 8.0),
    Subtype('PE', 'FRH_B', 'PE', 'Proyecto de extensión y responsabilidad social en CTeI', 5, 100.0, 10.0),
    Subtype('AP_A', 'FRH_A', 'AP_PROG', 'Apoyo a creación de programa doctorado', 5, 100.0, 10.0),
    Subtype('AP_B', 'FRH_B', 'AP_PROG', 'Apoyo a creación de programa maestría', 5, 100.0, 8.0),
    Subtype('AP_C', 'FRH_A', 'AP_CUR', 'Apoyo a creación de cursos doctorado', 5, 100.0, 5.0),
    Subtype('AP_D', 'FRH_B', 'AP_CUR', 'Apoyo a creación de cursos maestría', 5, 100.0, 3.0),
    Subtype('APO', 'FRH_B', 'APO', 'Acompañamiento y asesoría Programa Ondas', 5, 30.0, 10.0),
)}


class DataError(ValueError):
    """Dato inválido o relación que rompería la integridad."""


class BackendError(DataError):
    """La comunicación falló; no se puede continuar la operación como un rechazo de fila."""


class ListNode:
    def __init__(self, value: dict[str, str]):
        self.value = value
        self.prev: ListNode | None = None
        self.next: ListNode | None = None


class DoublyLinkedList:
    """Lista fuente de verdad para entidades; índices auxiliares no la reemplazan."""

    def __init__(self) -> None:
        self.head: ListNode | None = None
        self.tail: ListNode | None = None
        self.length = 0
        self.by_id: dict[str, ListNode] = {}

    def append(self, value: dict[str, str]) -> None:
        key = value["id"]
        if key in self.by_id:
            raise DataError(f"ID duplicado: {key}")
        node = ListNode(value)
        node.prev = self.tail
        if self.tail:
            self.tail.next = node
        else:
            self.head = node
        self.tail = node
        self.by_id[key] = node
        self.length += 1

    def get(self, key: str) -> dict[str, str] | None:
        node = self.by_id.get(key)
        return node.value if node else None

    def remove(self, key: str) -> dict[str, str]:
        node = self.by_id.pop(key, None)
        if node is None:
            raise DataError(f"No existe ID {key}")
        if node.prev:
            node.prev.next = node.next
        else:
            self.head = node.next
        if node.next:
            node.next.prev = node.prev
        else:
            self.tail = node.prev
        self.length -= 1
        return node.value

    def __iter__(self) -> Iterator[dict[str, str]]:
        node = self.head
        while node:
            yield node.value
            node = node.next

    def backwards(self) -> Iterator[dict[str, str]]:
        node = self.tail
        while node:
            yield node.value
            node = node.prev


class RelationNode:
    def __init__(self, left: str, right: str, value: dict[str, str]):
        self.left = left
        self.right = right
        self.value = value
        self.next_left: RelationNode | None = None
        self.next_right: RelationNode | None = None


class MultiList:
    """Cada vínculo pertenece simultáneamente a dos cadenas de recorrido."""

    def __init__(self, left_field: str, right_field: str) -> None:
        self.left_field = left_field
        self.right_field = right_field
        self.left_heads: dict[str, RelationNode] = {}
        self.right_heads: dict[str, RelationNode] = {}
        self.by_pair: dict[tuple[str, str], RelationNode] = {}

    def add(self, row: dict[str, str]) -> None:
        left, right = row[self.left_field], row[self.right_field]
        pair = (left, right)
        if pair in self.by_pair:
            raise DataError(f"Relación duplicada: {left} / {right}")
        node = RelationNode(left, right, row)
        node.next_left = self.left_heads.get(left)
        node.next_right = self.right_heads.get(right)
        self.left_heads[left] = node
        self.right_heads[right] = node
        self.by_pair[pair] = node

    def get(self, left: str, right: str) -> dict[str, str] | None:
        node = self.by_pair.get((left, right))
        return node.value if node else None

    def by_left(self, left: str) -> Iterator[dict[str, str]]:
        node = self.left_heads.get(left)
        while node:
            yield node.value
            node = node.next_left

    def by_right(self, right: str) -> Iterator[dict[str, str]]:
        node = self.right_heads.get(right)
        while node:
            yield node.value
            node = node.next_right

    def remove(self, left: str, right: str) -> dict[str, str]:
        node = self.by_pair.pop((left, right), None)
        if not node:
            raise DataError("No existe esa relación")
        previous = None
        current = self.left_heads[left]
        while current is not node:
            previous, current = current, current.next_left
        if previous:
            previous.next_left = node.next_left
        elif node.next_left:
            self.left_heads[left] = node.next_left
        else:
            del self.left_heads[left]
        previous = None
        current = self.right_heads[right]
        while current is not node:
            previous, current = current, current.next_right
        if previous:
            previous.next_right = node.next_right
        elif node.next_right:
            self.right_heads[right] = node.next_right
        else:
            del self.right_heads[right]
        return node.value

    def __iter__(self) -> Iterator[dict[str, str]]:
        for left in self.left_heads:
            yield from self.by_left(left)


class StackNode:
    def __init__(self, value: str, next_node: StackNode | None = None):
        self.value, self.next = value, next_node


class LinkedStack:
    def __init__(self, limit: int = 30) -> None:
        self.top: StackNode | None = None
        self.length = 0
        self.limit = limit

    def push(self, value: str) -> None:
        if len(value.encode("utf-8")) > MAX_HISTORY_BYTES:
            raise DataError("La acción supera el límite de 64 MiB del historial")
        self.top = StackNode(value, self.top)
        self.length += 1
        if self.length > self.limit:
            node = self.top
            for _ in range(self.limit - 1):
                node = node.next  # type: ignore[assignment]
            node.next = None
            self.length = self.limit

    def peek(self) -> str | None:
        return self.top.value if self.top else None

    def pop(self) -> str | None:
        if not self.top:
            return None
        value = self.top.value
        self.top = self.top.next
        self.length -= 1
        return value

    def clear(self) -> None:
        self.top, self.length = None, 0

    def __iter__(self) -> Iterator[str]:
        node = self.top
        while node:
            yield node.value
            node = node.next


class QueueNode:
    def __init__(self, value: dict[str, str]):
        self.value = value
        self.next: QueueNode | None = None


class LinkedQueue:
    def __init__(self) -> None:
        self.front: QueueNode | None = None
        self.rear: QueueNode | None = None
        self.length = 0

    def enqueue(self, value: dict[str, str]) -> None:
        node = QueueNode(value)
        if self.rear:
            self.rear.next = node
        else:
            self.front = node
        self.rear = node
        self.length += 1

    def prepend(self, value: dict[str, str]) -> None:
        node = QueueNode(value)
        node.next = self.front
        self.front = node
        if self.rear is None:
            self.rear = node
        self.length += 1

    def remove(self, job_id: str) -> None:
        previous = None
        node = self.front
        while node and node.value["id"] != job_id:
            previous, node = node, node.next
        if node is None:
            raise DataError("Trabajo de cola no encontrado")
        if previous:
            previous.next = node.next
        else:
            self.front = node.next
        if self.rear is node:
            self.rear = previous
        self.length -= 1

    def peek(self) -> dict[str, str] | None:
        return self.front.value if self.front else None

    def dequeue(self) -> dict[str, str] | None:
        if not self.front:
            return None
        value = self.front.value
        self.front = self.front.next
        if not self.front:
            self.rear = None
        self.length -= 1
        return value

    def clear(self) -> None:
        self.front, self.rear, self.length = None, None, 0

    def __iter__(self) -> Iterator[dict[str, str]]:
        node = self.front
        while node:
            yield node.value
            node = node.next


def clean_row(kind: str, values: dict[str, Any]) -> dict[str, str]:
    if kind not in ALL_FIELDS:
        raise DataError(f"Tipo desconocido: {kind}")
    extra = set(values) - set(ALL_FIELDS[kind])
    if extra:
        raise DataError(f"Columnas desconocidas: {', '.join(sorted(extra))}")
    row = {field: str(values.get(field, "") or "").strip() for field in ALL_FIELDS[kind]}
    row["activo"] = row["activo"] or "1"
    if row["activo"] not in ("0", "1"):
        raise DataError("Activo debe ser 0 o 1")
    for field in REQUIRED.get(kind, ()):
        if not row[field]:
            raise DataError(f"Falta {LABELS[field]}")
    for field in ALL_FIELDS[kind]:
        if field.endswith("_id") and not row[field]:
            raise DataError(f"Falta {LABELS[field]}")
    if kind == "productos":
        if row["fecha"] and not row["anio"]:
            row["anio"] = row["fecha"][:4]
        if row["anio"]:
            try:
                if not re.fullmatch(r"[0-9]{4}", row["anio"]):
                    raise ValueError("Formato de año")
                year = int(row["anio"])
            except ValueError as exc:
                raise DataError("Año inválido") from exc
            if not 1900 <= year <= date.today().year + 1:
                raise DataError("Año fuera de rango")
    if kind == "productos":
        category = row["categoria"]
        if category and category not in MODEL_2024:
            raise DataError(f"Categoría desconocida: {category}")
        if row["categoria_origen"] not in ("", *CATEGORY_ORIGINS):
            raise DataError("Origen de categoría: estimada o manual")
        row["categoria_origen"] = (row["categoria_origen"] or "manual") if category else ""
        # El repositorio recalcula este campo; aquí solo se admite un valor conocido.
        # «pendiente» procede de carpetas anteriores a la validación automática.
        row["validacion"] = row["validacion"] or "rechazado"
        if row["validacion"] not in (*VALIDATION_STATES, "pendiente"):
            raise DataError("Validación: validado o rechazado")
    if kind == "autorias" and row["orden"]:
        significant = row["orden"].lstrip("0")
        if (not re.fullmatch(r"[0-9]+", row["orden"]) or not significant or len(significant) > 19
                or (len(significant) == 19 and significant > "9223372036854775807")):
            raise DataError("Orden de autor debe ser positivo")
    for field in ("fecha", "fecha_creacion", "inicio", "fin"):
        if field in row and row[field]:
            try:
                if not re.fullmatch(r"[0-9]{4}-[0-9]{2}-[0-9]{2}", row[field]):
                    raise ValueError("Formato de fecha")
                date.fromisoformat(row[field])
            except ValueError as exc:
                raise DataError(f"{LABELS[field]} debe ser AAAA-MM-DD") from exc
    if kind == "productos" and row["fecha"] and row["anio"]:
        if row["fecha"][:4] != row["anio"]:
            raise DataError("El año y la fecha del producto no coinciden")
    if kind in ("planes", "membresias") and row["inicio"] and row["fin"]:
        if row["inicio"] > row["fin"]:
            raise DataError("La fecha final precede la inicial")
    for field, value in row.items():
        if len(value.encode("utf-8")) > MAX_FIELD_LEN:
            raise DataError(f"El campo {LABELS.get(field, field)} supera el tamaño permitido")
    return row


def _legacy_row(kind: str, values: dict[str, str]) -> dict[str, str]:
    """Adapta filas antiguas, también del historial, sin alterar los datos de origen."""
    row = values.copy()
    if kind == "grupos":
        row.pop("sigla", None)
    if kind == "productos":
        row.setdefault("validacion", "pendiente")
        # Solo las categorías anteriores a v3 se trasladan a la observación.
        # Los snapshots v5 incluyen categoria_origen y conservan su clasificación.
        if "categoria_origen" not in row:
            category = (row.pop("categoria", "") or "").strip()
            if category:
                note = f"Categoría registrada antes de la versión 3: {category}"
                row["observacion"] = f"{row.get('observacion', '').strip()}\n{note}".strip()
    return row


def _doi_text(value: str) -> str:
    text = value.strip().lower()
    return text[len("https://doi.org/"):] if text.startswith("https://doi.org/") else text


def _valid_doi(doi: str) -> bool:
    """10.NNNN/sufijo sin espacios; se evalúa igual que en C++ (sin expresiones regulares)."""
    if not doi.startswith("10."):
        return False
    prefix, slash, suffix = doi[3:].partition("/")
    return (bool(slash) and 4 <= len(prefix) <= 9 and prefix.isascii() and prefix.isdigit()
            and bool(suffix) and not any(ch in " \t\n\r\f\v" for ch in suffix))


def _source_doi(detail: str) -> tuple[str, str]:
    """DOI tal como lo publica la ficha (deriva el ID) y DOI válido que se guarda, o vacío.

    Repara un DOI partido por un espacio («10.5539/ gjhs.v11n9p34») y descarta uno truncado.
    """
    match = re.search(r"\bDOI:\s*(10\.\S+)", detail, re.IGNORECASE)
    if not match:
        return "", ""
    published = match.group(1).rstrip(".,;)")
    split = re.search(r"\bDOI:\s*(10\.\d{4,9}/)\s+([^\s,;]+)", detail, re.IGNORECASE)
    for candidate in ((split.group(1) + split.group(2).rstrip(".,;)")) if split else "", published):
        if candidate and _valid_doi(_doi_text(candidate)):
            return published, candidate
    return published, ""


def _verifiable_url(url: str) -> bool:
    text = url.strip().lower()
    return any(text.startswith(scheme) and len(text) > len(scheme) for scheme in ("http://", "https://"))


def auto_reason_codes(reason: str) -> list[str] | None:
    """Códigos de reglas de un motivo generado automáticamente; None si el motivo es manual."""
    if not reason.startswith(AUTO_REASON_PREFIX):
        return None
    return [code for code in reason[len(AUTO_REASON_PREFIX):].split(",") if code in RULE_LABELS]


class Repository:
    """Operaciones del dominio. Las entidades viven en listas; relaciones en multilistas."""

    _external_fields = {"grupos": "codigo_gruplac", "investigadores": "codigo_cvlac", "productos": "doi"}

    def __init__(self) -> None:
        self.entities = {name: DoublyLinkedList() for name in ENTITY_FIELDS}
        self.relations = {
            name: MultiList(ends[0], ends[1]) for name, ends in RELATION_ENDS.items()
        }
        self.history = LinkedStack()
        self.queue = LinkedQueue()
        self.issues: dict[str, tuple[str, ...]] = {}
        self._batch_changes: list[list[dict[str, str]]] | None = None
        self._batch_bytes = 0
        self._external_index: dict[str, dict[str, str]] = {kind: {} for kind in self._external_fields}
        self._indicator_population: dict[tuple[Any, ...], tuple[dict[str, float], int]] = {}
        self.dirty = False

    def catalog(self) -> list[dict[str, Any]]:
        return [{**subtype._asdict(), "peso_global": subtype.peso_global}
                for subtype in MODEL_2024.values()]

    def rows(self, kind: str) -> list[dict[str, str]]:
        source = self.entities.get(kind) or self.relations.get(kind)
        if source is None:
            raise DataError(f"Tipo desconocido: {kind}")
        rows = [row.copy() for row in source]
        if kind in RELATION_ENDS:
            left, right = RELATION_ENDS[kind][:2]
            rows.sort(key=lambda row: (row[left], row[right]))
        return rows

    @staticmethod
    def _matches(row: dict[str, str], needle: str) -> bool:
        return not needle or needle in " ".join(value for _, value in sorted(row.items()) if value).casefold()

    def _page_references(self, kind: str, rows: list[dict[str, str]], side: str = "") -> dict[str, Any]:
        left, right, left_kind, right_kind = RELATION_ENDS[kind]
        ends = ((left, left_kind), (right, right_kind)) if not side else (
            ((right, right_kind),) if side == "left" else ((left, left_kind),))
        return {entity_kind: [self.get(entity_kind, key).copy() for key in dict.fromkeys(row[field] for row in rows)]
                for field, entity_kind in ends}

    def page(self, kind: str, offset: int = 0, limit: int = PAGE_SIZE,
             query: str = "", *, active: str = "", group_id: str = "",
             include_references: bool = False) -> dict[str, Any]:
        source = self.entities.get(kind) or self.relations.get(kind)
        if source is None:
            raise DataError(f"Tipo desconocido: {kind}")
        if active not in ("", "0", "1"):
            raise DataError("Activo debe ser 0 o 1")
        if group_id and kind != "planes":
            raise DataError("El filtro de grupo solo se admite para planes")
        if include_references and kind not in RELATION_ENDS:
            raise DataError("La búsqueda de extremos solo se admite para relaciones")
        needle = query.casefold().strip()

        def matches(row: dict[str, str]) -> bool:
            if (active and row["activo"] != active) or (group_id and row["grupo_id"] != group_id):
                return False
            if self._matches(row, needle):
                return True
            if include_references:
                left, right, left_kind, right_kind = RELATION_ENDS[kind]
                return (self._matches(self.get(left_kind, row[left]), needle)
                        or self._matches(self.get(right_kind, row[right]), needle))
            return False

        rows = (row for row in source if matches(row))
        if kind in RELATION_ENDS:
            left, right = RELATION_ENDS[kind][:2]
            matches = sorted(rows, key=lambda row: (row[left], row[right]))
            result = {"rows": [row.copy() for row in matches[offset:offset + limit]], "total": len(matches)}
            if include_references:
                result["references"] = self._page_references(kind, result["rows"])
            return result
        page_rows = []
        total = 0
        for row in rows:
            if offset <= total < offset + limit:
                page_rows.append(row.copy())
            total += 1
        return {"rows": page_rows, "total": total}

    def summary(self) -> dict[str, Any]:
        return {"active_groups": sum(row["activo"] == "1" for row in self.entities["grupos"]),
                "active_people": sum(row["activo"] == "1" for row in self.entities["investigadores"])}

    def field_usage(self, kind: str) -> dict[str, int]:
        """Registros con valor por campo; la interfaz oculta columnas que nadie ha llenado."""
        source = self.entities.get(kind) or self.relations.get(kind)
        if source is None:
            raise DataError(f"Tipo desconocido: {kind}")
        usage = dict.fromkeys(ALL_FIELDS[kind], 0)
        for row in source:
            for field, value in row.items():
                if value:
                    usage[field] += 1
        return usage

    # Validación automática: un producto es «validado» si cumple todas las reglas.
    # Depende de sus vínculos, así que se recalcula al cambiar el producto, sus
    # autorías o grupos, o el estado de un investigador o grupo enlazado.
    def _rule_failures(self, product: dict[str, str]) -> tuple[str, ...]:
        failed = []
        if not product["anio"]:
            failed.append("anio")
        if not product["tipologia"]:
            failed.append("tipologia")
        has_author = any(link["activo"] == "1" and (person := self.get("investigadores", link["investigador_id"]))
                         and person["activo"] == "1" for link in self.relations["autorias"].by_left(product["id"]))
        has_group = any(link["activo"] == "1" and (group := self.get("grupos", link["grupo_id"]))
                        and group["activo"] == "1" for link in self.relations["grupos_productos"].by_right(product["id"]))
        if not has_author and not has_group:
            failed.append("vinculo")
        doi = _doi_text(product["doi"])
        if doi and not _valid_doi(doi):
            failed.append("doi")
        if not _verifiable_url(product["url"]) and not (doi and _valid_doi(doi)):
            failed.append("fuente")
        return tuple(failed)

    def _revalidate(self, product_ids: Any) -> None:
        self._indicator_population.clear()
        for product_id in set(product_ids):
            product = self.get("productos", product_id)
            if product is None:
                self.issues.pop(product_id, None)
                continue
            failed = self._rule_failures(product)
            self.issues[product_id] = failed
            product["validacion"] = "rechazado" if failed else "validado"

    def _affected_products(self, kind: str, row: dict[str, str]) -> list[str]:
        if kind == "productos":
            return [row["id"]]
        if kind in ("autorias", "grupos_productos"):
            return [row["producto_id"]]
        if kind == "investigadores":
            return [link["producto_id"] for link in self.relations["autorias"].by_right(row["id"])]
        if kind == "grupos":
            return [link["producto_id"] for link in self.relations["grupos_productos"].by_left(row["id"])]
        return []

    def product_issues(self, product_id: str) -> list[str]:
        if self.get("productos", product_id) is None:
            raise DataError("Producto inexistente")
        return list(self.issues.get(product_id, ()))

    def get(self, kind: str, key: str | tuple[str, str]) -> dict[str, str] | None:
        if kind in ENTITY_FIELDS:
            return self.entities[kind].get(str(key))
        if kind in RELATION_FIELDS and isinstance(key, tuple):
            return self.relations[kind].get(*key)
        return None

    def related(self, kind: str, key: str, side: str) -> list[dict[str, str]]:
        if kind not in RELATION_ENDS:
            raise DataError(f"Relación desconocida: {kind}")
        if side not in ("left", "right"):
            raise DataError("Extremo desconocido")
        relation = self.relations[kind]
        return list(relation.by_left(key) if side == "left" else relation.by_right(key))

    def related_page(self, kind: str, key: str, side: str, offset: int = 0,
                     limit: int = PAGE_SIZE, active_only: bool = False,
                     query: str = "", *, include_references: bool = False) -> dict[str, Any]:
        if kind not in RELATION_ENDS:
            raise DataError(f"Relación desconocida: {kind}")
        if side not in ("left", "right"):
            raise DataError("Extremo desconocido")
        relation = self.relations[kind]
        source = relation.by_left(key) if side == "left" else relation.by_right(key)
        needle = query.casefold().strip()
        left, right, left_kind, right_kind = RELATION_ENDS[kind]
        other_field, other_kind = (right, right_kind) if side == "left" else (left, left_kind)
        rows = []
        total = 0
        for row in source:
            if active_only and row["activo"] != "1":
                continue
            if (needle and not self._matches(row, needle)
                    and not self._matches(self.get(other_kind, row[other_field]), needle)):
                continue
            if offset <= total < offset + limit:
                rows.append(row.copy())
            total += 1
        result = {"rows": rows, "total": total}
        if include_references:
            result["references"] = self._page_references(kind, rows, side)
        return result

    def history_size(self) -> int:
        return self.history.length

    def suggest_id(self, kind: str) -> str:
        return new_id({"grupos": "G", "investigadores": "I",
                       "productos": "P", "planes": "PL"}[kind])

    def snapshot(self) -> str:
        state = {kind: self.rows(kind) for kind in ALL_FIELDS}
        state["cola_validacion"] = [job.copy() for job in self.queue]
        return json.dumps(state, ensure_ascii=False, separators=(",", ":"))

    def _restore(self, snapshot: str) -> None:
        state = json.loads(snapshot)
        fresh = Repository()
        for kind in ENTITY_FIELDS:
            for row in state[kind]:
                fresh.create(kind, _legacy_row(kind, row), remember=False)
        for kind in RELATION_FIELDS:
            for row in state[kind]:
                fresh.create(kind, row, remember=False)
        for job in state.get("cola_validacion", []):
            fresh._validate_job(job)
            fresh.queue.enqueue(job)
        self.entities, self.relations, self.queue = fresh.entities, fresh.relations, fresh.queue
        self.issues = fresh.issues
        self._external_index = fresh._external_index
        self._indicator_population.clear()
        self.dirty = True

    @staticmethod
    def _undo_row(op: str, kind: str, row: dict[str, str]) -> dict[str, str]:
        return {**row, "__op": op, "__kind": kind}

    def _remember(self, changes: list[dict[str, str]]) -> None:
        if not changes:
            return
        state = json.dumps({"__undo": changes}, ensure_ascii=False, separators=(",", ":"))
        size = len(state.encode("utf-8"))
        if size + self._batch_bytes > MAX_HISTORY_BYTES:
            raise DataError("La acción supera el límite de 64 MiB del historial")
        if self._batch_changes is None:
            self.history.push(state)
        else:
            self._batch_changes.append(changes)
            self._batch_bytes += size
        self.dirty = True

    def begin_batch(self) -> None:
        if self._batch_changes is not None:
            raise DataError("Ya hay una importación en curso")
        self._batch_changes = []
        self._batch_bytes = 0

    def end_batch(self) -> None:
        if self._batch_changes is None:
            raise DataError("No hay una importación en curso")
        changes = [row for operation in reversed(self._batch_changes) for row in operation]
        self._batch_changes = None
        self._batch_bytes = 0
        self._remember(changes)

    @contextmanager
    def batch(self):
        self.begin_batch()
        try:
            yield
        finally:
            self.end_batch()

    def _outside_batch(self) -> None:
        if self._batch_changes is not None:
            raise DataError("Termine la importación antes de cambiar el archivo o el historial")

    @staticmethod
    def _decode_delta(snapshot: str) -> list[dict[str, str]] | None:
        state = json.loads(snapshot)
        if "__undo" not in state:
            return None
        if set(state) != {"__undo"} or not isinstance(state["__undo"], list) or not state["__undo"]:
            raise DataError("Historial de cambios inválido")
        for change in state["__undo"]:
            if not isinstance(change, dict) or any(not isinstance(value, str) for value in change.values()):
                raise DataError("Historial de cambios inválido")
            op, kind = change.get("__op"), change.get("__kind")
            if op == "queue_remove":
                if not change.get("id"):
                    raise DataError("Historial de cola inválido")
            elif op == "queue_prepend":
                if not {"id", "producto_id", "motivo", "creado"} <= set(change):
                    raise DataError("Historial de cola inválido")
            elif op in ("create", "replace") and kind in ALL_FIELDS:
                clean_row(kind, _legacy_row(kind, {key: value for key, value in change.items()
                                                   if not key.startswith("__")}))
            elif op == "delete" and kind in ALL_FIELDS:
                keys = ("id",) if kind in ENTITY_FIELDS else RELATION_ENDS[kind][:2]
                if any(not change.get(key) for key in keys):
                    raise DataError("Historial sin clave")
            else:
                raise DataError("Operación de historial desconocida")
        return state["__undo"]

    def _validate_delta(self, changes: list[dict[str, str]]) -> None:
        # Simula solo las claves afectadas. Un historial incoherente no puede deshacer media acción.
        pending: dict[str, dict[Any, dict[str, str] | None]] = {kind: {} for kind in ALL_FIELDS}
        jobs = {job["id"]: job for job in self.queue}
        external_owners: dict[tuple[str, str], str | None] = {}

        def row_key(kind: str, row: dict[str, str]) -> str | tuple[str, str]:
            return row["id"] if kind in ENTITY_FIELDS else tuple(row[name] for name in RELATION_ENDS[kind][:2])

        def get(kind: str, key: Any) -> dict[str, str] | None:
            return pending[kind][key] if key in pending[kind] else self.get(kind, key)

        def rows(kind: str) -> Iterator[dict[str, str]]:
            source = self.entities[kind] if kind in ENTITY_FIELDS else self.relations[kind]
            for row in source:
                effective = get(kind, row_key(kind, row))
                if effective is not None:
                    yield effective
            for key, row in pending[kind].items():
                if row is not None and self.get(kind, key) is None:
                    yield row

        for change in changes:
            op, kind = change["__op"], change["__kind"]
            if op == "queue_remove":
                if jobs.pop(change["id"], None) is None:
                    raise DataError("Trabajo de historial no encontrado")
                continue
            if op == "queue_prepend":
                job = {name: change[name] for name in QUEUE_FIELDS}
                if (not job["id"] or not job["creado"] or job["id"] in jobs
                        or get("productos", job["producto_id"]) is None):
                    raise DataError("Trabajo de historial inválido")
                jobs[job["id"]] = job
                continue
            row = _legacy_row(kind, {name: value for name, value in change.items() if not name.startswith("__")})
            if op != "delete":
                row = clean_row(kind, row)
            key = row_key(kind, row)
            current = get(kind, key)
            if (op == "create" and current is not None) or (op != "create" and current is None):
                raise DataError("Historial incompatible con los registros actuales")
            if op == "delete":
                if kind == "grupos" and any(plan["grupo_id"] == key for plan in rows("planes")):
                    raise DataError("El grupo del historial aún tiene planes")
                for rel, (left, right, left_kind, right_kind) in RELATION_ENDS.items():
                    endpoint = left if kind == left_kind else right if kind == right_kind else None
                    if endpoint and any(link[endpoint] == key for link in rows(rel)):
                        raise DataError("El registro del historial aún tiene relaciones")
                if kind == "productos" and any(job["producto_id"] == key for job in jobs.values()):
                    raise DataError("El producto del historial aún tiene revisiones")
                external = self._external_fields.get(kind)
                if external:
                    normalized = self._normalized_external(kind, current[external])
                    if normalized:
                        external_owners[kind, normalized] = None
                pending[kind][key] = None
                continue
            if kind == "planes" and get("grupos", row["grupo_id"]) is None:
                raise DataError("El grupo del plan de historial no existe")
            if kind in RELATION_ENDS:
                left, right, left_kind, right_kind = RELATION_ENDS[kind]
                if get(left_kind, row[left]) is None or get(right_kind, row[right]) is None:
                    raise DataError("La relación del historial apunta a una entidad inexistente")
            external = self._external_fields.get(kind)
            if external:
                normalized = self._normalized_external(kind, row[external])
                previous = self._normalized_external(kind, current[external]) if current is not None else ""
                if normalized:
                    identity = (kind, normalized)
                    owner = (external_owners[identity] if identity in external_owners
                             else self._external_index[kind].get(normalized))
                    if owner is not None and owner != key:
                        raise DataError("Identidad externa duplicada en el historial")
                if previous and previous != normalized:
                    external_owners[kind, previous] = None
                if normalized:
                    external_owners[kind, normalized] = key
            pending[kind][key] = row

    def _apply_delta(self, changes: list[dict[str, str]]) -> None:
        self._validate_delta(changes)
        self.dirty = True
        for change in changes:
            op, kind = change["__op"], change["__kind"]
            if op == "queue_remove":
                self.queue.remove(change["id"])
            elif op == "queue_prepend":
                job = {name: change[name] for name in QUEUE_FIELDS}
                self._validate_job(job)
                self.queue.prepend(job)
            else:
                row = _legacy_row(kind, {key: value for key, value in change.items()
                                         if not key.startswith("__")})
                key = row["id"] if kind in ENTITY_FIELDS else tuple(row[name] for name in RELATION_ENDS[kind][:2])
                if op == "delete":
                    self.delete(kind, key, remember=False)
                elif op == "create":
                    self.create(kind, row, remember=False)
                else:
                    current = self.get(kind, key)
                    if current is None:
                        raise DataError("Registro de historial no encontrado")
                    row = clean_row(kind, row)
                    self._check_links(kind, row)
                    if kind in ENTITY_FIELDS:
                        self._unique_external(kind, row, str(key))
                        self._reindex_external(kind, current, row)
                    current.update(row)
                    self._revalidate(self._affected_products(kind, current))
        self.dirty = True

    def undo(self) -> bool:
        self._outside_batch()
        previous = self.history.peek()
        if previous is None:
            return False
        changes = self._decode_delta(previous)
        if changes is None:
            self._restore(previous)
        else:
            self._apply_delta(changes)
        self.history.pop()
        return True

    def clear_history(self) -> None:
        self._outside_batch()
        self.history.clear()
        self.dirty = True

    def _check_links(self, kind: str, row: dict[str, str]) -> None:
        if kind == "planes" and not self.get("grupos", row["grupo_id"]):
            raise DataError("El grupo del plan no existe")
        if kind in RELATION_ENDS:
            left, right, left_kind, right_kind = RELATION_ENDS[kind]
            if not self.get(left_kind, row[left]) or not self.get(right_kind, row[right]):
                raise DataError("La relación apunta a una entidad inexistente")

    @staticmethod
    def _normalized_external(kind: str, value: str) -> str:
        return _doi_text(value) if kind == "productos" else value.strip().lower()

    def _unique_external(self, kind: str, row: dict[str, str], old_id: str = "") -> None:
        field = self._external_fields.get(kind)
        if field:
            normalized = self._normalized_external(kind, row[field])
            owner = self._external_index[kind].get(normalized) if normalized else None
            if owner is not None and owner != old_id:
                raise DataError(f"{LABELS[field]} ya registrado en {owner}")

    def _reindex_external(self, kind: str, old: dict[str, str] | None,
                          new: dict[str, str] | None) -> None:
        field = self._external_fields.get(kind)
        if not field:
            return
        previous = self._normalized_external(kind, old[field]) if old is not None else ""
        current = self._normalized_external(kind, new[field]) if new is not None else ""
        if previous == current:
            return
        if current:
            self._external_index[kind][current] = new["id"]
        if previous:
            del self._external_index[kind][previous]

    def find_external(self, kind: str, value: str) -> dict[str, str] | None:
        if kind not in self._external_fields:
            raise DataError("Tipo sin identidad externa")
        needle = self._normalized_external(kind, value)
        owner = self._external_index[kind].get(needle) if needle else None
        row = self.get(kind, owner) if owner is not None else None
        return row.copy() if row is not None else None

    def create(self, kind: str, values: dict[str, Any], remember: bool = True) -> dict[str, str]:
        row = clean_row(kind, values)
        self._check_links(kind, row)
        if kind in ENTITY_FIELDS:
            self._unique_external(kind, row)
            if self.get(kind, row["id"]):
                raise DataError("ID duplicado")
        else:
            left, right = RELATION_ENDS[kind][:2]
            if self.get(kind, (row[left], row[right])):
                raise DataError("Relación duplicada")
        if remember:
            fields = ("id",) if kind in ENTITY_FIELDS else RELATION_ENDS[kind][:2]
            self._remember([self._undo_row("delete", kind, {field: row[field] for field in fields})])
        if kind in ENTITY_FIELDS:
            self.entities[kind].append(row)
            self._reindex_external(kind, None, row)
        else:
            self.relations[kind].add(row)
        self._revalidate(self._affected_products(kind, row))
        return row

    def update(self, kind: str, key: str | tuple[str, str], values: dict[str, Any], remember: bool = True) -> dict[str, str]:
        current = self.get(kind, key)
        if current is None:
            raise DataError("Registro no encontrado")
        updated_values = {**current, **values}
        if kind == "productos" and "categoria" in values and "categoria_origen" not in values:
            category = str(values["categoria"] or "").strip()
            if category != current["categoria"]:
                updated_values["categoria_origen"] = "manual" if category else ""
        updated = clean_row(kind, updated_values)
        fixed = ("id",) if kind in ENTITY_FIELDS else RELATION_ENDS[kind][:2]
        if any(updated[field] != current[field] for field in fixed):
            raise DataError("Los identificadores son estables; cree otra relación si cambian")
        self._check_links(kind, updated)
        if kind in ENTITY_FIELDS:
            self._unique_external(kind, updated, current["id"])
        if remember:
            self._remember([self._undo_row("replace", kind, current)])
        if kind in ENTITY_FIELDS:
            self._reindex_external(kind, current, updated)
        current.update(updated)
        self._revalidate(self._affected_products(kind, current))
        return current

    def delete(self, kind: str, key: str | tuple[str, str], remember: bool = True) -> None:
        row = self.get(kind, key)
        if not row:
            raise DataError("Registro no encontrado")
        if kind in ENTITY_FIELDS:
            if kind == "grupos" and any(p["grupo_id"] == key for p in self.entities["planes"]):
                raise DataError("El grupo tiene planes; elimínelos o desactívelo")
            for rel_kind, (left, right, left_kind, right_kind) in RELATION_ENDS.items():
                field = left if kind == left_kind else right if kind == right_kind else None
                if field and any(rel[field] == key for rel in self.relations[rel_kind]):
                    raise DataError("Hay relaciones asociadas; elimínelas o desactive el registro")
            if kind == "productos" and any(job["producto_id"] == key for job in self.queue):
                raise DataError("Hay revisiones pendientes para el producto")
        if remember:
            self._remember([self._undo_row("create", kind, row)])
        affected = self._affected_products(kind, row)
        if kind in ENTITY_FIELDS:
            self._reindex_external(kind, row, None)
            self.entities[kind].remove(str(key))
        else:
            self.relations[kind].remove(*key)
        self._revalidate(affected)

    def toggle(self, kind: str, key: str | tuple[str, str]) -> None:
        row = self.get(kind, key)
        if not row:
            raise DataError("Registro no encontrado")
        self.update(kind, key, {"activo": "0" if row["activo"] == "1" else "1"})

    def product_scope(self, view: str = "Todos", selected_id: str = "") -> list[dict[str, str]]:
        if view == "Grupo":
            ids = {r["producto_id"] for r in self.relations["grupos_productos"].by_left(selected_id) if r["activo"] == "1"} if selected_id else set()
        elif view == "Investigador":
            ids = {r["producto_id"] for r in self.relations["autorias"].by_right(selected_id) if r["activo"] == "1"} if selected_id else set()
        elif view == "Producto":
            ids = {selected_id} if selected_id else set()
        else:
            ids = None
        return [p for p in self.entities["productos"] if p["activo"] == "1" and (ids is None or p["id"] in ids)]

    def statistics(self, view: str = "Todos", selected_id: str = "", start: int | None = None,
                   end: int | None = None, status: str = "",
                   offset: int = 0, limit: int | None = None) -> dict[str, Any]:
        # Hipercubo lógico: cada producto es un hecho; grupo e investigador llegan
        # por multilistas, y año, tipología y validación son dimensiones.
        # La vista elige grupo/investigador/producto; los filtros recortan el conjunto
        # y estos conteos se calculan bajo demanda. No hay cubo OLAP materializado
        # ni filtro simultáneo por grupo e investigador.
        if view not in ("Todos", "Grupo", "Investigador", "Producto"):
            raise DataError("Vista desconocida")
        if start is not None and end is not None and start > end:
            raise DataError("El año inicial supera al final")
        if status and status not in VALIDATION_STATES:
            raise DataError("Validación: validado o rechazado")
        products = []
        total = 0
        tallies: dict[str, dict[str, int]] = {name: {} for name in ("anio", "tipologia", "categoria", "validacion")}
        rules: dict[str, int] = {}
        for row in self.product_scope(view, selected_id):
            year = int(row["anio"]) if row["anio"] else None
            if start is not None and (year is None or year < start):
                continue
            if end is not None and (year is None or year > end):
                continue
            if status and row["validacion"] != status:
                continue
            if total >= offset and (limit is None or len(products) < limit):
                products.append(row.copy())
            total += 1
            for name, counts in tallies.items():
                key = row[name] or "Sin dato"
                counts[key] = counts.get(key, 0) + 1
            for code in self.issues.get(row["id"], ()):
                rules[code] = rules.get(code, 0) + 1
        return {"productos": products, "total": total,
                "por_anio": tallies["anio"], "por_tipologia": tallies["tipologia"],
                "por_categoria": tallies["categoria"],
                "por_validacion": tallies["validacion"], "por_regla": rules}


    def _indicator_raw(self, kind: str, key: str, corte: int, inicio: int | None,
                       status: str, groups_of: dict[str, set[str]]) -> dict[str, Any]:
        """Valores sin normalizar de una entidad, recorriendo sus vínculos activos."""
        result: dict[str, Any] = {
            "valores": dict.fromkeys((*MODEL_CLASSES, "COHE", "COOP"), 0.0),
            "productos": 0, "sin_categoria": 0, "estimadas": 0, "subtipos": [],
        }
        period_cap = corte - inicio + 1 if inicio is not None else None
        if kind == "grupos":
            created = self.get(kind, key)["fecha_creacion"]
            if created:
                age = corte - int(created[:4]) + 1
                period_cap = min(period_cap, age) if period_cap is not None else age
            links = self.relations["grupos_productos"].by_left(key)
        else:
            links = self.relations["autorias"].by_right(key)
        if period_cap is not None and period_cap < 1:
            return result

        uncategorized_start = max(corte - 9, inicio) if inicio is not None else corte - 9
        empty_groups: set[str] = set()
        researcher_groups = groups_of.get(key, empty_groups)
        counts: dict[str, int] = {}
        new_knowledge = authors = group_links = 0
        for link in links:
            if link["activo"] != "1":
                continue
            product = self.get("productos", link["producto_id"])
            if product is None or product["activo"] != "1" or not product["anio"]:
                continue
            if status and product["validacion"] != status:
                continue
            year = int(product["anio"])
            code = product["categoria"]
            if not code:
                if uncategorized_start <= year <= corte:
                    result["sin_categoria"] += 1
                continue
            subtype = MODEL_2024.get(code)
            if subtype is None:
                continue
            period = min(subtype.ventana, period_cap) if period_cap is not None else subtype.ventana
            if not corte - period + 1 <= year <= corte:
                continue
            counts[code] = counts.get(code, 0) + 1
            result["productos"] += 1
            if product["categoria_origen"] == "estimada":
                result["estimadas"] += 1
            if subtype.clase not in NEW_KNOWLEDGE_CLASSES:
                continue
            new_knowledge += 1
            product_groups = {key} if kind == "grupos" else set()
            for authorship in self.relations["autorias"].by_left(product["id"]):
                if authorship["activo"] != "1":
                    continue
                researcher_id = authorship["investigador_id"]
                researcher = self.get("investigadores", researcher_id)
                if researcher is None or researcher["activo"] != "1":
                    continue
                author_groups = groups_of.get(researcher_id, empty_groups)
                if kind == "grupos":
                    belongs = key in author_groups
                else:
                    # S contiene al investigador y a los miembros de cualquiera de sus grupos activos.
                    belongs = researcher_id == key or not researcher_groups.isdisjoint(author_groups)
                if belongs:
                    authors += 1
                product_groups.update(author_groups)
            group_links += len(product_groups)

        for code, subtype in MODEL_2024.items():
            count = counts.get(code, 0)
            if not count:
                continue
            period = min(subtype.ventana, period_cap) if period_cap is not None else subtype.ventana
            frequency = math.log(count / period + 1)
            value = frequency * subtype.peso_global
            result["valores"][subtype.clase] += value
            result["subtipos"].append({"codigo": code, "clase": subtype.clase, "cantidad": count,
                                       "periodo": period, "lambda": frequency,
                                       "peso_global": subtype.peso_global, "valor": value})
        if new_knowledge:
            result["valores"]["COHE"] = max(0.0, authors / new_knowledge - 1)
            result["valores"]["COOP"] = max(0.0, group_links / new_knowledge - 1)
        return result

    def indicator(self, kind: str, key: str, start: int | None = None,
                  end: int | None = None, status: str = "") -> dict[str, Any]:
        """Indicador 2024 relativo al máximo de la población activa del mismo tipo."""
        if kind not in ("grupos", "investigadores"):
            raise DataError("Tipo de indicador: grupos o investigadores")
        if start is not None and end is not None and start > end:
            raise DataError("El año inicial supera al final")
        if status and status not in VALIDATION_STATES:
            raise DataError("Validación: validado o rechazado")
        entity = self.get(kind, key)
        if entity is None:
            raise DataError("Registro no encontrado")
        if entity["activo"] != "1":
            raise DataError("Registro inactivo")
        cutoff = end if end is not None else date.today().year

        # Índice auxiliar de esta consulta; las listas siguen siendo la fuente de verdad.
        active_groups = {row["id"] for row in self.entities["grupos"] if row["activo"] == "1"}
        groups_of: dict[str, set[str]] = {}
        for membership in self.relations["membresias"]:
            if membership["activo"] != "1" or membership["grupo_id"] not in active_groups:
                continue
            researcher_id = membership["investigador_id"]
            researcher = self.get("investigadores", researcher_id)
            if researcher is not None and researcher["activo"] == "1":
                groups_of.setdefault(researcher_id, set()).add(membership["grupo_id"])

        raw = self._indicator_raw(kind, key, cutoff, start, status, groups_of)
        cache_key = (kind, start, end, status, cutoff)
        cached = self._indicator_population.get(cache_key)
        if cached is None:
            maxima = raw["valores"].copy()
            population = 1
            for other in self.entities[kind]:
                if other["activo"] != "1" or other["id"] == key:
                    continue
                population += 1
                values = self._indicator_raw(kind, other["id"], cutoff, start, status, groups_of)["valores"]
                for name in maxima:
                    maxima[name] = max(maxima[name], values[name])
            if len(self._indicator_population) >= 16:
                del self._indicator_population[next(iter(self._indicator_population))]
            self._indicator_population[cache_key] = (maxima, population)
        else:
            maxima, population = cached

        def term(name: str) -> dict[str, float]:
            value, maximum = raw["valores"][name], maxima[name]
            index = value / maximum if maximum else 0.0
            coefficient = MODEL_COEFFICIENTS[name]
            return {"coeficiente": coefficient, "valor": value, "maximo": maximum,
                    "indice": index, "aporte": coefficient * index}

        classes = [{"clase": name, **term(name)} for name in MODEL_CLASSES]
        cohesion, cooperation = term("COHE"), term("COOP")
        return {"tipo": kind, "id": key, "corte": cutoff, "inicio": start, "estado": status,
                "poblacion": population, "productos": raw["productos"],
                "sin_categoria": raw["sin_categoria"], "estimadas": raw["estimadas"],
                "clases": classes, "cohesion": cohesion, "cooperacion": cooperation,
                "subtipos": raw["subtipos"],
                "indicador": sum(c["aporte"] for c in classes) + cohesion["aporte"] + cooperation["aporte"]}

    def queue_rows(self) -> list[dict[str, str]]:
        return list(self.queue)
    def queue_page(self, offset: int = 0, limit: int = PAGE_SIZE) -> dict[str, Any]:
        rows = []
        for index, job in enumerate(self.queue):
            if offset <= index < offset + limit:
                rows.append(job.copy())
            if index >= offset + limit:
                break
        return {"rows": rows, "total": self.queue.length}
    def queue_front(self) -> dict[str, str] | None:
        return self.queue.peek()
    def queue_size(self) -> int:
        return self.queue.length
    def _validate_job(self, job: dict[str, str]) -> None:
        if set(job) != {"id", "producto_id", "motivo", "creado"}:
            raise DataError("Formato inválido de la cola")
        if not self.get("productos", job["producto_id"]):
            raise DataError("La cola referencia un producto inexistente")
        if not job["id"] or not job["creado"]:
            raise DataError("Trabajo de cola incompleto")
    @staticmethod
    def _new_job(product_id: str, reason: str) -> dict[str, str]:
        return {"id": uuid.uuid4().hex[:12], "producto_id": product_id,
                "motivo": reason.strip(), "creado": datetime.now().isoformat(timespec="seconds")}
    def enqueue_review(self, product_id: str, reason: str) -> None:
        if not self.get("productos", product_id):
            raise DataError("Producto inexistente")
        if any(job["producto_id"] == product_id for job in self.queue):
            raise DataError("El producto ya está en la cola")
        job = self._new_job(product_id, reason)
        self._remember([self._undo_row("queue_remove", "", {"id": job["id"]})])
        self.queue.enqueue(job)
    def enqueue_rejected(self) -> int:
        """Encola los productos activos rechazados que aún no esperan revisión (una acción)."""
        queued = {job["producto_id"] for job in self.queue}
        targets = [row["id"] for row in self.entities["productos"]
                   if row["activo"] == "1" and row["validacion"] == "rechazado" and row["id"] not in queued]
        if targets:
            jobs = [self._new_job(product_id, AUTO_REASON_PREFIX + ",".join(self.issues.get(product_id, ())))
                    for product_id in targets]
            self._remember([self._undo_row("queue_remove", "", {"id": job["id"]}) for job in jobs])
            for job in jobs:
                self.queue.enqueue(job)
        return len(targets)
    def process_review(self, observation: str) -> dict[str, str]:
        """Cierra la revisión del frente con una nota; la validación sigue siendo automática."""
        job = self.queue.peek()
        if not job:
            raise DataError("La cola está vacía")
        if not observation.strip():
            raise DataError("Debe registrar una observación")
        current = self.get("productos", job["producto_id"])
        if current is None:
            raise DataError("El producto ya no existe")
        clean_row("productos", {**current, "observacion": observation})
        self._remember([self._undo_row("replace", "productos", current),
                        self._undo_row("queue_prepend", "", job)])
        self.update("productos", job["producto_id"], {"observacion": observation}, remember=False)
        self.queue.dequeue()
        return {**job, "validacion": current["validacion"]}
    def discard_review(self) -> dict[str, str]:
        if not self.queue.peek():
            raise DataError("La cola está vacía")
        self._remember([self._undo_row("queue_prepend", "", self.queue.peek())])
        return self.queue.dequeue()  # type: ignore[return-value]

def _storage_cell(value: str) -> str:
    return "'" + value if value.startswith(("=", "+", "-", "@", "\t", "\r", "'")) else value


def _csv_lines(stream: Any, path: Path, allow_blank: bool) -> Iterator[str]:
    quoted = False
    for number, line in enumerate(stream, start=1):
        if not allow_blank and not quoted and line in ("\n", "\r", "\r\n"):
            raise DataError(f"{path.name}, línea {number}: fila CSV vacía")
        # csv.strict comprueba cierres y escapes, pero acepta comillas en campos sin entrecomillar.
        for match in re.finditer('"', line):
            position = match.start()
            if not quoted and position and line[position - 1] not in ',\r\n"':
                raise DataError(f"{path.name}, línea {number}: comilla CSV mal ubicada")
            quoted = not quoted
        yield line


@contextmanager
def _csv_input(path: Path, *, allow_blank: bool = True):
    with path.open("r", encoding="utf-8-sig", newline="") as stream:
        try:
            yield _csv_lines(stream, path, allow_blank)
        except csv.Error as exc:
            raise DataError(f"{path.name}: CSV inválido: {exc}") from exc


def _csv_read(path: Path, fields: tuple[str, ...], *, version: str = SCHEMA_VERSION,
              encoded: bool = False) -> list[dict[str, str]]:
    # Las carpetas v1 a v3 usan el encabezado de productos anterior a v5.
    # Los CSV sueltos aceptan ambos encabezados, también con un nombre de archivo distinto.
    legacy = version in ("1", "2")
    with _csv_input(path, allow_blank=False) as stream:
        reader = csv.DictReader(stream, strict=True)
        actual = set(reader.fieldnames or ())
        product_fields = fields == ENTITY_FIELDS["productos"]
        expected_fields = PRODUCT_FIELDS_V4 if product_fields and version in FOLDER_VERSIONS else fields
        expected = set(expected_fields)
        v4_products = product_fields and actual == set(PRODUCT_FIELDS_V4)
        # v1: grupos con «sigla» y productos sin «validacion»; v1 y v2: productos con «categoria».
        old_groups = version == "1" and path.stem == "grupos" and actual == expected | {"sigla"}
        old_products = legacy and path.stem == "productos" and (
            actual == expected | {"categoria"}
            or (version == "1" and actual in ((expected | {"categoria"}) - {"validacion"},
                                              expected - {"validacion"})))
        if (reader.fieldnames is None or len(reader.fieldnames) != len(actual)
                or not (actual == expected or v4_products or old_groups or old_products)):
            raise DataError(f"Encabezados incorrectos en {path.name}; se esperan: {', '.join(expected_fields)}")
        rows = []
        for number, row in enumerate(reader, start=2):
            if None in row or any(value is None for value in row.values()):
                raise DataError(f"Fila {number} malformada en {path.name}")
            if len(rows) >= MAX_CSV_ROWS:
                raise DataError(f"Demasiadas filas en {path.name}")
            if encoded:
                row = {key: value[1:] if value.startswith("'") else value for key, value in row.items()}
            rows.append(_legacy_row(path.stem, row) if legacy else row)
        return rows


@contextmanager
def _csv_output(path: Path):
    with path.open("w", encoding="utf-8", newline="") as stream:
        yield csv.writer(stream, lineterminator="\n")
        stream.flush()
        os.fsync(stream.fileno())


def _csv_write(path: Path, fields: tuple[str, ...], rows: list[dict[str, str]],
               *, spreadsheet: bool = False) -> None:
    encode = _neutralize_cell if spreadsheet else _storage_cell
    with _csv_output(path) as writer:
        writer.writerow(fields)
        writer.writerows([encode(row.get(field, "")) for field in fields] for row in rows)


def _cell_text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value)


def _check_local_file(path: Path, *, archive: bool = False) -> None:
    try:
        size = path.stat().st_size
    except OSError as exc:
        raise DataError(f"No se puede abrir el archivo: {exc}") from exc
    if size > MAX_LOCAL_FILE:
        raise DataError("El archivo supera 32 MB")
    if archive:
        try:
            with zipfile.ZipFile(path) as package:
                members = package.infolist()
                if len(members) > 10_000 or sum(item.file_size for item in members) > MAX_ARCHIVE_UNPACKED:
                    raise DataError("El documento comprimido supera los límites de extracción")
        except (zipfile.BadZipFile, zipfile.LargeZipFile, OSError) as exc:
            raise DataError("El archivo no es un documento Office válido") from exc


def _xlsx_to_csv(source: Path, target: Path) -> Path:
    """Convierte la primera hoja de un .xlsx a CSV respetando los límites de importación."""
    _check_local_file(source, archive=True)
    try:
        import openpyxl
    except ImportError as exc:
        raise DataError("Para importar Excel instale openpyxl (pip install openpyxl)") from exc
    try:
        workbook = openpyxl.load_workbook(source, read_only=True, data_only=True)
    except Exception as exc:
        raise DataError(f"No se pudo leer el Excel: {exc}") from exc
    try:
        sheet = workbook.active
        rows = sheet.iter_rows(values_only=True)
        header = next(rows, None)
        header_cells = [_cell_text(cell) for cell in (header or ())]
        if not any(header_cells):
            raise DataError("El Excel no tiene encabezados")
        if any(not name for name in header_cells):
            raise DataError("El Excel tiene columnas sin encabezado")
        if any(len(name) > MAX_FIELD_LEN for name in header_cells):
            raise DataError("El Excel tiene un encabezado demasiado largo")
        with target.open("w", encoding="utf-8", newline="") as stream:
            writer = csv.writer(stream, lineterminator="\n")
            writer.writerow(header_cells)
            count = 1
            for record in rows:
                count += 1
                if count > MAX_CSV_ROWS:
                    raise DataError("El Excel tiene demasiadas filas")
                cells = [_cell_text(cell) for cell in record]
                if any(len(cell) > MAX_FIELD_LEN for cell in cells):
                    raise DataError("El Excel tiene un valor demasiado largo")
                writer.writerow(cells)
    except DataError:
        raise
    except Exception as exc:
        raise DataError(f"No se pudo leer el Excel: {exc}") from exc
    finally:
        workbook.close()
    return target


def _neutralize_cell(value: str) -> str:
    if value and value[0] in "=+-@\t\r":
        return "'" + value
    return value


def export_repository_spreadsheet(repo: Repository | CppRepository, directory: Path) -> None:
    """Exporta los datos neutralizados para abrirlos en hojas de cálculo sin fórmulas."""
    directory = Path(directory)
    directory.mkdir(parents=True, exist_ok=True)
    _csv_write(directory / "manifest.csv", MANIFEST_FIELDS,
               [{"version": SCHEMA_VERSION, "guardado": datetime.now().isoformat(timespec="seconds")}], spreadsheet=True)
    for kind, fields in ALL_FIELDS.items():
        _csv_write(directory / f"{kind}.csv", fields, repo.rows(kind), spreadsheet=True)


def is_folder_data(path: Path) -> bool:
    """Carpeta de las versiones 1 a 3, indicada por la carpeta o por su manifest.csv."""
    path = Path(path)
    return path.is_dir() or path.name == "manifest.csv"


def backup_of(path: Path) -> Path:
    """Copia del guardado anterior: .backup junto al archivo (o dentro de una carpeta anterior)."""
    path = Path(path)
    if is_folder_data(path):
        return (path if path.is_dir() else path.parent) / ".backup"
    return path.parent / ".backup" / path.name


def save_repository(repo: Repository, path: Path) -> None:
    """Escribe el archivo único en un temporal y lo sustituye de una vez; conserva la copia anterior."""
    repo._outside_batch()
    path = Path(path)
    if path.is_dir():
        raise DataError("Indique un archivo .csv, no una carpeta")
    path.parent.mkdir(parents=True, exist_ok=True)
    tables = {"manifest": [{"version": SCHEMA_VERSION, "guardado": datetime.now().isoformat(timespec="seconds")}],
              **{kind: repo.rows(kind) for kind in ALL_FIELDS},
              "cola_validacion": list(repo.queue),
              "historial": [{"orden": str(index), "snapshot": state} for index, state in enumerate(repo.history)]}
    temporary = path.with_name(f".{path.name}.{os.getpid()}.tmp")
    try:
        with _csv_output(temporary) as writer:
            for section, fields in DATA_SECTIONS.items():
                writer.writerow([f"#{section}"])
                writer.writerow(fields)
                writer.writerows([_storage_cell(row.get(field, "")) for field in fields] for row in tables[section])
        if path.exists():
            backup = backup_of(path)
            backup.parent.mkdir(exist_ok=True)
            shutil.copy2(path, backup)
        os.replace(temporary, path)
    finally:
        temporary.unlink(missing_ok=True)
    repo.dirty = False


def read_data_sections(path: Path) -> Iterator[tuple[str, int, dict[str, str]]]:
    """Recorre el archivo único en orden y entrega (sección, línea, fila).

    Cada sección empieza con una fila de una sola celda «#nombre», sigue su encabezado y
    luego sus filas hasta la siguiente marca. Ninguna tabla tiene una sola columna, así que
    una fila de datos nunca se confunde con una marca.
    """
    path = Path(path)
    expected = iter(DATA_SECTIONS)
    section, header, count, previous = "", None, 0, 0
    version = ""
    with _csv_input(path) as stream:
        reader = csv.reader(stream, strict=True)
        for record in reader:
            line, previous = previous + 1, reader.line_num
            where = f"{path.name}, línea {line}"
            if not record or record == [""]:
                continue
            if len(record) == 1 and record[0].startswith("#"):
                if section and header is None:
                    raise DataError(f"{where}: falta el encabezado de la sección #{section}")
                wanted = next(expected, None)
                if wanted is None:
                    raise DataError(f"{where}: sección sobrante {record[0]}")
                if record[0] != f"#{wanted}":
                    raise DataError(f"{where}: se esperaba la sección #{wanted}")
                section, header, count = wanted, None, 0
                continue
            if not section:
                raise DataError(f"{where}: el archivo debe comenzar con #manifest")
            if header is None:
                fields = PRODUCT_FIELDS_V4 if section == "productos" and version == "4" else DATA_SECTIONS[section]
                if len(record) != len(fields) or set(record) != set(fields):
                    raise DataError(f"{where}: encabezados incorrectos en #{section}; se esperan: {', '.join(fields)}")
                header = record
                continue
            if len(record) != len(header):
                raise DataError(f"{where}: fila malformada en #{section}")
            count += 1
            if count > MAX_CSV_ROWS:
                raise DataError(f"{where}: demasiadas filas en #{section}")
            row = {name: value[1:] if value.startswith("'") else value
                   for name, value in zip(header, record)}
            if section == "manifest":
                version = row["version"]
            yield section, line, row
    missing = next(expected, None)
    if missing is not None:
        raise DataError(f"{path.name}: falta la sección #{missing}")
    if header is None:
        raise DataError(f"{path.name}: falta el encabezado de la sección #{section}")


def _load_history(repo: Repository, entries: list[dict[str, str]]) -> None:
    if len(entries) > 30:
        raise DataError("Historial supera 30 acciones")
    for entry in reversed(entries):
        if len(entry["snapshot"].encode("utf-8")) > MAX_HISTORY_BYTES:
            raise DataError("El historial supera 64 MiB")
        # Cada snapshot se verifica antes de incorporarlo al historial.
        validator = Repository()
        try:
            if validator._decode_delta(entry["snapshot"]) is None:
                validator._restore(entry["snapshot"])
        except (KeyError, TypeError, ValueError, json.JSONDecodeError) as exc:
            raise DataError("Historial dañado") from exc
        repo.history.push(entry["snapshot"])


def _load_file(path: Path) -> Repository:
    """Crea los registros a medida que los lee: las entidades llegan antes que sus relaciones."""
    repo = Repository()
    manifests = 0
    job_ids: set[str] = set()
    history: list[dict[str, str]] = []
    for section, line, row in read_data_sections(path):
        try:
            if section == "manifest":
                manifests += 1
                if manifests > 1:
                    raise DataError("Manifest con más de una fila")
                if row["version"] not in FILE_VERSIONS:
                    raise DataError("Versión de datos incompatible")
            elif section in ALL_FIELDS:
                repo.create(section, row, remember=False)
            elif section == "cola_validacion":
                repo._validate_job(row)
                if row["id"] in job_ids:
                    raise DataError(f"Trabajo de cola duplicado: {row['id']}")
                job_ids.add(row["id"])
                repo.queue.enqueue(row)
            else:
                if len(history) >= 30:
                    raise DataError("El historial supera 30 acciones")
                history.append(row)
        except DataError as exc:
            raise DataError(f"{path.name}, línea {line} (#{section}): {exc}") from exc
    if manifests != 1:
        raise DataError(f"{path.name}: falta la fila de #manifest")
    _load_history(repo, history)
    repo.dirty = False
    return repo


def _load_folder(directory: Path) -> Repository:
    """Carpeta de las versiones 1 a 3 (un CSV por tabla); se guarda luego como archivo único."""
    manifest = _csv_read(directory / "manifest.csv", MANIFEST_FIELDS, encoded=True)
    if len(manifest) != 1 or manifest[0]["version"] not in FOLDER_VERSIONS:
        raise DataError("Versión de datos incompatible")
    version = manifest[0]["version"]
    legacy = version == "1"
    repo = Repository()
    for kind in ENTITY_FIELDS:
        for number, row in enumerate(_csv_read(directory / f"{kind}.csv", ENTITY_FIELDS[kind],
                                               version=version, encoded=not legacy), start=2):
            try:
                repo.create(kind, row, remember=False)
            except DataError as exc:
                raise DataError(f"{kind}.csv, fila {number}: {exc}") from exc
    for kind in RELATION_FIELDS:
        for number, row in enumerate(_csv_read(directory / f"{kind}.csv", RELATION_FIELDS[kind],
                                               encoded=not legacy), start=2):
            try:
                repo.create(kind, row, remember=False)
            except DataError as exc:
                raise DataError(f"{kind}.csv, fila {number}: {exc}") from exc
    job_ids: set[str] = set()
    queue_path = directory / "cola_validacion.csv"
    if not queue_path.exists() and not legacy:
        raise DataError("Falta cola_validacion.csv")
    jobs = _csv_read(queue_path, QUEUE_FIELDS, encoded=not legacy) if queue_path.exists() else []
    for job in jobs:
        repo._validate_job(job)
        if job["id"] in job_ids:
            raise DataError(f"Trabajo de cola duplicado: {job['id']}")
        job_ids.add(job["id"])
        repo.queue.enqueue(job)
    _load_history(repo, _csv_read(directory / "historial.csv", HISTORY_FIELDS, encoded=not legacy))
    repo.dirty = False
    return repo


def load_repository(path: Path) -> Repository:
    """Abre el archivo único o, para convertirla, una carpeta anterior (o su manifest.csv)."""
    path = Path(path)
    if is_folder_data(path):
        return _load_folder(path if path.is_dir() else path.parent)
    return _load_file(path)


def merge_csv(repo: Repository, path: Path, kind: str) -> tuple[int, list[str]]:
    """Importa filas válidas en lote y conserva una sola acción para deshacer."""
    if kind not in ALL_FIELDS:
        raise DataError("Tipo de CSV desconocido")
    rows = _csv_read(Path(path), ALL_FIELDS[kind])
    accepted, errors = 0, []
    with repo.batch():
        for number, row in enumerate(rows, start=2):
            try:
                repo.create(kind, row)
                accepted += 1
            except DataError as exc:
                errors.append(f"Fila {number}: {exc}")
    return accepted, errors


def preview_csv_merge(repo: Repository, path: Path, kind: str) -> tuple[int, int, list[str]]:
    """Comprueba en una copia todos los aciertos/rechazos antes de mutar datos."""
    if kind not in ALL_FIELDS:
        raise DataError("Tipo de CSV desconocido")
    rows = _csv_read(Path(path), ALL_FIELDS[kind])
    staged = Repository()
    staged._restore(repo.snapshot())
    accepted, errors = 0, []
    for number, row in enumerate(rows, start=2):
        try:
            staged.create(kind, row, remember=False)
            accepted += 1
        except DataError as exc:
            errors.append(f"Fila {number}: {exc}")
    return len(rows), accepted, errors


def bundled_windows_cpp(root: Path) -> Path | None:
    """Usa el binario incluido solo si coincide con las fuentes y con su huella propia."""
    binary = root / "bin" / "windows" / "pea_cpp.exe"
    fingerprint = binary.with_suffix(".source-sha256")
    integrity = root / "bin" / "windows" / "pea_cpp.exe.sha256"
    if not binary.is_file() or not fingerprint.is_file() or not integrity.is_file():
        return None
    try:
        digest = hashlib.sha256()
        for source in (root / "CMakeLists.txt", root / "src/cpp/Taller2_REMR.cpp"):
            digest.update(source.relative_to(root).as_posix().encode("utf-8"))
            digest.update(b"\0")
            digest.update(source.read_bytes().replace(b"\r\n", b"\n"))
        if fingerprint.read_text(encoding="ascii").strip() != digest.hexdigest():
            return None
        binary_digest = hashlib.sha256(binary.read_bytes()).hexdigest()
        if integrity.read_text(encoding="ascii").strip() != binary_digest:
            return None
        return binary
    except (OSError, UnicodeError):
        pass
    return None


def cpp_executable(explicit: Path | None = None) -> Path:
    """Localiza o compila el ejecutable que atiende a la interfaz Tkinter."""
    root = Path(__file__).resolve().parents[2]
    if explicit is not None:
        binary = Path(explicit).expanduser().resolve()
        if not binary.is_file():
            raise DataError(f"No existe el backend C++: {binary}")
        return binary
    candidates = [root / "build" / "mingw" / "pea_cpp.exe",
                  root / "build" / "default" / "Release" / "pea_cpp.exe",
                  root / "build" / "default" / "Release" / "pea_cpp",
                  root / "build" / "default" / "pea_cpp.exe",
                  root / "build" / "default" / "pea_cpp",
                  root / "build" / "Release" / "pea_cpp.exe",
                  root / "build" / "pea_cpp.exe",
                  root / "build" / "pea_cpp"]
    available = [item for item in candidates if item.is_file()]
    binary = max(available, key=lambda item: item.stat().st_mtime) if available else None
    newest_source = max((root / "CMakeLists.txt").stat().st_mtime,
                        (root / "src" / "cpp" / "Taller2_REMR.cpp").stat().st_mtime)
    if binary is not None and binary.stat().st_mtime >= newest_source:
        return binary
    if os.name == "nt":
        bundled = bundled_windows_cpp(root)
        if bundled is not None:
            return bundled
    if binary is None or binary.stat().st_mtime < newest_source:
        if shutil.which("cmake") is None:
            raise DataError("CMake no está instalado. Instálelo para compilar el backend C++.")
        mingw = bool(os.name == "nt" and shutil.which("g++") and shutil.which("mingw32-make"))
        build_dir = root / "build" / ("mingw" if mingw else "default")
        configure = ["cmake", "-S", str(root), "-B", str(build_dir)]
        if mingw:
            configure += ["-G", "MinGW Makefiles"]
        for command in (configure, ["cmake", "--build", str(build_dir), "--config", "Release"]):
            try:
                completed = subprocess.run(command, cwd=root, text=True, encoding="utf-8",
                                           stdout=subprocess.PIPE, stderr=subprocess.STDOUT,
                                           check=False)
            except OSError as exc:
                raise DataError(f"No se pudo iniciar CMake: {exc}") from exc
            if completed.returncode:
                raise DataError("No se pudo compilar el backend C++:\n" + completed.stdout[-3000:])
        available = [item for item in candidates if item.is_file()]
        binary = max(available, key=lambda item: item.stat().st_mtime) if available else None
    if binary is None:
        raise DataError("CMake terminó sin generar pea_cpp; revise la carpeta build.")
    return binary


CPP_PROTOCOL = "6"


class CppRepository:
    """Adaptador local: Tkinter envía operaciones y el proceso C++ conserva el estado."""

    def __init__(self, binary: Path | None = None, *, timeout: float = 60.0) -> None:
        if not math.isfinite(timeout) or timeout <= 0:
            raise DataError("El plazo del backend debe ser positivo y finito")
        executable = cpp_executable(binary)
        root = Path(__file__).resolve().parents[2]
        try:
            self.process = subprocess.Popen(
                [str(executable), "--api"], cwd=root,
                stdin=subprocess.PIPE, stdout=subprocess.PIPE, stderr=subprocess.PIPE,
                text=True, encoding="utf-8", bufsize=1,
            )
        except OSError as exc:
            raise DataError(f"No se pudo iniciar el backend C++: {exc}") from exc
        self.dirty = False
        self._history_size = 0
        self._queue_size = 0
        self.timeout = timeout
        self._closed = False
        self._request_lock = threading.Lock()
        self._requests: queue.Queue[Any] = queue.Queue()
        self._worker = threading.Thread(target=self._io_loop, daemon=True)
        self._worker.start()
        try:
            hello = self._request("ping")
            if hello != {"backend": "cpp", "protocol": CPP_PROTOCOL}:
                raise BackendError("El backend C++ usa un protocolo incompatible")
        except Exception:
            self.close()
            raise

    def _io_loop(self) -> None:
        assert self.process.stdin is not None and self.process.stdout is not None
        try:
            while True:
                work = self._requests.get()
                if work is None:
                    return
                payload, replies = work
                try:
                    self.process.stdin.write(payload)
                    self.process.stdin.flush()
                    replies.put(self.process.stdout.readline())
                except (OSError, UnicodeError) as exc:
                    replies.put(exc)
                    return
        finally:
            # Solo el trabajador cierra sus streams: close() también espera si otro hilo está leyendo.
            for stream in (self.process.stdin, self.process.stdout, self.process.stderr):
                if stream is not None:
                    try:
                        stream.close()
                    except OSError:
                        pass  # El error de comunicación ya fue entregado al solicitante.

    def _stop(self) -> None:
        if self._closed:
            return
        self._closed = True
        deadline = time.monotonic() + 1.0
        if self.process.poll() is None:
            try:
                self.process.kill()
            except ProcessLookupError:
                pass  # El proceso terminó entre poll() y kill().
        try:
            self.process.wait(timeout=max(0.0, deadline - time.monotonic()))
        except subprocess.TimeoutExpired:
            pass  # El plazo de limpieza no debe bloquear el error que recibe la interfaz.
        self._requests.put(None)
        self._worker.join(timeout=max(0.0, deadline - time.monotonic()))

    def _request(self, action: str, *, _timeout: float | None = None, **fields: str) -> Any:
        timeout = self.timeout if _timeout is None else _timeout
        deadline = time.monotonic() + timeout
        if not self._request_lock.acquire(timeout=timeout):
            raise BackendError("El backend C++ sigue ocupado con otra operación; espere a que termine")
        try:
            if self._closed or self.process.poll() is not None:
                self._stop()
                raise BackendError(f"El backend C++ se cerró (código {self.process.returncode})")
            payload = json.dumps({"action": action, **fields}, ensure_ascii=False, separators=(",", ":")) + "\n"
            replies: queue.Queue[Any] = queue.Queue(maxsize=1)
            self._requests.put((payload, replies))
            try:
                line = replies.get(timeout=max(0.0, deadline - time.monotonic()))
            except queue.Empty as exc:
                self._stop()
                raise BackendError("Se agotó el plazo de respuesta del backend C++; vuelva a abrir el programa. "
                                "La operación no se reintentó; abra el último archivo guardado.") from exc
            if isinstance(line, Exception) or not line:
                self._stop()
                raise BackendError(f"No se pudo comunicar con el backend C++: {line or 'sin respuesta'}")
            try:
                response = json.loads(line)
                state = response["state"]
                if not isinstance(response["ok"], bool) or not isinstance(state["dirty"], bool):
                    raise TypeError("Estado de protocolo inválido")
                if any(type(state[name]) is not int or state[name] < 0 for name in ("history_size", "queue_size")):
                    raise TypeError("Contadores de protocolo inválidos")
                self.dirty = state["dirty"]
                self._history_size = state["history_size"]
                self._queue_size = state["queue_size"]
                result = response["result"] if response["ok"] else response["error"]
            except (KeyError, TypeError, ValueError) as exc:
                self._stop()
                raise BackendError("Respuesta inválida del backend C++") from exc
            if not response["ok"]:
                raise DataError(result)
            return result
        finally:
            self._request_lock.release()

    @staticmethod
    def _keys(key: str | tuple[str, str]) -> dict[str, str]:
        return {"key": key[0], "second": key[1]} if isinstance(key, tuple) else {"key": key, "second": ""}

    def catalog(self) -> list[dict[str, Any]]:
        return self._request("catalog")

    def rows(self, kind: str) -> list[dict[str, str]]:
        return self._request("rows", kind=kind)

    def page(self, kind: str, offset: int = 0, limit: int = PAGE_SIZE,
             query: str = "", *, active: str = "", group_id: str = "",
             include_references: bool = False) -> dict[str, Any]:
        return self._request("page", kind=kind, offset=str(offset), limit=str(limit), query=query,
                             active=active, group_id=group_id, references="1" if include_references else "0")

    def summary(self) -> dict[str, Any]:
        return self._request("summary")

    def field_usage(self, kind: str) -> dict[str, int]:
        return self._request("field_usage", kind=kind)

    def product_issues(self, product_id: str) -> list[str]:
        return self._request("product_issues", key=product_id)

    def get(self, kind: str, key: str | tuple[str, str]) -> dict[str, str] | None:
        return self._request("get", kind=kind, **self._keys(key))

    def find_external(self, kind: str, value: str) -> dict[str, str] | None:
        return self._request("find_external", kind=kind, value=value)

    @contextmanager
    def batch(self):
        self._request("begin_batch")
        try:
            yield
        finally:
            if not self._closed:
                self._request("end_batch")

    def related(self, kind: str, key: str, side: str) -> list[dict[str, str]]:
        return self._request("related", kind=kind, key=key, side=side)

    def related_page(self, kind: str, key: str, side: str, offset: int = 0,
                     limit: int = PAGE_SIZE, active_only: bool = False,
                     query: str = "", *, include_references: bool = False) -> dict[str, Any]:
        return self._request("related_page", kind=kind, key=key, side=side,
                             offset=str(offset), limit=str(limit), active="1" if active_only else "0", query=query,
                             references="1" if include_references else "0")

    def create(self, kind: str, values: dict[str, str]) -> dict[str, str]:
        return self._request("create", kind=kind, values=json.dumps(values, ensure_ascii=False))

    def update(self, kind: str, key: str | tuple[str, str], values: dict[str, str]) -> dict[str, str]:
        return self._request("update", kind=kind, values=json.dumps(values, ensure_ascii=False),
                             **self._keys(key))

    def delete(self, kind: str, key: str | tuple[str, str]) -> None:
        self._request("delete", kind=kind, **self._keys(key))

    def toggle(self, kind: str, key: str | tuple[str, str]) -> None:
        self._request("toggle", kind=kind, **self._keys(key))

    def statistics(self, view: str = "Todos", selected_id: str = "", start: int | None = None,
                   end: int | None = None, status: str = "",
                   offset: int = 0, limit: int = PAGE_SIZE) -> dict[str, Any]:
        return self._request("statistics", view=view, selected=selected_id,
                             start="" if start is None else str(start),
                             end="" if end is None else str(end),
                             status=status, offset=str(offset), limit=str(limit))

    def indicator(self, kind: str, key: str, start: int | None = None,
                  end: int | None = None, status: str = "") -> dict[str, Any]:
        return self._request("indicator", kind=kind, key=key,
                             start="" if start is None else str(start),
                             end="" if end is None else str(end), status=status)

    def history_size(self) -> int:
        return self._history_size

    def suggest_id(self, kind: str) -> str:
        return self._request("new_id", kind=kind)

    def undo(self) -> bool:
        return self._request("undo")

    def clear_history(self) -> None:
        self._request("clear_history")

    def reset(self) -> None:
        self._request("reset")

    def load(self, path: Path) -> None:
        self._request("load", path=str(Path(path).resolve()))

    def save(self, path: Path) -> None:
        self._request("save", path=str(Path(path).resolve()))

    def preview_csv(self, path: Path, kind: str) -> tuple[int, int, list[str]]:
        result = self._request("preview_csv", path=str(Path(path).resolve()), kind=kind)
        return result["total"], result["accepted"], result["errors"]

    def import_csv(self, path: Path, kind: str) -> tuple[int, list[str]]:
        result = self._request("import_csv", path=str(Path(path).resolve()), kind=kind)
        return result["accepted"], result["errors"]

    def queue_rows(self) -> list[dict[str, str]]:
        return self._request("queue_rows")
    def queue_page(self, offset: int = 0, limit: int = PAGE_SIZE) -> dict[str, Any]:
        return self._request("queue_page", offset=str(offset), limit=str(limit))
    def queue_front(self) -> dict[str, str] | None:
        return self._request("queue_front")
    def queue_size(self) -> int:
        return self._queue_size
    def enqueue_review(self, product_id: str, reason: str) -> None:
        self._request("enqueue_review", product_id=product_id, reason=reason)
    def enqueue_rejected(self) -> int:
        return self._request("enqueue_rejected")
    def process_review(self, observation: str) -> dict[str, str]:
        return self._request("process_review", observation=observation)
    def discard_review(self) -> dict[str, str]:
        return self._request("discard_review")

    def close(self) -> None:
        if not self._closed and self.process.poll() is None:
            try:
                self._request("shutdown", _timeout=min(self.timeout, 1.0))
                self.process.wait(timeout=1)
            except (DataError, OSError, subprocess.TimeoutExpired):
                self._stop()
                return
        self._stop()


class VisibleText(HTMLParser):
    BLOCKS = {"tr", "td", "th", "h1", "h2", "h3", "h4", "p", "li", "br", "div", "section", "article", "title"}

    def __init__(self) -> None:
        super().__init__(convert_charrefs=True)
        self.parts: list[str] = []
        self.suppressed = 0

    def handle_starttag(self, tag: str, attrs: list[tuple[str, str | None]]) -> None:
        if tag in ("script", "style"):
            self.suppressed += 1
        elif tag in self.BLOCKS:
            self.parts.append("\n")

    def handle_endtag(self, tag: str) -> None:
        if tag in ("script", "style"):
            self.suppressed = max(0, self.suppressed - 1)
        elif tag in self.BLOCKS:
            self.parts.append("\n")

    def handle_data(self, data: str) -> None:
        if not self.suppressed:
            self.parts.append(data)

    def lines(self) -> list[str]:
        return [re.sub(r"\s+", " ", part).strip() for part in "".join(self.parts).splitlines() if part.strip()]


class LinkedTableRows(HTMLParser):
    """Conserva celdas y enlaces de tablas HTML para distinguir nombres de códigos CvLAC."""

    def __init__(self) -> None:
        super().__init__(convert_charrefs=True)
        self.rows: list[list[tuple[str, str]]] = []
        self._row: list[tuple[str, str]] | None = None
        self._cell: list[str] | None = None
        self._href = ""

    def handle_starttag(self, tag: str, attrs: list[tuple[str, str | None]]) -> None:
        if tag == "tr":
            self._row = []
        elif tag in ("td", "th") and self._row is not None:
            self._cell = []
            self._href = ""
        elif tag == "a" and self._cell is not None:
            self._href = dict(attrs).get("href") or self._href

    def handle_data(self, data: str) -> None:
        if self._cell is not None:
            self._cell.append(data)

    def handle_endtag(self, tag: str) -> None:
        if tag in ("td", "th") and self._cell is not None:
            if self._row is not None:
                self._row.append((re.sub(r"\s+", " ", "".join(self._cell)).strip(), self._href))
            self._cell = None
        elif tag == "tr" and self._row is not None:
            if self._row:
                self.rows.append(self._row)
            self._row = None


class WebPreview(NamedTuple):
    url: str
    format: str
    title: str
    lines: tuple[str, ...]
    metadata: dict[str, str]
    suggested_kind: str | None = None
    suggested_row: dict[str, str] | None = None
    csv_bytes: bytes | None = None
    tables: tuple[tuple[tuple[str, ...], ...], ...] = ()
    related_products: tuple[dict[str, str], ...] = ()
    related_members: tuple[tuple[dict[str, str], dict[str, str]], ...] = ()
    related_plan: dict[str, str] | None = None
    related_authorships: tuple[dict[str, str], ...] = ()
    diagnostics: tuple[str, ...] = ()


class GrupLACProduct(NamedTuple):
    row: dict[str, str]
    authors: tuple[str, ...]
    section: str


class PageText(VisibleText):
    """Recoge texto visible y metadatos explícitos, sin interpretar el tema del sitio."""

    META_KEYS = {
        "description", "og:description", "og:title", "twitter:title", "author", "date", "article:published_time",
        "citation_title", "citation_doi", "citation_date", "citation_publication_date",
        "citation_author", "citation_journal_title", "dc.title", "dc.date",
    }

    def __init__(self) -> None:
        super().__init__()
        self.metadata: dict[str, str] = {}
        self.title_parts: list[str] = []
        self.heading_parts: list[str] = []
        self._title = False
        self._heading = False
        self._json_ld = False
        self._json_parts: list[str] = []
        self.schemas: list[dict[str, Any]] = []
        self.tables: list[tuple[tuple[str, ...], ...]] = []
        self._table_level = 0
        self._table_rows: list[tuple[str, ...]] = []
        self._table_row: list[str] | None = None
        self._table_cell: list[str] | None = None

    def handle_starttag(self, tag: str, attrs: list[tuple[str, str | None]]) -> None:
        attributes = dict(attrs)
        if tag == "meta":
            key = (attributes.get("name") or attributes.get("property") or "").casefold()
            value = (attributes.get("content") or "").strip()
            if key in self.META_KEYS and value and key not in self.metadata:
                self.metadata[key] = value[:1000]
        if tag == "title":
            self._title = True
        if tag == "h1" and not self.heading_parts:
            self._heading = True
        if tag == "script" and attributes.get("type", "").casefold() == "application/ld+json":
            self._json_ld = True
            self._json_parts = []
        if tag == "table":
            self._table_level += 1
            if self._table_level == 1:
                self._table_rows = []
        elif self._table_level == 1 and tag == "tr":
            self._table_row = []
        elif self._table_level == 1 and tag in ("td", "th") and self._table_row is not None:
            self._table_cell = []
        super().handle_starttag(tag, attrs)

    def handle_endtag(self, tag: str) -> None:
        if tag == "script" and self._json_ld:
            self._json_ld = False
            try:
                schema = json.loads("".join(self._json_parts))
            except (ValueError, TypeError, RecursionError):
                schema = None
            self._collect_schema(schema)
        if tag == "title":
            self._title = False
        if tag == "h1":
            self._heading = False
        if self._table_level == 1 and tag in ("td", "th") and self._table_cell is not None:
            if self._table_row is not None and len(self._table_row) < 30:
                self._table_row.append(re.sub(r"\s+", " ", "".join(self._table_cell)).strip()[:500])
            self._table_cell = None
        elif self._table_level == 1 and tag == "tr" and self._table_row is not None:
            if self._table_row and len(self._table_rows) < 100:
                self._table_rows.append(tuple(self._table_row))
            self._table_row = None
        if tag == "table" and self._table_level:
            if self._table_level == 1 and self._table_rows and len(self.tables) < 5:
                self.tables.append(tuple(self._table_rows))
            self._table_level -= 1
        super().handle_endtag(tag)

    def handle_data(self, data: str) -> None:
        if self._json_ld:
            self._json_parts.append(data)
        if self._title:
            self.title_parts.append(data)
        if self._heading:
            self.heading_parts.append(data)
        if self._table_cell is not None:
            self._table_cell.append(data)
        super().handle_data(data)

    def _collect_schema(self, item: Any, depth: int = 0) -> None:
        if depth > 64:
            return
        if isinstance(item, list):
            for value in item:
                self._collect_schema(value, depth + 1)
        elif isinstance(item, dict):
            if isinstance(item.get("@graph"), list):
                self._collect_schema(item["@graph"], depth + 1)
            if "@type" in item:
                self.schemas.append(item)


def _web_source(url: str) -> str:
    return f"{urlparse(url).hostname}, consultado {date.today().isoformat()}"


def _redact_url(url: str, keep_params: tuple[str, ...] = ("nro", "cod_rh")) -> str:
    """Elimina parámetros de consulta sensibles; conserva solo los necesarios para reabrir."""
    parsed = urlparse(url)
    if parsed.scheme not in ("http", "https"):
        return url
    params = parse_qs(parsed.query)
    kept = {key: params[key][0] for key in keep_params if key in params}
    query = urlencode(sorted(kept.items()))
    return parsed._replace(query=query, fragment="").geturl()


def _publication_year(value: str) -> str:
    match = re.search(r"(?<!\d)(?:19|20)\d{2}(?!\d)", value)
    if match and 1900 <= int(match.group()) <= date.today().year + 1:
        return match.group()
    return ""


def _is_scholarly_article(item: dict[str, Any]) -> bool:
    declared = item.get("@type")
    types = [declared] if isinstance(declared, str) else declared if isinstance(declared, list) else []
    return any(isinstance(value, str) and value.rsplit("/", 1)[-1] == "ScholarlyArticle" for value in types)


def parse_web_page(text: str, url: str, format: str = "HTML") -> WebPreview:
    """Previsualiza cualquier página de texto; solo propone campos con etiquetas fiables."""
    metadata: dict[str, str] = {}
    suggestion: tuple[str, dict[str, str]] | None = None
    tables: tuple[tuple[tuple[str, ...], ...], ...] = ()
    related_products: tuple[dict[str, str], ...] = ()
    related_members: tuple[tuple[dict[str, str], dict[str, str]], ...] = ()
    related_plan: dict[str, str] | None = None
    related_authorships: tuple[dict[str, str], ...] = ()
    diagnostics: list[str] = []
    if format == "HTML":
        parser = PageText()
        parser.feed(text)
        lines = parser.lines()
        metadata = parser.metadata.copy()
        tables = tuple(parser.tables)
        if tables:
            metadata["Tablas HTML"] = str(len(tables))
        title = re.sub(r"\s+", " ", "".join(parser.title_parts)).strip()
        heading = re.sub(r"\s+", " ", "".join(parser.heading_parts)).strip()
        title = title or metadata.get("og:title", "") or metadata.get("twitter:title", "") or heading or metadata.get("citation_title", "")
        scienti = urlparse(url).hostname == "scienti.minciencias.gov.co"
        if scienti:
            try:
                suggestion = parse_scienti_html(text, url, diagnostics=diagnostics)
            except DataError as exc:
                diagnostics.append(f"Perfil Scienti: {exc}")
            if suggestion and suggestion[0] == "grupos":
                members: list[tuple[dict[str, str], dict[str, str]]] = []
                products: list[GrupLACProduct] = []
                try:
                    members = parse_gruplac_members(text, url, diagnostics=diagnostics)
                    related_members = tuple(members)
                except DataError as exc:
                    diagnostics.append(f"Integrantes GrupLAC: {exc}")
                try:
                    products = parse_gruplac_products(text, url, diagnostics=diagnostics)
                    related_products = tuple(item.row for item in products)
                except DataError as exc:
                    diagnostics.append(f"Productos GrupLAC: {exc}")
                try:
                    related_plan = parse_gruplac_plan(text, url, suggestion[1]["objetivos"])
                except DataError as exc:
                    diagnostics.append(f"Plan GrupLAC: {exc}")
                try:
                    related_authorships = tuple(match_gruplac_authors(products, members))
                except DataError as exc:
                    diagnostics.append(f"Autorías GrupLAC: {exc}")
            elif suggestion and suggestion[0] == "investigadores":
                try:
                    products = parse_cvlac_products(text, url, diagnostics=diagnostics)
                    related_products = tuple(item.row for item in products)
                    related_authorships = tuple(clean_row("autorias", {
                        "producto_id": item.row["id"], "investigador_id": suggestion[1]["id"],
                        "rol": "Producto listado en CvLAC"}) for item in products)
                except DataError as exc:
                    diagnostics.append(f"Productos CvLAC: {exc}")
        if suggestion is None and not scienti:
            article = next((item for item in parser.schemas if _is_scholarly_article(item)), None)
            article_title = metadata.get("citation_title", "")
            if article and not article_title and isinstance(article.get("headline") or article.get("name"), str):
                article_title = article.get("headline") or article.get("name")
            if article_title:
                published = metadata.get("citation_publication_date") or metadata.get("citation_date") or (
                    article.get("datePublished", "") if article else "")
                doi = metadata.get("citation_doi") or (article.get("doi", "") if article else "")
                if article:
                    for key, value in (("schema.org: título", article_title),
                                       ("schema.org: fecha", published), ("schema.org: DOI", doi)):
                        if value and key not in metadata:
                            metadata[key] = str(value)[:1000]
                row = {"id": new_id("P"), "titulo": article_title, "anio": _publication_year(str(published)),
                       "doi": str(doi) if isinstance(doi, str) else "", "url": _redact_url(url), "fuente": _web_source(url)}
                suggestion = ("productos", row)
        if not title:
            title = lines[0] if lines else urlparse(url).hostname or url
    elif format == "CSV":
        reader = csv.reader(io.StringIO(text))
        try:
            columns = next(reader)
        except StopIteration:
            columns = []
        metadata["Columnas CSV"] = ", ".join(columns[:30]) or "Archivo vacío"
        lines = [f"Encabezados: {metadata['Columnas CSV']}"]
        for index, row in enumerate(reader, start=1):
            if index > 100:
                lines.append("… Vista previa limitada a 100 filas")
                break
            lines.append(f"Fila {index}: " + " | ".join(row[:30]))
        title = "CSV: " + (urlparse(url).path.rsplit("/", 1)[-1] or urlparse(url).hostname or url)
    else:
        lines = [line.strip() for line in text.splitlines() if line.strip()]
        title = lines[0][:200] if lines else urlparse(url).path.rsplit("/", 1)[-1] or url
    if not lines and not metadata:
        lines = ["No se encontró texto legible en la respuesta. El sitio podría requerir JavaScript o acceso especial."]
    if suggestion:
        kind, row = suggestion
    else:
        kind, row = None, None
    return WebPreview(url, format, title[:250], tuple(lines[:300]), metadata, kind, row, None, tables,
                      related_products, related_members, related_plan, related_authorships, tuple(diagnostics))


def _resolve_public_addrs(host: str, port: int) -> list[str]:
    try:
        addresses = socket.getaddrinfo(host, port, type=socket.SOCK_STREAM)
    except (OSError, ValueError) as exc:
        raise DataError(f"No se pudo resolver el sitio: {exc}") from exc
    resolved: list[str] = []
    for item in addresses:
        address = ipaddress.ip_address(item[4][0])
        if not address.is_global:
            raise DataError("No se permiten direcciones locales o privadas")
        if item[4][0] not in resolved:
            resolved.append(item[4][0])
    if not resolved:
        raise DataError("No se pudo resolver el sitio")
    return resolved


def _validate_public_url(url: str) -> str:
    try:
        parsed = urlparse(url.strip())
        host = parsed.hostname
        port = parsed.port
    except ValueError as exc:
        raise DataError("URL inválida") from exc
    if parsed.scheme not in ("http", "https") or not host or parsed.username or parsed.password:
        raise DataError("Indique una URL pública HTTP o HTTPS, sin usuario ni contraseña")
    if host.casefold() == "localhost" or host.casefold().endswith((".localhost", ".local")):
        raise DataError("No se permiten direcciones locales o privadas")
    if port is not None and port not in (80, 443):
        raise DataError("Solo se permiten los puertos 80 y 443")
    try:
        literal_address = ipaddress.ip_address(host)
    except ValueError:
        literal_address = None
    if literal_address is not None and not literal_address.is_global:
        raise DataError("No se permiten direcciones locales o privadas")
    _resolve_public_addrs(host, port or (443 if parsed.scheme == "https" else 80))
    return url.strip()


class _PinnedHTTPConnection(http.client.HTTPConnection):
    def connect(self) -> None:
        address = _resolve_public_addrs(self.host, self.port)[0]
        self.sock = socket.create_connection((address, self.port), self.timeout, self.source_address)


class _PinnedHTTPSConnection(http.client.HTTPSConnection):
    def connect(self) -> None:
        address = _resolve_public_addrs(self.host, self.port)[0]
        self.sock = socket.create_connection((address, self.port), self.timeout, self.source_address)
        if self._tunnel_host:
            self._tunnel()
        self.sock = self._context.wrap_socket(self.sock, server_hostname=self._tunnel_host or self.host)


class _PinnedHTTPHandler(HTTPHandler):
    def http_open(self, req: Request) -> Any:
        return self.do_open(_PinnedHTTPConnection, req)


class _PinnedHTTPSHandler(HTTPSHandler):
    def https_open(self, req: Request) -> Any:
        return self.do_open(_PinnedHTTPSConnection, req, context=self._context)


class PublicRedirects(HTTPRedirectHandler):
    def redirect_request(self, req: Request, fp: Any, code: int, msg: str,
                         headers: Any, newurl: str) -> Request | None:
        _validate_public_url(newurl)
        return super().redirect_request(req, fp, code, msg, headers, newurl)


def _build_public_opener() -> Any:
    return build_opener(ProxyHandler({}), _PinnedHTTPHandler(),
                        _PinnedHTTPSHandler(), PublicRedirects())


def fetch_web_page(url: str) -> WebPreview:
    """Descarga una sola URL pública y previsualiza HTML, texto, CSV o PDF de texto."""
    url = _validate_public_url(url)
    request = Request(url, headers={"User-Agent": "PEA-i-UPC/1.0 (academic project)",
                                    "Accept": "text/html,text/plain,text/csv,application/pdf;q=0.9,*/*;q=0.1"})
    try:
        opener = _build_public_opener()
        with opener.open(request, timeout=15) as response:
            final_url = _validate_public_url(response.geturl())
            content_type = response.headers.get_content_type().casefold()
            charset = response.headers.get_content_charset()
            content = response.read(8_000_001)
    except (OSError, TimeoutError) as exc:
        raise DataError(f"No se pudo descargar la URL: {exc}") from exc
    if len(content) > 8_000_000:
        raise DataError("La respuesta supera 8 MB; use un archivo CSV local")
    if content.startswith(b"%PDF") or content_type == "application/pdf":
        with tempfile.TemporaryDirectory() as temporary:
            path = Path(temporary) / "pagina.pdf"
            path.write_bytes(content)
            text = extract_pdf_text(path)
            preview = parse_web_page(text, final_url, "PDF de texto")
            if urlparse(final_url).hostname == "scienti.minciencias.gov.co":
                try:
                    diagnostics: list[str] = []
                    kind, row = parse_scienti_pdf(path, final_url, text=text, diagnostics=diagnostics)
                    preview = preview._replace(suggested_kind=kind, suggested_row=row,
                                               diagnostics=tuple(diagnostics))
                except DataError as exc:
                    preview = preview._replace(diagnostics=tuple(diagnostics) + (f"Perfil Scienti PDF: {exc}",))
            return preview
    csv_file = content_type in ("text/csv", "application/csv", "application/vnd.ms-excel") or urlparse(final_url).path.lower().endswith(".csv")
    html_file = content_type in ("text/html", "application/xhtml+xml") or content[:500].lstrip().lower().startswith((b"<!doctype html", b"<html"))
    if not (csv_file or html_file or content_type.startswith("text/")):
        raise DataError(f"Formato web no legible: {content_type}")
    if not charset and html_file:
        match = re.search(rb"charset\s*=\s*['\"]?([A-Za-z0-9_-]+)", content[:5000], re.IGNORECASE)
        charset = match.group(1).decode("ascii") if match else None
    try:
        decoded = content.decode(charset or "utf-8-sig", errors="replace")
    except LookupError as exc:
        raise DataError(f"Codificación desconocida: {charset}") from exc
    preview = parse_web_page(decoded, final_url, "CSV" if csv_file else "HTML" if html_file else "Texto")
    return preview._replace(csv_bytes=content) if csv_file else preview


def chart_values(values: dict[str, int]) -> dict[str, int]:
    """Los campos sin valor cuentan en el total, pero no forman una barra de datos."""
    return {key: value for key, value in values.items() if key.strip().casefold() != "sin dato" and value > 0}


_GRUPLAC_LEVELS = ("A1", "A", "B", "C", "D", "Reconocido")


def _gruplac_category(value: str) -> str:
    """Reduce el texto de 'Clasificación' a un nivel válido (A1..D, Reconocido) o vacío."""
    token = value.split(" ")[0].strip()
    return token if token in _GRUPLAC_LEVELS else ""


def parse_scienti_html(html: str, url: str, *, diagnostics: list[str] | None = None) -> tuple[str, dict[str, str]]:
    """Extrae campos identificables; evita atribuir productos a un perfil ambiguo."""
    parsed = urlparse(url)
    if parsed.scheme != "https" or parsed.hostname != "scienti.minciencias.gov.co":
        raise DataError("Solo se admiten URL públicas HTTPS de Scienti Minciencias")
    parser = VisibleText()
    parser.feed(html)
    lines = parser.lines()
    if not lines:
        raise DataError("La página no contiene texto legible")
    code = parse_qs(parsed.query)
    if "/gruplac/" in parsed.path and code.get("nro"):
        if not any(_name_key(line) == "datos basicos" for line in lines):
            raise DataError("No se identificó la sección Datos básicos de la ficha GrupLAC")
        candidates = [line for line in lines if line.lower() not in ("gruplac - plataforma scienti - colombia", "datos básicos")]
        name = next((line for line in candidates if line in lines[:15] and len(line) > 2 and not line.lower().startswith(("inicio", "menú", "buscar"))), "")
        if not name:
            raise DataError("No se identificó el nombre del grupo")
        text = "\n".join(lines)
        def after(label: str) -> str:
            match = re.search(rf"(?im)^{re.escape(label)}\s*(?:[:|])?\s*(.+)$", text)
            if match:
                return match.group(1).strip()
            try:
                index = next(i for i, line in enumerate(lines) if line.casefold() == label.casefold())
                return lines[index + 1] if index + 1 < len(lines) else ""
            except StopIteration:
                return ""
        def block(label: str, stop: str) -> str:
            try:
                first = next(i for i, line in enumerate(lines) if line.casefold().startswith(label.casefold()))
            except StopIteration:
                return ""
            collected = [lines[first].split(":", 1)[1].strip()] if ":" in lines[first] else []
            for index in range(first + 1, len(lines)):
                line = lines[index]
                if line.casefold().startswith(stop.casefold()):
                    break
                collected.append(line)
            return "\n".join(part for part in collected if part)
        formed = after("Año y mes de formación")
        formation_date = ""
        formation_match = re.fullmatch(r"(\d{4})\s*-\s*(\d{1,2})", formed)
        if formation_match:
            try:
                formation_date = date(int(formation_match[1]), int(formation_match[2]), 1).isoformat()
            except ValueError:
                if diagnostics is not None:
                    diagnostics.append(f"Perfil Scienti: fecha de formación inválida: {formed[:100]}")
        elif formed and diagnostics is not None:
            diagnostics.append(f"Perfil Scienti: no se pudo interpretar la fecha de formación: {formed[:100]}")
        group = {
            "id": f"G-{code['nro'][0]}", "nombre": name,
            "fecha_creacion": formation_date,
            "codigo_gruplac": code["nro"][0], "responsable": after("Líder"),
            "categoria": _gruplac_category(after("Clasificación")),
            "descripcion": block("Estado del arte", "Objetivos"),
            "objetivos": block("Objetivos", "Retos"),
            "vision": after("Visión"),
            "lineas": block("Líneas de investigación declaradas por el grupo", "Integrantes del grupo"),
            "url": _redact_url(url), "fuente": f"Scienti, consultado {date.today().isoformat()}; formación: {formed}",
        }
        return "grupos", clean_row("grupos", group)
    if "/cvlac/" in parsed.path and code.get("cod_rh"):
        text = "\n".join(lines)
        match = re.search(r"(?im)^Nombre\s+([^\n]+)$", text)
        if not match:
            raise DataError("No se identificó el nombre del investigador")
        name = match.group(1).strip()
        category = ""
        for index, line in enumerate(lines):
            if line.casefold() == "categoría" and index + 1 < len(lines):
                candidate = lines[index + 1].split(" con vigencia")[0].strip()
                if _name_key(candidate) not in ("nombre", "nombre en citaciones", "nacionalidad", "sexo") \
                        and not candidate.casefold().startswith(("producción ", "formación ", "experiencia ", "trabajos dirigidos/")):
                    category = candidate
                break
        if not category:
            found = re.search(r"Categoría Investigador\s+([^\n]+)", text, re.IGNORECASE)
            category = found.group(1).split(" con vigencia")[0].strip() if found else ""
        researcher = {
            "id": f"I-{code['cod_rh'][0]}", "nombre": name,
            "codigo_cvlac": code["cod_rh"][0],
            "categoria": category,
            "url": _redact_url(url), "fuente": f"Scienti, consultado {date.today().isoformat()}",
        }
        return "investigadores", clean_row("investigadores", researcher)
    raise DataError("La URL debe ser una ficha GrupLAC o CvLAC con su código")


def _require_gruplac_url(url: str) -> str:
    parsed = urlparse(url)
    if parsed.scheme != "https" or parsed.hostname != "scienti.minciencias.gov.co" or "/gruplac/" not in parsed.path:
        raise DataError("La fuente debe ser una ficha GrupLAC HTTPS")
    code = parse_qs(parsed.query).get("nro", [""])[0]
    if not code:
        raise DataError("La ficha GrupLAC no tiene código de grupo")
    return code


def _name_key(value: str) -> str:
    plain = "".join(char for char in unicodedata.normalize("NFKD", value) if not unicodedata.combining(char))
    return re.sub(r"\s+", " ", plain).casefold().strip()


def parse_gruplac_members(html: str, url: str, *, diagnostics: list[str] | None = None) -> list[tuple[dict[str, str], dict[str, str]]]:
    """Importa el censo explícito; conserva meses textuales sin inventar un día ISO."""
    code = _require_gruplac_url(url)
    parser = LinkedTableRows()
    parser.feed(html)
    first = next((i + 1 for i, row in enumerate(parser.rows)
                  if len(row) >= 4 and _name_key(row[0][0]) == "nombre"
                  and _name_key(row[1][0]) == "vinculacion"
                  and "inicio" in _name_key(row[3][0])), None)
    if first is None:
        return []
    members: list[tuple[dict[str, str], dict[str, str]]] = []
    seen: set[str] = set()
    for index in range(first, len(parser.rows)):
        row = parser.rows[index]
        if len(row) < 4:
            if row and re.match(r"^\d+\.-", row[0][0]):
                if diagnostics is not None:
                    diagnostics.append(f"Integrantes GrupLAC, {row[0][0][:80]}: extracción incompleta, faltan celdas del censo")
                continue
            if members:
                break
            continue
        match = re.match(r"^\d+\.-\s*(.+)$", row[0][0])
        if not match:
            if members:
                break
            continue
        name = match.group(1).strip()
        href = row[0][1]
        profile = urlparse(href)
        profile_code = ""
        if profile.scheme == "https" and profile.hostname == "scienti.minciencias.gov.co" and "/cvlac/" in profile.path:
            profile_code = parse_qs(profile.query).get("cod_rh", [""])[0]
        person_id = ("I-" + profile_code) if profile_code else (
            "I-GR-" + hashlib.sha1((code + "|" + _name_key(name)).encode("utf-8")).hexdigest()[:12].upper())
        if person_id in seen:
            continue
        period = row[3][0].strip()
        role = row[1][0].strip() or "Integrante"
        ended = bool(re.search(r"\s-\s(?:19|20)\d{2}/\d{1,2}\s*$", period))
        try:
            person = clean_row("investigadores", {
                "id": person_id, "nombre": name, "codigo_cvlac": profile_code,
                "url": _redact_url(href) if profile_code else _redact_url(url),
                "fuente": f"GrupLAC, censo de integrantes, consultado {date.today().isoformat()}; vinculación: {period}",
            })
            membership = clean_row("membresias", {
                "grupo_id": "G-" + code, "investigador_id": person_id,
                "rol": f"{role} · {period}" if period else role,
                "activo": "0" if ended else "1",
            })
        except DataError as exc:
            if diagnostics is None:
                raise
            diagnostics.append(f"Integrantes GrupLAC, {person_id[:100]} ({name[:80]}): {exc}")
            continue
        seen.add(person_id)
        members.append((person, membership))
    return members


def parse_gruplac_plan(html: str, url: str, objectives: str = "") -> dict[str, str] | None:
    code = _require_gruplac_url(url)
    parser = VisibleText()
    parser.feed(html)
    lines = parser.lines()
    try:
        first = lines.index("Plan Estratégico") + 1
    except ValueError:
        return None
    last = next((i for i in range(first, len(lines)) if lines[i].startswith("Estado del arte")), first)
    activities = "\n".join(lines[first:last]).strip()
    if not activities and not objectives:
        return None
    return clean_row("planes", {
        "id": "PL-" + code, "grupo_id": "G-" + code, "nombre": "Plan Estratégico",
        "objetivo": objectives, "actividad": activities,
    })


class GrupLACSection(NamedTuple):
    heading: str
    stop: str
    family: str
    year_rule: str
    only_categorized: bool


# Las secciones nuevas requieren una categoría reconocida; las anteriores conservan todo lo importable.
# Encabezado, siguiente encabezado, familia, regla de año, categoría estimada por defecto.
_GRUPLAC_NEW_SECTIONS = (
    ("Libros de divulgación y/o Compilación de divulgación", "Manuales y Guías Especializadas", "Producción bibliográfica", "publication", "LIB_DIV"),
    ("Manuales y Guías Especializadas", "Otros Libros publicados", "Producción bibliográfica", "publication", "MAN_GUI"),
    ("Notas científicas", "PRODUCCIÓN TÉCNICA Y TECNOLÓGICA", "Producción bibliográfica", "publication", "N_D"),
    ("Conceptos técnicos", "Diseños industriales", "Producción técnica y tecnológica", "publication", "CT"),
    ("Diseños industriales", "Esquemas de trazados de circuito integrado", "Producción técnica y tecnológica", "publication", "DI_B"),
    ("Esquemas de trazados de circuito integrado", "Informes técnicos", "Producción técnica y tecnológica", "publication", "ECI"),
    ("Informes técnicos", "Innovaciones en Procesos y Procedimientos", "Producción técnica y tecnológica", "publication", "INF"),
    ("Innovaciones en Procesos y Procedimientos", "Innovaciones generadas en la Gestión Empresarial", "Producción técnica y tecnológica", "publication", "IPP"),
    ("Innovaciones generadas en la Gestión Empresarial", "Nuevas variedades animal", "Producción técnica y tecnológica", "publication", "IG_B2"),
    ("Nuevas variedades animal", "Poblaciones mejoradas de razas pecuarias", "Producción técnica y tecnológica", "publication", "VA_A"),
    ("Poblaciones mejoradas de razas pecuarias", "Nuevas variedades vegetal", "Producción técnica y tecnológica", "publication", "VA_B"),
    ("Nuevas variedades vegetal", "Nuevos registros científicos", "Producción técnica y tecnológica", "publication", "VV_B4"),
    ("Nuevos registros científicos", "Plantas piloto", "Producción técnica y tecnológica", "publication", "NRC_B"),
    ("Plantas piloto", "Productos nutracéuticos", "Producción técnica y tecnológica", "publication", "PP"),
    ("Productos nutracéuticos", "Otros productos tecnológicos", "Producción técnica y tecnológica", "publication", "PN"),
    ("Prototipos", "Regulaciones y Normas", "Producción técnica y tecnológica", "publication", "PI"),
    ("Regulaciones y Normas", "Protocolos de vigilancia epidemiológica", "Producción técnica y tecnológica", "publication", "RNT"),
    ("Protocolos de vigilancia epidemiológica", "Reglamentos técnicos", "Producción técnica y tecnológica", "publication", "PVE"),
    ("Reglamentos técnicos", "Guias de práctica clínica", "Producción técnica y tecnológica", "publication", "RNT"),
    ("Guias de práctica clínica", "Proyectos de ley", "Producción técnica y tecnológica", "publication", "RNPC"),
    ("Proyectos de ley", "Signos distintivos", "Producción técnica y tecnológica", "publication", "RNPL"),
    ("Signos distintivos", "Softwares", "Producción técnica y tecnológica", "publication", "SD"),
    ("Publicaciones editoriales no especializadas", "Producciones de contenido digital - Audiovisual", "Divulgación pública de la ciencia", "range", "PEE_C2"),
    ("Producciones de contenido digital - Audiovisual", "Producciones de contenido digital - Sonoro", "Divulgación pública de la ciencia", "range", "PCD_C2"),
    ("Producciones de contenido digital - Sonoro", "Producciones de contenido digital - Recursos gráficos", "Divulgación pública de la ciencia", "range", "PCD_C2"),
    ("Producciones de contenido digital - Recursos gráficos", "Divulgación Pública de la Ciencia producción de estrategias y contenidos transmedia", "Divulgación pública de la ciencia", "range", "PCD_C2"),
    ("Divulgación Pública de la Ciencia producción de estrategias y contenidos transmedia", "Desarrollo web", "Divulgación pública de la ciencia", "range", "TRM_C2"),
    ("Desarrollo web", "Procesos de apropiación social del Conocimiento para el fortalecimiento o solución de asuntos de interés social", "Divulgación pública de la ciencia", "range", "DW_C2"),
    ("Procesos de apropiación social del Conocimiento para el fortalecimiento o solución de asuntos de interés social", "Productos de apropiación social del conocimiento resultado del trabajo conjunto entre un Centro de Ciencia y un grupo de investigación", "Apropiación social del conocimiento", "range", "FIS"),
    ("Productos de apropiación social del conocimiento resultado del trabajo conjunto entre un Centro de Ciencia y un grupo de investigación", "Proceso de Apropiación Social del Conocimiento para la generación de insumos de política pública y normatividad", "Apropiación social del conocimiento", "range", "TCCG_D"),
    ("Proceso de Apropiación Social del Conocimiento para la generación de insumos de política pública y normatividad", "Proceso de apropiación social del Conocimiento para el fortalecimiento de cadenas productivas", "Apropiación social del conocimiento", "range", "GPP_C"),
    ("Proceso de apropiación social del Conocimiento para el fortalecimiento de cadenas productivas", "Consultorías científico-tecnológicas", "Apropiación social del conocimiento", "range", "FCP_C"),
    ("Consultorías científico-tecnológicas", "Ediciones", "Divulgación pública de la ciencia", "range", "CON_CT"),
    ("Ediciones", "Eventos Científicos", "Divulgación pública de la ciencia", "range", "ERL"),
    ("Eventos Científicos", "Informes de investigación", "Divulgación pública de la ciencia", "range", "EC_B"),
    ("Informes de investigación", "Nuevas secuencias genéticas", "Divulgación pública de la ciencia", "range", "IFI"),
    ("Nuevas secuencias genéticas", "Redes de Conocimiento Especializado", "Divulgación pública de la ciencia", "range", "NSG"),
    ("Redes de Conocimiento Especializado", "Generaciónes de contenido de audio", "Divulgación pública de la ciencia", "range", "RC_B"),
    ("Generaciónes de contenido de audio", "Generación de Contenido Impreso", "Divulgación pública de la ciencia", "range", "GC"),
    ("Generación de Contenido Impreso", "Generación de Contenido Multimedia", "Divulgación pública de la ciencia", "range", "GC"),
    ("Generación de Contenido Multimedia", "Generación de Contenido Virtual", "Divulgación pública de la ciencia", "range", "GC"),
    ("Generación de Contenido Virtual", "Estrategias de Comunicación del Conocimiento", "Divulgación pública de la ciencia", "range", "GC"),
    ("Obras o productos", "Industrias creativas y culturales", "Producción en arte, arquitectura y diseño", "range", "AAD_C"),
    ("Industrias creativas y culturales", "Eventos Artísticos", "Producción en arte, arquitectura y diseño", "range", "ICC_B"),
    ("Eventos Artísticos", "Talleres de Creación", "Producción en arte, arquitectura y diseño", "range", "ECA_B"),
    ("Talleres de Creación", "ACTIVIDADES DE FORMACIÓN", "Producción en arte, arquitectura y diseño", "range", "TC_C"),
    ("Asesorías al Programa Ondas", "Curso de Corta Duración Dictados", "Formación de recurso humano", "range", "APO"),
    ("Trabajos dirigidos/turorías", "ACTIVIDADES COMO EVALUADOR", "Formación de recurso humano", "range", ""),
    ("Proyectos", "", "Proyectos", "range", ""),
    ("Programa académico de doctorado", "Programa académico de maestría", "Formación y extensión", "range", "AP_A"),
    ("Programa académico de maestría", "Otro programa académico", "Formación y extensión", "range", "AP_B"),
    ("Curso de doctorado", "Curso de maestría", "Formación y extensión", "range", "AP_C"),
    ("Curso de maestría", "Curso especializado de extensión", "Formación y extensión", "range", "AP_D"),
)
GRUPLAC_SECTIONS = (
    GrupLACSection("Artículos publicados", "Libros publicados", "Producción bibliográfica", "publication", False),
    GrupLACSection("Libros publicados", "Capítulos de libro publicados", "Producción bibliográfica", "publication", False),
    GrupLACSection("Capítulos de libro publicados", "Documentos de trabajo", "Producción bibliográfica", "publication", False),
    GrupLACSection("Documentos de trabajo", "Otra publicación divulgativa", "Producción bibliográfica", "publication", False),
    GrupLACSection("Otra publicación divulgativa", "Otros artículos publicados", "Producción bibliográfica", "publication", False),
    GrupLACSection("Otros artículos publicados", "Libros de formación", "Producción bibliográfica", "publication", False),
    GrupLACSection("Libros de formación", "Libros de divulgación y/o Compilación de divulgación", "Producción bibliográfica", "publication", False),
    GrupLACSection("Otros Libros publicados", "Traducciones", "Producción bibliográfica", "publication", False),
    GrupLACSection("Otros productos tecnológicos", "Prototipos", "Producción técnica y tecnológica", "publication", False),
    GrupLACSection("Softwares", "Empresas de base tecnológica", "Producción técnica y tecnológica", "publication", False),
    GrupLACSection("Empresas de base tecnológica", "APROPIACIÓN SOCIAL Y DIVULGACIÓN PÚBLICA DE LA CIENCIA", "Producción técnica y tecnológica", "publication", False),
    *(GrupLACSection(heading, stop, family, rule, True) for heading, stop, family, rule, _ in _GRUPLAC_NEW_SECTIONS),
)
CATEGORY_ESTIMATES = {
    ("Artículos publicados", "Publicado en revista especializada"): "ART_D",
    ("Artículos publicados", "Corto (Resumen)"): "N_D",
    ("Libros publicados", "Libro resultado de investigación"): "LIB_C",
    ("*", "Libro pedagógico y/o de divulgación"): "LIB_DIV",
    ("Capítulos de libro publicados", "Capítulo de libro"): "CAP_LIB_C",
    ("Documentos de trabajo", "*"): "WP",
    ("Otra publicación divulgativa", "*"): "PEE_C2",
    # GrupLAC lista la prensa y la divulgación entre «Otros artículos publicados».
    ("Otros artículos publicados", "Revista de divulgación"): "PEE_C2",
    ("Otros artículos publicados", "Periódico de noticias"): "PEE_C2",
    ("Otros artículos publicados", "Columna de opinión"): "PEE_C2",
    ("Libros de formación", "*"): "LIB_FOR3",
    ("Softwares", "*"): "SF",
    ("Empresas de base tecnológica", "*"): "EBT_B",
    ("Trabajos dirigidos/turorías", "Trabajos de grado de pregrado"): "TP_B",
    ("Trabajos dirigidos/turorías", "Pregrado"): "TP_B",
    ("Trabajos dirigidos/turorías", "Trabajo de grado de maestría o especialidad clínica"): "TM_B",
    ("Trabajos dirigidos/turorías", "Maestría"): "TM_B",
    ("Trabajos dirigidos/turorías", "Tesis de doctorado"): "TD_B",
    ("Trabajos dirigidos/turorías", "Doctorado"): "TD_B",
    ("Proyectos", "Investigación y desarrollo"): "PID_C",
    ("Proyectos", "Investigación, desarrollo e innovación"): "PF_B",
    ("Proyectos", "Investigación-creación"): "PIC_C",
    ("Proyectos", "Extensión y responsabilidad social CTI"): "PE",
    **{(heading, "*"): code for heading, _, _, _, code in _GRUPLAC_NEW_SECTIONS if code},
}
PEOPLE_PREFIXES = ("Autores:", "Tutor(es)/Cotutor(es):")


def estimate_category(section: str, tipologia: str) -> str:
    for key in ((section, tipologia), (section, "*"), ("*", tipologia)):
        if key in CATEGORY_ESTIMATES:
            return CATEGORY_ESTIMATES[key]
    return ""


def product_identity(title: str, year: str, doi: str, kind: str, article: bool) -> str:
    identity = "doi:" + doi.casefold() if doi else "title:" + _name_key(title) + "|" + year
    if not doi and not article:
        identity += "|" + _name_key(kind)
    return identity


def product_id(identity: str) -> str:
    return "P-GR-" + hashlib.sha1(identity.encode("utf-8")).hexdigest()[:12].upper()


def _combine_product(previous: GrupLACProduct, current: GrupLACProduct) -> GrupLACProduct:
    authors = list(previous.authors)
    known = {_name_key(name) for name in authors}
    for name in current.authors:
        if _name_key(name) not in known:
            authors.append(name)
            known.add(_name_key(name))
    row = previous.row
    if not row["categoria"] and current.row["categoria"]:
        row = {**row, "categoria": current.row["categoria"], "categoria_origen": "estimada"}
    return previous._replace(row=row, authors=tuple(authors))


def parse_gruplac_products(html: str, url: str, *, diagnostics: list[str] | None = None) -> list[GrupLACProduct]:
    """Extrae secciones con tipo, título y año explícitos; une fichas repetidas por DOI o título."""
    _require_gruplac_url(url)
    parser = VisibleText()
    parser.feed(html)
    lines = parser.lines()
    by_identity: dict[str, GrupLACProduct] = {}
    headings = {part for section in GRUPLAC_SECTIONS for part in (section.heading, section.stop) if part}
    for section, stop, family, year_rule, only_categorized in GRUPLAC_SECTIONS:
        try:
            first = lines.index(section) + 1
        except ValueError:
            continue
        last = next((i for i in range(first, len(lines)) if lines[i] in headings), len(lines)) if stop else len(lines)
        for index in range(first, last):
            match = re.match(r"^\d+\.(?:-\s*|\s+)([^:]+?)\s*:\s*(.+)$", lines[index])
            if not match:
                if diagnostics is not None and re.match(r"^\d+\.(?:-\s*|\s+)", lines[index]):
                    diagnostics.append(f"Productos GrupLAC, {section}, {lines[index][:80]}: tipo o título no identificable")
                continue
            kind, title = match.group(1).strip(), match.group(2).strip()
            category = estimate_category(section, kind)
            if only_categorized and not category:
                if diagnostics is not None:
                    diagnostics.append(f"Productos GrupLAC, {section}, {title[:80]}: tipología sin categoría admitida ({kind[:80]})")
                continue
            if not title:
                continue
            details: list[str] = []
            authors: list[str] = []
            for cursor in range(index + 1, last):
                line = lines[cursor]
                if re.match(r"^\d+\.(?:-\s*|\s+)", line):
                    break
                prefix = next((prefix for prefix in PEOPLE_PREFIXES if line.startswith(prefix)), None)
                if prefix:
                    authors.extend(name.strip() for name in line[len(prefix):].split(",") if name.strip())
                    break
                details.append(line)
            detail = " ".join(details)
            if year_rule == "range":
                years = re.findall(r"(?<!\d)(?:19|20)\d{2}(?!\d)", " ".join(details[:2]))
                year = years[-1] if years else ""
            else:
                year_match = re.search(r"(?<![\d-])((?:19|20)\d{2})(?=\s*(?:[,.;]|vol:|$))", detail)
                year = year_match.group(1) if year_match else ""
            if not _publication_year(year):
                if diagnostics is not None:
                    diagnostics.append(f"Productos GrupLAC, {section}, {title[:80]}: extracción incompleta, falta un año válido")
                continue
            published, doi = _source_doi(detail)
            identity = product_identity(title, year, published, kind, section == "Artículos publicados")
            note = f"; DOI incompleto en la fuente: {published}" if published and not doi else ""
            try:
                row = clean_row("productos", {
                    "id": product_id(identity), "titulo": title, "anio": year,
                    "familia": family, "tipologia": kind,
                    "categoria": category, "categoria_origen": "estimada" if category else "",
                    "doi": doi, "url": _redact_url(url),
                    "fuente": f"GrupLAC, sección {section}, consultado {date.today().isoformat()}{note}",
                })
            except DataError as exc:
                if diagnostics is None:
                    raise
                diagnostics.append(f"Productos GrupLAC, {section}, {title[:80]}: {exc}")
                continue
            entry = GrupLACProduct(row, tuple(authors), section)
            previous = by_identity.get(identity)
            by_identity[identity] = _combine_product(previous, entry) if previous else entry
    return list(by_identity.values())


def parse_gruplac_articles(html: str, url: str) -> list[dict[str, str]]:
    """Compatibilidad: artículos de revista de la sección propia de GrupLAC."""
    return [entry.row for entry in parse_gruplac_products(html, url) if entry.section == "Artículos publicados"]


def match_gruplac_authors(products: list[GrupLACProduct],
                          members: list[tuple[dict[str, str], dict[str, str]]]) -> list[dict[str, str]]:
    by_name = {_name_key(person["nombre"]): person["id"] for person, _ in members}
    relations: list[dict[str, str]] = []
    seen: set[tuple[str, str]] = set()
    for product in products:
        for author in product.authors:
            researcher_id = by_name.get(_name_key(author))
            pair = (product.row["id"], researcher_id or "")
            if not researcher_id or pair in seen:
                continue
            seen.add(pair)
            relations.append(clean_row("autorias", {"producto_id": pair[0], "investigador_id": pair[1],
                                                   "rol": "Autor listado en GrupLAC"}))
    return relations


# Campos que una ficha Scienti vuelve a traer al importarla de nuevo; solo esos se actualizan,
# y solo con valores no vacíos. Las claves, la validación calculada, la observación de revisión
# y la URL de un producto (cambia según el grupo consultado) nunca se reemplazan.
REFRESH_FIELDS = {
    "grupos": ("nombre", "responsable", "categoria", "descripcion", "objetivos", "vision", "lineas", "url", "fecha_creacion"),
    "investigadores": ("nombre", "categoria", "url"),
    "planes": ("nombre", "objetivo", "actividad"),
    "productos": ("familia", "tipologia", "categoria", "categoria_origen"),
    "membresias": ("rol", "activo"),
}


def _import_counts() -> dict[str, Any]:
    counts: dict[str, Any] = dict.fromkeys(("grupos", "investigadores", "membresias", "planes", "productos",
                                         "grupos_productos", "autorias", "existentes", "actualizados"), 0)
    counts["errors"] = []
    return counts


def _merge_record(repository: Repository | CppRepository, kind: str, row: dict[str, str],
                  key: str | tuple[str, str], counts: dict[str, Any],
                  extra: dict[str, Any]) -> str | tuple[str, str] | None:
    existing = repository.get(kind, key)
    external = {"grupos": "codigo_gruplac", "investigadores": "codigo_cvlac", "productos": "doi"}.get(kind)
    if existing is None and external and row.get(external):
        existing = repository.find_external(kind, row[external])
        if existing is not None:
            key = existing["id"]
            row = {**row, "id": key}
    if existing is not None:
        if kind == "grupos" and existing["codigo_gruplac"].strip().lower() != row["codigo_gruplac"].strip().lower():
            counts["errors"].append(f"{kind} {key}: ID ocupado por otro grupo")
            return None
        if kind == "investigadores" and existing["codigo_cvlac"].strip().lower() != row["codigo_cvlac"].strip().lower():
            counts["errors"].append(f"{kind} {key}: ID ocupado por otro perfil")
            return None
        if kind == "productos":
            same_doi = bool(_doi_text(existing["doi"]) and row.get("doi")
                            and _doi_text(existing["doi"]) == _doi_text(row["doi"]))
            # Las fichas repiten un título con o sin punto final; el ID no cambia por eso.
            same_title_year = (_name_key(existing["titulo"]).rstrip(" .,;:") == _name_key(row["titulo"]).rstrip(" .,;:")
                               and existing["anio"] == row["anio"])
            if not same_doi and not same_title_year:
                counts["errors"].append(f"{kind} {key}: ID ocupado por otro producto")
                return None
        changes = {name: row[name] for name in REFRESH_FIELDS.get(kind, ())
                   if row.get(name) and row[name] != existing.get(name, "")}
        if kind == "productos" and existing["categoria_origen"] == "manual":
            changes.pop("categoria", None)
            changes.pop("categoria_origen", None)
        # Un DOI guardado con formato inválido se reemplaza por el de la fuente, aunque esta ya no traiga ninguno.
        if (kind == "productos" and existing["doi"] and not _valid_doi(_doi_text(existing["doi"]))
                and row.get("doi", "") != existing["doi"]):
            changes["doi"] = row.get("doi", "")
        elif kind == "productos" and "categoria" in changes:
            # La categoría y su origen se envían juntos aunque el origen no cambie.
            changes["categoria_origen"] = row.get("categoria_origen", "")
        if not changes:
            counts["existentes"] += 1
            return key
        # La fecha de consulta cambia en cada descarga: la fuente solo se renueva con el contenido.
        if row.get("fuente"):
            changes["fuente"] = row["fuente"]
        try:
            repository.update(kind, key, changes, **extra)
        except BackendError:
            raise
        except (DataError, OSError, ValueError) as exc:
            counts["errors"].append(f"{kind} {key}: {exc}")
            return key
        counts["actualizados"] += 1
        return key
    try:
        repository.create(kind, row, **extra)
    except BackendError:
        raise
    except (DataError, OSError, ValueError) as exc:
        counts["errors"].append(f"{kind} {key}: {exc}")
        return None
    counts[kind] += 1
    return key


def import_gruplac_preview(repository: Repository | CppRepository, preview: WebPreview,
                           reviewed_group: dict[str, str] | None = None, *,
                           remember: bool = True) -> dict[str, Any]:
    """Incorpora una ficha: crea lo nuevo y actualiza lo existente con lo que diga la fuente.

    remember=False (solo con Repository) omite el historial de deshacer en cargas masivas.
    """
    if preview.suggested_kind != "grupos" or not preview.suggested_row:
        raise DataError("La vista previa no contiene una ficha de grupo")
    source_group = preview.suggested_row
    group = clean_row("grupos", reviewed_group or source_group)
    if group["codigo_gruplac"] != source_group["codigo_gruplac"]:
        raise DataError("El código GrupLAC revisado no coincide con la fuente")
    group_id = group["id"]
    counts = _import_counts()
    extra: dict[str, Any] = {} if remember else {"remember": False}

    def add(kind: str, row: dict[str, str], key: str | tuple[str, str]) -> str | tuple[str, str] | None:
        return _merge_record(repository, kind, row, key, counts, extra)

    with repository.batch() if remember else nullcontext():
        group_id = add("grupos", group, group_id)
        if group_id is None:
            return counts
        available_people: dict[str, str] = {}
        for person, _ in preview.related_members:
            canonical = add("investigadores", person, person["id"])
            if canonical is not None:
                available_people[person["id"]] = canonical
        if preview.related_plan:
            plan = {**preview.related_plan, "grupo_id": group_id}
            add("planes", plan, plan["id"])
        for person, membership in preview.related_members:
            canonical = available_people.get(person["id"])
            if canonical is not None:
                row = {**membership, "grupo_id": group_id, "investigador_id": canonical}
                add("membresias", row, (group_id, canonical))
        available_products: dict[str, str] = {}
        for product in preview.related_products:
            canonical = add("productos", product, product["id"])
            if canonical is not None:
                available_products[product["id"]] = canonical
        for canonical in dict.fromkeys(available_products.values()):
            row = {"grupo_id": group_id, "producto_id": canonical,
                   "origen": "GrupLAC: producto listado en la ficha"}
            add("grupos_productos", row, (group_id, canonical))
        for authorship in preview.related_authorships:
            product_id = available_products.get(authorship["producto_id"])
            researcher_id = available_people.get(authorship["investigador_id"])
            if product_id is not None and researcher_id is not None:
                row = {**authorship, "producto_id": product_id, "investigador_id": researcher_id}
                add("autorias", row, (product_id, researcher_id))
    return counts


CVLAC_SECTIONS = {
    "Artículo": "Artículos publicados", "Libro": "Libros publicados",
    "Capítulo de libro": "Capítulos de libro publicados", "Capitulo de libro": "Capítulos de libro publicados",
    "Documento de trabajo (Working Paper)": "Documentos de trabajo",
    "Otro artículo publicado": "Otra publicación divulgativa", "Softwares": "Softwares",
    "Innovación de proceso o procedimiento": "Innovaciones en Procesos y Procedimientos",
    "Empresas de base tecnológica": "Empresas de base tecnológica",
    "Empresa de base tecnológica": "Empresas de base tecnológica",
    "Trabajos de grado de pregrado": "Trabajos dirigidos/turorías",
    "Trabajo de grado de maestría o especialidad clínica": "Trabajos dirigidos/turorías",
}


def parse_cvlac_products(html: str, url: str, *, diagnostics: list[str] | None = None) -> list[GrupLACProduct]:
    """Lee productos declarados en CvLAC; no incorpora eventos ni proyectos."""
    parsed = urlparse(url)
    if parsed.scheme != "https" or parsed.hostname != "scienti.minciencias.gov.co" or \
            "/cvlac/" not in parsed.path or not parse_qs(parsed.query).get("cod_rh"):
        raise DataError("La fuente debe ser una ficha CvLAC HTTPS con código de investigador")

    class CvlacText(VisibleText):
        def handle_endtag(self, tag: str) -> None:
            super().handle_endtag(tag)
            if tag == "blockquote":
                self.parts.append("\n__CVLAC_END__\n")

    parser = CvlacText()
    parser.feed(html)
    lines = parser.lines()
    header = re.compile(r"^(Producción bibliográfica|Producción técnica|Trabajos dirigidos/Tutorías)\s+-\s+(.+)$")
    headers = [(index, match) for index, line in enumerate(lines) if (match := header.match(line))]
    by_identity: dict[str, GrupLACProduct] = {}
    for position, (index, match) in enumerate(headers):
        parts = re.split(r"\s+-\s+", match[2])
        section = CVLAC_SECTIONS.get(parts[0])
        if section is None:
            continue
        kind = parts[-1]
        last = headers[position + 1][0] if position + 1 < len(headers) else len(lines)
        last = next((i for i in range(index + 1, last) if lines[i] == "__CVLAC_END__"), last)
        details = lines[index + 1:last]
        cursor = 0
        author_lines = []
        while cursor < len(details) and details[cursor].endswith(",") and ":" not in details[cursor] \
                and not details[cursor].startswith((".", ",")):
            author_lines.append(details[cursor])
            cursor += 1
        quoted = next((found for line in details if (found := re.search(r'["“](.+?)["”]', line))), None)
        if quoted:
            title = quoted[1].strip()
        elif cursor < len(details) and details[cursor].startswith(("Nombre comercial:", "contrato/registro:", ". En:")):
            # En productos técnicos el título lleva coma final y precede a los metadatos.
            title = author_lines.pop().rstrip(",").strip() if len(author_lines) > 1 else ""
        else:
            title = details[cursor].rstrip(",").strip() if cursor < len(details) else ""
        if not title:
            if diagnostics is not None:
                diagnostics.append(f"Productos CvLAC, {parts[0]}, entrada {position + 1}: extracción incompleta, falta el título")
            continue
        detail = " ".join(details)
        year_match = re.search(r"(?<!\d)((?:19|20)\d{2})(?=[,.])", detail)
        if not year_match or not _publication_year(year_match[1]):
            if diagnostics is not None:
                diagnostics.append(f"Productos CvLAC, {parts[0]}, {title[:80]}: extracción incompleta, falta un año válido")
            continue
        year = year_match[1]
        published, doi = _source_doi(detail)
        note = f"; DOI incompleto en la fuente: {published}" if published and not doi else ""
        category = estimate_category(section, kind)
        identity = product_identity(title, year, published, kind, section == "Artículos publicados")
        family = {"Trabajos dirigidos/Tutorías": "Formación de recurso humano",
                  "Producción técnica": "Producción técnica y tecnológica"}.get(match[1], match[1])
        try:
            row = clean_row("productos", {
                "id": product_id(identity), "titulo": title, "anio": year, "tipologia": kind,
                "familia": family,
                "categoria": category, "categoria_origen": "estimada" if category else "",
                "doi": doi, "url": _redact_url(url),
                "fuente": f"CvLAC, sección {parts[0]}, consultado {date.today().isoformat()}{note}",
            })
        except DataError as exc:
            if diagnostics is None:
                raise
            diagnostics.append(f"Productos CvLAC, {parts[0]}, {title[:80]}: {exc}")
            continue
        authors = tuple(name.strip() for line in author_lines for name in line.split(",") if name.strip())
        entry = GrupLACProduct(row, authors, section)
        previous = by_identity.get(identity)
        by_identity[identity] = _combine_product(previous, entry) if previous else entry
    return list(by_identity.values())


def import_cvlac_preview(repository: Repository | CppRepository, preview: WebPreview,
                         reviewed_person: dict[str, str] | None = None, *,
                         remember: bool = True) -> dict[str, Any]:
    """Fusiona el perfil, sus productos y autorías; conserva los vínculos existentes."""
    if preview.suggested_kind != "investigadores" or not preview.suggested_row:
        raise DataError("La vista previa no contiene una ficha de investigador")
    person = clean_row("investigadores", reviewed_person or preview.suggested_row)
    if person["codigo_cvlac"] != preview.suggested_row["codigo_cvlac"]:
        raise DataError("El código CvLAC revisado no coincide con la fuente")
    counts = _import_counts()
    extra: dict[str, Any] = {} if remember else {"remember": False}
    with repository.batch() if remember else nullcontext():
        person_id = _merge_record(repository, "investigadores", person, person["id"], counts, extra)
        if person_id is None:
            return counts
        for product in preview.related_products:
            product_id = _merge_record(repository, "productos", product, product["id"], counts, extra)
            if product_id is None:
                continue
            key = (product_id, person_id)
            # «existentes» cuenta perfil y productos; una autoría presente no se vuelve a importar.
            if repository.get("autorias", key) is None:
                authorship = {"producto_id": key[0], "investigador_id": key[1], "rol": "Producto listado en CvLAC"}
                _merge_record(repository, "autorias", authorship, key, counts, extra)
    return counts


def extract_docx_html(path: Path) -> str:
    """Convierte un .docx a fragmento HTML con mammoth (instalación opcional)."""
    _check_local_file(path, archive=True)
    try:
        import mammoth
    except ImportError as exc:
        raise DataError("Para importar Word instale mammoth (pip install mammoth)") from exc
    try:
        with path.open("rb") as stream:
            html = mammoth.convert_to_html(stream).value
            if len(html) > MAX_PDF_TEXT:
                raise DataError("El documento Word produce demasiado texto")
            return html
    except DataError:
        raise
    except Exception as exc:
        raise DataError(f"No se pudo leer el documento Word: {exc}") from exc


def extract_pdf_text(path: Path) -> str:
    """Admite PDF de texto si pdftotext (Poppler) está instalado."""
    _check_local_file(path)
    if shutil.which("pdftotext") is None:
        raise DataError("Para importar PDF instale Poppler (pdftotext) o use CSV")

    def limit_resources() -> None:
        import resource
        resource.setrlimit(resource.RLIMIT_AS, (MAX_PDF_MEMORY, MAX_PDF_MEMORY))
        resource.setrlimit(resource.RLIMIT_CPU, (PDF_TIMEOUT, PDF_TIMEOUT))

    try:
        process = subprocess.Popen(
            ["pdftotext", "-layout", str(path), "-"],
            stdout=subprocess.PIPE, stderr=subprocess.DEVNULL, text=True,
            start_new_session=(os.name != "nt"),
            preexec_fn=limit_resources if os.name != "nt" else None,
        )
    except (OSError, ValueError) as exc:
        raise DataError(f"No se pudo leer el PDF: {exc}") from exc
    chunks: list[str] = []
    failures: list[Exception] = []

    def collect() -> None:
        total = 0
        try:
            assert process.stdout is not None
            while chunk := process.stdout.read(65536):
                total += len(chunk)
                if total > MAX_PDF_TEXT:
                    raise DataError("El PDF produce demasiado texto")
                chunks.append(chunk)
        except (OSError, UnicodeError, DataError) as exc:
            failures.append(exc)

    reader = threading.Thread(target=collect, daemon=True)
    reader.start()
    reader.join(PDF_TIMEOUT)
    if reader.is_alive() or failures:
        process.kill()
        process.wait()
        reader.join(timeout=1)
        if process.stdout is not None:
            process.stdout.close()
        if reader.is_alive():
            raise DataError("Se agotó el tiempo al leer el PDF")
        if failures:
            raise DataError(f"No se pudo leer el PDF: {failures[0]}") from failures[0]
        raise DataError("Se agotó el tiempo al leer el PDF")
    try:
        process.wait(timeout=2)
    except subprocess.TimeoutExpired as exc:
        process.kill()
        process.wait()
        raise DataError("Se agotó el tiempo al leer el PDF") from exc
    finally:
        if process.stdout is not None:
            process.stdout.close()
    if process.returncode != 0:
        raise DataError("No se pudo leer el PDF")
    text = "".join(chunks)
    if not text.strip():
        raise DataError("El PDF no contiene texto seleccionable; requiere OCR")
    return text


def parse_scienti_pdf(path: Path, source_url: str, *, text: str | None = None,
                      diagnostics: list[str] | None = None) -> tuple[str, dict[str, str]]:
    """Procesa PDFs impresos de GrupLAC/CvLAC con URL original conocida."""
    if text is None:
        text = extract_pdf_text(path)
    # El exportador PDF presenta texto en bloques; reutilizamos la misma extracción
    # de campos, con saltos de línea preservados para no inventar asociaciones.
    from html import escape
    html = "<p>" + "</p><p>".join(escape(line) for line in text.splitlines()) + "</p>"
    return parse_scienti_html(html, source_url, diagnostics=diagnostics)


def new_id(prefix: str) -> str:
    return prefix + "-" + uuid.uuid4().hex[:8].upper()


# La interfaz usa CustomTkinter; el núcleo de datos, los scripts y las pruebas no lo necesitan.
CTK_REQUIREMENT = "customtkinter==6.0.0"


def _interactive() -> bool:
    return bool(sys.stdin and sys.stdin.isatty())


def _bootstrap_ask(text: str) -> bool:
    """Pregunta antes de instalar: en consola con input; con el lanzador .pyw, con un aviso del sistema."""
    if _interactive():
        return input(f"{text} [S/n] ").strip().casefold() in ("", "s", "si", "sí", "y", "yes")
    # Todavía no existe ninguna ventana CustomTkinter: única ventana Tk clásica del programa.
    import tkinter as tk
    from tkinter import messagebox
    root = tk.Tk()
    root.withdraw()
    try:
        return messagebox.askyesno("PEA-i", text, parent=root)
    finally:
        root.destroy()


def _bootstrap_notice(text: str) -> None:
    if sys.stderr:
        print(text, file=sys.stderr)
    if _interactive():
        return
    import tkinter as tk
    from tkinter import messagebox
    root = tk.Tk()
    root.withdraw()
    try:
        messagebox.showerror("PEA-i", text, parent=root)
    finally:
        root.destroy()


def ensure_customtkinter(*, ask: Callable[[str], bool] | None = None,
                         run: Callable[..., subprocess.CompletedProcess] = subprocess.run) -> bool:
    """True si CustomTkinter puede importarse, instalándolo con pip cuando falta y el usuario acepta."""
    if importlib.util.find_spec("customtkinter") is not None:
        return True
    manual = f"{Path(sys.executable).name} -m pip install {CTK_REQUIREMENT}"
    question = (f"PEA-i necesita la biblioteca CustomTkinter ({CTK_REQUIREMENT}) para su interfaz. "
                "¿Instalarla ahora con pip? Se necesita Internet una sola vez.")
    if not (ask or _bootstrap_ask)(question):
        _bootstrap_notice(f"PEA-i no puede abrir su interfaz sin CustomTkinter. Instálela con:\n{manual}")
        return False
    in_virtualenv = sys.prefix != sys.base_prefix
    command = [sys.executable, "-m", "pip", "install", CTK_REQUIREMENT] + ([] if in_virtualenv else ["--user"])
    result = run(command, capture_output=True, text=True)
    if result.returncode != 0:
        output = f"{result.stdout or ''}\n{result.stderr or ''}"
        if "externally-managed-environment" in output:
            _bootstrap_notice(
                "El Python del sistema no admite instalar paquetes con pip (PEP 668). Use un entorno virtual "
                "desde la carpeta del proyecto:\n"
                "python3 -m venv .venv\n"
                ".venv/bin/pip install -r requirements.txt\n"
                ".venv/bin/python src/python/Taller2_REMR.py\n"
                "En Windows use .venv\\Scripts\\pip y .venv\\Scripts\\python.")
        else:
            tail = "\n".join(output.strip().splitlines()[-20:])
            _bootstrap_notice(f"No se pudo instalar CustomTkinter.\n{tail}\n\nInstálelo manualmente con:\n{manual}")
        return False
    if not in_virtualenv:
        # Un directorio de usuario recién creado no está aún en sys.path de este proceso.
        site.addsitedir(site.getusersitepackages())
    importlib.invalidate_caches()
    return importlib.util.find_spec("customtkinter") is not None


def open_gui_repository(*, python_backend: bool = False,
                        cpp_binary: Path | None = None) -> tuple[Repository | CppRepository, str]:
    """Abre el motor disponible y devuelve un aviso si se usó el respaldo Python."""
    if python_backend:
        return Repository(), ""
    try:
        return CppRepository(cpp_binary), ""
    except DataError as exc:
        if cpp_binary is not None:
            raise
        return Repository(), str(exc).splitlines()[0]


# «Sierra y sabana»: verde institucional UPC; el amarillo solo marca lo pendiente.
LIGHT_THEME = {
    "page": "#f3f6f3", "surface": "#ffffff", "field": "#ffffff",
    "sidebar": "#163a2e", "sidebar_fg": "#e3ece7", "sidebar_muted": "#a9c2b6",
    "sidebar_hover": "#1f4a3b", "ink": "#1c2a24", "muted": "#5b6c64",
    "line": "#dce4df", "line_strong": "#c3cfc8", "hover": "#eaf0ec",
    "primary": "#1e6b50", "primary_hover": "#185a43", "primary_pressed": "#124636",
    "on_primary": "#ffffff", "primary_soft": "#dcefe5", "stripe": "#f7faf8",
    "thumb": "#c3cfc8", "thumb_active": "#9fb0a7",
    "chart_grid": "#e7ece9", "chart_track": "#eef3f0", "chart_older": "#8fb8a5",
    "pending": "#bb7e0e", "rejected": "#b5472f", "unknown": "#87938d",
    "error": "#b42318", "disabled": "#9fb0a7",
    # Clases del modelo 2024 en los gráficos de anillo y de mosaico.
    "class_top": "#1e6b50", "class_a": "#4a9d78", "class_b": "#97c9ac", "class_asc": "#c48a1c",
    "class_dpc": "#2f78a0", "class_frh_a": "#8659a6", "class_frh_b": "#c0634a", "class_none": "#bcc6c1",
}

DARK_THEME = {
    "page": "#111714", "surface": "#18201c", "field": "#1d2621",
    "sidebar": "#0b100e", "sidebar_fg": "#e3ece7", "sidebar_muted": "#8fa79b",
    "sidebar_hover": "#16201b", "ink": "#e3ece7", "muted": "#94a69d",
    "line": "#28332e", "line_strong": "#3a4842", "hover": "#222c27",
    "primary": "#4db38c", "primary_hover": "#2f8a68", "primary_pressed": "#236b51",
    "on_primary": "#ffffff", "primary_soft": "#1f3a30", "stripe": "#1b241f",
    "thumb": "#34423b", "thumb_active": "#4a5b52",
    "chart_grid": "#232d28", "chart_track": "#202a25", "chart_older": "#3f7d65",
    "pending": "#e6b04a", "rejected": "#e07a62", "unknown": "#7d8c85",
    "error": "#ffb4ab", "disabled": "#4f5f57",
    "class_top": "#4db38c", "class_a": "#7ccaa6", "class_b": "#b7e2ca", "class_asc": "#e6b04a",
    "class_dpc": "#62abd6", "class_frh_a": "#b78fd6", "class_frh_b": "#e48b73", "class_none": "#56645d",
}
# Botón primario: en oscuro el verde claro sirve para trazos, no como fondo de texto blanco.
DARK_THEME["button"] = "#2a7a5c"
LIGHT_THEME["button"] = LIGHT_THEME["primary"]
STATUS_COLORS = {"validado": "primary", "rechazado": "rejected", "Sin dato": "unknown"}


def build_gui(root: Any, repository: Repository | CppRepository,
              data_file: Path | None = None, *, load_on_start: bool = True) -> Any:
    """Construye la ventana; permite verificar sus flujos con un repositorio aislado."""
    import tkinter as tk
    import tkinter.font as tkfont
    from tkinter import TclError, filedialog, ttk

    from types import SimpleNamespace
    import weakref

    import customtkinter as ctk

    catalog = repository.catalog()
    catalog_by_code = {row["codigo"]: row for row in catalog}

    def decimal(value: float, digits: int | None = None) -> str:
        return (f"{value:g}" if digits is None else f"{value:.{digits}f}").replace(".", ",")

    available_fonts = set(tkfont.families(root))

    def first_font(*names: str) -> str | None:
        return next((name for name in names if name in available_fonts), None)

    BODY_FONT = first_font("Segoe UI Variable Text", "Segoe UI", "Inter", "Noto Sans", "DejaVu Sans") or "TkDefaultFont"

    # Campos, menús y diálogos estándar usan las fuentes con nombre de Tk.
    for name in ("TkDefaultFont", "TkTextFont", "TkMenuFont", "TkHeadingFont", "TkCaptionFont"):
        try:
            tkfont.nametofont(name, root=root).configure(family=BODY_FONT, size=10)
        except tk.TclError:
            pass

    TITLES = {"grupos": "Grupos", "investigadores": "Investigadores", "productos": "Productos",
              "planes": "Planes", "membresias": "Integrantes", "autorias": "Autorías",
              "grupos_productos": "Grupos y productos"}
    SINGULAR = {"grupos": "Grupo", "investigadores": "Investigador", "productos": "Producto", "planes": "Plan"}

    def record_name(row: dict[str, str]) -> str:
        return row.get("nombre") or row.get("titulo") or row["id"]

    def display_value(field: str, value: str) -> str:
        if field == "activo":
            return "Activo" if value == "1" else "Inactivo"
        if field == "validacion":
            return value.capitalize()
        return value or "Sin dato"

    def reason_text(reason: str, short: bool = False) -> str:
        codes = auto_reason_codes(reason)
        if codes is None:
            return reason or "Sin motivo registrado"
        rules = "; ".join(RULE_LABELS[code] for code in codes) or "rechazado"
        return rules if short else "Validación automática: " + rules

    def cell_value(field: str, value: str) -> str:
        """Texto de una celda de tabla: una línea y un guion discreto para lo vacío."""
        return " ".join(display_value(field, value).split()) if value or field == "activo" else "—"

    def category_tone(row: dict[str, str]) -> str:
        if not row.get("categoria"):
            return "unknown"
        return "pending" if row.get("categoria_origen") == "estimada" else "primary"

    # Kit CustomTkinter: tokens de color, tipografía y componentes compartidos por todas las pantallas.
    ctk.set_default_color_theme("green")
    T = {name: (LIGHT_THEME[name], DARK_THEME[name]) for name in LIGHT_THEME}

    def color(name: str) -> str:
        """Color del token en la apariencia actual; los lienzos no cambian de tema por sí solos."""
        return T[name][1 if ctk.get_appearance_mode() == "Dark" else 0]

    def mix(first: str, second: str, amount: float) -> str:
        start = [int(first[i:i + 2], 16) for i in (1, 3, 5)]
        end = [int(second[i:i + 2], 16) for i in (1, 3, 5)]
        return "#" + "".join(f"{round(a + (b - a) * amount):02x}" for a, b in zip(start, end))

    def soft(tone: str) -> tuple[str, str]:
        """Fondo tenue de un tono para chips y etiquetas de estado."""
        return (mix(LIGHT_THEME["surface"], LIGHT_THEME[tone], 0.14),
                mix(DARK_THEME["surface"], DARK_THEME[tone], 0.24))

    def number(value: int) -> str:
        return f"{value:,}".replace(",", ".")

    FAMILY = first_font("Segoe UI Variable Text", "Segoe UI", "Inter", "Roboto") or "Roboto"
    FONTS = {name: ctk.CTkFont(family=FAMILY, size=size, weight=weight) for name, (size, weight) in {
        "display": (30, "bold"), "h1": (22, "bold"), "h2": (16, "bold"),
        "body": (13, "normal"), "small": (12, "normal"), "caption": (11, "normal")}.items()}

    def button(parent: Any, text: str, command: Callable[[], Any], variant: str = "secondary",
               **options: Any) -> Any:
        styles = {
            "primary": {"fg_color": T["button"], "hover_color": T["primary_pressed"], "text_color": T["on_primary"]},
            "secondary": {"fg_color": "transparent", "hover_color": T["hover"], "text_color": T["ink"],
                          "border_width": 1, "border_color": T["line_strong"]},
            "danger": {"fg_color": T["rejected"], "text_color": T["on_primary"],
                       "hover_color": (mix(LIGHT_THEME["rejected"], "#000000", 0.15),
                                       mix(DARK_THEME["rejected"], "#000000", 0.15))},
            "ghost": {"fg_color": "transparent", "hover_color": T["sidebar_hover"],
                      "text_color": T["sidebar_fg"], "anchor": "w"},
        }[variant]
        return ctk.CTkButton(parent, text=text, command=command,
                             **{"height": 34, "corner_radius": 8, "font": FONTS["body"],
                                "text_color_disabled": T["disabled"], **styles, **options})

    class Card(ctk.CTkFrame):
        """Superficie con borde fino; franja opcional del color de estado y título."""
        def __init__(self, parent: Any, title: str | None = None, stripe: str | None = None, **options: Any):
            super().__init__(parent, fg_color=T["surface"], border_color=T["line"], border_width=1,
                             corner_radius=12, **options)
            if stripe:
                ctk.CTkFrame(self, height=3, corner_radius=0, fg_color=T[stripe]).pack(fill="x", padx=12, pady=(1, 0))
            if title:
                ctk.CTkLabel(self, text=title, font=FONTS["h2"], text_color=T["ink"], anchor="w").pack(
                    fill="x", padx=16, pady=(12, 0))
            self.body = ctk.CTkFrame(self, fg_color="transparent")
            self.body.pack(fill="both", expand=True, padx=16, pady=(8, 14))

    def segmented(parent: Any, variable: Any, values: list[str], command: Callable[[str], Any]) -> Any:
        return ctk.CTkSegmentedButton(parent, variable=variable, values=values, command=command,
                                      height=34, corner_radius=8, font=FONTS["body"], text_color=T["ink"],
                                      fg_color=T["field"], selected_color=T["primary_soft"],
                                      selected_hover_color=T["hover"], unselected_color=T["field"],
                                      unselected_hover_color=T["hover"])

    def tabview(parent: Any) -> Any:
        return ctk.CTkTabview(parent, fg_color=T["page"], text_color=T["ink"], segmented_button_font=FONTS["body"],
                              segmented_button_fg_color=T["field"], segmented_button_selected_color=T["primary_soft"],
                              segmented_button_selected_hover_color=T["hover"],
                              segmented_button_unselected_color=T["field"],
                              segmented_button_unselected_hover_color=T["hover"])

    class Kpi(Card):
        def __init__(self, parent: Any, caption: str, tone: str, variable: Any, **options: Any):
            super().__init__(parent, stripe=tone, **options)
            self.value_label = ctk.CTkLabel(self.body, textvariable=variable, font=FONTS["display"],
                                            text_color=T["ink"], anchor="w")
            self.value_label.pack(fill="x")
            ctk.CTkLabel(self.body, text=caption, font=FONTS["caption"], text_color=T["muted"], anchor="w").pack(fill="x")

    class Chip(ctk.CTkLabel):
        def __init__(self, parent: Any, text: str, tone: str, **options: Any):
            super().__init__(parent, text=text, font=FONTS["caption"], corner_radius=10, height=22, padx=8, **options)
            self.set(text, tone)

        def set(self, text: str, tone: str) -> None:
            self.configure(text=text, text_color=T[tone], fg_color=soft(tone))

    class MenuButton(ctk.CTkOptionMenu):
        """Menú de cabecera: muestra siempre su nombre y ejecuta la acción elegida."""
        def __init__(self, parent: Any, label: str, items: list[tuple[str, Callable[[], Any]]], **options: Any):
            self._label = label
            self._actions = dict(items)
            super().__init__(parent, values=[name for name, _ in items], command=self.choose,
                             dynamic_resizing=False, width=options.pop("width", 118), height=32, corner_radius=8,
                             font=FONTS["body"], dropdown_font=FONTS["body"], fg_color=T["page"],
                             button_color=T["page"], button_hover_color=T["hover"], text_color=T["ink"],
                             dropdown_fg_color=T["surface"], dropdown_hover_color=T["hover"],
                             dropdown_text_color=T["ink"], **options)
            self.set(label)

        def choose(self, choice: str) -> None:
            self.set(self._label)
            action = self._actions.get(choice)
            if action:
                action()

    class SearchField(ctk.CTkEntry):
        """Campo de búsqueda: avisa una vez cuando se deja de escribir."""
        def __init__(self, parent: Any, variable: Any, on_change: Callable[[], Any], delay: int = 250,
                     **options: Any):
            placeholder = options.pop("placeholder_text", "Buscar…")
            super().__init__(parent, textvariable=variable, height=34, corner_radius=8, font=FONTS["body"],
                             fg_color=T["field"], border_color=T["line_strong"], text_color=T["ink"], **options)
            self._on_change, self._delay, self._timer = on_change, delay, None
            self._variable = variable
            # CTkEntry no muestra placeholder_text junto a textvariable. La pista es solo visual.
            self._focused = False
            self.hint = ctk.CTkLabel(self, text=placeholder, font=FONTS["body"], text_color=T["muted"],
                                     fg_color=T["field"], anchor="w", height=24)
            self.hint.bind("<Button-1>", lambda _event: self.focus_set())
            self.bind("<FocusIn>", lambda _event: self._focus_hint(True), add="+")
            self.bind("<FocusOut>", lambda _event: self._focus_hint(False), add="+")
            self._trace = variable.trace_add("write", lambda *_: self._schedule())
            self._show_hint()

        def _focus_hint(self, focused: bool) -> None:
            self._focused = focused
            self._show_hint()

        def _show_hint(self) -> None:
            if not self._variable.get() and not self._focused:
                self.hint.place(x=12, rely=0.5, anchor="w")
            else:
                self.hint.place_forget()

        def _schedule(self) -> None:
            self._show_hint()
            if self._timer:
                self.after_cancel(self._timer)
            self._timer = self.after(self._delay, self._fire)

        def _fire(self) -> None:
            self._timer = None
            self._on_change()

        def cancel_pending(self) -> None:
            if self._timer:
                self.after_cancel(self._timer)
                self._timer = None

        def destroy(self) -> None:
            if self._timer:
                self.after_cancel(self._timer)
            try:
                self._variable.trace_remove("write", self._trace)
            except TclError:
                pass
            super().destroy()

    class Pager(ctk.CTkFrame):
        def __init__(self, parent: Any, on_move: Callable[[int], Any], **options: Any):
            super().__init__(parent, fg_color="transparent", **options)
            self.previous = button(self, "Anterior", lambda: on_move(-1), width=96)
            self.previous.pack(side="left")
            self.label = ctk.CTkLabel(self, text="", font=FONTS["small"], text_color=T["muted"])
            self.label.pack(side="left", padx=12)
            self.next = button(self, "Siguiente", lambda: on_move(1), width=96)
            self.next.pack(side="left")

        def show(self, offset: int, shown: int, total: int) -> None:
            self.label.configure(text=f"{offset + 1}–{offset + shown} de {number(total)}" if total else "Sin resultados")
            self.previous.configure(state="normal" if offset > 0 else "disabled")
            self.next.configure(state="normal" if offset + shown < total else "disabled")

    class EmptyState(ctk.CTkFrame):
        def __init__(self, parent: Any, title: str, hint: str = "", **options: Any):
            super().__init__(parent, fg_color="transparent", **options)
            ctk.CTkLabel(self, text=title, font=FONTS["h2"], text_color=T["muted"]).pack(pady=(24, 4))
            if hint:
                ctk.CTkLabel(self, text=hint, font=FONTS["small"], text_color=T["muted"], wraplength=420).pack()

    def reason_chips(parent: Any, reason: str | None, per_row: int = 3) -> list[Any]:
        """Motivo de una revisión: un chip por regla incumplida, o uno con el motivo escrito a mano."""
        for child in parent.winfo_children():
            child.destroy()
        if reason is None:
            return []
        codes = auto_reason_codes(reason)
        if codes is None:
            text = reason.strip() or "Sin motivo registrado"
            texts, tone = [text if len(text) <= 70 else text[:69] + "…"], "pending"
        else:
            texts, tone = [RULE_LABELS[code] for code in codes] or ["Rechazado"], "rejected"
        chips = [Chip(parent, text, tone) for text in texts]
        for index, chip in enumerate(chips):
            chip.grid(row=index // per_row, column=index % per_row, sticky="w", padx=(0, 6), pady=(0, 4))
        return chips

    def appearance(pair: tuple[str, str]) -> str:
        return pair[1 if ctk.get_appearance_mode() == "Dark" else 0]

    def palette() -> dict[str, str]:
        """Colores de la apariencia actual para lo que se dibuja en lienzos."""
        return DARK_THEME if ctk.get_appearance_mode() == "Dark" else LIGHT_THEME

    def scale_of(widget: Any) -> float:
        """Escala de CustomTkinter; una ventana Tk clásica (aún sin migrar) usa la escala global."""
        try:
            return ctk.ScalingTracker.get_widget_scaling(widget)
        except KeyError:
            return ctk.ScalingTracker.widget_scaling

    def chart_font(widget: Any, points: int, weight: str = "normal") -> tuple[Any, ...]:
        """Fuente en píxeles escalados (un lienzo no se escala solo); points se expresa en puntos tipográficos."""
        size = -round(points * 4 / 3 * scale_of(widget))
        return (FAMILY, size, "bold") if weight != "normal" else (FAMILY, size)

    def wheel_units(event: Any) -> int:
        if getattr(event, "num", None) in (4, 5):
            return -1 if event.num == 4 else 1
        return -int(event.delta / 120) if abs(event.delta) >= 120 else -1 if event.delta > 0 else 1

    scrollbar_updates: set[Any] = set()

    def set_scrollbar(scroll: Any, first: str, last: str, *, autohide: bool = False) -> None:
        """CTkScrollbar procesa eventos al dibujar: evita actualizaciones anidadas durante cambios de escala."""
        if scroll in scrollbar_updates:
            return
        scrollbar_updates.add(scroll)
        try:
            values = (float(first), float(last))
            if scroll.get() != values:
                scroll.set(*values)
            if autohide:
                needed = values[0] > 0 or values[1] < 1
                if needed != bool(scroll.winfo_manager()):
                    scroll.grid() if needed else scroll.grid_remove()
        finally:
            scrollbar_updates.remove(scroll)

    def cancel_widget_timers(window: Any) -> None:
        """Cancela callbacks del árbol de widgets antes de destruir sus comandos Tcl."""
        names: set[str] = set()
        pending = [window]
        while pending:
            widget = pending.pop()
            names.update(getattr(widget, "_tclCommands", None) or ())
            pending.extend(widget.winfo_children())
        for job in window.tk.call("after", "info"):
            try:
                script = str(window.tk.call("after", "info", job)[0])
            except TclError:
                continue
            if any(name in script for name in names):
                # after_cancel borraría también el comando, que destroy() elimina después.
                window.tk.call("after", "cancel", job)

    # Lienzos que se vuelven a dibujar al cambiar entre apariencia clara y oscura.
    canvas_registry: weakref.WeakSet[Any] = weakref.WeakSet()

    class ChartCanvas(ctk.CTkCanvas):
        """Lienzo de gráfico sobre la superficie; se redibuja al cambiar de tamaño o de apariencia."""
        def __init__(self, parent: Any, draw: Callable[[Any], Any] | None = None, **options: Any):
            super().__init__(parent, highlightthickness=0, borderwidth=0, bg=color("surface"), **options)
            self.draw = draw
            canvas_registry.add(self)
            self.bind("<Configure>", lambda _event: self.redraw())
            for sequence in ("<MouseWheel>", "<Button-4>", "<Button-5>"):
                self.bind(sequence, self._wheel)

        def _wheel(self, event: Any) -> str | None:
            axis = "x" if event.state & 0x1 else "y"
            view = self.xview() if axis == "x" else self.yview()
            if view[1] - view[0] < 1:
                command = self.xview_scroll if axis == "x" else self.yview_scroll
                command(wheel_units(event), "units")
                return "break"
            return None

        def redraw(self) -> None:
            self.configure(bg=color("surface"))
            if self.draw is not None:
                self.draw(self)

        def destroy(self) -> None:
            canvas_registry.discard(self)
            super().destroy()

    class DataTable(ctk.CTkFrame):
        """Tabla dibujada en un lienzo: solo las filas visibles, textos recortados con «…» y chips de estado.

        Ofrece el subconjunto de ttk.Treeview que usan las pantallas (insert, delete, item, selection_set…),
        teclado, rueda del ratón y una ayuda emergente con el texto completo de una celda recortada.
        """
        ROW, HEADER, OVERSCAN, PAD = 34, 36, 5, 12

        def __init__(self, parent: Any, columns: tuple[str, ...], *, headings: dict[str, str],
                     widths: dict[str, int], anchors: dict[str, str] | None = None,
                     chips: dict[str, Callable[[str], str | None]] | None = None,
                     row_chips: dict[str, Callable[[str, str], str | None]] | None = None,
                     tooltips: dict[str, Callable[[str, str], str]] | None = None,
                     height_rows: int = 12, open_on_click: bool = False, **options: Any):
            super().__init__(parent, fg_color=T["surface"], border_color=T["line"], border_width=1,
                             corner_radius=10, **options)
            self._columns = tuple(columns)
            self._display = self._columns
            self._index = {name: position for position, name in enumerate(self._columns)}
            self._headings = {name: headings.get(name, name) for name in self._columns}
            self._widths = {name: widths.get(name, 120) for name in self._columns}
            self._anchors = {name: (anchors or {}).get(name, "w") for name in self._columns}
            self._chips = dict(chips or {})
            self._row_chips = dict(row_chips or {})
            self._tooltips = dict(tooltips or {})
            self._height_rows = height_rows
            self.open_on_click = open_on_click
            self.empty_text = "Sin registros"
            self._order: list[str] = []
            self._items: dict[str, dict[str, tuple[Any, ...]]] = {}
            self._selected: str | None = None
            self._hover: int | None = None
            self._pressed: int | None = None
            self._select_callbacks: list[Callable[[str], Any]] = []
            self._open_callbacks: list[Callable[[str], Any]] = []
            self._col_x: list[int] = []
            self._col_w: list[int] = []
            self._content_width = 1
            self._fitted: dict[tuple[str, int, str], str] = {}
            self._font_scale = 0.0
            self._pending: str | None = None
            self._tip_job: str | None = None
            self._tip_cell: tuple[int, str] | None = None
            self.tooltip: Any = None
            self.tooltip_label: Any = None
            self._counter = 0
            self.grid_columnconfigure(0, weight=1)
            # Sin barra vertical, el lienzo deja libre el borde derecho del marco, como el izquierdo.
            self.grid_columnconfigure(1, minsize=6)
            self.grid_rowconfigure(1, weight=1)
            self.header = ctk.CTkCanvas(self, highlightthickness=0, borderwidth=0)
            self.body = ctk.CTkCanvas(self, highlightthickness=0, borderwidth=0, takefocus=1)
            self.vbar = ctk.CTkScrollbar(self, orientation="vertical", height=0, command=self.body.yview)
            self.hbar = ctk.CTkScrollbar(self, orientation="horizontal", width=0, command=self.body.xview)
            self.header.grid(row=0, column=0, sticky="ew", padx=(6, 0), pady=(6, 0))
            self.body.grid(row=1, column=0, sticky="nsew", padx=(6, 0), pady=(0, 6))
            self.vbar.grid(row=1, column=1, sticky="ns", padx=(0, 4), pady=(0, 6))
            self.hbar.grid(row=2, column=0, sticky="ew", padx=(6, 0), pady=(0, 4))
            self.vbar.grid_remove()
            self.hbar.grid_remove()
            self.body.configure(yscrollcommand=self._yscrolled, xscrollcommand=self._xscrolled)
            for sequence, handler in (("<Button-1>", self._press), ("<ButtonRelease-1>", self._release),
                                      ("<Double-Button-1>", self._double), ("<Motion>", self._motion),
                                      ("<Leave>", self._leave), ("<Key>", self._key),
                                      ("<Configure>", lambda _event: self._schedule()),
                                      ("<FocusIn>", lambda _event: self.configure(border_color=T["primary"])),
                                      ("<FocusOut>", lambda _event: self.configure(border_color=T["line"]))):
                self.body.bind(sequence, handler)
            for widget in (self.body, self.header):
                for sequence in ("<MouseWheel>", "<Shift-MouseWheel>", "<Button-4>", "<Button-5>",
                                 "<Shift-Button-4>", "<Shift-Button-5>"):
                    widget.bind(sequence, self._wheel)
            canvas_registry.add(self)
            self.redraw()

        # Subconjunto de ttk.Treeview
        def insert(self, parent: str, index: Any, iid: str | None = None, values: Any = (), tags: Any = (),
                   **_options: Any) -> str:
            if parent != "":
                raise TclError("DataTable no admite filas anidadas")
            if iid is None:
                self._counter += 1
                iid = f"I{self._counter:04d}"
            if iid in self._items:
                raise TclError(f"Item {iid} already exists")
            self._items[iid] = {"values": tuple(values), "tags": tuple(tags)}
            self._order.insert(len(self._order) if index == "end" else int(index), iid)
            self._schedule()
            return iid

        def delete(self, *items: str) -> None:
            selected = self._selected
            for iid in items:
                if iid in self._items:
                    del self._items[iid]
                    self._order.remove(iid)
                    if self._selected == iid:
                        self._selected = None
            self._tip_cell = None
            self._hide_tip()
            if selected and self._selected is None:
                self._notify_selection()
            self._schedule()

        def get_children(self, _item: str = "") -> tuple[str, ...]:
            return tuple(self._order)

        def exists(self, iid: str) -> bool:
            return iid in self._items

        def _entry(self, iid: str) -> dict[str, tuple[Any, ...]]:
            if iid not in self._items:
                raise TclError(f"Item {iid} not found")
            return self._items[iid]

        def item(self, iid: str, option: str | None = None, **changes: Any) -> Any:
            entry = self._entry(iid)
            if changes:
                entry.update({name: tuple(value) for name, value in changes.items() if name in ("values", "tags")})
                self._schedule()
                return None
            data = {"text": "", "values": entry["values"], "tags": entry["tags"]}
            return data if option is None else data[option]

        def selection(self) -> tuple[str, ...]:
            return (self._selected,) if self._selected else ()

        def selection_set(self, items: Any) -> None:
            iid = items if isinstance(items, str) else (tuple(items)[0] if items else None)
            if iid is None:
                self._selected = None
                self._draw_rows()
                self._notify_selection()
                return
            self._entry(iid)
            self._select(iid)
            self.see(iid)

        def focus(self, item: str | None = None) -> str | None:
            if item is None:
                return self._selected or ""
            self._entry(item)
            self._selected = item
            self._draw_rows()
            return None

        def see(self, iid: str) -> None:
            self._flush()
            index, row = self._order.index(iid), self._row_h()
            total = max(len(self._order) * row, 1)
            top, height = self.body.canvasy(0), max(self.body.winfo_height(), row)
            if index * row < top:
                self.body.yview_moveto(index * row / total)
            elif (index + 1) * row > top + height:
                self.body.yview_moveto(((index + 1) * row - height) / total)

        def identify_row(self, y: int) -> str:
            self._flush()
            index = self._row_at(y)
            return self._order[index] if index is not None else ""

        def bbox(self, iid: str, column: str | None = None) -> tuple[int, int, int, int] | str:
            self._flush()
            index, row = self._order.index(iid), self._row_h()
            y = index * row - round(self.body.canvasy(0))
            if y + row <= 0 or y >= self.body.winfo_height():
                return ""
            if column is None:
                return (-round(self.body.canvasx(0)), y, self._content_width, row)
            position = self._display.index(column)
            return (self._col_x[position] - round(self.body.canvasx(0)), y, self._col_w[position], row)

        def heading(self, column: str, option: str | None = None, text: str | None = None, **_options: Any) -> Any:
            if text is not None:
                self._headings[column] = text
                self._schedule()
                return None
            data = {"text": self._headings[column]}
            return data if option is None else data[option]

        def column(self, column: str, option: str | None = None, width: int | None = None,
                   anchor: str | None = None, **_options: Any) -> Any:
            if width is not None or anchor is not None:
                if width is not None:
                    self._widths[column] = width
                if anchor is not None:
                    self._anchors[column] = anchor
                self._schedule()
                return None
            data = {"width": self._widths[column], "anchor": self._anchors[column]}
            return data if option is None else data[option]

        def tag_configure(self, *_args: Any, **_options: Any) -> None:
            """Las franjas alternas son automáticas; se acepta por compatibilidad con Treeview."""

        def configure(self, require_redraw: bool = False, **options: Any) -> None:
            if "displaycolumns" in options:
                shown = options.pop("displaycolumns")
                self._display = self._columns if tuple(shown) in (("#all",), ("#", "a", "l", "l")) \
                    else tuple(shown)
                self._schedule()
            if options or require_redraw:
                super().configure(require_redraw=require_redraw, **options)

        def cget(self, attribute: str) -> Any:
            if attribute == "displaycolumns":
                return self._display
            if attribute == "columns":
                return self._columns
            return super().cget(attribute)

        def __getitem__(self, attribute: str) -> Any:
            return self.cget(attribute)

        # Selección y apertura
        def on_select(self, callback: Callable[[str], Any]) -> None:
            self._select_callbacks.append(callback)

        def on_open(self, callback: Callable[[str], Any]) -> None:
            self._open_callbacks.append(callback)

        def click_row(self, iid: str, double: bool = False) -> None:
            self._select(iid)
            if double or self.open_on_click:
                self._open(iid)

        def column_width(self, column: str) -> int:
            self._flush()
            return self._col_w[self._display.index(column)]

        def _select(self, iid: str) -> None:
            changed = iid != self._selected
            self._selected = iid
            self._draw_rows()
            if changed:
                self._notify_selection()

        def _notify_selection(self) -> None:
            for callback in self._select_callbacks:
                callback(self._selected or "")
            if self.winfo_exists():
                self.body.event_generate("<<TreeviewSelect>>")

        def _open(self, iid: str) -> None:
            for callback in self._open_callbacks:
                callback(iid)
            if self.winfo_exists():
                self.body.event_generate("<<TableOpen>>")

        # Dibujo
        def _scale(self) -> float:
            return scale_of(self)

        def _row_h(self) -> int:
            return round(self.ROW * self._scale())

        def _schedule(self) -> None:
            if self._pending is None and self.winfo_exists():
                self._pending = self.after_idle(self.redraw)

        def _flush(self) -> None:
            if self._pending is not None:
                self.after_cancel(self._pending)
                self.redraw()

        def redraw(self) -> None:
            self._pending = None
            scale = self._scale()
            if scale != self._font_scale:
                self._font_scale = scale
                self._fitted.clear()
                self._font = tkfont.Font(root=self, family=FAMILY, size=-round(12 * scale))
                self._head_font = tkfont.Font(root=self, family=FAMILY, size=-round(12 * scale), weight="bold")
                self._chip_font = tkfont.Font(root=self, family=FAMILY, size=-round(11 * scale))
                row = self._row_h()
                self.header.configure(height=round(self.HEADER * scale))
                self.body.configure(height=self._height_rows * row, yscrollincrement=row)
            for canvas in (self.header, self.body):
                canvas.configure(bg=color("surface"))
            self._layout()
            self._draw_header()
            self._draw_rows()

        def _layout(self) -> None:
            scale = self._scale()
            available = max(1, self.body.winfo_width())
            base = [max(1, round(self._widths[name] * scale)) for name in self._display]
            total = sum(base)
            widths = base
            if base and total < available:
                widths = [width + (available - total) * width // total for width in base]
                widths[-1] += available - sum(widths)
            self._col_x, self._col_w, left = [], [], 0
            for width in widths:
                self._col_x.append(left)
                self._col_w.append(width)
                left += width
            self._content_width = max(left, 1)
            height = max(len(self._order) * self._row_h(), 1)
            self.body.configure(scrollregion=(0, 0, self._content_width, height))
            self.header.configure(scrollregion=(0, 0, self._content_width, round(self.HEADER * scale)))

        def _fit(self, text: str, width: int, font: Any) -> str:
            key = (text, width, str(font))
            if key not in self._fitted:
                if len(self._fitted) > 20000:
                    self._fitted.clear()
                if width <= 0:
                    fitted = ""
                elif font.measure(text) <= width:
                    fitted = text
                else:
                    low, high = 0, len(text)
                    while low < high:
                        middle = (low + high + 1) // 2
                        if font.measure(text[:middle] + "…") <= width:
                            low = middle
                        else:
                            high = middle - 1
                    fitted = text[:low].rstrip() + "…"
                self._fitted[key] = fitted
            return self._fitted[key]

        def _cell_text(self, index: int, column: str) -> str:
            values = self._items[self._order[index]]["values"]
            position = self._index[column]
            return " ".join(str(values[position]).split()) if position < len(values) else ""

        def cell_tone(self, iid: str, column: str) -> str | None:
            values = self._entry(iid)["values"]
            text = str(values[self._index[column]]) if self._index[column] < len(values) else ""
            if column in self._row_chips:
                return self._row_chips[column](iid, text)
            return self._chips[column](text) if column in self._chips and text else None

        def cell_tooltip(self, iid: str, column: str) -> str:
            values = self._entry(iid)["values"]
            text = str(values[self._index[column]]) if self._index[column] < len(values) else ""
            return self._tooltips[column](iid, text) if column in self._tooltips else ""

        def _draw_header(self) -> None:
            scale, canvas = self._scale(), self.header
            canvas.delete("all")
            height, pad = round(self.HEADER * scale), round(self.PAD * scale)
            canvas.create_line(0, height - 1, self._content_width, height - 1, fill=color("line"))
            for name, left, width in zip(self._display, self._col_x, self._col_w):
                anchor = self._anchors[name]
                x = left + pad if anchor == "w" else left + width - pad
                canvas.create_text(x, height // 2, text=self._fit(self._headings[name], width - 2 * pad, self._head_font),
                                   anchor=anchor, font=self._head_font, fill=color("muted"))

        def _draw_rows(self) -> None:
            canvas, scale, row = self.body, self._scale(), self._row_h()
            canvas.delete("row", "empty")
            if not self._order:
                canvas.create_text(max(canvas.winfo_width(), 1) // 2, round(row * 1.5), text=self.empty_text,
                                   font=self._font, fill=color("muted"), tags=("empty",))
                return
            top = canvas.canvasy(0)
            height = max(canvas.winfo_height(), row)
            first = max(0, int(top // row) - self.OVERSCAN)
            last = min(len(self._order), int((top + height) // row) + 1 + self.OVERSCAN)
            pad, chip_height, chip_pad = round(self.PAD * scale), round(22 * scale), round(8 * scale)
            for index in range(first, last):
                iid, y = self._order[index], index * row
                fill = ("primary_soft" if iid == self._selected else "hover" if index == self._hover
                        else "stripe" if index % 2 else "surface")
                canvas.create_rectangle(0, y, self._content_width, y + row, fill=color(fill), outline="", tags=("row",))
                if iid == self._selected:
                    canvas.create_rectangle(0, y, round(3 * scale), y + row, fill=color("primary"), outline="",
                                            tags=("row",))
                for name, left, width in zip(self._display, self._col_x, self._col_w):
                    text, anchor = self._cell_text(index, name), self._anchors[name]
                    tone = self.cell_tone(iid, name)
                    if tone:
                        label = self._fit(text, width - 2 * pad - 2 * chip_pad, self._chip_font)
                        pill = self._chip_font.measure(label) + 2 * chip_pad
                        x1 = left + pad if anchor == "w" else left + width - pad - pill
                        y1 = y + (row - chip_height) / 2
                        backdrop = color("surface") if iid == self._selected else appearance(soft(tone))
                        self._pill(x1, y1, x1 + pill, y1 + chip_height, backdrop)
                        canvas.create_text(x1 + pill / 2, y + row / 2, text=label, font=self._chip_font,
                                           fill=color(tone), tags=("row", "chip_text"))
                    else:
                        x = left + pad if anchor == "w" else left + width - pad
                        canvas.create_text(x, y + row / 2, text=self._fit(text, width - 2 * pad, self._font),
                                           anchor=anchor, font=self._font, fill=color("ink"), tags=("row", "cell"))

        def _pill(self, x1: float, y1: float, x2: float, y2: float, fill: str) -> None:
            radius = (y2 - y1) / 2
            points = (x1 + radius, y1, x2 - radius, y1, x2, y1, x2, y1 + radius, x2, y2 - radius, x2, y2,
                      x2 - radius, y2, x1 + radius, y2, x1, y2, x1, y2 - radius, x1, y1 + radius, x1, y1)
            self.body.create_polygon(points, smooth=True, fill=fill, outline="", tags=("row", "chip"))

        # Desplazamiento
        def _yscrolled(self, first: str, last: str) -> None:
            set_scrollbar(self.vbar, first, last)
            needed = len(self._order) * self._row_h() > self.body.winfo_height()
            if needed != bool(self.vbar.winfo_manager()):
                self.vbar.grid() if needed else self.vbar.grid_remove()
            self._draw_rows()

        def _xscrolled(self, first: str, last: str) -> None:
            set_scrollbar(self.hbar, first, last)
            # Las columnas se expanden al ancho disponible. Ese ancho anterior puede quedar obsoleto
            # al aparecer la barra vertical: decide por los mínimos, no por el scrollregion transitorio.
            minimum = sum(round(self._widths[name] * self._scale()) for name in self._display)
            needed = minimum > self.body.winfo_width()
            if needed != bool(self.hbar.winfo_manager()):
                self.hbar.grid() if needed else self.hbar.grid_remove()
            self.header.xview_moveto(self.body.xview()[0])

        def _wheel(self, event: Any) -> str:
            step = -3 if getattr(event, "num", None) == 4 else 3 if getattr(event, "num", None) == 5 else (
                -3 if event.delta > 0 else 3)
            if event.state & 0x1:
                self.body.xview_scroll(step, "units")
            else:
                self.body.yview_scroll(step, "units")
            return "break"

        # Ratón y teclado
        def _row_at(self, y: int) -> int | None:
            index = int(self.body.canvasy(y) // self._row_h())
            return index if 0 <= index < len(self._order) else None

        def _column_at(self, x: int) -> str | None:
            position = self.body.canvasx(x)
            return next((name for name, left, width in zip(self._display, self._col_x, self._col_w)
                         if left <= position < left + width), None)

        def _press(self, event: Any) -> None:
            self.body.focus_set()
            self._pressed = self._row_at(event.y)
            if self._pressed is not None:
                self._select(self._order[self._pressed])

        def _release(self, event: Any) -> None:
            index = self._row_at(event.y)
            if self.open_on_click and index is not None and index == self._pressed:
                self._open(self._order[index])
            self._pressed = None

        def _double(self, event: Any) -> None:
            index = self._row_at(event.y)
            if index is not None and not self.open_on_click:
                self._open(self._order[index])

        def _key(self, event: Any) -> str | None:
            if not self._order:
                return None
            if event.keysym in ("Return", "KP_Enter"):
                if self._selected:
                    self._open(self._selected)
                return "break"
            index = self._order.index(self._selected) if self._selected in self._items else -1
            page = max(1, self.body.winfo_height() // self._row_h() - 1)
            target = {"Up": index - 1, "Down": index + 1, "Home": 0, "End": len(self._order) - 1,
                      "Prior": index - page, "Next": index + page}.get(event.keysym)
            if target is None:
                return None
            iid = self._order[max(0, min(len(self._order) - 1, target))]
            self._select(iid)
            self.see(iid)
            return "break"

        def _motion(self, event: Any) -> None:
            index = self._row_at(event.y)
            if index != self._hover:
                self._hover = index
                self._draw_rows()
            column = self._column_at(event.x)
            cell = (index, column) if index is not None and column else None
            if cell == self._tip_cell:
                return
            self._hide_tip()
            self._tip_cell = cell
            if cell:
                text = self._cell_text(*cell)
                hint = self.cell_tooltip(self._order[index], column)
                width = self._col_w[self._display.index(column)] - 2 * round(self.PAD * self._scale())
                if hint or self._fit(text, width, self._font) != text:
                    self._tip_job = self.after(600, lambda: self._show_tip(hint or text))

        def _leave(self, _event: Any) -> None:
            self._hover = None
            self._tip_cell = None
            self._hide_tip()
            self._draw_rows()

        def _show_tip(self, text: str) -> None:
            self._tip_job = None
            tip = ctk.CTkToplevel(self)
            tip.withdraw()
            tip.overrideredirect(True)
            label = ctk.CTkLabel(tip, text=text, wraplength=480, justify="left", font=FONTS["small"],
                                 fg_color=T["ink"], text_color=T["surface"], corner_radius=6, padx=10, pady=6)
            label.pack()
            tip.geometry(f"+{self.winfo_pointerx() + 14}+{self.winfo_pointery() + 18}")
            tip.deiconify()
            tip.lift()
            self.tooltip, self.tooltip_label = tip, label

        def _hide_tip(self) -> None:
            if self._tip_job:
                self.after_cancel(self._tip_job)
                self._tip_job = None
            if self.tooltip is not None:
                self.tooltip.destroy()
                self.tooltip = self.tooltip_label = None

        def destroy(self) -> None:
            canvas_registry.discard(self)
            if self._pending is not None:
                self.after_cancel(self._pending)
                self._pending = None
            self._hide_tip()
            super().destroy()

    INDICATOR_TERMS = ("Nc.Top", "Nc.A", "Nc.B", "AP", "DPC", "Fr.A", "Fr.B", "Cohe", "Coop")

    def draw_indicator(canvas: Any, result: dict[str, Any] | None) -> None:
        """Una barra por término: el fondo es el coeficiente (máximo posible) y el relleno, el aporte."""
        canvas.delete("all")
        if not result:
            return
        c, scale = palette(), scale_of(canvas)

        def px(value: float) -> int:
            return round(value * scale)
        width = max(px(360), canvas.winfo_width())
        left, right = px(96), width - px(64)
        terms = [*result["clases"], result["cohesion"], result["cooperacion"]]
        unit = (right - left) / max(term["coeficiente"] for term in terms)
        canvas.create_text(0, 0, anchor="nw", width=width, font=chart_font(canvas, 9), fill=c["muted"],
                           text="Aporte al indicador · la barra gris es el peso máximo de cada término")
        for index, (label, term) in enumerate(zip(INDICATOR_TERMS, terms)):
            y = px(36) + index * px(24)
            canvas.create_text(0, y, anchor="w", text=f"{label} · {decimal(term['coeficiente'])}",
                               font=chart_font(canvas, 9), fill=c["ink"])
            canvas.create_rectangle(left, y - px(5), left + term["coeficiente"] * unit, y + px(5),
                                    fill=c["chart_track"], outline="", tags=("indicator:track",))
            if term["aporte"] > 0:
                canvas.create_rectangle(left, y - px(5), left + max(px(2), term["aporte"] * unit), y + px(5),
                                        fill=c["primary"], outline="", tags=("indicator:fill",))
            canvas.create_text(right + px(10), y, anchor="w", text=decimal(term["aporte"], 3),
                               font=chart_font(canvas, 9), fill=c["muted"])

    CLASS_TONES = {"TOP": "class_top", "A": "class_a", "B": "class_b", "ASC": "class_asc", "DPC": "class_dpc",
                   "FRH_A": "class_frh_a", "FRH_B": "class_frh_b"}
    CLASS_LABELS = {"TOP": "Nuevo conocimiento TOP", "A": "Nuevo conocimiento A", "B": "Nuevo conocimiento B",
                    "ASC": "Apropiación social", "DPC": "Divulgación pública", "FRH_A": "Formación RH A",
                    "FRH_B": "Formación RH B"}

    def class_counts(categories: dict[str, int]) -> dict[str, int]:
        """Productos por clase del modelo a partir del conteo por código de categoría."""
        counts: dict[str, int] = {}
        for code, count in categories.items():
            category = catalog_by_code.get(code)
            name = category["clase"] if category else "Sin categoría"
            counts[name] = counts.get(name, 0) + count
        return {name: counts[name] for name in (*MODEL_CLASSES, "Sin categoría") if counts.get(name)}

    def class_color(name: str) -> str:
        return color(CLASS_TONES.get(name, "class_none"))

    def text_on(fill: str) -> str:
        """Tinta legible sobre un relleno: blanca sobre los colores oscuros."""
        red, green, blue = (int(fill[index:index + 2], 16) for index in (1, 3, 5))
        return "#ffffff" if 0.299 * red + 0.587 * green + 0.114 * blue < 150 else LIGHT_THEME["ink"]

    def percent(part: int, total: int) -> str:
        return decimal(100 * part / total, 1) + " %" if total else "0 %"

    def chart_start(canvas: Any, title: str, subtitle: str = "") -> tuple[dict[str, str], Callable[[float], int], int, int]:
        """Limpia el lienzo, escribe título y subtítulo; devuelve colores, escala, ancho y alto."""
        canvas.delete("all")
        c, scale = palette(), scale_of(canvas)

        def px(value: float) -> int:
            return round(value * scale)
        canvas.create_text(px(12), px(14), anchor="w", text=title, font=chart_font(canvas, 10, "strong"), fill=c["ink"])
        if subtitle:
            canvas.create_text(px(12), px(34), anchor="w", text=subtitle, font=chart_font(canvas, 9), fill=c["muted"])
        return c, px, max(canvas.winfo_width(), px(320)), max(canvas.winfo_height(), px(200))

    def chart_empty(canvas: Any, text: str, width: int, height: int) -> None:
        scale = scale_of(canvas)
        canvas.create_text(round(12 * scale), round(80 * scale), anchor="w", text=text, font=chart_font(canvas, 10),
                           fill=palette()["muted"])
        canvas.configure(scrollregion=(0, 0, width, height))

    def draw_donut(canvas: Any, title: str, subtitle: str, items: list[tuple[str, str, int, str]], caption: str,
                   empty: str = "Sin productos para estos filtros") -> None:
        """Anillo: una porción por elemento (clave, rótulo, cantidad, color) y su leyenda a la derecha."""
        c, px, width, height = chart_start(canvas, title, subtitle)
        items = [item for item in items if item[2] > 0]
        total = sum(item[2] for item in items)
        if not total:
            chart_empty(canvas, empty, width, height)
            return
        top = px(54)
        size = max(px(120), min(height - top - px(14), round(width * 0.42), px(250)))
        # En lienzos anchos, el anillo y su leyenda (hasta 440 px) quedan centrados.
        x0 = max(px(18), (width - size - px(30) - px(440)) // 2)
        y0 = top + max(0, (height - top - px(14) - size) // 2)
        start = 90.0
        for key, _label, count, fill in items:
            tags = ("donut:slice", f"slice:{key}", f"count:{count}")
            if len(items) == 1:
                canvas.create_oval(x0, y0, x0 + size, y0 + size, fill=fill, outline="", tags=tags)
                break
            extent = -360 * count / total
            canvas.create_arc(x0, y0, x0 + size, y0 + size, start=start, extent=extent, style="pieslice", fill=fill,
                              outline=c["surface"], width=max(1, px(2)), tags=tags)
            start += extent
        inset = size * 0.2
        canvas.create_oval(x0 + inset, y0 + inset, x0 + size - inset, y0 + size - inset, fill=c["surface"], outline="")
        middle_x, middle_y = x0 + size / 2, y0 + size / 2
        canvas.create_text(middle_x, middle_y - px(8), text=number(total), font=chart_font(canvas, 16, "strong"),
                           fill=c["ink"])
        canvas.create_text(middle_x, middle_y + px(14), text=caption, font=chart_font(canvas, 8), fill=c["muted"])
        left = x0 + size + px(30)
        right = min(width - px(14), left + px(440))
        # Con poco espacio la leyenda muestra solo la clave (por ejemplo «DPC»).
        full = right - left >= px(300)
        row = max(px(18), min(px(28), (height - top - px(8)) // len(items)))
        y = y0 + max(row / 2, (size - row * (len(items) - 1)) / 2)
        for key, label, count, fill in items:
            canvas.create_oval(left, y - px(5), left + px(10), y + px(5), fill=fill, outline="")
            canvas.create_text(left + px(18), y, anchor="w", text=label if full else key, font=chart_font(canvas, 9),
                               fill=c["ink"])
            canvas.create_text(right, y, anchor="e", text=f"{number(count)} · {percent(count, total)}",
                               font=chart_font(canvas, 8), fill=c["muted"])
            y += row
        canvas.configure(scrollregion=(0, 0, width, max(height, round(y - row / 2 + px(8)))))

    def squarify(values: list[float], x: float, y: float, width: float, height: float) -> list[tuple[float, ...]]:
        """Mosaico «squarified» (Bruls et al.): rectángulos casi cuadrados con área proporcional al valor."""
        total = sum(values)
        areas = [value * width * height / total for value in values]
        rectangles: list[tuple[float, ...]] = []

        def worst(row: list[float], side: float) -> float:
            size = sum(row)
            return max(max(side * side * area / (size * size), size * size / (side * side * area)) for area in row)
        while areas:
            side = max(min(width, height), 1e-9)
            row = [areas[0]]
            while len(row) < len(areas) and worst(row + [areas[len(row)]], side) <= worst(row, side):
                row.append(areas[len(row)])
            size = sum(row)
            if width >= height:
                column = size / max(height, 1e-9)
                top = y
                for area in row:
                    rectangles.append((x, top, column, area / max(column, 1e-9)))
                    top += area / max(column, 1e-9)
                x, width = x + column, width - column
            else:
                band = size / max(width, 1e-9)
                left = x
                for area in row:
                    rectangles.append((left, y, area / max(band, 1e-9), band))
                    left += area / max(band, 1e-9)
                y, height = y + band, height - band
            areas = areas[len(row):]
        return rectangles

    def draw_treemap(canvas: Any, title: str, subtitle: str, items: list[tuple[str, int, str]], limit: int = 30,
                     empty: str = "Sin productos para estos filtros") -> None:
        """Mosaico: un rectángulo por elemento (rótulo, cantidad, color) con área proporcional a la cantidad."""
        c, px, width, height = chart_start(canvas, title, subtitle)
        items = sorted((item for item in items if item[1] > 0), key=lambda item: (-item[1], item[0]))
        if not items:
            chart_empty(canvas, empty, width, height)
            return
        if len(items) > limit:
            rest = items[limit - 1:]
            items = items[:limit - 1] + [(f"Otras {len(rest)}", sum(item[1] for item in rest), c["class_none"])]
        total = sum(item[1] for item in items)
        legend = [name for name in (*MODEL_CLASSES, "Sin categoría")
                  if any(fill == class_color(name) for _label, _count, fill in items)]
        x = px(12)
        for name in legend:
            canvas.create_rectangle(x, px(52), x + px(10), px(62), fill=class_color(name), outline="")
            label = canvas.create_text(x + px(14), px(57), anchor="w", text=name, font=chart_font(canvas, 8),
                                       fill=c["muted"])
            x = canvas.bbox(label)[2] + px(14)
        top = px(72) if legend else px(52)
        cells = squarify([float(item[1]) for item in items], px(12), top, width - px(24), height - top - px(12))
        for (label, count, fill), (left, cell_top, cell_width, cell_height) in zip(items, cells):
            canvas.create_rectangle(left, cell_top, left + cell_width, cell_top + cell_height, fill=fill,
                                    outline=c["surface"], width=max(1, px(2)),
                                    tags=("treemap:cell", f"cell:{label}", f"count:{count}"))
            # Cada rótulo aparece solo si cabe entero en su celda.
            ink, bottom = text_on(fill), cell_top + px(6)
            for text, font in ((label, chart_font(canvas, 9, "strong")),
                               (f"{number(count)} · {percent(count, total)}", chart_font(canvas, 8))):
                item = canvas.create_text(left + px(7), bottom, anchor="nw", text=text, font=font, fill=ink)
                box = canvas.bbox(item)
                if box[2] > left + cell_width - px(3) or box[3] > cell_top + cell_height - px(3):
                    canvas.delete(item)
                    break
                bottom = box[3] + px(1)
        canvas.configure(scrollregion=(0, 0, width, height))

    def draw_trend(canvas: Any, values: dict[str, int], bounds: tuple[int | None, int | None] = (None, None)) -> None:
        """Línea de productos por año (área bajo la curva) y su media móvil de tres años."""
        numeric = {int(key): value for key, value in values.items() if key.isdigit()}
        missing = values.get("Sin dato", 0)
        first = bounds[0] if bounds[0] is not None else min(numeric, default=0)
        last = bounds[1] if bounds[1] is not None else max(numeric, default=0)
        cutoff = date.today().year - 20
        older = sum(count for year, count in numeric.items() if first <= year <= min(cutoff, last))
        years = list(range(max(first, cutoff + 1), last + 1))
        notes = [f"{number(sum(values.values()))} productos", f"{number(missing)} sin año"]
        if older:
            notes.append(f"{number(older)} de {cutoff} o antes (fuera de la línea)")
        c, px, width, height = chart_start(canvas, "Tendencia anual · productos por año", " · ".join(notes))
        if len(years) < 2:
            chart_empty(canvas, "Se necesitan al menos dos años para trazar la tendencia", width, height)
            return
        counts = [numeric.get(year, 0) for year in years]
        average = [sum(counts[max(0, index - 2):index + 1]) / len(counts[max(0, index - 2):index + 1])
                   for index in range(len(counts))]
        left, right, top, bottom = px(48), width - px(18), px(88), height - px(36)
        maximum = max(1, max(counts))
        for tick in sorted({round(maximum * step / 4) for step in range(5)}):
            y = bottom - (bottom - top) * tick / maximum
            canvas.create_line(left, y, right, y, fill=c["chart_grid"] if tick else c["line_strong"])
            canvas.create_text(left - px(8), y, anchor="e", text=number(tick), font=chart_font(canvas, 8), fill=c["muted"])
        step = (right - left) / (len(years) - 1)
        points = [(left + index * step, bottom - (bottom - top) * count / maximum) for index, count in enumerate(counts)]
        canvas.create_polygon(left, bottom, *[value for point in points for value in point], right, bottom,
                              fill=appearance(soft("primary")), outline="", tags=("trend:area",))
        canvas.create_line(*[value for point in points for value in point], fill=c["primary"], width=max(2, px(2.5)),
                           smooth=False, tags=("trend:line",))
        canvas.create_line(*[value for index, mean in enumerate(average)
                             for value in (left + index * step, bottom - (bottom - top) * mean / maximum)],
                           fill=c["pending"], width=max(1, px(2)), dash=(6, 4), tags=("trend:average",))
        label_every = max(1, math.ceil(px(40) / step))
        for index, ((x, y), year, count) in enumerate(zip(points, years, counts)):
            radius = px(3.5)
            canvas.create_oval(x - radius, y - radius, x + radius, y + radius, fill=c["surface"], outline=c["primary"],
                               width=max(1, px(2)), tags=("trend:point", f"year:{year}", f"count:{count}"))
            if index % label_every == 0 or index == len(years) - 1:
                canvas.create_text(x, bottom + px(16), text=str(year), font=chart_font(canvas, 8), fill=c["muted"])
        peak = max(range(len(counts)), key=counts.__getitem__)
        if counts[peak]:
            canvas.create_text(points[peak][0], points[peak][1] - px(12), text=f"máx. {number(counts[peak])}",
                               font=chart_font(canvas, 8, "strong"), fill=c["ink"])
        legend_x = width - px(14)
        for text, fill, dash in (("Media móvil 3 años", c["pending"], (6, 4)), ("Productos del año", c["primary"], None)):
            label = canvas.create_text(legend_x, px(56), anchor="e", text=text, font=chart_font(canvas, 8), fill=c["muted"])
            box = canvas.bbox(label)
            canvas.create_line(box[0] - px(26), px(56), box[0] - px(6), px(56), fill=fill, width=max(2, px(2)), dash=dash)
            legend_x = box[0] - px(40)
        canvas.configure(scrollregion=(0, 0, width, height))

    def draw_gauge(canvas: Any, value: float | None, left: int, top: int, size: int) -> None:
        """Semicírculo de 0 a 10 con el valor del indicador en el centro."""
        c, scale = palette(), scale_of(canvas)
        ring = max(4, round(size * 0.12))
        box = (left + ring / 2, top + ring / 2, left + size - ring / 2, top + size - ring / 2)
        canvas.create_arc(*box, start=180, extent=-180, style="arc", outline=c["chart_track"], width=ring,
                          tags=("gauge:track",))
        if value is None:
            return
        canvas.create_arc(*box, start=180, extent=-180 * min(max(value, 0), 10) / 10, style="arc", outline=c["primary"],
                          width=ring, tags=("gauge:value",))
        middle, points = left + size / 2, max(12, round(size / scale / 9))
        canvas.create_text(middle, top + size * 0.34, text=decimal(value, 2), font=chart_font(canvas, points, "strong"),
                           fill=c["ink"])
        canvas.create_text(middle, top + size * 0.34 + round(points * 1.6 * scale), text="de 10",
                           font=chart_font(canvas, 8), fill=c["muted"])
        for mark, anchor in ((0, "n"), (10, "n")):
            canvas.create_text(left + ring / 2 if mark == 0 else left + size - ring / 2, top + size / 2 + round(6 * scale),
                               anchor=anchor, text=str(mark), font=chart_font(canvas, 8), fill=c["muted"])

    def draw_radar(canvas: Any, result: dict[str, Any], middle_x: float, middle_y: float, radius: float) -> None:
        """Radar de los nueve términos: índice relativo (aporte / coeficiente) de 0 a 1 en cada eje."""
        c = palette()
        terms = [*result["clases"], result["cohesion"], result["cooperacion"]]
        angles = [math.pi / 2 - 2 * math.pi * index / len(terms) for index in range(len(terms))]

        def point(angle: float, fraction: float) -> tuple[float, float]:
            return middle_x + radius * fraction * math.cos(angle), middle_y - radius * fraction * math.sin(angle)
        for fraction in (0.25, 0.5, 0.75, 1):
            canvas.create_polygon(*[value for angle in angles for value in point(angle, fraction)],
                                  fill="", outline=c["chart_grid"] if fraction < 1 else c["line_strong"])
        for angle, label in zip(angles, INDICATOR_TERMS):
            canvas.create_line(middle_x, middle_y, *point(angle, 1), fill=c["chart_grid"], tags=("radar:axis",))
            x, y = point(angle, 1.16)
            canvas.create_text(x, y, text=label, font=chart_font(canvas, 8), fill=c["muted"],
                               anchor="center" if abs(math.cos(angle)) < 0.3 else "w" if math.cos(angle) > 0 else "e")
        values = [min(1.0, term["aporte"] / term["coeficiente"]) if term["coeficiente"] else 0.0 for term in terms]
        shape = [value for angle, fraction in zip(angles, values) for value in point(angle, fraction)]
        canvas.create_polygon(*shape, fill=appearance(soft("primary")), outline=c["primary"],
                              width=max(1, round(2 * scale_of(canvas))), tags=("radar:value",))
        for index in range(0, len(shape), 2):
            canvas.create_oval(shape[index] - 3, shape[index + 1] - 3, shape[index] + 3, shape[index + 1] + 3,
                               fill=c["primary"], outline="")

    def draw_indicator_profile(canvas: Any, result: dict[str, Any] | None) -> None:
        """Valor del indicador en un semicírculo y perfil de sus términos en un radar."""
        c, px, width, height = chart_start(canvas, "Perfil del indicador",
                                           "Semicírculo: valor de 0 a 10 · radar: índice relativo de cada término (0 a 1)")
        if result is None:
            chart_empty(canvas, "Seleccione un grupo o investigador activo en los filtros del Dashboard para ver su indicador.",
                        width, height)
            return
        size = min(px(240), max(px(150), round(width * 0.3)))
        draw_gauge(canvas, result["indicador"], px(24), px(64), size)
        canvas.create_text(px(24) + size / 2, px(64) + size / 2 + px(30), text=f"Corte {result['corte']}",
                           font=chart_font(canvas, 9), fill=c["muted"])
        radius = max(px(60), min((height - px(110)) / 2, (width - size - px(140)) / 2.6))
        draw_radar(canvas, result, px(24) + size + (width - size - px(24)) / 2, px(56) + (height - px(56)) / 2, radius)
        canvas.configure(scrollregion=(0, 0, width, height))

    class IndicatorPanel(Card):
        """Indicador de producción: valor sobre 10, contexto en chips y una barra por término de la fórmula."""
        def __init__(self, parent: Any, on_details: Callable[[], Any], pack_options: dict[str, Any] | None = None,
                     **options: Any):
            super().__init__(parent, **options)
            self._pack_options = pack_options or {"fill": "x"}
            self.result: dict[str, Any] | None = None
            self.value = ctk.StringVar(master=self, value="")
            self.note = ctk.StringVar(master=self, value="")
            self.body.grid_columnconfigure(1, weight=1)
            summary = ctk.CTkFrame(self.body, fg_color="transparent")
            summary.grid(row=0, column=0, sticky="nw", padx=(0, 28))
            ctk.CTkLabel(summary, text="INDICADOR DE PRODUCCIÓN · MINCIENCIAS 2024", font=FONTS["caption"],
                         text_color=T["muted"], anchor="w").pack(anchor="w")
            # El valor (sobre 10) se dibuja en el semicírculo; value conserva el texto para quien lo consulte.
            self.gauge = ChartCanvas(summary, draw=self._draw_gauge, width=round(220 * scale_of(parent)),
                                     height=round(136 * scale_of(parent)))
            self.gauge.pack(anchor="w", pady=(6, 4))
            self.entity = ctk.CTkLabel(summary, text="", font=FONTS["h2"], text_color=T["ink"], anchor="w",
                                       justify="left", wraplength=300)
            self.entity.pack(anchor="w", pady=(0, 10))
            chips = ctk.CTkFrame(summary, fg_color="transparent")
            chips.pack(anchor="w")
            self.chips = [Chip(chips, "", "muted") for _ in range(4)]
            for index, chip in enumerate(self.chips):
                chip.grid(row=index // 2, column=index % 2, sticky="w", padx=(0, 6), pady=(0, 6))
            ctk.CTkLabel(summary, textvariable=self.note, font=FONTS["caption"], text_color=T["muted"], anchor="w",
                         justify="left", wraplength=300).pack(anchor="w", pady=(2, 12))
            self.details_button = button(summary, "Detalle por subtipo", on_details)
            self.details_button.pack(anchor="w")
            self.canvas = ChartCanvas(self.body, draw=lambda canvas: draw_indicator(canvas, self.result),
                                      height=round(250 * scale_of(parent)), width=360)
            self.canvas.grid(row=0, column=1, sticky="nsew")

        def _draw_gauge(self, canvas: Any) -> None:
            canvas.delete("all")
            scale = scale_of(canvas)
            draw_gauge(canvas, self.result["indicador"] if self.result else None, round(4 * scale), round(4 * scale),
                       round(212 * scale))

        def show(self, result: dict[str, Any] | None, entity_label: str) -> None:
            self.result = result
            if result is None:
                self.value.set("")
                self.note.set("")
                self.entity.configure(text="")
                self.pack_forget()
                return
            population = "grupos activos" if result["tipo"] == "grupos" else "investigadores activos"
            self.value.set(decimal(result["indicador"], 2))
            self.entity.configure(text=entity_label)
            for chip, (text, tone) in zip(self.chips, (
                    (f"Corte {result['corte']}", "primary"),
                    (f"{number(result['productos'])} contados", "primary"),
                    (f"{number(result['estimadas'])} estimadas", "pending" if result["estimadas"] else "muted"),
                    (f"{number(result['sin_categoria'])} sin categoría",
                     "unknown" if result["sin_categoria"] else "muted"))):
                chip.set(text, tone)
            self.note.set(f"Índices relativos al máximo entre {number(result['poblacion'])} {population} cargados; "
                          "no es una clasificación oficial de Minciencias.")
            if not self.winfo_manager():
                self.pack(**self._pack_options)
            self.canvas.redraw()
            self.gauge.redraw()

    dialog_state: dict[str, Any] = {"last": None}

    class Dialog(ctk.CTkToplevel):
        """Ventana modal centrada sobre su origen; Escape la cierra y devuelve el foco modal anterior."""
        def __init__(self, parent: Any, title: str, width: int = 460, height: int | None = None):
            super().__init__(parent)
            self.withdraw()
            self.title(title)
            self.configure(fg_color=T["page"])
            try:
                # X11: los gestores de ventanas en mosaico dejan flotar los diálogos con su tamaño pedido.
                self.wm_attributes("-type", "dialog")
            except TclError:
                pass
            self._origin, self._size, self._previous_grab = parent, (width, height), None
            self.buttons: dict[str, Any] = {}
            self.body = ctk.CTkFrame(self, fg_color="transparent")
            self.body.pack(fill="both", expand=True, padx=22, pady=(20, 10))
            self.footer = ctk.CTkFrame(self, fg_color="transparent")
            self.footer.pack(fill="x", padx=22, pady=(0, 18))
            self.bind("<Escape>", lambda _event: self.close())
            self.protocol("WM_DELETE_WINDOW", self.close)

        def present(self, wait: bool = True, focus: Any = None) -> None:
            dialog_state["last"] = self
            self._previous_grab = self.grab_current()
            host = self._origin.winfo_toplevel()
            # Un diálogo transitorio de una ventana oculta también quedaría oculto.
            if host.winfo_viewable():
                self.transient(host)
            self._place(host)
            self.deiconify()
            self.lift()
            try:
                self.wait_visibility()
                self.grab_set()
            except TclError:
                if self.winfo_exists():
                    self.after(20, self.grab_set)
            # Mientras se espera la ventana, un evento pendiente puede haberla cerrado ya.
            if not self.winfo_exists():
                return
            (focus or self).focus_force()
            if wait:
                self.wait_window()

        def _place(self, host: Any) -> None:
            """Tamaño del contenido (o el fijado) y posición centrada sobre la ventana de origen."""
            self.update_idletasks()
            width = max(self._size[0], self._reverse_window_scaling(self.winfo_reqwidth()))
            height = self._size[1] or self._reverse_window_scaling(self.winfo_reqheight())
            width = round(min(width, self._reverse_window_scaling(host.winfo_screenwidth() - 40)))
            height = round(min(height, self._reverse_window_scaling(host.winfo_screenheight() - 80)))
            pixels_w, pixels_h = self._apply_window_scaling(width), self._apply_window_scaling(height)
            x = max(0, min(host.winfo_rootx() + (host.winfo_width() - pixels_w) // 2, host.winfo_screenwidth() - pixels_w))
            y = max(0, min(host.winfo_rooty() + (host.winfo_height() - pixels_h) // 2, host.winfo_screenheight() - pixels_h))
            self.geometry(f"{width}x{height}+{x}+{y}")

        def _cancel_timers(self) -> None:
            """CTkTextbox repite un temporizador propio; al cerrar se cancela para no llamar a un widget destruido."""
            cancel_widget_timers(self)

        def close(self) -> None:
            previous = self._previous_grab
            if self.winfo_exists():
                self._cancel_timers()
                self.destroy()
            if previous is not None and previous.winfo_exists():
                try:
                    previous.grab_set()
                except TclError:
                    pass

    def show_columns(table: DataTable, columns: tuple[str, ...], usage: dict[str, int],
                     keep: tuple[str, ...] = ()) -> None:
        """Oculta las columnas que ningún registro ha llenado; reaparecen al tener datos."""
        shown = [name for name in columns
                 if name in keep or name.startswith("_") or usage.get(name, 1) > 0]
        if tuple(table.cget("displaycolumns")) != tuple(shown):
            table.configure(displaycolumns=shown)

    class ChoiceDialog(Dialog):
        def __init__(self, parent: tk.Misc, title: str, prompt: str,
                     options: dict[str, str], default: str = ""):
            super().__init__(parent, title, 530)
            self.result: str | None = None
            self.options = options
            ctk.CTkLabel(self.body, text=prompt, wraplength=480, justify="left", anchor="w",
                         font=FONTS["body"], text_color=T["ink"]).pack(fill="x", pady=(0, 12))
            self.value = tk.StringVar(value=next((label for label, value in options.items() if value == default), next(iter(options))))
            self.box = ctk.CTkOptionMenu(self.body, variable=self.value, values=list(options), height=34,
                                         font=FONTS["body"], dropdown_font=FONTS["body"], fg_color=T["field"],
                                         button_color=T["field"], button_hover_color=T["hover"], text_color=T["ink"],
                                         dropdown_fg_color=T["surface"], dropdown_text_color=T["ink"],
                                         dropdown_hover_color=T["hover"])
            self.box.pack(fill="x")
            button(self.footer, "Continuar", self.accept, variant="primary", width=112).pack(side="right")
            button(self.footer, "Cancelar", self.close, width=112).pack(side="right", padx=(0, 8))
            self.bind("<Return>", lambda _event: self.accept())
            self.present()

        def accept(self) -> None:
            self.result = self.options[self.value.get()]
            self.close()

    # Fichas públicas de ejemplo del enunciado del Parcial 2.
    EXAMPLE_URLS = {
        "Ejemplo GrupLAC": "https://scienti.minciencias.gov.co/gruplac/jsp/visualiza/visualizagr.jsp?nro=00000000002099",
        "Ejemplo CvLAC": "https://scienti.minciencias.gov.co/cvlac/visualizador/generarCurriculoCv.do?cod_rh=0000494917",
    }

    class UrlDialog(Dialog):
        """Consulta de una URL pública: la descarga corre en segundo plano con el diálogo abierto.

        result es la URL consultada y preview su vista previa; ambos quedan en None si se cancela.
        app se conserva para los consumidores de kit; la consulta no depende de la aplicación.
        """
        def __init__(self, parent: Any, app: "App | None" = None, initial: str = "", start: bool = False):
            super().__init__(parent, "Consultar URL pública", 600)
            self.result: str | None = None
            self.preview: WebPreview | None = None
            self._results: queue.Queue[WebPreview | Exception] | None = None
            ctk.CTkLabel(self.body, text="Pegue una URL pública HTTP o HTTPS: una ficha GrupLAC o CvLAC, o la página "
                         "de un artículo.", font=FONTS["body"], text_color=T["ink"], anchor="w", justify="left",
                         wraplength=550).pack(fill="x", pady=(0, 8))
            self.entry = ctk.CTkEntry(self.body, height=36, corner_radius=8, font=FONTS["body"], fg_color=T["field"],
                                      border_color=T["line_strong"], text_color=T["ink"], placeholder_text="https://…")
            if initial:
                self.entry.insert(0, initial)
            self.entry.pack(fill="x")
            examples = ctk.CTkFrame(self.body, fg_color="transparent")
            examples.pack(fill="x", pady=(8, 0))
            self.examples = {}
            for label, url in EXAMPLE_URLS.items():
                control = ctk.CTkButton(examples, text=label, command=lambda value=url: self._use(value), width=0,
                                        height=24, corner_radius=12, font=FONTS["caption"], fg_color=soft("primary"),
                                        hover_color=T["primary_soft"], text_color=T["primary"])
                control.pack(side="left", padx=(0, 6))
                self.examples[label] = control
            self.progress = ctk.CTkProgressBar(self.body, mode="indeterminate", height=6, fg_color=T["chart_track"],
                                               progress_color=T["primary"])
            self.error = tk.StringVar()
            self.error_label = ctk.CTkLabel(self.body, textvariable=self.error, font=FONTS["small"],
                                            text_color=T["rejected"], anchor="w", justify="left", wraplength=550)
            self.error_label.pack(fill="x", pady=(8, 0))
            self.consult = button(self.footer, "Consultar", self.accept, variant="primary", width=120)
            self.consult.pack(side="right")
            button(self.footer, "Cancelar", self.close, width=112).pack(side="right", padx=(0, 8))
            self.bind("<Return>", lambda _event: self.accept())
            if start and initial:
                self.after(50, self.accept)
            self.present(focus=self.entry)

        def _use(self, url: str) -> None:
            self.entry.delete(0, "end")
            self.entry.insert(0, url)
            self.error.set("")

        def _busy(self, busy: bool) -> None:
            for control in (self.consult, self.entry, *self.examples.values()):
                control.configure(state="disabled" if busy else "normal")
            if busy:
                self.progress.pack(fill="x", pady=(12, 0), before=self.error_label)
                self.progress.start()
            else:
                self.progress.stop()
                self.progress.pack_forget()

        def accept(self) -> None:
            if self._results is not None:
                return
            url = self.entry.get().strip()
            if not url:
                self.error.set("Escriba o pegue una URL.")
                return
            self.error.set("")
            self._busy(True)
            results = self._results = queue.Queue(maxsize=1)

            def fetch() -> None:
                try:
                    results.put(fetch_web_page(url))
                except Exception as exc:
                    results.put(exc)
            threading.Thread(target=fetch, daemon=True).start()
            self.after(100, lambda: self._check(url))

        def _check(self, url: str) -> None:
            try:
                result = self._results.get_nowait()  # type: ignore[union-attr]
            except queue.Empty:
                self.after(100, lambda: self._check(url))
                return
            self._results = None
            self._busy(False)
            if isinstance(result, Exception):
                self.error.set(str(result))
                return
            self.result, self.preview = url, result
            self.close()

    class RecordPicker(Dialog):
        """Consulta paginada: el límite por página no limita las opciones disponibles."""
        def __init__(self, parent: tk.Misc, repo: Repository | CppRepository,
                     kind: str, current: str = ""):
            super().__init__(parent, "Seleccionar " + TITLES[kind].lower(), 760, 560)
            self.minsize(540, 400)
            self.repo, self.kind = repo, kind
            self.result: str | None = None
            self.offset = 0
            self.search = tk.StringVar()
            self.search.trace_add("write", self._search_changed)
            ctk.CTkLabel(self.body, text="Buscar nombre o ID", font=FONTS["h2"],
                         text_color=T["ink"], anchor="w").pack(fill="x", pady=(0, 8))
            self.search_entry = SearchField(self.body, self.search, self.refresh, placeholder_text="Buscar…")
            self.search_entry.pack(fill="x", pady=(0, 12))
            self.table = DataTable(self.body, ("nombre", "id", "activo"),
                                   headings={"nombre": "Nombre / título", "id": "ID", "activo": "Estado"},
                                   widths={"nombre": 330, "id": 180, "activo": 100}, height_rows=9,
                                   chips={"activo": lambda value: "primary" if value == "Activo" else "muted"})
            self.table.pack(fill="both", expand=True)
            self.pager = Pager(self.body, self.move)
            self.pager.pack(fill="x", pady=(10, 0))
            self.previous, self.next = self.pager.previous, self.pager.next
            self.count = tk.StringVar()
            self.choose = button(self.footer, "Seleccionar", self.accept, variant="primary", width=120)
            self.choose.pack(side="right")
            button(self.footer, "Cancelar", self.close, width=112).pack(side="right", padx=(0, 8))
            self.table.on_select(lambda _iid: self.choose.configure(state="normal" if self.table.selection() else "disabled"))
            self.table.on_open(lambda _iid: self.accept())
            self.refresh()
            if current and self.table.exists(current):
                self.table.selection_set(current)
            self.present(focus=self.search_entry)

        def _search_changed(self, *_: Any) -> None:
            self.offset = 0

        def refresh(self) -> None:
            self.search_entry.cancel_pending()
            page = self.repo.page(self.kind, self.offset, PAGE_SIZE, self.search.get())
            if self.offset and self.offset >= page["total"]:
                self.offset = max(0, (page["total"] - 1) // PAGE_SIZE * PAGE_SIZE)
                page = self.repo.page(self.kind, self.offset, PAGE_SIZE, self.search.get())
            self.table.delete(*self.table.get_children())
            for row in page["rows"]:
                self.table.insert("", "end", iid=row["id"],
                                  values=(record_name(row), row["id"], display_value("activo", row["activo"])))
            self.pager.show(self.offset, len(page["rows"]), page["total"])
            self.count.set(self.pager.label.cget("text"))
            self.choose.configure(state="disabled")

        def move(self, direction: int) -> None:
            self.offset = max(0, self.offset + direction * PAGE_SIZE)
            self.refresh()

        def accept(self) -> None:
            selected = self.table.selection()
            if selected:
                self.result = selected[0]
                self.close()


    class ReferenceInput(ctk.CTkFrame):
        def __init__(self, parent: tk.Misc, repo: Repository | CppRepository, kind: str, value: str = "",
                     variable: tk.StringVar | None = None, command: Callable[[], Any] | None = None):
            super().__init__(parent, fg_color="transparent")
            self.repo, self.kind, self.command = repo, kind, command
            self.value = variable if variable is not None else tk.StringVar()
            self.label = tk.StringVar()
            self._trace = self.value.trace_add("write", lambda *_: self._update_label())
            self.columnconfigure(0, weight=1)
            self.entry = ctk.CTkEntry(self, textvariable=self.label, state="readonly", height=34,
                                      font=FONTS["body"], fg_color=T["field"], text_color=T["ink"],
                                      border_color=T["line_strong"], corner_radius=8)
            self.entry.grid(row=0, column=0, sticky="ew")
            self.button = button(self, "Elegir…", self.select, width=88)
            self.button.grid(row=0, column=1, padx=(8, 0))
            self.set(value)

        def get(self) -> str:
            return self.value.get()

        def set(self, value: str) -> None:
            self.value.set(value)

        def _update_label(self) -> None:
            value = self.value.get()
            row = self.repo.get(self.kind, value) if value else None
            self.label.set(f"{record_name(row)} · {value}" if row else value)

        def select(self) -> None:
            picker = RecordPicker(self, self.repo, self.kind, self.get())
            if picker.result:
                self.set(picker.result)
                if self.command:
                    self.command()

        def set_enabled(self, enabled: bool) -> None:
            self.button.configure(state="normal" if enabled else "disabled")
            self.entry.configure(state="readonly" if enabled else "disabled")

        def destroy(self) -> None:
            self.value.trace_remove("write", self._trace)
            super().destroy()

    class CategoryCatalog(ctk.CTkFrame):
        """Catálogo de referencia compartido por el selector y la página del modelo."""
        def __init__(self, parent: Any):
            super().__init__(parent, fg_color="transparent")
            self.search = tk.StringVar()
            self.class_filter = tk.StringVar(value="Todas")
            self.search_entry = SearchField(self, self.search, self.refresh, placeholder_text="Buscar código o descripción…")
            self.search_entry.pack(fill="x", pady=(0, 8))
            self.class_control = segmented(self, self.class_filter, ["Todas", *MODEL_CLASSES], lambda _value: self.refresh())
            self.class_control.pack(fill="x", pady=(0, 12))
            columns = ("codigo", "descripcion", "clase", "peso_relativo", "peso_global", "ventana")
            self.table = DataTable(self, columns,
                                   headings={"codigo": "Código", "descripcion": "Descripción", "clase": "Clase",
                                             "peso_relativo": "Peso relativo", "peso_global": "Puntos", "ventana": "Ventana (años)"},
                                   widths={"codigo": 120, "descripcion": 300, "clase": 85,
                                           "peso_relativo": 110, "peso_global": 100, "ventana": 120},
                                   anchors={field: "e" for field in columns[3:]}, chips={"clase": lambda _value: "primary"},
                                   height_rows=10)
            self.table.empty_text = "Sin categorías para estos filtros."
            self.table.pack(fill="both", expand=True)
            self.count = ctk.CTkLabel(self, text="", font=FONTS["caption"], text_color=T["muted"], anchor="w")
            self.count.pack(side="bottom", fill="x", pady=(6, 0), before=self.table)
            self.refresh()

        def refresh(self) -> None:
            self.search_entry.cancel_pending()
            current = self.table.selection()
            needle = self.search.get().strip().casefold()
            rows = [row for row in catalog if (self.class_filter.get() == "Todas" or row["clase"] == self.class_filter.get())
                    and (not needle or needle in f"{row['codigo']} {row['descripcion']}".casefold())]
            self.table.delete(*self.table.get_children())
            for row in rows:
                self.table.insert("", "end", iid=row["codigo"], values=(
                    row["codigo"], row["descripcion"], row["clase"], decimal(row["peso_relativo"]),
                    decimal(row["peso_global"]), row["ventana"]))
            if current and self.table.exists(current[0]):
                self.table.selection_set(current[0])
            self.count.configure(text=f"{number(len(rows))} de {number(len(catalog))} categorías")

    class CategoryPicker(Dialog):
        """Selector autónomo; app se conserva por compatibilidad con los consumidores de kit."""
        def __init__(self, parent: Any, app: "App | None" = None, current: str = ""):
            super().__init__(parent, "Seleccionar categoría · Modelo 2024", 880, 620)
            self.minsize(640, 440)
            self.result: str | None = None
            self.catalog_view = CategoryCatalog(self.body)
            self.catalog_view.pack(fill="both", expand=True)
            self.search = self.catalog_view.search
            self.class_filter = self.catalog_view.class_filter
            self.class_control = self.catalog_view.class_control
            self.table = self.catalog_view.table
            self.choose = button(self.footer, "Seleccionar", self.accept, variant="primary", width=120, state="disabled")
            self.choose.pack(side="right")
            button(self.footer, "Cancelar", self.close, width=112).pack(side="right", padx=(0, 8))
            button(self.footer, "Sin categoría", self.clear, width=120).pack(side="left")
            self.table.on_select(lambda _iid: self._selection_changed())
            self.table.on_open(lambda _iid: self.accept())
            if current and self.table.exists(current):
                self.table.selection_set(current)
            self.present(focus=self.catalog_view.search_entry)

        def _selection_changed(self) -> None:
            self.choose.configure(state="normal" if self.table.selection() else "disabled")

        def refresh(self) -> None:
            self.catalog_view.refresh()
            self._selection_changed()

        def accept(self) -> None:
            selected = self.table.selection()
            if selected:
                self.result = selected[0]
                self.close()

        def clear(self) -> None:
            self.result = ""
            self.close()

    class CategoryInput(ctk.CTkFrame):
        def __init__(self, parent: Any, app: "App | None" = None, value: str = ""):
            super().__init__(parent, fg_color="transparent")
            self.app = app
            self.value = tk.StringVar()
            self.label = tk.StringVar()
            self.info = tk.StringVar()
            self.grid_columnconfigure(0, weight=1)
            self.entry = ctk.CTkEntry(self, textvariable=self.label, state="readonly", height=34,
                                      font=FONTS["body"], fg_color=T["field"], text_color=T["ink"],
                                      border_color=T["line_strong"], corner_radius=8)
            self.entry.grid(row=0, column=0, sticky="ew")
            self.button = button(self, "Elegir…", self.select, width=88)
            self.button.grid(row=0, column=1, padx=(8, 0))
            self.info_label = ctk.CTkLabel(self, textvariable=self.info, font=FONTS["caption"], text_color=T["muted"],
                                          anchor="w", justify="left", wraplength=440)
            self.info_label.grid(row=1, column=0, columnspan=2, sticky="ew", pady=(4, 0))
            self.bind("<Configure>", lambda event: self.info_label.configure(wraplength=max(200, event.width / scale_of(self) - 8)), add="+")
            self.set(value)

        def get(self) -> str:
            return self.value.get()

        def set(self, value: str) -> None:
            code = value.partition(" — ")[0].strip().upper()
            self.value.set(code)
            category = catalog_by_code.get(code)
            self.label.set(f"{code} — {category['descripcion']}" if category else code)
            self.info.set(f"Clase {category['clase']} · peso relativo {decimal(category['peso_relativo'])} · "
                          f"{decimal(category['peso_global'])} puntos · ventana {category['ventana']} años"
                          if category else "Sin categoría asignada")

        def select(self) -> None:
            picker = CategoryPicker(self, self.app, self.get())
            if picker.result is not None:
                self.set(picker.result)

        def set_enabled(self, enabled: bool) -> None:
            self.entry.configure(state="readonly" if enabled else "disabled")
            self.button.configure(state="normal" if enabled else "disabled")


    class RecordDialog(Dialog):
        def __init__(self, parent: tk.Misc, kind: str, initial: dict[str, str] | None = None,
                     suggested_id: str = "", external: bool = False,
                     repo: Repository | CppRepository | None = None,
                     on_save: Callable[[dict[str, str]], Any] | None = None,
                     is_new: bool = False, locked_fields: tuple[str, ...] = ()):
            title = ("Revisar " if external else "Crear " if is_new or initial is None else "Editar ") + TITLES[kind].lower()
            super().__init__(parent, title, 720, 760)
            self.minsize(540, 420)
            self.result: dict[str, str] | None = None
            self.kind, self.on_save = kind, on_save
            self.inputs: dict[str, Any] = {}
            self.error = tk.StringVar()
            self.scroller = ctk.CTkScrollableFrame(self.body, fg_color="transparent", corner_radius=0,
                                                   scrollbar_button_color=T["line_strong"],
                                                   scrollbar_button_hover_color=T["muted"])
            self.scroller.pack(fill="both", expand=True)
            self.scroll_canvas = self.scroller._parent_canvas
            self.scroll_canvas.configure(yscrollcommand=lambda first, last: set_scrollbar(self.scroller._scrollbar, first, last))
            sections = {
                "grupos": (("Identificación", ("id", "nombre", "codigo_gruplac", "fecha_creacion", "unidad", "responsable", "categoria", "activo")),
                           ("Perfil de investigación", ("descripcion", "objetivos", "mision", "vision", "lineas")),
                           ("Procedencia", ("url", "fuente"))),
                "investigadores": (("Información personal pública", ("id", "nombre", "afiliacion", "contacto", "activo")),
                                    ("Clasificación", ("codigo_cvlac", "categoria")), ("Procedencia", ("url", "fuente"))),
                "productos": (("Producto", ("id", "titulo", "anio", "fecha", "doi", "activo")),
                               ("Clasificación y revisión", ("familia", "tipologia", "categoria", "observacion")),
                               ("Procedencia", ("url", "fuente"))),
                "planes": (("Plan del grupo", ("id", "grupo_id", "nombre", "inicio", "fin", "activo")),
                            ("Objetivos y actividades", ("objetivo", "indicador", "meta", "actividad"))),
            }.get(kind, (("Asociación", ALL_FIELDS[kind]),))
            self.section_cards = []
            for title, fields in sections:
                card = Card(self.scroller, title=title)
                card.pack(fill="x", padx=(0, 8), pady=(0, 12))
                self.section_cards.append(card)
                card.body.grid_columnconfigure(1, weight=1)
                for index, field in enumerate(fields):
                    required = field in REQUIRED.get(kind, ()) or field.endswith("_id")
                    label = LABELS.get(field, field).replace(" (ID)", "") + (" *" if required else "")
                    ctk.CTkLabel(card.body, text=label, font=FONTS["small"], text_color=T["muted"],
                                 justify="left", anchor="w", wraplength=112, width=112).grid(
                                     row=index, column=0, sticky="nw", padx=(0, 16), pady=(8, 4))
                    value = initial.get(field, "") if initial else ""
                    if field == "activo" and not initial:
                        value = "1"
                    if field == "id" and not value:
                        value = suggested_id
                    if field in LONG_FIELDS:
                        widget = ctk.CTkTextbox(card.body, height=80, width=220, wrap="word", font=FONTS["body"],
                                               fg_color=T["field"], text_color=T["ink"], border_width=1,
                                               border_color=T["line_strong"], corner_radius=8,
                                               scrollbar_button_color=T["line_strong"], scrollbar_button_hover_color=T["muted"])
                        widget.insert("1.0", value)
                    elif field == "activo":
                        state = tk.StringVar(value=display_value(field, value or "1"))
                        widget = segmented(card.body, state, ["Activo", "Inactivo"], lambda _value: None)
                    elif kind == "productos" and field == "categoria":
                        widget = CategoryInput(card.body, value=value)
                        self.category_info = widget.info
                    elif field.endswith("_id") and repo is not None:
                        ref_kind = {"grupo_id": "grupos", "investigador_id": "investigadores", "producto_id": "productos"}[field]
                        widget = ReferenceInput(card.body, repo, ref_kind, value)
                    else:
                        widget = ctk.CTkEntry(card.body, width=220, height=34, font=FONTS["body"],
                                              fg_color=T["field"], text_color=T["ink"], border_color=T["line_strong"],
                                              corner_radius=8,
                                              placeholder_text="AAAA-MM-DD" if field in ("fecha", "fecha_creacion", "inicio", "fin") else "")
                        widget.insert(0, value)
                    widget.grid(row=index, column=1, sticky="ew", pady=4)
                    stable = initial and not external and not is_new and (field == "id" or field in RELATION_ENDS.get(kind, ())[:2])
                    if stable or field in locked_fields:
                        if isinstance(widget, (ReferenceInput, CategoryInput)):
                            widget.set_enabled(False)
                        elif isinstance(widget, (ctk.CTkTextbox, ctk.CTkSegmentedButton)):
                            widget.configure(state="disabled")
                        else:
                            widget.configure(state="readonly")
                    self.inputs[field] = widget
            help_text = "* Campo obligatorio. Fechas: AAAA-MM-DD. Ctrl+Enter guarda."
            if kind == "productos":
                help_text += ("\nLa validación es automática y se recalcula al guardar: requiere año, tipología, "
                              "un autor o un grupo activos, y una URL o un DOI válido.")
            ctk.CTkLabel(self.scroller, text=help_text, font=FONTS["caption"], text_color=T["muted"],
                         justify="left", anchor="w", wraplength=580).pack(fill="x", padx=(0, 8), pady=(0, 8))
            self.error_label = ctk.CTkLabel(self.footer, textvariable=self.error, font=FONTS["small"],
                                            text_color=T["rejected"], anchor="w", justify="left", wraplength=650)
            self.error_label.pack(fill="x", pady=(0, 8))
            self.buttons["save"] = button(self.footer, "Guardar", self.accept, variant="primary", width=112)
            self.buttons["save"].pack(side="right")
            self.buttons["cancel"] = button(self.footer, "Cancelar", self.close, width=112)
            self.buttons["cancel"].pack(side="right", padx=(0, 8))
            self.bind("<Control-Return>", lambda _event: self.accept())
            self.error_label.bind("<Configure>", lambda event: self.error_label.configure(
                wraplength=max(220, event.width / scale_of(self) - 10)), add="+")
            self.present(focus=self.inputs.get("nombre") or self.inputs.get("titulo"))

        def accept(self) -> None:
            result = {}
            for field, widget in self.inputs.items():
                result[field] = (widget.get("1.0", "end-1c") if isinstance(widget, ctk.CTkTextbox) else widget.get()).strip()
                if field == "activo":
                    result[field] = "1" if result[field] == "Activo" else "0"
            try:
                result = clean_row(self.kind, result)
                if self.kind == "productos":
                    result.pop("categoria_origen", None)
                if self.on_save:
                    self.on_save(result)
            except (DataError, OSError, ValueError) as exc:
                self.error.set(str(exc))
                return
            self.result = result
            self.close()


    class ReviewDialog(Dialog):
        """Revisión del frente de la cola: corregir el producto y cerrar con una nota."""
        def __init__(self, parent: Any, app: "App", job: dict[str, str]):
            super().__init__(parent, "Revisar siguiente producto", 680, 600)
            self.minsize(500, 460)
            self.app, self.job = app, job
            self.result: dict[str, str] | None = None
            product = app.repo.get("productos", job["producto_id"])
            ctk.CTkLabel(self.body, text="SIGUIENTE EN LA COLA", font=FONTS["caption"], text_color=T["muted"],
                         anchor="w").pack(fill="x")
            ctk.CTkLabel(self.body, text=record_name(product), font=FONTS["h2"], text_color=T["ink"], anchor="w",
                         justify="left", wraplength=620).pack(fill="x", pady=(2, 8))
            reasons = ctk.CTkFrame(self.body, fg_color="transparent")
            reasons.pack(fill="x")
            self.reason_chips = reason_chips(reasons, job["motivo"])
            state = Card(self.body, title="Estado del producto")
            state.pack(fill="x", pady=12)
            state.body.grid_columnconfigure(0, weight=1)
            self.validation_chip = Chip(state.body, "", "muted")
            self.validation_chip.grid(row=0, column=0, sticky="w")
            button(state.body, "Editar producto", self.edit_product, width=136).grid(
                row=0, column=1, rowspan=2, sticky="ne", padx=(12, 0))
            self.state_text = tk.StringVar()
            self.state_label = ctk.CTkLabel(state.body, textvariable=self.state_text, font=FONTS["body"],
                                            text_color=T["ink"], anchor="w", justify="left", wraplength=470)
            self.state_label.grid(row=1, column=0, sticky="ew", pady=(8, 0))
            ctk.CTkLabel(self.body, text="Nota de la revisión *", font=FONTS["small"], text_color=T["muted"],
                         anchor="w").pack(fill="x")
            self.observation = ctk.CTkTextbox(self.body, height=110, wrap="word", font=FONTS["body"],
                                              fg_color=T["field"], text_color=T["ink"], border_width=1,
                                              border_color=T["line_strong"], corner_radius=8,
                                              scrollbar_button_color=T["line_strong"],
                                              scrollbar_button_hover_color=T["muted"])
            self.observation.pack(fill="both", expand=True, pady=(4, 0))
            self.error = tk.StringVar()
            ctk.CTkLabel(self.footer, textvariable=self.error, font=FONTS["small"], text_color=T["rejected"],
                         anchor="w", justify="left", wraplength=620).pack(fill="x", pady=(0, 8))
            button(self.footer, "Ver ficha", lambda: app.open_detail("productos", product["id"], self),
                   width=112).pack(side="left")
            self.buttons["accept"] = button(self.footer, "Cerrar revisión", self.accept, variant="primary", width=136)
            self.buttons["accept"].pack(side="right")
            self.buttons["cancel"] = button(self.footer, "Cancelar", self.close, width=112)
            self.buttons["cancel"].pack(side="right", padx=(0, 8))
            self.bind("<Control-Return>", lambda _event: self.accept())
            self._update_state()
            self.present(focus=self.observation)

        def _update_state(self) -> None:
            product = self.app.repo.get("productos", self.job["producto_id"])
            issues = self.app.repo.product_issues(self.job["producto_id"])
            validated = product["validacion"] == "validado"
            self.validation_chip.set(display_value("validacion", product["validacion"]),
                                     STATUS_COLORS.get(product["validacion"], "unknown"))
            if validated:
                self.state_text.set("Validado: el producto cumple todas las reglas.")
            else:
                self.state_text.set("Rechazado. Corrija el producto o explique en la nota por qué no es posible:\n"
                                    + "\n".join("• " + RULE_LABELS[code] for code in issues))
            self.state_label.configure(text_color=T["ink"] if validated else T["rejected"])

        def edit_product(self) -> None:
            self.app.edit_record("productos", key=self.job["producto_id"], parent=self)
            self._update_state()

        def accept(self) -> None:
            try:
                self.result = self.app.repo.process_review(self.observation.get("1.0", "end-1c"))
            except (DataError, OSError, ValueError) as exc:
                self.error.set(str(exc))
                return
            self.close()

    class EntityTab(ctk.CTkFrame):
        def __init__(self, parent: tk.Misc, app: "App", kind: str):
            super().__init__(parent, fg_color=T["page"], corner_radius=0)
            self.app, self.kind = app, kind
            self.search = tk.StringVar()
            self.state_filter = tk.StringVar(value="Todos")
            self.offset = 0
            self._rows: dict[str, dict[str, str]] = {}
            body = ctk.CTkFrame(self, fg_color="transparent")
            body.pack(fill="both", expand=True, padx=24, pady=20)
            heading = ctk.CTkFrame(body, fg_color="transparent")
            heading.pack(fill="x", pady=(0, 14))
            ctk.CTkLabel(heading, text=TITLES[kind], font=FONTS["h1"], text_color=T["ink"]).pack(side="left")
            self.count_chip = Chip(heading, "0 registros", "muted")
            self.count_chip.pack(side="left", padx=12)
            button(heading, "Nuevo", self.create, variant="primary", width=112).pack(side="right")
            toolbar = ctk.CTkFrame(body, fg_color="transparent")
            toolbar.pack(fill="x")
            self.search_entry = SearchField(toolbar, self.search, self.refresh, placeholder_text="Buscar nombre, ID o datos…", width=180)
            self.search_entry.pack(side="left", fill="x", expand=True, padx=(0, 12))
            self.search.trace_add("write", lambda *_: setattr(self, "offset", 0))
            self.state_control = segmented(toolbar, self.state_filter, ["Todos", "Activos", "Inactivos"],
                                           lambda _value: self._filter_changed())
            self.state_control.pack(side="left")
            self.source_button = None
            if kind in ("grupos", "investigadores"):
                self.source_button = button(toolbar, "Consultar fuente", self.open_source, width=132, state="disabled")
                self.source_button.pack(side="left", padx=(12, 0))
            actions = ctk.CTkFrame(body, fg_color="transparent")
            actions.pack(fill="x", pady=(12, 0))
            self.selection_buttons = []
            for label, action in (("Abrir ficha", self.open_detail), ("Editar", self.edit),
                                  ("Desactivar", self.toggle), ("Eliminar", self.delete)):
                control = button(actions, label, action, variant="danger" if label == "Eliminar" else "secondary",
                                 width=112, state="disabled")
                control.pack(side="left", padx=(0, 8))
                self.selection_buttons.append(control)
                if label == "Desactivar":
                    self.toggle_button = control
            self.visible = {
                "grupos": ("nombre", "categoria", "responsable", "activo", "id"),
                "investigadores": ("nombre", "afiliacion", "contacto", "activo", "id"),
                "productos": ("titulo", "anio", "tipologia", "categoria", "validacion", "activo", "id"),
                "planes": ("grupo_id", "nombre", "objetivo", "inicio", "fin", "activo", "id"),
            }[kind]
            widths = {field: (280 if field in ("titulo", "nombre", "grupo_id", "objetivo") else
                              160 if field in ("responsable", "tipologia", "afiliacion", "contacto", "id") else
                              60 if field == "anio" else 96) for field in self.visible}
            self.table = DataTable(body, self.visible,
                                   headings={field: "Estado" if field == "activo" else LABELS[field].replace(" (ID)", "")
                                             for field in self.visible}, widths=widths, anchors={"anio": "e"},
                                   chips={"activo": lambda value: "primary" if value == "Activo" else "muted",
                                          "validacion": lambda value: STATUS_COLORS.get(value.lower(), "unknown"),
                                          "categoria": lambda value: "unknown" if value == "—" else "primary"},
                                   row_chips={"categoria": lambda iid, _value: category_tone(self._rows.get(iid, {}))},
                                   tooltips={"categoria": self._category_hint}, height_rows=12, open_on_click=kind == "productos")
            self.table.empty_text = "Sin registros. Cree uno nuevo o cambie la búsqueda."
            self.table.on_open(lambda _iid: self.open_detail())
            self.table.on_select(lambda _iid: self._selection_changed())
            # El paginador se empaqueta antes que la tabla para no quedar recortado en ventanas bajas.
            self.pager = Pager(body, self._move)
            self.pager.pack(side="bottom", fill="x", pady=(10, 0))
            self.table.pack(fill="both", expand=True, pady=(12, 0))
            self.previous, self.next = self.pager.previous, self.pager.next
            self.page_label = tk.StringVar()

        def _category_hint(self, iid: str, text: str) -> str:
            row = self._rows.get(iid, {})
            category = catalog_by_code.get(row.get("categoria", "")) if self.kind == "productos" else None
            if row.get("categoria_origen") == "estimada":
                return text + " · Estimada por tipología" + (f"\n{category['descripcion']}" if category else "")
            return f"{text} — {category['descripcion']}" if category else ""

        def _filter_changed(self) -> None:
            self.offset = 0
            self.refresh()


        def selected(self) -> str | None:
            selection = self.table.selection()
            return selection[0] if selection else None

        def _selection_changed(self) -> None:
            key = self.selected()
            for button in self.selection_buttons:
                button.configure(state="normal" if key else "disabled")
            self.selection_buttons[-1].configure(fg_color=T["rejected"] if key else T["field"],
                                                  border_width=0 if key else 1, border_color=T["line_strong"],
                                                  text_color_disabled=T["disabled"])
            if self.source_button is not None:
                self.source_button.configure(state="normal" if key else "disabled")
            row = self.app.record(self.kind, key) if key else None
            self.toggle_button.configure(text="Activar" if row and row["activo"] == "0" else "Desactivar")

        def _move(self, direction: int) -> None:
            self.offset = max(0, self.offset + direction * PAGE_SIZE)
            self.refresh()

        def refresh(self) -> None:
            self.search_entry.cancel_pending()
            current = self.selected()
            self.table.delete(*self.table.get_children())
            page = self.app.entity_page(self.kind, self.offset, self.search.get(), self.state_filter.get())
            if self.offset and self.offset >= page["total"]:
                self.offset = max(0, (page["total"] - 1) // PAGE_SIZE * PAGE_SIZE)
                page = self.app.entity_page(self.kind, self.offset, self.search.get(), self.state_filter.get())
            self._rows = {row["id"]: row for row in page["rows"]}
            self.count_chip.set(f"{number(page['total'])} registros", "muted")
            for index, row in enumerate(page["rows"]):
                self.app._record_cache[(self.kind, row["id"])] = row
                values = [self.app.reference_label("grupos", row[field]) if field == "grupo_id" else cell_value(field, row[field])
                          for field in self.visible]
                self.table.insert("", "end", iid=row["id"], values=values,
                                  tags=("stripe",) if index % 2 else ())
            self.pager.show(self.offset, len(page["rows"]), page["total"])
            self.page_label.set(self.pager.label.cget("text"))
            show_columns(self.table, self.visible, self.app.usage(self.kind), keep=("nombre", "titulo", "activo", "id"))
            if current and self.table.exists(current):
                self.table.selection_set(current)
            self._selection_changed()

        def open_detail(self) -> None:
            key = self.selected()
            if key:
                self.app.open_detail(self.kind, key, self)

        def create(self) -> None:
            self.app.edit_record(self.kind, parent=self)

        def edit(self) -> None:
            key = self.selected()
            if not key:
                return
            self.app.edit_record(self.kind, key=key, parent=self)

        def open_source(self) -> None:
            key = self.selected()
            if key:
                row = self.app.repo.get(self.kind, key)
                if row and row["url"]:
                    self.app.import_url(row["url"])
                else:
                    self.app.inform("Sin fuente", "Este registro no tiene URL de origen", parent=self)

        def toggle(self) -> None:
            key = self.selected()
            if key:
                self.app.mutate(lambda: self.app.repo.toggle(self.kind, key), changed={self.kind})

        def delete(self) -> None:
            key = self.selected()
            row = self.app.record(self.kind, key) if key else None
            if row and self.app.confirm("Eliminar", f"¿Eliminar {record_name(row)} definitivamente?\nPuede deshacer este cambio antes de vaciar el historial.", parent=self):
                self.app.mutate(lambda: self.app.repo.delete(self.kind, key), changed={self.kind})

    class RelationTab(ctk.CTkFrame):
        def __init__(self, parent: tk.Misc, app: "App", kind: str):
            super().__init__(parent, fg_color=T["page"], corner_radius=0)
            self.app, self.kind = app, kind
            self.offset = 0
            self.search = tk.StringVar()
            body = ctk.CTkFrame(self, fg_color="transparent")
            body.pack(fill="both", expand=True, padx=8, pady=12)
            search_bar = ctk.CTkFrame(body, fg_color="transparent")
            search_bar.pack(fill="x", pady=(0, 12))
            self.search_entry = SearchField(search_bar, self.search, self.refresh,
                                             placeholder_text="Buscar nombres, IDs o datos del vínculo…")
            self.search_entry.pack(side="left", fill="x", expand=True, padx=(0, 12))
            button(search_bar, "Nuevo vínculo", self.create, variant="primary", width=132).pack(side="right")
            self.search.trace_add("write", lambda *_: setattr(self, "offset", 0))
            actions = ctk.CTkFrame(body, fg_color="transparent")
            actions.pack(fill="x")
            self.selection_buttons = []
            for label, action in (("Editar", self.edit), ("Desactivar", self.toggle), ("Desvincular", self.delete)):
                control = button(actions, label, action, variant="danger" if label == "Desvincular" else "secondary",
                                 width=112, state="disabled")
                control.pack(side="left", padx=(0, 8))
                self.selection_buttons.append(control)
            self.toggle_button = self.selection_buttons[1]
            self._rows: dict[str, dict[str, str]] = {}
            self.fields = RELATION_FIELDS[kind]
            self.table = DataTable(body, self.fields,
                                   headings={field: "Estado" if field == "activo" else LABELS.get(field, field).replace(" (ID)", "")
                                             for field in self.fields},
                                   widths={field: 260 if field.endswith("_id") else 130 for field in self.fields},
                                   anchors={"orden": "e"}, height_rows=12,
                                   chips={"activo": lambda value: "primary" if value == "Activo" else "muted"})
            self.table.empty_text = "Sin vínculos. Cree uno nuevo o cambie la búsqueda."
            self.table.on_open(lambda _iid: self.edit())
            self.table.on_select(lambda _iid: self._selection_changed())
            self.pager = Pager(body, self._move)
            self.pager.pack(side="bottom", fill="x", pady=(10, 0))
            self.table.pack(fill="both", expand=True, pady=(12, 0))
            self.previous, self.next = self.pager.previous, self.pager.next
            self.page_label = tk.StringVar()

        def _selection_changed(self) -> None:
            selection = self.table.selection()
            for control in self.selection_buttons:
                control.configure(state="normal" if selection else "disabled")
            self.selection_buttons[-1].configure(fg_color=T["rejected"] if selection else T["field"],
                                                  border_width=0 if selection else 1, border_color=T["line_strong"],
                                                  text_color_disabled=T["disabled"])
            row = self._rows.get(selection[0], {}) if selection else {}
            self.toggle_button.configure(text="Activar" if row.get("activo") == "0" else "Desactivar")


        def _move(self, direction: int) -> None:
            self.offset = max(0, self.offset + direction * PAGE_SIZE)
            self.refresh()

        def selected(self) -> tuple[str, str] | None:
            selection = self.table.selection()
            if not selection:
                return None
            return tuple(json.loads(selection[0]))  # type: ignore[return-value]

        def refresh(self) -> None:
            self.search_entry.cancel_pending()
            current = self.table.selection()
            self.table.delete(*self.table.get_children())
            left, right = RELATION_ENDS[self.kind][:2]
            page = self.app.relationship_page(self.kind, self.offset, self.search.get())
            if self.offset and self.offset >= page["total"]:
                self.offset = max(0, (page["total"] - 1) // PAGE_SIZE * PAGE_SIZE)
                page = self.app.relationship_page(self.kind, self.offset, self.search.get())
            left_kind, right_kind = RELATION_ENDS[self.kind][2:]
            self._rows.clear()
            for index, row in enumerate(page["rows"]):
                pair = json.dumps((row[left], row[right]))
                self._rows[pair] = row
                values = [self.app.reference_label(left_kind if field == left else right_kind, row[field])
                          if field in (left, right) else cell_value(field, row[field]) for field in self.fields]
                self.table.insert("", "end", iid=pair, values=values,
                                  tags=("stripe",) if index % 2 else ())
            self.pager.show(self.offset, len(page["rows"]), page["total"])
            self.page_label.set(self.pager.label.cget("text"))
            show_columns(self.table, self.fields, self.app.usage(self.kind), keep=(left, right, "activo"))
            if current and self.table.exists(current[0]):
                self.table.selection_set(current[0])
            self._selection_changed()

        def create(self) -> None:
            self.app.edit_record(self.kind, parent=self)

        def edit(self) -> None:
            key = self.selected()
            if not key:
                return
            self.app.edit_record(self.kind, key=key, parent=self)

        def toggle(self) -> None:
            key = self.selected()
            if key:
                self.app.mutate(lambda: self.app.repo.toggle(self.kind, key), changed={self.kind})

        def delete(self) -> None:
            key = self.selected()
            if key and self.app.confirm("Desvincular", "¿Eliminar esta relación?", parent=self):
                self.app.mutate(lambda: self.app.repo.delete(self.kind, key), changed={self.kind})

    class ContextTab(ctk.CTkFrame):
        def __init__(self, parent: Any, detail: "DetailDialog", source: str, side: str,
                     target_kind: str, columns: tuple[str, ...]):
            super().__init__(parent, fg_color="transparent")
            self.detail, self.app = detail, detail.app
            self.source, self.side, self.target_kind = source, side, target_kind
            self.columns = columns
            self.offset = 0
            self.search = tk.StringVar()
            self.state_filter = tk.StringVar(value="Todos")
            self.items: dict[str, tuple[dict[str, str], dict[str, str]]] = {}
            toolbar = ctk.CTkFrame(self, fg_color="transparent")
            toolbar.pack(fill="x", pady=(4, 10))
            self.search_entry = SearchField(toolbar, self.search, self.refresh, placeholder_text="Buscar en esta ficha…",
                                            width=180)
            self.search_entry.pack(side="left", fill="x", expand=True, padx=(0, 12))
            self.search.trace_add("write", lambda *_: setattr(self, "offset", 0))
            self.state_control = segmented(toolbar, self.state_filter, ["Todos", "Activos"],
                                           lambda _value: self._filter_changed())
            self.state_control.pack(side="left")
            button(toolbar, "Nuevo plan" if source == "planes" else "Vincular…", self.create, variant="primary",
                   width=120).pack(side="left", padx=(12, 0))
            actions = ctk.CTkFrame(self, fg_color="transparent")
            actions.pack(fill="x", pady=(0, 10))
            self.selection_buttons = []
            for label, command in (("Abrir ficha", self.open_record), ("Editar registro", self.edit_record),
                                    ("Editar vínculo" if source != "planes" else "Editar plan", self.edit_link),
                                    ("Desactivar", self.toggle),
                                    ("Desvincular" if source != "planes" else "Eliminar plan", self.delete)):
                if source == "planes" and label == "Editar registro":
                    continue
                control = button(actions, label, command, width=112, state="disabled",
                                 variant="danger" if label in ("Desvincular", "Eliminar plan") else "secondary")
                control.pack(side="left", padx=(0, 8))
                self.selection_buttons.append(control)
                if label == "Desactivar":
                    self.toggle_button = control
            state_tone = lambda value: "primary" if value == "Activo" else "muted"
            self.table = DataTable(self, columns, height_rows=8, open_on_click=target_kind == "productos",
                                   headings={field: {"_link_active": "Vínculo", "_record_active": "Registro"}.get(
                                       field, LABELS.get(field, field).replace(" (ID)", "")) for field in columns},
                                   widths={field: 280 if field in ("nombre", "titulo") else
                                           65 if field in ("anio", "orden") else 110 for field in columns},
                                   anchors={"anio": "e", "orden": "e"},
                                   chips={"_link_active": state_tone, "_record_active": state_tone,
                                          "validacion": lambda value: STATUS_COLORS.get(value.lower(), "unknown"),
                                          "categoria": lambda value: "unknown" if value == "—" else "primary"},
                                   row_chips={"categoria": lambda iid, _value: category_tone(
                                       self.items.get(iid, ({}, {}))[0])})
            self.table.empty_text = "Sin datos asociados. Vincule un registro o cambie la búsqueda."
            self.table.on_select(lambda _iid: self._selection_changed())
            self.table.on_open(lambda _iid: self.open_record())
            # El paginador se empaqueta antes que la tabla para no quedar recortado en ventanas bajas.
            self.pager = Pager(self, self.move)
            self.pager.pack(side="bottom", fill="x", pady=(10, 0))
            self.table.pack(fill="both", expand=True)
            self.previous, self.next = self.pager.previous, self.pager.next
            self.count = tk.StringVar()

        def _filter_changed(self) -> None:
            self.offset = 0
            self.refresh()

        def move(self, direction: int) -> None:
            self.offset = max(0, self.offset + direction * PAGE_SIZE)
            self.refresh()

        def selected(self) -> tuple[dict[str, str], dict[str, str]] | None:
            selected = self.table.selection()
            return self.items.get(selected[0]) if selected else None

        def _selection_changed(self) -> None:
            item = self.selected()
            for control in self.selection_buttons:
                control.configure(state="normal" if item else "disabled")
            self.selection_buttons[-1].configure(fg_color=T["rejected"] if item else T["field"],
                                                  border_width=0 if item else 1, border_color=T["line_strong"],
                                                  text_color_disabled=T["disabled"])
            self.toggle_button.configure(text="Activar" if item and item[1]["activo"] == "0" else "Desactivar")

        def refresh(self) -> None:
            self.search_entry.cancel_pending()
            selected = self.table.selection()
            page = self.detail.page_for(self, self.offset, self.search.get())
            if self.offset and self.offset >= page["total"]:
                self.offset = max(0, (page["total"] - 1) // PAGE_SIZE * PAGE_SIZE)
                page = self.detail.page_for(self, self.offset, self.search.get())
            self.items.clear()
            self.table.delete(*self.table.get_children())
            for index, (row, link) in enumerate(page["rows"]):
                self.items[row["id"]] = (row, link)
                combined = {**row, **link, "_link_active": link["activo"], "_record_active": row["activo"]}
                values = [cell_value("activo" if field.startswith("_") else field, combined.get(field, ""))
                          for field in self.columns]
                self.table.insert("", "end", iid=row["id"], values=values, tags=("stripe",) if index % 2 else ())
            usage = {**self.app.usage(self.target_kind), **self.app.usage(self.source)}
            show_columns(self.table, self.columns, usage, keep=("nombre", "titulo"))
            self.pager.show(self.offset, len(page["rows"]), page["total"])
            self.count.set(self.pager.label.cget("text"))
            if selected and self.table.exists(selected[0]):
                self.table.selection_set(selected[0])
            self._selection_changed()

        def create(self) -> None:
            field = "grupo_id" if self.source == "planes" else RELATION_ENDS[self.source][0 if self.side == "left" else 1]
            self.app.edit_record(self.source, defaults={field: self.detail.key}, parent=self.detail, locked_fields=(field,))

        def edit_link(self) -> None:
            item = self.selected()
            if item:
                row, link = item
                key = row["id"] if self.source == "planes" else tuple(link[field] for field in RELATION_ENDS[self.source][:2])
                self.app.edit_record(self.source, key=key, parent=self.detail)

        def edit_record(self) -> None:
            item = self.selected()
            if item:
                self.app.edit_record(self.target_kind, key=item[0]["id"], parent=self.detail)

        def open_record(self) -> None:
            item = self.selected()
            if item:
                self.app.open_detail(self.target_kind, item[0]["id"], self.detail)

        def toggle(self) -> None:
            item = self.selected()
            if item:
                row, link = item
                key = row["id"] if self.source == "planes" else tuple(link[field] for field in RELATION_ENDS[self.source][:2])
                self.app.mutate(lambda: self.app.repo.toggle(self.source, key), changed={self.source})

        def delete(self) -> None:
            item = self.selected()
            if item and self.app.confirm("Eliminar asociación" if self.source != "planes" else "Eliminar plan",
                                            f"¿Eliminar {record_name(item[0])} de esta ficha?", parent=self.detail):
                row, link = item
                key = row["id"] if self.source == "planes" else tuple(link[field] for field in RELATION_ENDS[self.source][:2])
                self.app.mutate(lambda: self.app.repo.delete(self.source, key), changed={self.source})

    class DetailDialog(Dialog):
        """Ficha de un registro: resumen, vínculos con sus acciones, indicador y estadísticas."""
        # Ventanas del indicador: años hasta el corte (el año actual); None usa todos los años.
        INDICATOR_WINDOWS = {"Completa": None, "Últimos 5 años": 5, "Últimos 2 años": 2}

        def __init__(self, parent: Any, app: "App", kind: str, key: str, tab: str | None = None):
            super().__init__(parent, "Ficha · " + TITLES[kind], 1040, 700)
            self.minsize(720, 500)
            self.app, self.kind, self.key = app, kind, key
            self.body.pack_configure(pady=(16, 8))
            header = ctk.CTkFrame(self.body, fg_color="transparent")
            header.pack(fill="x", pady=(0, 6))
            header.grid_columnconfigure(0, weight=1)
            self.eyebrow = ctk.CTkLabel(header, text="", font=FONTS["caption"], text_color=T["muted"], anchor="w")
            self.eyebrow.grid(row=0, column=0, sticky="w")
            self.title_label = ctk.CTkLabel(header, text="", font=FONTS["h1"], text_color=T["ink"], anchor="w",
                                            justify="left", wraplength=780)
            self.title_label.grid(row=1, column=0, sticky="ew")
            self.state_chips = ctk.CTkFrame(header, fg_color="transparent")
            self.state_chips.grid(row=2, column=0, sticky="w", pady=(6, 0))
            button(header, "Editar ficha", lambda: app.edit_record(kind, key=key, parent=self), variant="primary",
                   width=120).grid(row=0, column=1, rowspan=3, sticky="ne", padx=(12, 0))
            header.bind("<Configure>", lambda event: self.title_label.configure(
                wraplength=max(260, event.width / scale_of(header) - 150)), add="+")
            self.tabs = tabview(self.body)
            self.tabs.pack(fill="both", expand=True)
            summary_tab = self.tabs.add("Resumen")
            self.summary = ctk.CTkScrollableFrame(summary_tab, fg_color="transparent", corner_radius=0,
                                                  scrollbar_button_color=T["line_strong"],
                                                  scrollbar_button_hover_color=T["muted"])
            self.summary.pack(fill="both", expand=True)
            self.summary._parent_canvas.configure(yscrollcommand=lambda first, last: set_scrollbar(
                self.summary._scrollbar, first, last))
            self.summary.grid_columnconfigure(1, weight=1)
            self._wrap = 700
            self._wrapped: list[Any] = []
            summary_tab.bind("<Configure>", self._summary_width, add="+")
            # La clasificación de productos se muestra también en sus vínculos.
            specifications = {
                "grupos": (("Integrantes", "membresias", "left", "investigadores", ("nombre", "categoria", "rol", "inicio", "fin", "_link_active", "_record_active")),
                           ("Planes", "planes", "left", "planes", ("nombre", "inicio", "fin", "objetivo", "_record_active")),
                           ("Productos", "grupos_productos", "left", "productos", ("titulo", "anio", "tipologia", "categoria", "validacion", "_link_active", "_record_active"))),
                "investigadores": (("Grupos", "membresias", "right", "grupos", ("nombre", "rol", "inicio", "fin", "_link_active", "_record_active")),
                                    ("Productos", "autorias", "right", "productos", ("titulo", "anio", "tipologia", "categoria", "validacion", "orden", "rol", "_link_active", "_record_active"))),
                "productos": (("Autores", "autorias", "left", "investigadores", ("nombre", "orden", "rol", "_link_active", "_record_active")),
                               ("Grupos", "grupos_productos", "right", "grupos", ("nombre", "origen", "_link_active", "_record_active"))),
                "planes": (),
            }
            self.link_tabs = {}
            for title, source, side, target, columns in specifications[kind]:
                tab_view = ContextTab(self.tabs.add(title), self, source, side, target, columns)
                tab_view.pack(fill="both", expand=True)
                self.link_tabs[source] = tab_view
            self.indicator_panel = None
            if kind in ("grupos", "investigadores"):
                indicator = self.tabs.add("Indicador")
                self.indicator_window = tk.StringVar(value="Completa")
                window_bar = ctk.CTkFrame(indicator, fg_color="transparent")
                window_bar.pack(fill="x", pady=(4, 12))
                ctk.CTkLabel(window_bar, text="Ventana", font=FONTS["small"], text_color=T["muted"]).pack(
                    side="left", padx=(0, 10))
                self.indicator_control = segmented(window_bar, self.indicator_window, list(self.INDICATOR_WINDOWS),
                                                   lambda _value: self.refresh_indicator())
                self.indicator_control.pack(side="left")
                self.indicator_message = ctk.CTkLabel(indicator, text="", font=FONTS["body"], text_color=T["muted"],
                                                      anchor="w", justify="left", wraplength=700)
                self.indicator_panel = IndicatorPanel(
                    indicator, lambda: app.show_indicator_details(self.indicator_panel.result, parent=self),
                    bg_color=T["page"])
            self.stats_canvas = None
            if kind != "planes":
                stats = self.tabs.add("Estadísticas")
                self.stats_note = tk.StringVar()
                ctk.CTkLabel(stats, textvariable=self.stats_note, font=FONTS["small"], text_color=T["muted"],
                             anchor="w", justify="left", wraplength=800).pack(fill="x", pady=(4, 8))
                button(stats, "Explorar en Dashboard con filtros", self.show_statistics, width=240).pack(
                    side="bottom", anchor="e", pady=(10, 0))
                charts = ctk.CTkFrame(stats, fg_color="transparent")
                charts.pack(fill="both", expand=True)
                charts.grid_columnconfigure(0, weight=3, uniform="statistics")
                charts.grid_columnconfigure(1, weight=2, uniform="statistics")
                charts.grid_rowconfigure(0, weight=1)
                panel, self.stats_canvas = app._chart_panel(charts, 280, draw=lambda _canvas: self._draw_statistics())
                panel.grid(row=0, column=0, sticky="nsew", padx=(0, 12))
                panel, self.class_canvas = app._chart_panel(charts, 280, draw=self._draw_classes)
                panel.grid(row=0, column=1, sticky="nsew")
            if kind == "productos":
                button(self.footer, "Enviar a revisión", lambda: app.enqueue(key, parent=self), width=140).pack(
                    side="left", padx=(0, 8))
            elif kind == "planes":
                button(self.footer, "Abrir grupo", lambda: app.open_detail(
                    "grupos", app.repo.get(kind, key)["grupo_id"], self), width=120).pack(side="left", padx=(0, 8))
            self.source_button = button(self.footer, "Consultar fuente", self.open_source, width=140)
            if kind != "planes":
                self.source_button.pack(side="left")
            button(self.footer, "Cerrar", self.close, width=112).pack(side="right")
            self.refresh()
            if tab is not None:
                try:
                    self.tabs.set(tab)
                except ValueError:
                    pass
            app.details.add(self)
            self.present(wait=False)

        def page_for(self, tab: ContextTab, offset: int, query: str) -> dict[str, Any]:
            active_only = tab.state_filter.get() == "Activos"
            if tab.source == "planes":
                page = self.app.repo.page("planes", offset, PAGE_SIZE, query,
                                          active="1" if active_only else "", group_id=self.key)
                return {"total": page["total"], "rows": [(row, row) for row in page["rows"]]}
            page = self.app.repo.related_page(tab.source, self.key, tab.side, offset, PAGE_SIZE,
                                             active_only, query, include_references=True)
            self.app.cache_references(page)
            field = RELATION_ENDS[tab.source][1 if tab.side == "left" else 0]
            return {"total": page["total"],
                    "rows": [(self.app.record(tab.target_kind, link[field]), link) for link in page["rows"]]}

        def _summary_width(self, event: Any) -> None:
            self._wrap = max(240, round(event.width / scale_of(self)) - 230)
            for label in self._wrapped:
                label.configure(wraplength=self._wrap)

        def _summary_row(self, line: int, label: str, value: str, tone: str = "ink") -> None:
            ctk.CTkLabel(self.summary, text=label, font=FONTS["small"], text_color=T["muted"], anchor="nw",
                         justify="left", width=170, wraplength=160).grid(row=line, column=0, sticky="nw",
                                                                         padx=(0, 16), pady=(8, 0))
            shown = ctk.CTkLabel(self.summary, text=value, font=FONTS["body"], text_color=T[tone], anchor="nw",
                                 justify="left", wraplength=self._wrap)
            shown.grid(row=line, column=1, sticky="ew", pady=(8, 0))
            self._wrapped.append(shown)

        def _fill_summary(self, row: dict[str, str]) -> None:
            for child in self.summary.winfo_children():
                child.destroy()
            self._wrapped = []
            # Solo campos con datos; los vacíos se resumen en una línea. La categoría de
            # un investigador se consulta desde los integrantes de su grupo; la de un producto tiene su tarjeta.
            hidden = {"investigadores": {"categoria"}, "productos": {"categoria", "categoria_origen"}}.get(self.kind, set())
            missing = []
            line = 0
            for field in ENTITY_FIELDS[self.kind]:
                if field in hidden:
                    continue
                if field != "validacion" and not row[field]:
                    missing.append(LABELS[field])
                    continue
                if field == "validacion":
                    issues = self.app.repo.product_issues(self.key)
                    self._summary_row(line, LABELS[field], "Validado: cumple todas las reglas" if not issues else
                                      "Rechazado" + "".join("\n• " + RULE_LABELS[code] for code in issues),
                                      "ink" if not issues else "rejected")
                else:
                    self._summary_row(line, LABELS[field], display_value(field, row[field]))
                line += 1
            if self.kind == "productos":
                self._category_card(row).grid(row=line, column=0, columnspan=2, sticky="ew", pady=(16, 0))
                line += 1
            if missing:
                note = ctk.CTkLabel(self.summary, text="Sin registrar: " + ", ".join(missing), font=FONTS["caption"],
                                    text_color=T["muted"], anchor="w", justify="left", wraplength=self._wrap + 180)
                note.grid(row=line, column=0, columnspan=2, sticky="ew", pady=(16, 4))

        def _category_card(self, row: dict[str, str]) -> Any:
            card = Card(self.summary, title="Clasificación del modelo", bg_color=T["page"])
            category = catalog_by_code.get(row["categoria"])
            estimated = row["categoria_origen"] == "estimada"
            top = ctk.CTkFrame(card.body, fg_color="transparent")
            top.pack(fill="x")
            self.category_chip = Chip(top, row["categoria"] or "Sin categoría", category_tone(row))
            self.category_chip.pack(side="left")
            description = ctk.CTkLabel(top, text=category["descripcion"] if category else
                                       "Sin categoría del modelo 2024" if not row["categoria"] else "Código fuera del catálogo",
                                       font=FONTS["body"], text_color=T["ink"], anchor="w", justify="left",
                                       wraplength=self._wrap)
            description.pack(side="left", fill="x", expand=True, padx=(10, 0))
            self._wrapped.append(description)
            facts = ctk.CTkFrame(card.body, fg_color="transparent")
            facts.pack(fill="x", pady=(10, 0))
            values = [("Origen", "Estimada por tipología" if estimated else "Manual" if row["categoria"] else "—")]
            if category:
                values = [("Clase", category["clase"]), ("Peso relativo", decimal(category["peso_relativo"])),
                          ("Valor en puntos", decimal(category["peso_global"])),
                          ("Ventana de observación", f"{category['ventana']} años"), *values]
            for index, (label, value) in enumerate(values):
                cell = ctk.CTkFrame(facts, fg_color="transparent")
                cell.grid(row=index // 3, column=index % 3, sticky="w", padx=(0, 28), pady=(0, 8))
                ctk.CTkLabel(cell, text=label, font=FONTS["caption"], text_color=T["muted"], anchor="w").pack(anchor="w")
                ctk.CTkLabel(cell, text=value, font=FONTS["body"], text_color=T["ink"], anchor="w").pack(anchor="w")
            if estimated:
                ctk.CTkLabel(card.body, text="Confirme o cambie la categoría con Editar ficha.", font=FONTS["small"],
                             text_color=T["pending"], anchor="w").pack(fill="x")
            return card

        def refresh(self, changed: set[str] | None = None) -> None:
            row = self.app.repo.get(self.kind, self.key)
            if row is None:
                self.destroy()
                return
            if changed is None or self.kind in changed:
                self.eyebrow.configure(text=f"{SINGULAR[self.kind].upper()} · {self.key}")
                self.title_label.configure(text=record_name(row))
                chips = [(display_value("activo", row["activo"]), "primary" if row["activo"] == "1" else "muted")]
                if self.kind == "productos":
                    chips.append((display_value("validacion", row["validacion"]),
                                  STATUS_COLORS.get(row["validacion"], "unknown")))
                elif self.kind == "grupos" and row["categoria"]:
                    chips.append((f"Categoría {row['categoria']}", "primary"))
                for child in self.state_chips.winfo_children():
                    child.destroy()
                for text, tone in chips:
                    Chip(self.state_chips, text, tone).pack(side="left", padx=(0, 6))
                self._fill_summary(row)
                self.source_button.configure(state="normal" if row.get("url") else "disabled")
            for tab in self.link_tabs.values():
                if changed is None or tab.source in changed or tab.target_kind in changed:
                    tab.refresh()
            analysis_changed = changed is None or bool(changed & (set(ALL_FIELDS) - {"planes"}))
            if analysis_changed:
                self.refresh_indicator()
            if analysis_changed and self.stats_canvas is not None:
                self.statistics = self.app.repo.statistics({"grupos": "Grupo", "investigadores": "Investigador", "productos": "Producto"}[self.kind], self.key, limit=0)
                self.stats_note.set(f"{self.statistics['total']} productos activos únicos de esta ficha · Todos los años. Los vínculos inactivos se excluyen.")
                self.stats_canvas.redraw()
                self.class_canvas.redraw()

        def refresh_indicator(self) -> None:
            """Indicador de la ficha en la ventana elegida; solo los registros activos tienen indicador."""
            if self.indicator_panel is None:
                return
            row = self.app.repo.get(self.kind, self.key)
            years = self.INDICATOR_WINDOWS[self.indicator_window.get()]
            end = date.today().year if years else None
            result, message = None, ""
            if row is None or row["activo"] != "1":
                message = "El indicador se calcula solo para registros activos. Active la ficha para verlo."
            else:
                try:
                    result = self.app.indicator(self.kind, self.key, end - years + 1 if years else None, end)
                except DataError as exc:
                    message = str(exc)
            self.indicator_panel.show(result, record_name(row) if result else "")
            self.indicator_message.configure(text=message)
            if message:
                self.indicator_message.pack(fill="x", pady=(8, 0))
            else:
                self.indicator_message.pack_forget()

        def _draw_statistics(self) -> None:
            if self.stats_canvas is not None and hasattr(self, "statistics"):
                self.app._draw_histogram(self.stats_canvas, self.statistics["por_anio"])

        def _draw_classes(self, canvas: Any) -> None:
            counts = class_counts(self.statistics["por_categoria"]) if hasattr(self, "statistics") else {}
            draw_donut(canvas, "Por clase del modelo 2024", "", [(key, key, count, class_color(key))
                                                               for key, count in counts.items()], "productos")

        def show_statistics(self) -> None:
            kind, key = self.kind, self.key
            self.destroy()
            self.app.view.set({"grupos": "Grupo", "investigadores": "Investigador", "productos": "Producto"}[kind])
            self.app._scope_changed(key)
            self.app._show_page("dashboard")

        def open_source(self) -> None:
            row = self.app.repo.get(self.kind, self.key)
            self.destroy()
            self.app.import_url(row["url"])

        def destroy(self) -> None:
            self.app.details.discard(self)
            previous = self._previous_grab
            cancel_widget_timers(self)
            super().destroy()
            try:
                if previous is not None and previous.winfo_exists():
                    previous.grab_set()
            except TclError:
                pass

    class App:
        def __init__(self, root: Any, initial_file: Path | None,
                     repository: Repository | CppRepository):
            if load_on_start:
                root.withdraw()
            self.root = root
            self.repo = repository
            self.backend_name = "C++" if isinstance(repository, CppRepository) else "Python"
            self.theme = "light"
            # Archivo donde se guarda; None si los datos aún no tienen uno (vacío, carpeta anterior o copia).
            self.data_file: Path | None = None
            self.initial_file = initial_file
            self._url_busy = False
            self.details: set[DetailDialog] = set()
            self._record_cache: dict[tuple[str, str], dict[str, str]] = {}
            self._usage_cache: dict[str, dict[str, int]] = {}
            self._indicator_cache: dict[tuple[Any, ...], dict[str, Any]] = {}
            self._year_timer: str | None = None
            self.status = ctk.StringVar(master=root, value="Iniciando PEA-i…")
            self.workspace = ctk.StringVar(master=root)
            self.appearance = ctk.StringVar(master=root, value="light")
            self.workspace_caption = "Espacio vacío"
            self.root.title("PEA-i UPC — Investigación")
            self.root.geometry("1280x850")
            self.root.minsize(1024, 680)
            self._apply_theme("light")
            self.root.configure(fg_color=T["page"])
            self.menu_expanded = True
            self._compact_header: bool | None = None
            self.sidebar = ctk.CTkFrame(root, fg_color=T["sidebar"], corner_radius=0, width=220)
            self.sidebar.pack(side="left", fill="y")
            self.sidebar.pack_propagate(False)
            self.main = ctk.CTkFrame(root, fg_color=T["page"], corner_radius=0)
            self.main.pack(side="left", fill="both", expand=True)
            self.header = ctk.CTkFrame(self.main, fg_color="transparent", corner_radius=0)
            self.header.pack(side="top", fill="x", padx=24, pady=(16, 12))
            self.header.grid_columnconfigure(1, weight=1)
            self.burger = button(self.header, "☰", self._toggle_menu, width=34)
            self.burger.grid(row=0, column=0, sticky="w", padx=(0, 14))
            brand_bar = ctk.CTkFrame(self.header, fg_color="transparent")
            brand_bar.grid(row=0, column=1, sticky="w", padx=(0, 12))
            self.brand = ctk.CTkLabel(brand_bar, text="PEA-i", font=FONTS["h1"], text_color=T["ink"], anchor="w")
            self.brand.pack(side="left")
            self.brand_caption = ctk.CTkLabel(brand_bar, text="Universidad Popular del Cesar", font=FONTS["small"],
                                              text_color=T["muted"], anchor="w")
            self.brand_caption.pack(side="left", padx=(14, 0), pady=(6, 0))
            self.header_actions = ctk.CTkFrame(self.header, fg_color="transparent")
            self.header_actions.grid(row=0, column=2, sticky="e")
            self._menu()
            self.save_button = button(self.header_actions, "Guardar como…", lambda: self.save(),
                                      variant="primary", width=136, height=38)
            self.save_button.pack(side="left", padx=(4, 0))
            self.header.bind("<Configure>", self._header_width)
            ctk.CTkFrame(self.main, fg_color=T["line"], height=1, corner_radius=0).pack(side="top", fill="x")
            self.status_bar = ctk.CTkFrame(self.main, fg_color=T["surface"], corner_radius=0)
            self.status_bar.pack(side="bottom", fill="x")
            ctk.CTkFrame(self.status_bar, fg_color=T["line"], height=1, corner_radius=0).pack(fill="x")
            # Archivo abierto y su estado a la derecha; el último aviso a la izquierda.
            self.workspace_state = Chip(self.status_bar, "Guardado", "muted")
            self.workspace_state.pack(side="right", padx=(8, 24), pady=6)
            self.workspace_label = ctk.CTkLabel(self.status_bar, textvariable=self.workspace, font=FONTS["caption"],
                                               text_color=T["ink"], anchor="e")
            self.workspace_label.pack(side="right")
            self.status_label = ctk.CTkLabel(self.status_bar, textvariable=self.status, font=FONTS["caption"],
                                            text_color=T["muted"], anchor="w", justify="left")
            self.status_label.pack(side="left", fill="x", expand=True, padx=(24, 16), pady=8)
            self.status_bar.bind("<Configure>", lambda event: self.status_label.configure(
                wraplength=max(200, round(event.width / scale_of(self.status_bar)) - 360)))

            sidebar_footer = ctk.CTkFrame(self.sidebar, fg_color="transparent")
            sidebar_footer.pack(side="bottom", fill="x", padx=20, pady=(12, 20))
            self.theme_button = ctk.CTkSwitch(
                sidebar_footer, text="Modo oscuro", variable=self.appearance, onvalue="dark", offvalue="light",
                command=lambda: self.set_appearance(self.appearance.get()), font=FONTS["small"],
                text_color=T["sidebar_fg"], fg_color=T["thumb"], progress_color=T["primary"],
                button_color=T["sidebar_fg"], button_hover_color=T["sidebar_muted"])
            self.theme_button.pack(anchor="w")
            self.sidebar_items = ctk.CTkFrame(self.sidebar, fg_color=T["sidebar"], corner_radius=0)
            self.sidebar_items.pack(fill="both", expand=True, pady=(12, 0))
            self.content = ctk.CTkFrame(self.main, fg_color=T["page"], corner_radius=0)
            self.content.pack(side="top", fill="both", expand=True)
            self.content.rowconfigure(0, weight=1)
            self.content.columnconfigure(0, weight=1)

            self.loading_layer = ctk.CTkFrame(root, fg_color=T["page"], corner_radius=0)
            self.loading = Card(self.loading_layer)
            ctk.CTkLabel(self.loading.body, text="Cargando datos UPC…", font=FONTS["h2"],
                         text_color=T["ink"]).pack(padx=24, pady=(12, 16))
            self.loading.progress = ctk.CTkProgressBar(self.loading.body, mode="indeterminate", width=280,
                                                     fg_color=T["chart_track"], progress_color=T["primary"])
            self.loading.progress.pack(padx=24, pady=(0, 12))
            if load_on_start:
                # La primera pintura incluye el aviso de carga; las páginas se construyen detrás.
                self._set_loading(True)
                self.root.deiconify()
                self.root.update_idletasks()

            self.pages: dict[str, tk.Widget] = {}
            self.dashboard = ctk.CTkFrame(self.content, fg_color=T["page"], corner_radius=0)
            self._build_dashboard()
            self.graphs = ctk.CTkFrame(self.content, fg_color=T["page"], corner_radius=0)
            self._build_graphs()
            self.entity_tabs: dict[str, EntityTab] = {}
            for kind in ("grupos", "investigadores", "productos", "planes"):
                self.entity_tabs[kind] = EntityTab(self.content, self, kind)
            relations_host = ctk.CTkFrame(self.content, fg_color=T["page"], corner_radius=0)
            ctk.CTkLabel(relations_host, text="Relaciones", font=FONTS["h1"], text_color=T["ink"],
                         anchor="w").pack(fill="x", padx=24, pady=(20, 0))
            ctk.CTkLabel(relations_host, text="Vínculos entre investigadores, grupos y productos.", font=FONTS["small"],
                         text_color=T["muted"], anchor="w").pack(fill="x", padx=24, pady=(2, 12))
            relations = self.relations_view = ctk.CTkTabview(
                relations_host, fg_color=T["page"], text_color=T["ink"],
                segmented_button_font=FONTS["body"], segmented_button_fg_color=T["field"],
                segmented_button_selected_color=T["primary_soft"], segmented_button_selected_hover_color=T["hover"],
                segmented_button_unselected_color=T["field"], segmented_button_unselected_hover_color=T["hover"])
            relations.pack(fill="both", expand=True, padx=16, pady=(0, 10))
            self.relation_tabs: dict[str, RelationTab] = {}
            for kind, title in (("membresias", "Integrantes"), ("autorias", "Autorías"),
                                ("grupos_productos", "Grupos / productos")):
                tab = RelationTab(relations.add(title), self, kind)
                self.relation_tabs[kind] = tab
                tab.pack(fill="both", expand=True)
            self.queue_tab = ctk.CTkFrame(self.content, fg_color=T["page"], corner_radius=0)
            self._build_queue()
            self.model_page = ctk.CTkFrame(self.content, fg_color=T["page"], corner_radius=0)
            self._build_model_page()

            self.nav_buttons = {}
            self.nav_markers = {}
            self.nav_rows = {}
            self._nav_font = FONTS["body"]
            self._nav_active_font = ctk.CTkFont(family=FAMILY, size=FONTS["body"].cget("size"), weight="bold")
            for group, items in (
                    ("Análisis", (("dashboard", "Dashboard", self.dashboard),
                                  ("graphs", "Gráficos", self.graphs))),
                    ("Registros", (("grupos", "Grupos", self.entity_tabs["grupos"]),
                                   ("investigadores", "Investigadores", self.entity_tabs["investigadores"]),
                                   ("productos", "Productos", self.entity_tabs["productos"]),
                                   ("planes", "Planes", self.entity_tabs["planes"]),
                                   ("relaciones", "Relaciones", relations_host))),
                    ("Modelo", (("modelo", "Modelo 2024", self.model_page),)),
                    ("Seguimiento", (("cola", "Cola de revisión", self.queue_tab),))):
                ctk.CTkLabel(self.sidebar_items, text=group, font=FONTS["caption"],
                             text_color=T["sidebar_muted"], anchor="w").pack(fill="x", padx=20, pady=(12, 2))
                for page_id, title, page in items:
                    self.pages[page_id] = page
                    row = ctk.CTkFrame(self.sidebar_items, fg_color="transparent", corner_radius=0)
                    row.pack(fill="x", padx=10, pady=1)
                    control = button(row, title, lambda p=page_id: self._show_page(p), variant="ghost",
                                     width=0, font=self._nav_font)
                    control.pack(fill="x", padx=(8, 0))
                    self.nav_buttons[page_id] = control
                    self.nav_rows[page_id] = row
                    self.nav_markers[page_id] = ctk.CTkFrame(row, width=3, height=24, corner_radius=0,
                                                           fg_color=T["primary"])
            self.queue_badge = Chip(self.nav_rows["cola"], "", "pending")
            self.queue_badge.bind("<Button-1>", lambda _event: self._show_page("cola"))
            for page in self.pages.values():
                page.grid(row=0, column=0, sticky="nsew")

            self.refresh()
            self._show_page("dashboard")
            self.root.protocol("WM_DELETE_WINDOW", self.close)
            self._startup_timer = self.root.after(80, self.start_choice) if load_on_start else None
            if load_on_start:
                self._set_loading(True)
            self.root.bind("<Destroy>", self._destroyed, add="+")

        def _destroyed(self, event: tk.Event) -> None:
            if event.widget is self.root:
                for timer in (self._year_timer, self._startup_timer):
                    if timer:
                        self.root.after_cancel(timer)

        kit = SimpleNamespace(T=T, color=color, FONTS=FONTS, button=button, Card=Card, Kpi=Kpi, Chip=Chip,
                              MenuButton=MenuButton, SearchField=SearchField, Pager=Pager,
                              EmptyState=EmptyState, Dialog=Dialog, DataTable=DataTable,
                              ChartCanvas=ChartCanvas, IndicatorPanel=IndicatorPanel, palette=palette,
                              RecordPicker=RecordPicker, ReferenceInput=ReferenceInput, ChoiceDialog=ChoiceDialog,
                              RecordDialog=RecordDialog, CategoryInput=CategoryInput, CategoryPicker=CategoryPicker,
                              UrlDialog=UrlDialog)

        @property
        def _last_dialog(self) -> Dialog | None:
            return dialog_state["last"]

        def _dialog_parent(self, parent: Any = None) -> Any:
            if parent is not None:
                return parent
            holder = self.root.grab_current()
            return holder.winfo_toplevel() if holder is not None else self.root

        def _message(self, title: str, text: str, buttons: list[tuple[str, str, Any, str]], default: Any,
                     parent: Any = None, tone: str = "ink") -> Any:
            """Diálogo de aviso o pregunta; buttons = (clave, rótulo, valor devuelto, variante), el primero a la derecha."""
            dialog = Dialog(self._dialog_parent(parent), title, 460)
            ctk.CTkLabel(dialog.body, text=title, font=FONTS["h2"], text_color=T[tone], anchor="w").pack(
                fill="x", pady=(0, 8))
            if text.count("\n") > 12 or len(text) > 900:
                dialog.text_widget = ctk.CTkTextbox(dialog.body, width=540, height=280, wrap="word", font=FONTS["small"],
                                                    fg_color=T["surface"], text_color=T["ink"],
                                                    border_width=1, border_color=T["line"])
                dialog.text_widget.insert("1.0", text)
                dialog.text_widget.configure(state="disabled")
                dialog.text_widget.pack(fill="both", expand=True)
            else:
                ctk.CTkLabel(dialog.body, text=text, font=FONTS["body"], text_color=T["ink"], anchor="w",
                             justify="left", wraplength=420).pack(fill="x")
            result = {"value": default}

            def choose(value: Any) -> None:
                result["value"] = value
                dialog.close()
            for key, label, value, variant in buttons:
                widget = button(dialog.footer, label, lambda chosen=value: choose(chosen), variant=variant, width=112)
                widget.pack(side="right", padx=(8, 0))
                dialog.buttons[key] = widget
            dialog.bind("<Return>", lambda _event: dialog.buttons[buttons[0][0]].invoke())
            dialog.present()
            return result["value"]

        def inform(self, title: str, text: str, parent: Any = None) -> None:
            self._message(title, text, [("accept", "Aceptar", None, "primary")], None, parent)

        def alert(self, title: str, text: str, parent: Any = None) -> None:
            self._message(title, text, [("accept", "Aceptar", None, "primary")], None, parent, tone="error")

        def confirm(self, title: str, text: str, accept: str = "Sí", cancel: str = "No", parent: Any = None) -> bool:
            return self._message(title, text, [("accept", accept, True, "primary"),
                                               ("cancel", cancel, False, "secondary")], False, parent)

        def decide(self, title: str, text: str, yes: str = "Guardar", no: str = "No guardar",
                   cancel: str = "Cancelar", parent: Any = None) -> bool | None:
            """Sí, no o cancelar (None), como una pregunta de tres opciones."""
            return self._message(title, text, [("yes", yes, True, "primary"), ("no", no, False, "secondary"),
                                               ("cancel", cancel, None, "secondary")], None, parent)

        def ask_text(self, title: str, prompt: str, initial: str = "", parent: Any = None) -> str | None:
            dialog = Dialog(self._dialog_parent(parent), title, 520)
            ctk.CTkLabel(dialog.body, text=prompt, font=FONTS["body"], text_color=T["ink"], anchor="w",
                         justify="left", wraplength=470).pack(fill="x", pady=(0, 8))
            dialog.entry = ctk.CTkEntry(dialog.body, height=36, corner_radius=8, font=FONTS["body"],
                                        fg_color=T["field"], border_color=T["line_strong"], text_color=T["ink"])
            dialog.entry.insert(0, initial)
            dialog.entry.pack(fill="x")
            result: dict[str, str | None] = {"value": None}

            def accept() -> None:
                result["value"] = dialog.entry.get().strip()
                dialog.close()
            dialog.buttons["accept"] = button(dialog.footer, "Aceptar", accept, variant="primary", width=112)
            dialog.buttons["accept"].pack(side="right", padx=(8, 0))
            dialog.buttons["cancel"] = button(dialog.footer, "Cancelar", dialog.close, width=112)
            dialog.buttons["cancel"].pack(side="right", padx=(8, 0))
            dialog.bind("<Return>", lambda _event: accept())
            dialog.present(focus=dialog.entry)
            return result["value"]

        def _show_page(self, page_id: str) -> None:
            self.active_page = page_id
            for key, control in self.nav_buttons.items():
                selected = key == page_id
                control.configure(fg_color=T["sidebar_hover"] if selected else "transparent",
                                  font=self._nav_active_font if selected else self._nav_font)
                marker = self.nav_markers[key]
                if selected:
                    marker.place(x=0, rely=0.15, relheight=0.7)
                else:
                    marker.place_forget()
            self.pages[page_id].tkraise()
            if page_id == "graphs":
                self._redraw_charts()

        def _header_width(self, event: Any) -> None:
            compact = event.width < round(920 * scale_of(self.header))
            if compact != self._compact_header:
                self._compact_header = compact
                self.header_actions.grid(row=1 if compact else 0, column=0 if compact else 2,
                                         columnspan=3 if compact else 1, sticky="e",
                                         pady=(12, 0) if compact else 0)

        def _toggle_menu(self) -> None:
            if self.menu_expanded:
                self.sidebar.pack_forget()
            else:
                self.sidebar.pack(side="left", fill="y", before=self.main)
            self.menu_expanded = not self.menu_expanded

        def _show_status(self) -> Dialog:
            return self.show_about("Estado de datos")

        def _apply_theme(self, theme: str) -> None:
            # Los selectores de archivo Tk crean controles ttk aunque la aplicación use CustomTkinter.
            style = ttk.Style()
            if "clam" in style.theme_names():
                style.theme_use("clam")
            colors = LIGHT_THEME if theme == "light" else DARK_THEME
            self.theme = theme
            ctk.set_appearance_mode(theme)
            self.colors = colors
            c = colors

            def flat(name: str, background: str, border: str | None = None, **options: Any) -> None:
                """clam dibuja biseles con light/darkcolor; igualarlos al fondo deja el control plano."""
                style.configure(name, background=background, bordercolor=border or background,
                                lightcolor=background, darkcolor=background, **options)

            self.root.configure(background=c["page"])
            for pattern, value in (("*Toplevel.background", c["page"]),
                                   ("*TCombobox*Listbox.background", c["surface"]),
                                   ("*TCombobox*Listbox.foreground", c["ink"]),
                                   ("*TCombobox*Listbox.selectBackground", c["primary_soft"]),
                                   ("*TCombobox*Listbox.selectForeground", c["ink"]),
                                   ("*TCombobox*Listbox.font", (BODY_FONT, 10))):
                self.root.option_add(pattern, value)
            style.configure(".", font=(BODY_FONT, 10), background=c["page"], foreground=c["ink"],
                            troughcolor=c["page"], focuscolor=c["primary"], selectbackground=c["primary_soft"],
                            selectforeground=c["ink"], insertcolor=c["ink"])
            style.configure("TFrame", background=c["page"])
            style.configure("TLabel", background=c["page"], foreground=c["ink"])
            # Botones: secundarios con borde visible en ambos temas.
            flat("TButton", c["surface"], c["line_strong"], foreground=c["ink"], padding=(12, 6), relief="solid", width=-6)
            style.map("TButton", background=[("disabled", c["page"]), ("pressed", c["line"]), ("active", c["hover"])],
                      lightcolor=[("disabled", c["page"]), ("pressed", c["line"]), ("active", c["hover"])],
                      darkcolor=[("disabled", c["page"]), ("pressed", c["line"]), ("active", c["hover"])],
                      bordercolor=[("disabled", c["line"]), ("active", c["primary"])],
                      foreground=[("disabled", c["unknown"])])
            # Campos con anillo de foco.
            for name in ("TEntry", "TCombobox"):
                style.configure(name, fieldbackground=c["field"], foreground=c["ink"], bordercolor=c["line_strong"],
                                lightcolor=c["field"], darkcolor=c["field"], padding=(7, 5), arrowcolor=c["muted"],
                                background=c["field"], insertcolor=c["ink"])
                # Una entrada de solo lectura se distingue de una editable; una lista desplegable no.
                readonly = c["page"] if name == "TEntry" else c["field"]
                style.map(name, bordercolor=[("focus", c["primary"])], lightcolor=[("focus", c["primary"])],
                          fieldbackground=[("disabled", c["page"]), ("readonly", readonly)],
                          foreground=[("disabled", c["muted"]), ("readonly", c["ink"])],
                          background=[("active", c["hover"]), ("pressed", c["line"])],
                          arrowcolor=[("disabled", c["line_strong"]), ("active", c["primary"])],
                          selectbackground=[("readonly", c["field"]), ("!focus", c["field"])],
                          selectforeground=[("readonly", c["ink"]), ("!focus", c["ink"])])
            # Barras de desplazamiento finas, sin flechas.
            for orient, sticky in (("Vertical", "ns"), ("Horizontal", "ew")):
                style.layout(f"{orient}.TScrollbar", [(f"{orient}.Scrollbar.trough", {"sticky": sticky, "children": [
                    (f"{orient}.Scrollbar.thumb", {"expand": "1", "sticky": "nswe"})]})])
                # El selector Tk de archivos también usa estas barras.
                flat(f"{orient}.TScrollbar", c["thumb"], troughcolor=c["surface"], gripcount=0,
                     arrowsize=9, width=9, relief="flat", borderwidth=0)
                style.configure(f"{orient}.TScrollbar", troughcolor=c["surface"], bordercolor=c["surface"])
                style.map(f"{orient}.TScrollbar", background=[("active", c["thumb_active"]), ("pressed", c["thumb_active"])],
                          lightcolor=[("active", c["thumb_active"])], darkcolor=[("active", c["thumb_active"])])

        @property
        def canvases(self) -> list[Any]:
            """Gráficos y tablas dibujados en lienzos, vigentes en esta sesión."""
            return list(canvas_registry)

        def set_appearance(self, mode: str) -> None:
            """Claro u oscuro en CustomTkinter y en los lienzos."""
            self._apply_theme(mode)
            for detail in tuple(self.details):
                detail.refresh()
            for canvas in self.canvases:
                canvas.redraw()
            self.appearance.set(mode)

        def _toggle_theme(self) -> None:
            self.set_appearance("dark" if self.theme == "light" else "light")

        def _menu(self) -> None:
            self.menus = {}
            for key, label, width, items in (
                ("archivo", "Archivo", 92, (("Iniciar vacío", "new_workspace"),
                                            ("Abrir archivo de datos...", "open_file"),
                                            ("Abrir datos UPC", "open_default"), ("Guardar", "save"),
                                            ("Guardar como...", "save_as"),
                                            ("Exportar para hoja de cálculo...", "export_spreadsheet"),
                                            ("Salir", "close"))),
                ("importar", "Importar", 112, (("Hoja de cálculo (CSV/Excel)...", "import_csv"),
                                             ("URL pública...", "import_url"),
                                             ("PDF de texto GrupLAC/CvLAC...", "import_pdf"),
                                             ("Documento Word...", "import_docx"))),
                ("editar", "Editar", 84, (("Deshacer", "undo"), ("Vaciar historial", "clear_history"))),
                ("ayuda", "Ayuda", 84, (("Estado de datos", "_show_status"), ("Acerca de PEA-i", "_show_about"))),
            ):
                actions = [(text, lambda method=method: getattr(self, method)()) for text, method in items]
                self.menus[key] = MenuButton(self.header_actions, label, actions, width=width)
                self.menus[key].pack(side="left", padx=(0, 6))

            def shortcut(action: str) -> Callable[[Any], str]:
                def invoke(_event: Any) -> str:
                    getattr(self, action)()
                    return "break"
                return invoke
            for sequence, action in (("<Control-z>", "undo"), ("<Control-s>", "save"),
                                     ("<Control-Shift-S>", "save_as"), ("<Control-n>", "new_workspace")):
                self.root.bind(sequence, shortcut(action))

        def _show_about(self) -> Dialog:
            return self.show_about("Acerca de PEA-i")

        def show_about(self, title: str = "Acerca de PEA-i") -> Dialog:
            """Acerca de PEA-i y estado de los datos abiertos: motor, archivo, versiones y conteos."""
            dialog = Dialog(self.root, title, 560)
            ctk.CTkLabel(dialog.body, text="PEA-i", font=FONTS["h1"], text_color=T["ink"], anchor="w").pack(fill="x")
            ctk.CTkLabel(dialog.body, text="Programa Estadístico de Análisis de Investigación · Universidad Popular del Cesar",
                         font=FONTS["small"], text_color=T["muted"], anchor="w", justify="left", wraplength=500).pack(fill="x")
            chips = ctk.CTkFrame(dialog.body, fg_color="transparent")
            chips.pack(fill="x", pady=(8, 12))
            for text, tone in ((f"Motor: {self.backend_name}", "primary"), (f"Esquema {SCHEMA_VERSION}", "muted"),
                               (f"Protocolo {CPP_PROTOCOL}", "muted"),
                               ("Cambios sin guardar", "pending") if self.repo.dirty else ("Guardado", "muted")):
                Chip(chips, text, tone).pack(side="left", padx=(0, 6))
            facts = {"Motor de datos": self.backend_name,
                     "Archivo": str(self.data_file) if self.data_file else "Sin archivo (Guardar como crea uno)",
                     "Esquema de datos": SCHEMA_VERSION, "Protocolo C++": CPP_PROTOCOL}
            for kind in ALL_FIELDS:
                facts[TITLES[kind]] = number(self.repo.page(kind, limit=0)["total"])
            facts["Cola de revisión"] = number(self.repo.queue_size())
            facts["Acciones para deshacer"] = number(self.repo.history_size())
            card = Card(dialog.body, title="Datos abiertos", bg_color=T["page"])
            card.pack(fill="x")
            card.body.grid_columnconfigure(1, weight=1)
            for line, (label, value) in enumerate(facts.items()):
                ctk.CTkLabel(card.body, text=label, font=FONTS["small"], text_color=T["muted"], anchor="w").grid(
                    row=line, column=0, sticky="w", padx=(0, 16), pady=2)
                ctk.CTkLabel(card.body, text=value, font=FONTS["body"], text_color=T["ink"], anchor="w", justify="left",
                             wraplength=330).grid(row=line, column=1, sticky="ew", pady=2)
            ctk.CTkLabel(dialog.body, text="Último aviso: " + self.status.get(), font=FONTS["caption"],
                         text_color=T["muted"], anchor="w", justify="left", wraplength=500).pack(fill="x", pady=(10, 0))
            dialog.facts = facts
            button(dialog.footer, "Cerrar", dialog.close, variant="primary", width=112).pack(side="right")
            dialog.present(wait=False)
            return dialog

        def _build_dashboard(self) -> None:
            self.dashboard_page = ctk.CTkScrollableFrame(self.dashboard, fg_color=T["page"], corner_radius=0,
                                                        scrollbar_button_color=T["line_strong"],
                                                        scrollbar_button_hover_color=T["muted"])
            self.dashboard_page.pack(fill="both", expand=True)
            # Conserva el acceso al desplazamiento que usaban los flujos del Dashboard.
            self.dashboard_page.scroll_canvas = self.dashboard_page._parent_canvas
            self.dashboard_page.scroll_canvas.configure(yscrollcommand=lambda first, last: set_scrollbar(
                self.dashboard_page._scrollbar, first, last))
            body = ctk.CTkFrame(self.dashboard_page, fg_color="transparent")
            body.pack(fill="both", expand=True, padx=24, pady=20)
            title_row = ctk.CTkFrame(body, fg_color="transparent")
            title_row.pack(fill="x")
            ctk.CTkLabel(title_row, text="Panorama de investigación", font=FONTS["h1"],
                         text_color=T["ink"], anchor="w").pack(side="left")
            # Población activa de toda la institución, independiente de los filtros.
            self.institution_chips = [Chip(title_row, "", "primary") for _ in range(2)]
            for chip in reversed(self.institution_chips):
                chip.pack(side="right", padx=(6, 0))
            self.context = tk.StringVar()
            self.context_label = ctk.CTkLabel(body, textvariable=self.context, font=FONTS["small"],
                                             text_color=T["muted"], anchor="w", justify="left", wraplength=700)
            self.context_label.pack(fill="x", pady=(2, 14))
            filter_card = Card(body)
            filter_card.pack(fill="x", pady=(0, 14))
            filters = filter_card.body
            for col in range(2):
                filters.grid_columnconfigure(col, weight=1, uniform="filter")
            self.view = tk.StringVar(value="Todos")
            self.scope_id = tk.StringVar()
            self._scope_kind: str | None = None
            self.period = tk.StringVar(value="Todos")
            self.start_year = tk.StringVar()
            self.end_year = tk.StringVar()
            self.validation = tk.StringVar(value="Todos")

            def field_box(col: int, row: int, label: str, span: int = 1) -> Any:
                frame = ctk.CTkFrame(filters, fg_color="transparent")
                frame.grid(row=row, column=col, columnspan=span, sticky="ew", padx=(0, 16) if col == 0 and span == 1 else 0,
                           pady=4)
                ctk.CTkLabel(frame, text=label, font=FONTS["caption"], text_color=T["muted"],
                             anchor="w", height=20).pack(fill="x", pady=(0, 4))
                return frame

            segmented_style = {"height": 34, "corner_radius": 8, "font": FONTS["body"], "text_color": T["ink"],
                               "fg_color": T["field"], "selected_color": T["primary_soft"],
                               "selected_hover_color": T["hover"], "unselected_color": T["field"],
                               "unselected_hover_color": T["hover"]}
            self.view_control = ctk.CTkSegmentedButton(
                field_box(0, 0, "Vista", 2), variable=self.view,
                values=["Todos", "Grupo", "Investigador", "Producto"],
                command=lambda _value: self._scope_changed(), **segmented_style)
            self.view_control.pack(fill="x")
            self.scope_box = ReferenceInput(field_box(0, 1, "Nombre / título"), self.repo, "grupos",
                                            variable=self.scope_id, command=self.refresh_dashboard)
            self.scope_box.pack(fill="x")
            self.scope_box.set_enabled(False)
            self.period_box = ctk.CTkOptionMenu(
                field_box(1, 1, "Ventana de observación"), variable=self.period,
                values=["Todos", "Últimos 2 años", "Últimos 5 años", "Personalizado"],
                command=lambda _value: self.refresh_dashboard(), height=34, corner_radius=8,
                font=FONTS["body"], dropdown_font=FONTS["body"], dynamic_resizing=False,
                fg_color=T["field"], button_color=T["field"], button_hover_color=T["hover"],
                text_color=T["ink"], dropdown_fg_color=T["surface"],
                dropdown_text_color=T["ink"], dropdown_hover_color=T["hover"])
            self.period_box.pack(fill="x")
            self.validation_control = ctk.CTkSegmentedButton(
                field_box(0, 2, "Validación automática"), variable=self.validation,
                values=["Todos", "Validado", "Rechazado"], command=lambda _value: self.refresh_dashboard(),
                **segmented_style)
            self.validation_control.pack(fill="x")
            years = field_box(1, 2, "Años: desde / hasta")
            entry_style = {"width": 90, "height": 34, "corner_radius": 8, "font": FONTS["body"],
                           "fg_color": T["field"], "text_color": T["ink"], "border_color": T["line_strong"]}
            self.start_entry = ctk.CTkEntry(years, textvariable=self.start_year, **entry_style)
            self.start_entry.pack(side="left", fill="x", expand=True)
            ctk.CTkLabel(years, text="a", font=FONTS["small"], text_color=T["muted"], width=16).pack(side="left", padx=6)
            self.end_entry = ctk.CTkEntry(years, textvariable=self.end_year, **entry_style)
            self.end_entry.pack(side="left", fill="x", expand=True)
            for variable in (self.start_year, self.end_year):
                variable.trace_add("write", self._year_changed)
            for entry in (self.start_entry, self.end_entry):
                entry.bind("<Return>", lambda _e: self.refresh_dashboard())
            self.filter_error = tk.StringVar()
            self.filter_error_label = ctk.CTkLabel(filters, textvariable=self.filter_error, font=FONTS["small"],
                                                  text_color=T["rejected"], anchor="w", justify="left", wraplength=650)

            cards = ctk.CTkFrame(body, fg_color="transparent")
            cards.pack(fill="x", pady=(0, 8))
            self.card_values = {}
            # La franja superior usa el color de estado que repiten los gráficos.
            for index, (key, title, tone) in enumerate((("total", "Productos", "sidebar"),
                                                       ("validated", "Validados", "primary"),
                                                       ("rejected", "Rechazados", "rejected"),
                                                       ("undated", "Sin año", "unknown"))):
                cards.columnconfigure(index, weight=1, uniform="cards")
                value = tk.StringVar(value="0")
                self.card_values[key] = value
                Kpi(cards, title, tone, value).grid(row=0, column=index, sticky="nsew", padx=(0, 12) if index < 3 else 0)
            spacer = ctk.CTkFrame(body, fg_color="transparent", height=14)
            spacer.pack(fill="x")
            # Solo aparece cuando hay rechazos (sus causas) o datos sin año/tipología.
            self.note = tk.StringVar()
            self.note_label = ctk.CTkLabel(body, textvariable=self.note, font=FONTS["small"],
                                          text_color=T["rejected"], anchor="w", justify="left", wraplength=700)
            self._note_pack = {"fill": "x", "pady": (0, 14), "after": spacer}
            charts = ctk.CTkFrame(body, fg_color="transparent")
            self.dashboard_charts = charts
            charts.pack(fill="x", pady=(0, 18))
            charts.columnconfigure(0, weight=1, uniform="preview")
            charts.columnconfigure(1, weight=1, uniform="preview")
            # Solo aparece con la vista Grupo o Investigador; se ubica antes de los gráficos.
            self.indicator_panel = IndicatorPanel(body, self.show_indicator_details, bg_color=T["page"],
                                                  pack_options={"fill": "x", "pady": (0, 14), "before": charts})
            self.indicator_value = self.indicator_panel.value
            self.indicator_note = self.indicator_panel.note
            self.indicator_canvas = self.indicator_panel.canvas
            self._indicator_result: dict[str, Any] | None = None
            panels = []
            for name in ("dashboard_year_canvas", "dashboard_class_canvas", "dashboard_type_canvas",
                         "dashboard_category_canvas"):
                panel, canvas = self._chart_panel(charts, 260, draw=self._chart_draw(name))
                setattr(self, name, canvas)
                panels.append(panel)
            self.dashboard_chart_panels = tuple(panels)
            self._charts_stacked: bool | None = None
            self._layout_dashboard_charts(False)
            table_heading = ctk.CTkFrame(body, fg_color="transparent")
            table_heading.pack(fill="x", pady=(0, 6))
            ctk.CTkLabel(table_heading, text="Productos de la vista", font=FONTS["h2"],
                         text_color=T["ink"]).pack(side="left")
            button(table_heading, "Ver todas las distribuciones", lambda: self._show_page("graphs"), width=220).pack(side="right")
            columns = ("titulo", "anio", "tipologia", "categoria", "validacion")
            self.stats_table = DataTable(body, columns, headings={field: LABELS[field] for field in columns},
                                         widths={"titulo": 270, "anio": 60, "tipologia": 160, "categoria": 95, "validacion": 110},
                                         anchors={"anio": "e"},
                                         chips={"validacion": lambda value: STATUS_COLORS.get(value.lower(), "unknown"),
                                                "categoria": lambda value: "unknown" if value == "Sin dato" else "primary"},
                                         height_rows=8, open_on_click=True)
            self.stats_table.pack(fill="x")
            self.stats_table.on_open(lambda _iid: self._open_product(None))
            self.stats_pager = Pager(body, self._stats_move)
            self.stats_pager.pack(fill="x", pady=(8, 0))
            self.stats_previous = self.stats_pager.previous
            self.stats_next = self.stats_pager.next
            self.stats_page_label = tk.StringVar()
            self.stats_offset = 0
            self._stats_filter: tuple[Any, ...] | None = None
            self.current_stats: dict[str, Any] = {"por_anio": {}, "por_tipologia": {}, "por_categoria": {},
                                                  "por_validacion": {}, "por_regla": {}}
            body.bind("<Configure>", lambda event: self._dashboard_width(event.width), add="+")

        def _dashboard_width(self, width: int) -> None:
            logical_width = width / scale_of(self.dashboard_page)
            if abs(logical_width - getattr(self, "_dashboard_logical_width", 0)) < 1:
                return
            self._dashboard_logical_width = logical_width
            for label in (self.context_label, self.note_label, self.filter_error_label):
                label.configure(wraplength=max(260, logical_width - 32))
            self._layout_dashboard_charts(logical_width < 850)

        def _layout_dashboard_charts(self, stacked: bool) -> None:
            """Gráficos del Dashboard en una cuadrícula de 2×2, o en una columna si la ventana es estrecha."""
            if stacked == self._charts_stacked:
                return
            self._charts_stacked = stacked
            for index, panel in enumerate(self.dashboard_chart_panels):
                row, column = (index, 0) if stacked else divmod(index, 2)
                panel.grid(row=row, column=column, columnspan=2 if stacked else 1, sticky="nsew",
                           padx=0 if stacked or column else (0, 12), pady=(0, 12))

        def _chart_panel(self, parent: tk.Misc, height: int = 240,
                         draw: Callable[[Any], Any] | None = None) -> tuple[Card, ChartCanvas]:
            panel = Card(parent)
            panel.body.grid_columnconfigure(0, weight=1)
            panel.body.grid_rowconfigure(0, weight=1)
            canvas = ChartCanvas(panel.body, draw=draw, width=280, height=round(height * scale_of(parent)))
            scroll_style = {"fg_color": T["surface"], "button_color": T["line_strong"], "button_hover_color": T["muted"]}
            xscroll = ctk.CTkScrollbar(panel.body, orientation="horizontal", width=0, command=canvas.xview, **scroll_style)
            yscroll = ctk.CTkScrollbar(panel.body, orientation="vertical", height=0, command=canvas.yview, **scroll_style)
            canvas.configure(xscrollcommand=lambda first, last: set_scrollbar(xscroll, first, last, autohide=True),
                             yscrollcommand=lambda first, last: set_scrollbar(yscroll, first, last, autohide=True))
            canvas.grid(row=0, column=0, sticky="nsew")
            yscroll.grid(row=0, column=1, sticky="ns")
            xscroll.grid(row=1, column=0, sticky="ew")
            return panel, canvas

        def _chart_draw(self, name: str) -> Callable[[Any], None]:
            """Dibujo de cada gráfico del dashboard y de Gráficos con los resultados vigentes de los filtros."""
            def stats() -> dict[str, Any]:
                return getattr(self, "current_stats", {})

            def bounds() -> tuple[int | None, int | None]:
                filters = getattr(self, "_stats_filter", None)
                return (filters[2], filters[3]) if filters else (None, None)

            def classes(canvas: Any) -> None:
                counts = class_counts(stats().get("por_categoria", {}))
                draw_donut(canvas, "Productos por clase del modelo 2024", "Clase de la categoría asignada o estimada",
                           [(key, f"{key} · {CLASS_LABELS[key]}" if key in CLASS_LABELS else key, count, class_color(key))
                            for key, count in counts.items()], "productos")

            def categories(canvas: Any, limit: int) -> None:
                draw_treemap(canvas, "Productos por categoría" if limit > 12 else "Categorías más frecuentes",
                             "Área proporcional a la cantidad · color por clase del modelo",
                             [(code, count, class_color(catalog_by_code[code]["clase"]) if code in catalog_by_code
                               else color("class_none")) for code, count in stats().get("por_categoria", {}).items()],
                             limit=limit)

            def validation(canvas: Any) -> None:
                draw_donut(canvas, "Validación automática",
                           "Reglas: año, tipología, autor o grupo activo, URL o DOI válido y fuente",
                           [(state, display_value("validacion", state), count, color(STATUS_COLORS.get(state, "unknown")))
                            for state, count in sorted(stats().get("por_validacion", {}).items(), reverse=True)],
                           "productos")

            def indicator(canvas: Any) -> None:
                draw_indicator(canvas, self._indicator_result)
                scale = scale_of(canvas)
                if self._indicator_result is None:
                    canvas.create_text(round(12 * scale), round(12 * scale), anchor="nw",
                                       text="Aporte de cada término: disponible al elegir un grupo o investigador activo.",
                                       font=chart_font(canvas, 10), fill=color("muted"))
                canvas.configure(scrollregion=(0, 0, max(round(360 * scale), canvas.winfo_width()),
                                               max(round(250 * scale), canvas.winfo_height())))
            drawings: dict[str, Callable[[Any], None]] = {
                "year_canvas": lambda canvas: self._draw_histogram(canvas, stats().get("por_anio", {}), bounds()),
                "trend_canvas": lambda canvas: draw_trend(canvas, stats().get("por_anio", {}), bounds()),
                "type_canvas": lambda canvas: self._draw_distribution(canvas, stats().get("por_tipologia", {}),
                                                                      "Productos por tipología"),
                "dashboard_type_canvas": lambda canvas: self._draw_distribution(
                    canvas, stats().get("por_tipologia", {}), "Tipologías más frecuentes", 4),
                "class_canvas": classes, "dashboard_class_canvas": classes,
                "category_canvas": lambda canvas: categories(canvas, 30),
                "dashboard_category_canvas": lambda canvas: categories(canvas, 12),
                "validation_canvas": validation,
                "graph_profile_canvas": lambda canvas: draw_indicator_profile(canvas, self._indicator_result),
                "graph_indicator_canvas": indicator,
            }
            drawings["dashboard_year_canvas"] = drawings["year_canvas"]
            return drawings[name]

        def _build_graphs(self) -> None:
            body = ctk.CTkFrame(self.graphs, fg_color="transparent")
            body.pack(fill="both", expand=True, padx=24, pady=20)
            heading = ctk.CTkFrame(body, fg_color="transparent")
            heading.pack(fill="x")
            ctk.CTkLabel(heading, text="Distribuciones de productos", font=FONTS["h1"],
                         text_color=T["ink"]).pack(side="left")
            button(heading, "Cambiar filtros en Dashboard", lambda: self._show_page("dashboard"), width=220).pack(side="right")
            self.graph_context_label = ctk.CTkLabel(body, textvariable=self.context, font=FONTS["small"],
                                                    text_color=T["muted"], anchor="w", justify="left", wraplength=700)
            self.graph_context_label.pack(fill="x", pady=(2, 14))
            charts = self.graph_tabs = ctk.CTkTabview(
                body, fg_color=T["page"], text_color=T["ink"], border_width=0, corner_radius=8,
                segmented_button_fg_color=T["field"], segmented_button_selected_color=T["primary_soft"],
                segmented_button_selected_hover_color=T["hover"], segmented_button_unselected_color=T["field"],
                segmented_button_unselected_hover_color=T["hover"])
            charts.pack(fill="both", expand=True)
            # Siete clases de gráfico: histograma, línea, barras, anillo, mosaico, semicírculo y radar.
            for name, title in (("year_canvas", "Por año"), ("trend_canvas", "Tendencia"),
                                ("type_canvas", "Por tipología"), ("class_canvas", "Por clase"),
                                ("category_canvas", "Por categoría"), ("graph_indicator_canvas", "Indicador"),
                                ("validation_canvas", "Por validación")):
                tab = charts.add(title)
                if name == "graph_indicator_canvas":
                    summary = ctk.CTkFrame(tab, fg_color="transparent")
                    summary.pack(fill="x", pady=(0, 10))
                    self.graph_indicator_summary = ctk.CTkLabel(
                        summary, text="", font=FONTS["small"], text_color=T["muted"],
                        anchor="w", justify="left", wraplength=430)
                    self.graph_indicator_summary.pack(side="left", fill="x", expand=True, padx=(0, 16))
                    self.graph_indicator_details = button(summary, "Detalle por subtipo", self.show_indicator_details, width=160)
                    self.graph_indicator_details.pack(side="right")
                    profile, self.graph_profile_canvas = self._chart_panel(
                        tab, 270, draw=self._chart_draw("graph_profile_canvas"))
                    profile.pack(fill="x", pady=(0, 10))
                panel, canvas = self._chart_panel(tab, 250 if name == "graph_indicator_canvas" else 240,
                                                  draw=self._chart_draw(name))
                setattr(self, name, canvas)
                panel.pack(fill="both", expand=True)
            body.bind("<Configure>", lambda event: self.graph_context_label.configure(
                wraplength=max(260, event.width / scale_of(body) - 16)), add="+")

        def _build_model_page(self) -> None:
            self.model_scroll = ctk.CTkScrollableFrame(self.model_page, fg_color=T["page"], corner_radius=0,
                                                       scrollbar_button_color=T["line_strong"],
                                                       scrollbar_button_hover_color=T["muted"])
            self.model_scroll.pack(fill="both", expand=True)
            self.model_scroll._parent_canvas.configure(yscrollcommand=lambda first, last: set_scrollbar(
                self.model_scroll._scrollbar, first, last))
            body = ctk.CTkFrame(self.model_scroll, fg_color="transparent")
            body.pack(fill="both", expand=True, padx=24, pady=20)
            ctk.CTkLabel(body, text="Modelo de medición Minciencias 2024", font=FONTS["h1"],
                         text_color=T["ink"], anchor="w").pack(fill="x", pady=(0, 12))
            formula = Card(body)
            formula.pack(fill="x", pady=(0, 12))
            self.model_formula = ctk.CTkLabel(
                formula.body, text="Ind = 3,7·Nc.Top + 2,3·Nc.A + 0,4·Nc.B + 1,5·AP + 0,5·DPC + "
                                   "1·Fr.A + 0,2·Fr.B + 0,1·Cohe + 0,3·Coop",
                font=FONTS["h2"], text_color=T["ink"], anchor="w", justify="left", wraplength=850)
            self.model_formula.pack(fill="x", pady=(0, 8))
            self.model_definitions = ctk.CTkLabel(
                formula.body,
                text="λ = ln(n / periodo + 1): n es la cantidad contada y periodo la ventana efectiva en años.\n"
                     "Peso global = peso de alta calidad × peso relativo. Valor del subtipo = λ × peso global.\n"
                     "Índice = valor / máximo de la población activa con los mismos filtros (0 si el máximo es 0).",
                font=FONTS["caption"], text_color=T["muted"], anchor="w", justify="left", wraplength=850)
            self.model_definitions.pack(fill="x")
            formula.body.bind("<Configure>", lambda event: [label.configure(
                wraplength=max(260, event.width / scale_of(body) - 8))
                for label in (self.model_formula, self.model_definitions)], add="+")
            classes = ctk.CTkFrame(body, fg_color="transparent")
            classes.pack(fill="x", pady=(0, 12))
            self.model_class_chips = []
            for index, name in enumerate(MODEL_CLASSES):
                count = sum(row["clase"] == name for row in catalog)
                chip = Chip(classes, f"{name} · {decimal(MODEL_COEFFICIENTS[name])} · {count} categorías", "primary")
                chip.grid(row=index // 4, column=index % 4, sticky="w", padx=(0, 8), pady=(0, 6))
                self.model_class_chips.append(chip)
            self.model_catalog = CategoryCatalog(body)
            self.model_catalog.pack(fill="both", expand=True)
            self.model_table = self.model_catalog.table

        def _build_queue(self) -> None:
            body = ctk.CTkFrame(self.queue_tab, fg_color="transparent")
            body.pack(fill="both", expand=True, padx=24, pady=20)
            heading = ctk.CTkFrame(body, fg_color="transparent")
            heading.pack(fill="x")
            ctk.CTkLabel(heading, text="Revisión de productos", font=FONTS["h1"], text_color=T["ink"]).pack(side="left")
            button(heading, "Añadir producto…", self.enqueue, variant="primary", width=150).pack(side="right")
            button(heading, "Enviar rechazados", self.enqueue_rejected, width=150).pack(side="right", padx=(0, 8))
            caption = ctk.CTkLabel(body, text="La validación es automática. Los productos rechazados se revisan aquí en orden "
                                   "de llegada: corrija el producto o deje una nota que explique por qué no es posible.",
                                   font=FONTS["small"], text_color=T["muted"], anchor="w", justify="left", wraplength=680)
            caption.pack(fill="x", pady=(2, 14))
            # La siguiente revisión, destacada en una tarjeta con la franja de «pendiente».
            front = Card(body, title="Siguiente en la cola", stripe="pending")
            front.pack(fill="x")
            front.body.grid_columnconfigure(0, weight=1)
            self.queue_count = tk.StringVar()
            ctk.CTkLabel(front.body, textvariable=self.queue_count, font=FONTS["caption"], text_color=T["muted"],
                         anchor="w").grid(row=0, column=0, sticky="w")
            self.queue_front_label = tk.StringVar()
            self.queue_front_widget = ctk.CTkLabel(front.body, textvariable=self.queue_front_label, font=FONTS["h2"],
                                                   text_color=T["ink"], anchor="w", justify="left", wraplength=560)
            self.queue_front_widget.grid(row=1, column=0, sticky="ew", pady=(2, 8))
            self.queue_reasons = ctk.CTkFrame(front.body, fg_color="transparent")
            self.queue_reasons.grid(row=2, column=0, sticky="w")
            self.queue_reason_chips: list[Chip] = []
            controls = ctk.CTkFrame(front.body, fg_color="transparent")
            controls.grid(row=3, column=0, sticky="w", pady=(8, 0))
            self.review_button = button(controls, "Revisar siguiente", self.process_queue, variant="primary", width=150)
            self.review_button.pack(side="left")
            self.discard_button = button(controls, "Descartar siguiente", self.discard_queue, variant="danger", width=160)
            self.discard_button.pack(side="left", padx=(8, 0))
            body.bind("<Configure>", lambda event: [label.configure(wraplength=max(240, event.width / scale_of(body) - 40))
                                                    for label in (caption, self.queue_front_widget)], add="+")
            ctk.CTkLabel(body, text="En cola", font=FONTS["h2"], text_color=T["ink"], anchor="w").pack(fill="x", pady=(18, 6))
            # El ID del trabajo es interno (identifica la fila); no se muestra.
            self.queue_table = DataTable(body, ("id", "producto_id", "motivo", "creado"), height_rows=8,
                                         headings={"id": "ID", "producto_id": "Producto", "motivo": "Motivo",
                                                   "creado": "Recibido"},
                                         widths={"id": 100, "producto_id": 340, "motivo": 280, "creado": 150})
            self.queue_table.configure(displaycolumns=("producto_id", "motivo", "creado"))
            self.queue_table.empty_text = "No hay productos pendientes de revisión."
            self.queue_table.on_open(self._open_queued_product)
            self.queue_pager = Pager(body, self._queue_move)
            self.queue_pager.pack(side="bottom", fill="x", pady=(10, 0))
            self.queue_table.pack(fill="both", expand=True)
            self.queue_previous, self.queue_next = self.queue_pager.previous, self.queue_pager.next
            self.queue_page_label = tk.StringVar()
            self.queue_offset = 0
            self._queue_jobs: dict[str, dict[str, str]] = {}

        def _open_queued_product(self, job_id: str | None = None) -> None:
            selection = (job_id,) if job_id else self.queue_table.selection()
            job = self._queue_jobs.get(selection[0]) if selection else None
            if job:
                self.open_detail("productos", job["producto_id"])

        def _queue_move(self, direction: int) -> None:
            self.queue_offset = max(0, self.queue_offset + direction * PAGE_SIZE)
            self.refresh_queue()

        def refresh_queue(self) -> None:
            page = self.repo.queue_page(self.queue_offset, PAGE_SIZE)
            if self.queue_offset and self.queue_offset >= page["total"]:
                self.queue_offset = max(0, (page["total"] - 1) // PAGE_SIZE * PAGE_SIZE)
                page = self.repo.queue_page(self.queue_offset, PAGE_SIZE)
            self._queue_jobs = {job["id"]: job for job in page["rows"]}
            self.queue_table.delete(*self.queue_table.get_children())
            for index, job in enumerate(page["rows"]):
                self.queue_table.insert("", "end", iid=job["id"],
                                        values=(job["id"], self.reference_label("productos", job["producto_id"]), reason_text(job["motivo"], short=True), job["creado"].replace("T", " ")),
                                        tags=("stripe",) if index % 2 else ())
            self.queue_count.set(f"Pendientes: {number(page['total'])}")
            self.queue_badge.set(number(page["total"]) if page["total"] else "", "pending")
            if page["total"]:
                self.queue_badge.place(relx=0.96, rely=0.5, anchor="e")
                self.queue_badge.lift()
            else:
                self.queue_badge.place_forget()
            self.queue_pager.show(self.queue_offset, len(page["rows"]), page["total"])
            self.queue_page_label.set(self.queue_pager.label.cget("text"))
            front = self.repo.queue_front()
            self.queue_front_label.set(self.reference_label("productos", front["producto_id"]) if front
                                       else "Cola vacía. Añada un producto para revisarlo.")
            self.queue_reason_chips = reason_chips(self.queue_reasons, front["motivo"] if front else None)
            # Un CTkFrame vacío conserva su alto por defecto (200 px): sin motivo no ocupa lugar.
            self.queue_reasons.grid() if front else self.queue_reasons.grid_remove()
            self.review_button.configure(state="normal" if front else "disabled",
                                         fg_color=T["button"] if front else T["field"], border_width=0 if front else 1,
                                         border_color=T["line_strong"], text_color_disabled=T["disabled"])
            self.discard_button.configure(state="normal" if front else "disabled",
                                          fg_color=T["rejected"] if front else T["field"], border_width=0 if front else 1,
                                          border_color=T["line_strong"], text_color_disabled=T["disabled"])

        def enqueue(self, product_id: str | None = None, *, parent: tk.Misc | None = None) -> None:
            parent = parent or self.root
            if not self.repo.page("productos", limit=0)["total"]:
                self.inform("Sin productos", "Cree un producto primero", parent=self.root)
                return
            product_id = product_id or self.pick_record("productos", parent=parent)
            if not product_id:
                return
            reason = self.ask_text("Motivo", "Motivo de la revisión:", parent=parent)
            if reason is None:
                return
            self.mutate(lambda: self.repo.enqueue_review(product_id.strip(), reason), changed={"cola"})

        def enqueue_rejected(self) -> None:
            count = self.mutate(self.repo.enqueue_rejected, changed={"cola"})
            if count == 0:
                self.status.set("No hay productos rechazados fuera de la cola")
            elif count:
                self.status.set(f"{count} productos rechazados enviados a revisión. Guarde para conservarlos.")

        def process_queue(self) -> None:
            job = self.repo.queue_front()
            if not job:
                self.inform("Cola", "No hay revisiones pendientes", parent=self.root)
                return
            dialog = ReviewDialog(self.root, self, job)
            self.refresh({"productos", "cola"})
            if dialog.result:
                self.status.set(f"Revisión cerrada; el producto quedó {dialog.result['validacion']}. Guarde para conservarlo.")

        def discard_queue(self) -> None:
            job = self.repo.queue_front()
            if job and self.confirm("Descartar siguiente", f"¿Descartar revisión de {self.reference_label('productos', job['producto_id'])}?", parent=self.root):
                self.mutate(self.repo.discard_review, changed={"cola"})

        def _stats_move(self, direction: int) -> None:
            self.stats_offset = max(0, self.stats_offset + direction * PAGE_SIZE)
            self.refresh_dashboard()

        def _redraw_charts(self) -> None:
            for name in ("year_canvas", "dashboard_year_canvas", "trend_canvas", "type_canvas", "dashboard_type_canvas",
                         "class_canvas", "dashboard_class_canvas", "category_canvas", "dashboard_category_canvas",
                         "validation_canvas", "graph_profile_canvas", "graph_indicator_canvas"):
                canvas = getattr(self, name, None)
                if canvas is not None:
                    canvas.redraw()
            if getattr(self, "_indicator_result", None):
                self.indicator_canvas.redraw()

        def _draw_histogram(self, canvas: Any, values: dict[str, int],
                            bounds: tuple[int | None, int | None] = (None, None)) -> None:
            c, scale = palette(), scale_of(canvas)

            def px(value: float) -> int:
                return round(value * scale)
            canvas.delete("all")
            numeric = {int(key): value for key, value in values.items() if key.isdigit()}
            missing = values.get("Sin dato", 0)
            total = sum(values.values())
            canvas.create_text(px(12), px(14), anchor="w", text="Productos por año · histograma",
                               font=chart_font(canvas, 10, "strong"), fill=c["ink"])
            canvas.create_text(px(12), px(34), anchor="w", text=f"{total} productos · {missing} sin año",
                               font=chart_font(canvas, 9), fill=c["muted"])
            height = max(px(170), canvas.winfo_height())
            if not numeric:
                canvas.create_text(px(12), px(80), anchor="w",
                                   text="Sin años registrados" if total else "Sin productos para estos filtros",
                                   font=chart_font(canvas, 10), fill=c["muted"])
                canvas.configure(scrollregion=(0, 0, max(canvas.winfo_width(), px(280)), height))
                return
            first = bounds[0] if bounds[0] is not None else min(numeric)
            last = bounds[1] if bounds[1] is not None else max(numeric)
            cutoff = date.today().year - 20
            bins = []
            if first <= cutoff:
                count = sum(value for year, value in numeric.items() if first <= year <= min(cutoff, last))
                bins.append(("20 años\no más", count, f"older:{cutoff}"))
            bins.extend((str(year), numeric.get(year, 0), f"year:{year}") for year in range(max(first, cutoff + 1), last + 1))
            width = max(canvas.winfo_width(), px(320))
            left, right, top, bottom = px(46), width - px(14), px(74), height - px(40)
            canvas.create_text(px(12), px(52), anchor="w", text=f"Columna clara: 20 años o más ({cutoff} y anteriores)"
                               if bins[0][2].startswith("older:") else "Un año por columna",
                               font=chart_font(canvas, 8), fill=c["muted"])
            maximum = max(1, max(count for _, count, _ in bins))
            for value in sorted({round(maximum * tick / 4) for tick in range(5)}):
                y = bottom - (bottom - top) * value / maximum
                if value:
                    canvas.create_line(left, y, right, y, fill=c["chart_grid"])
                canvas.create_text(left - px(8), y, anchor="e", text=str(value),
                                   font=chart_font(canvas, 8), fill=c["muted"])
            step = (right - left) / len(bins)
            # Con pocas columnas la barra no ocupa todo el ancho.
            gap = max(min(px(3), step * 0.12), (step - px(64)) / 2)
            label_every = max(1, math.ceil(px(34) / step))
            for index, (label, count, tag) in enumerate(bins):
                x0, x1 = left + index * step, left + (index + 1) * step
                y = bottom - (bottom - top) * count / maximum
                fill = c["chart_older"] if tag.startswith("older:") else c["primary"]
                canvas.create_rectangle(x0 + gap, y, x1 - gap, bottom, fill=fill, outline="",
                                        tags=("bin", tag, f"count:{count}"))
                if count and (step >= px(28) or tag.startswith("older:")):
                    canvas.create_text((x0 + x1) / 2, y - px(8), text=str(count),
                                       font=chart_font(canvas, 8), fill=c["ink"])
                if (index % label_every == 0 and (index == 0 or len(bins) - 1 - index >= label_every)) or index == len(bins) - 1:
                    canvas.create_text((x0 + x1) / 2, bottom + px(16), text=label,
                                       font=chart_font(canvas, 8), fill=c["muted"])
            canvas.create_line(left, bottom, right, bottom, fill=c["line_strong"])
            canvas.configure(scrollregion=(0, 0, width, height))
            signature = (tuple(sorted(numeric.items())), first, last)
            if signature != getattr(canvas, "_hist_signature", None):
                canvas.xview_moveto(0.0)
                canvas._hist_signature = signature

        def show_indicator_details(self, result: dict[str, Any] | None = None, parent: Any = None) -> Dialog | None:
            """Subtipos del indicador indicado o, por defecto, del que muestran los filtros del Dashboard."""
            result = result if result is not None else self._indicator_result
            if result is None:
                return None
            dialog = Dialog(parent or self.root, "Detalle del indicador por subtipo", 880, 470)
            ctk.CTkLabel(dialog.body, text=f"Detalle por subtipo · corte {result['corte']}", font=FONTS["h2"],
                         text_color=T["ink"], anchor="w").pack(fill="x", pady=(0, 10))
            columns = ("codigo", "clase", "cantidad", "periodo", "lambda", "peso_global", "valor")
            table = DataTable(dialog.body, columns, height_rows=8,
                              headings={"codigo": "Código", "clase": "Clase", "cantidad": "Cantidad",
                                        "periodo": "Periodo (años)", "lambda": "λ", "peso_global": "Peso global",
                                        "valor": "Valor"},
                              widths={"codigo": 130, "clase": 80, "cantidad": 90, "periodo": 120, "lambda": 110,
                                      "peso_global": 120, "valor": 120},
                              anchors={name: "e" for name in columns[2:]})
            table.pack(fill="both", expand=True)
            for subtype in result["subtipos"]:
                table.insert("", "end", values=(subtype["codigo"], subtype["clase"], subtype["cantidad"],
                                                subtype["periodo"], decimal(subtype["lambda"], 4),
                                                decimal(subtype["peso_global"], 3), decimal(subtype["valor"], 3)))
            dialog.table = table
            button(dialog.footer, "Cerrar", dialog.close, variant="primary", width=112).pack(side="right")
            dialog.present(wait=False)
            return dialog

        def _draw_distribution(self, canvas: Any, values: dict[str, int], title: str,
                               limit: int | None = None) -> None:
            """Diagrama de barras horizontales, de la categoría más frecuente a la menos frecuente."""
            total = sum(values.values())
            c, px, width, _height = chart_start(canvas, title, f"Base: {number(total)} productos (incluye «Sin dato»)")
            if not total:
                chart_empty(canvas, "Sin productos para estos filtros", width, max(px(160), canvas.winfo_height()))
                return
            items = sorted(((name, count) for name, count in values.items() if count),
                           key=lambda item: (-item[1], item[0]))
            if limit is not None and len(items) > limit:
                known = [item for item in items if item[0] != "Sin dato"]
                shown = known[:limit]
                if len(known) > limit:
                    shown.append(("Otras tipologías (agrupadas)", sum(count for _, count in known[limit:])))
                if values.get("Sin dato"):
                    shown.append(("Sin dato", values["Sin dato"]))
                items = shown
            max_count = max(count for _, count in items)
            label_width = min(px(230), max(px(120), width * 0.38))
            left, right = label_width + px(26), width - px(100)
            y = px(58)
            for name, count in items:
                label = canvas.create_text(px(12), y, anchor="nw", text=name[:1].upper() + name[1:], width=label_width,
                                           font=chart_font(canvas, 9), fill=c["ink"])
                bbox = canvas.bbox(label)
                row_height = max(px(30), (bbox[3] - bbox[1]) + px(12)) if bbox else px(30)
                bar_y = y + row_height / 2 - px(8)
                fill = c["unknown" if name == "Sin dato" else "primary"]
                canvas.create_rectangle(left, bar_y, right, bar_y + px(10), fill=c["chart_track"], outline="")
                canvas.create_rectangle(left, bar_y, left + max(px(2), (right - left) * count / max_count),
                                        bar_y + px(10), fill=fill, outline="", tags=("bar", f"count:{count}"))
                canvas.create_text(right + px(10), bar_y + px(5), anchor="w", text=f"{number(count)} · {percent(count, total)}",
                                   font=chart_font(canvas, 8), fill=c["muted"])
                y += row_height
            canvas.configure(scrollregion=(0, 0, width, max(y + px(10), canvas.winfo_height())))

        def _scope_changed(self, selected: str | None = None) -> None:
            kind = {"Grupo": "grupos", "Investigador": "investigadores", "Producto": "productos"}.get(self.view.get())
            self.scope_box.repo = self.repo
            self.scope_box.kind = kind or "grupos"
            self.scope_box.set_enabled(bool(kind))
            if selected is not None and kind and self.repo.get(kind, selected):
                self.scope_id.set(selected)
            elif kind != self._scope_kind or (kind and not self.repo.get(kind, self.scope_id.get())):
                rows = self.repo.page(kind, limit=1)["rows"] if kind else []
                self.scope_id.set(rows[0]["id"] if rows else "")
            else:
                self.scope_box._update_label()
            self._scope_kind = kind
            self.refresh_dashboard()

        def _year_changed(self, *_: Any) -> None:
            if self._year_timer:
                self.root.after_cancel(self._year_timer)
            if self.period.get() == "Personalizado":
                self._year_timer = self.root.after(350, self.refresh_dashboard)

        def clear_filters(self) -> None:
            self.period.set("Todos")
            self.start_year.set("")
            self.end_year.set("")
            self.validation.set("Todos")
            self.view.set("Todos")
            self._scope_changed()

        def refresh_dashboard(self) -> None:
            if self._year_timer:
                self.root.after_cancel(self._year_timer)
                self._year_timer = None
            custom = self.period.get() == "Personalizado"
            for entry in (self.start_entry, self.end_entry):
                entry.configure(state="normal" if custom else "disabled")
            try:
                start = end = None
                period = self.period.get()
                if period in ("Últimos 2 años", "Últimos 5 años"):
                    years = 2 if period == "Últimos 2 años" else 5
                    start, end = date.today().year - years + 1, date.today().year
                elif period == "Personalizado":
                    start = int(self.start_year.get()) if self.start_year.get().strip() else None
                    end = int(self.end_year.get()) if self.end_year.get().strip() else None
                    if start is not None and end is not None and start > end:
                        raise DataError("El año inicial supera al final")
                    if any(value is not None and not 1900 <= value <= date.today().year + 1 for value in (start, end)):
                        raise DataError("Año fuera de rango")
                status = "" if self.validation.get() == "Todos" else self.validation.get().lower()
                filters = (self.view.get(), self.scope_id.get(), start, end, status)
                if filters != self._stats_filter:
                    self.stats_offset = 0
                    self._stats_filter = filters
                result = self.repo.statistics(*filters, offset=self.stats_offset, limit=PAGE_SIZE)
                if self.stats_offset and self.stats_offset >= result["total"]:
                    self.stats_offset = max(0, (result["total"] - 1) // PAGE_SIZE * PAGE_SIZE)
                    result = self.repo.statistics(*filters, offset=self.stats_offset, limit=PAGE_SIZE)
                indicator = None
                kind = {"Grupo": "grupos", "Investigador": "investigadores"}.get(self.view.get())
                if kind and self.scope_id.get():
                    entity = self.repo.get(kind, self.scope_id.get())
                    if entity is not None and entity["activo"] == "1":
                        indicator = self.indicator(kind, self.scope_id.get(), start, end, status)
            except (ValueError, DataError) as exc:
                self.filter_error.set(f"Filtro inválido: {exc}. Se conservan los últimos resultados válidos.")
                self.filter_error_label.grid(row=3, column=0, columnspan=2, sticky="ew", pady=(8, 0))
                self.status.set(f"Filtro inválido: {exc}")
                return
            self.filter_error.set("")
            self.filter_error_label.grid_remove()
            self.current_stats = result
            self._indicator_result = indicator
            self.indicator_panel.show(indicator, self.scope_box.label.get() if indicator is not None else "")
            self.graph_indicator_summary.configure(
                text=(f"Indicador: {self.indicator_value.get()} / 10 · Corte {indicator['corte']}\n"
                      + self.indicator_note.get()) if indicator is not None else "Indicador de producción · Minciencias 2024")
            self.graph_indicator_details.configure(state="normal" if indicator is not None else "disabled")
            active_groups = self.overview["active_groups"]
            active_people = self.overview["active_people"]
            scope = self.scope_box.label.get() if self.view.get() != "Todos" else "Toda la institución"
            window = f"{start or 'sin inicio'}–{end or 'sin fin'}" if start is not None or end is not None else "Todos los años"
            self.context.set(f"{self.view.get()}: {scope} · {window} · Validación: {status or 'todas'}")
            self.card_values["total"].set(f"{result['total']:,}".replace(",", "."))
            for key, field, value in (("validated", "por_validacion", "validado"),
                                       ("rejected", "por_validacion", "rechazado"),
                                       ("undated", "por_anio", "Sin dato")):
                self.card_values[key].set(f"{result[field].get(value, 0):,}".replace(",", "."))
            for chip, (count, singular, plural) in zip(self.institution_chips, (
                    (active_groups, "grupo activo", "grupos activos"),
                    (active_people, "investigador activo", "investigadores activos"))):
                chip.set(f"{number(count)} {singular if count == 1 else plural}", "primary")
            rules = sorted(result["por_regla"].items(), key=lambda item: -item[1])
            self.note.set("Rechazos por: " + ", ".join(f"{RULE_LABELS[code].lower()} ({number(count)})"
                                                       for code, count in rules) if rules else "")
            if rules and not self.note_label.winfo_manager():
                self.note_label.pack(**self._note_pack)
            elif not rules:
                self.note_label.pack_forget()
            self.stats_table.delete(*self.stats_table.get_children())
            for index, row in enumerate(result["productos"]):
                self.stats_table.insert("", "end", iid=row["id"],
                                        values=tuple(display_value(field, row[field]) for field in self.stats_table["columns"]),
                                        tags=("stripe",) if index % 2 else ())
            self.stats_pager.show(self.stats_offset, len(result["productos"]), result["total"])
            self.stats_page_label.set(self.stats_pager.label.cget("text"))
            self._redraw_charts()

        def _open_product(self, _event: tk.Event | None = None) -> None:
            selection = self.stats_table.selection()
            if selection:
                self.open_detail("productos", selection[0])

        def usage(self, kind: str) -> dict[str, int]:
            if kind not in self._usage_cache:
                self._usage_cache[kind] = self.repo.field_usage(kind)
            return self._usage_cache[kind]

        def refresh(self, changed: set[str] | None = None) -> None:
            affected = set(ALL_FIELDS) | {"cola"} if changed is None else set(changed)
            if affected & {"grupos", "investigadores", "productos", "autorias", "grupos_productos"}:
                affected.add("productos")  # La validación del producto depende de sus vínculos y extremos.
            if affected & set(ALL_FIELDS):
                self._record_cache.clear()
                for kind in affected:
                    self._usage_cache.pop(kind, None)
            analysis_changed = bool(affected & (set(ALL_FIELDS) - {"planes"}))
            if analysis_changed:
                self._indicator_cache.clear()
            for kind, tab in self.entity_tabs.items():
                if kind in affected or (kind == "planes" and "grupos" in affected):
                    tab.refresh()
            for kind, tab in self.relation_tabs.items():
                if kind in affected or affected.intersection(RELATION_ENDS[kind][2:]):
                    tab.refresh()
            if affected & {"cola", "productos"}:
                self.refresh_queue()
            if affected & {"grupos", "investigadores"}:
                self.overview = self.repo.summary()
            if analysis_changed:
                self._scope_changed()
            self.root.title("PEA-i UPC — Investigación" + (" *" if self.repo.dirty else ""))
            self.workspace.set(self.workspace_caption)
            self.workspace_state.set("Cambios sin guardar" if self.repo.dirty else "Guardado",
                                     "pending" if self.repo.dirty else "muted")
            self.save_button.configure(text="Guardar como…" if self.data_file is None else "Guardar")
            for detail in tuple(self.details):
                if detail.winfo_exists() and (changed is None or detail.kind in affected or analysis_changed
                        or any(tab.source in affected or tab.target_kind in affected
                               for tab in detail.link_tabs.values())):
                    detail.refresh(affected if changed is not None else None)

        def mutate(self, action: Callable[[], Any], *, changed: set[str] | None = None) -> Any:
            try:
                result = action()
            except (DataError, OSError, ValueError) as exc:
                self.alert("No se pudo completar", str(exc), parent=self.root)
                return None
            self.refresh(changed)
            self.status.set("Cambio en memoria. Guarde para conservarlo al cerrar.")
            return result

        def pick_record(self, kind: str, current: str = "", parent: tk.Misc | None = None) -> str | None:
            return RecordPicker(parent or self.root, self.repo, kind, current).result

        def record(self, kind: str, key: str) -> dict[str, str]:
            cache_key = (kind, key)
            if cache_key not in self._record_cache:
                self._record_cache[cache_key] = self.repo.get(kind, key) or {"id": key, "activo": "0"}
            return self._record_cache[cache_key]

        def cache_references(self, page: dict[str, Any]) -> None:
            for kind, rows in page["references"].items():
                for row in rows:
                    self._record_cache[(kind, row["id"])] = row

        def indicator(self, kind: str, key: str, start: int | None = None,
                      end: int | None = None, status: str = "") -> dict[str, Any]:
            cache_key = (kind, key, start, end, status, date.today().year)
            if cache_key not in self._indicator_cache:
                result = self.repo.indicator(kind, key, start, end, status)
                if len(self._indicator_cache) >= 64:
                    self._indicator_cache.pop(next(iter(self._indicator_cache)))
                self._indicator_cache[cache_key] = result
            return self._indicator_cache[cache_key]

        def open_detail(self, kind: str, key: str, parent: tk.Misc | None = None,
                        tab: str | None = None) -> DetailDialog | None:
            """Abre la ficha; tab elige una pestaña (por ejemplo «Indicador») si la ficha la tiene."""
            if self.repo.get(kind, key) is None:
                self.status.set("Este registro ya no existe")
                return None
            return DetailDialog(parent or self.root, self, kind, key, tab)

        def reference_label(self, kind: str, key: str) -> str:
            row = self.record(kind, key)
            return f"{record_name(row)} · {key}"

        def entity_page(self, kind: str, offset: int, query: str, state: str) -> dict[str, Any]:
            active = {"Todos": "", "Activos": "1", "Inactivos": "0"}[state]
            return self.repo.page(kind, offset, PAGE_SIZE, query, active=active)

        def relationship_page(self, kind: str, offset: int, query: str) -> dict[str, Any]:
            page = self.repo.page(kind, offset, PAGE_SIZE, query, include_references=True)
            self.cache_references(page)
            return {"total": page["total"], "rows": page["rows"]}

        def edit_record(self, kind: str, key: str | tuple[str, str] | None = None,
                        defaults: dict[str, str] | None = None, parent: tk.Misc | None = None,
                        external: bool = False, locked_fields: tuple[str, ...] = ()) -> bool:
            initial = self.repo.get(kind, key) if key is not None else defaults
            if key is not None and defaults:
                initial = {**(initial or {}), **defaults}
            on_save = (lambda row: self.repo.update(kind, key, row)) if key is not None else (
                lambda row: self.repo.create(kind, row))
            dialog = RecordDialog(parent or self.root, kind, initial, repo=self.repo, on_save=on_save,
                                  suggested_id=self.repo.suggest_id(kind) if kind in ENTITY_FIELDS else "",
                                  external=external, is_new=key is None, locked_fields=locked_fields)
            if dialog.result is None:
                return False
            self.refresh({kind})
            self.status.set("Cambio en memoria. Guarde para conservarlo al cerrar.")
            return True

        def _may_discard(self) -> bool:
            if not self.repo.dirty:
                return True
            choice = self.decide("Cambios sin guardar", "¿Guardar cambios antes de continuar?", parent=self.root)
            if choice is None:
                return False
            if choice:
                return self.save()
            return True

        def start_choice(self) -> None:
            if self._startup_timer is not None:
                self.root.after_cancel(self._startup_timer)
                self._startup_timer = None
            self._set_loading(True)
            self.root.update_idletasks()
            try:
                if self.initial_file is not None:
                    self._load(self.initial_file)
                elif DEFAULT_DATA_FILE.exists():
                    self._load(DEFAULT_DATA_FILE)
                else:
                    self.status.set(f"No se encontró {DEFAULT_DATA_FILE.name}; espacio vacío.")
            finally:
                self._set_loading(False)

        def _set_loading(self, visible: bool) -> None:
            if not self.loading.winfo_exists():
                return
            if visible:
                self.loading_layer.place(x=0, y=0, relwidth=1, relheight=1)
                self.loading_layer.lift()
                self.loading.place(relx=0.5, rely=0.5, anchor="center")
                self.loading.progress.start()
            else:
                self.loading.progress.stop()
                self.loading.place_forget()
                self.loading_layer.place_forget()

        def new_workspace(self) -> None:
            if self._may_discard():
                if isinstance(self.repo, CppRepository):
                    try:
                        self.repo.reset()
                    except BackendError as exc:
                        self.alert("Motor C++ no disponible", str(exc), parent=self.root)
                        return
                else:
                    self.repo = Repository()
                self.data_file = None
                self.workspace_caption = "Espacio vacío"
                self.refresh()
                self.status.set("Espacio vacío. Guarde en un archivo antes de salir.")

        def _load(self, path: Path) -> None:
            path = Path(path)
            recovered = False
            try:
                if isinstance(self.repo, CppRepository):
                    self.repo.load(path)
                else:
                    new_repo = load_repository(path)
            except BackendError as exc:
                self.alert("Motor C++ no disponible", str(exc), parent=self.root)
                return
            except (DataError, OSError, ValueError, UnicodeError, json.JSONDecodeError) as exc:
                backup = backup_of(path)
                if backup.exists() and self.confirm("Datos dañados", f"{exc}\n\n¿Abrir la copia anterior de .backup?", parent=self.root):
                    try:
                        if isinstance(self.repo, CppRepository):
                            self.repo.load(backup)
                        else:
                            new_repo = load_repository(backup)
                    except BackendError as backup_exc:
                        self.alert("Motor C++ no disponible", str(backup_exc), parent=self.root)
                        return
                    except (DataError, OSError, ValueError, UnicodeError, json.JSONDecodeError) as backup_exc:
                        self.alert("Copia inválida", str(backup_exc), parent=self.root)
                        return
                    path = backup
                    recovered = True
                else:
                    self.alert("No se pudo cargar", str(exc), parent=self.root)
                    return
            if not isinstance(self.repo, CppRepository):
                self.repo = new_repo
            # Una carpeta anterior o una copia recuperada se guardan en un archivo nuevo.
            folder = is_folder_data(path)
            self.data_file = None if folder or recovered else path
            self.workspace_caption = ("Copia recuperada" if recovered else
                                      "Carpeta anterior" if folder else path.name)
            self.refresh()
            self.status.set(f"Datos cargados de {path}" + (" (Guardar como crea el archivo único)"
                                                          if self.data_file is None else ""))

        def open_file(self) -> None:
            if not self._may_discard():
                return
            chosen = filedialog.askopenfilename(
                parent=self.root, title="Archivo de datos PEA-i (o manifest.csv de una carpeta anterior)",
                filetypes=[("Datos PEA-i", "*.csv"), ("Todos los archivos", "*")])
            if chosen:
                self._load(Path(chosen))

        def open_default(self) -> None:
            if self._may_discard():
                self._load(DEFAULT_DATA_FILE)

        def save(self) -> bool:
            if self.data_file is None:
                return self.save_as()
            try:
                if isinstance(self.repo, CppRepository):
                    self.repo.save(self.data_file)
                else:
                    save_repository(self.repo, self.data_file)
            except (OSError, DataError) as exc:
                self.alert("No se pudo guardar", str(exc), parent=self.root)
                return False
            self.refresh(set())
            self.status.set(f"Guardado en {self.data_file}")
            return True

        def save_as(self) -> bool:
            chosen = filedialog.asksaveasfilename(
                parent=self.root, title="Guardar datos en un archivo", defaultextension=".csv",
                initialfile=DEFAULT_DATA_FILE.name, filetypes=[("Datos PEA-i", "*.csv")])
            if not chosen:
                return False
            previous = self.data_file
            previous_caption = self.workspace_caption
            self.data_file = Path(chosen)
            self.workspace_caption = self.data_file.name
            if not self.save():
                self.data_file = previous
                self.workspace_caption = previous_caption
                return False
            return True

        def _keep_import(self, summary: str) -> None:
            """Lo importado se guarda de inmediato en el archivo abierto: los datos crecen con cada fuente."""
            if self.data_file is not None and self.repo.dirty:
                saved = self.save()
                self.status.set(f"{summary} " + (f"Guardado en {self.data_file.name}." if saved
                                                 else "No se pudo guardar; use Guardar como."))
                return
            self.status.set(f"{summary} Guarde para conservarlo." if self.repo.dirty else summary)

        def export_spreadsheet(self) -> None:
            folder = filedialog.askdirectory(parent=self.root, title="Carpeta para exportar a hoja de cálculo")
            if not folder:
                return
            try:
                export_repository_spreadsheet(self.repo, Path(folder))
            except (OSError, DataError) as exc:
                self.alert("No se pudo exportar", str(exc), parent=self.root)
                return
            self.status.set(f"Exportado para hoja de cálculo en {folder}")

        def close(self) -> None:
            if self._may_discard():
                self._set_loading(False)
                cancel_widget_timers(self.root)
                self.root.destroy()

        def undo(self) -> None:
            try:
                undone = self.repo.undo()
            except DataError as exc:
                self.alert("No se pudo deshacer", str(exc), parent=self.root)
                return
            if undone:
                self.refresh()
                self.status.set("Última acción deshecha")
            else:
                self.status.set("No hay acciones para deshacer")

        def clear_history(self) -> None:
            if self.repo.history_size() and self.confirm("Vaciar historial", "¿Eliminar las acciones para deshacer?", parent=self.root):
                self.mutate(self.repo.clear_history, changed=set())

        def import_csv(self) -> None:
            path = filedialog.askopenfilename(parent=self.root, title="Importar hoja de cálculo",
                                              filetypes=[("Hojas de cálculo", "*.csv *.xlsx"),
                                                         ("CSV", "*.csv"), ("Excel", "*.xlsx")])
            if not path:
                return
            path = Path(path)
            if path.suffix.lower() == ".xlsx":
                try:
                    with tempfile.TemporaryDirectory() as temporary:
                        self.import_csv_path(_xlsx_to_csv(path, Path(temporary) / "hoja.csv"))
                except (DataError, OSError) as exc:
                    self.alert("Excel inválido", str(exc), parent=self.root)
                return
            self.import_csv_path(path)

        def import_csv_path(self, path: Path) -> None:
            try:
                with path.open(encoding="utf-8-sig", newline="") as stream:
                    header = next(csv.reader(stream), [])
            except (OSError, UnicodeError, csv.Error) as exc:
                self.alert("CSV inválido", str(exc), parent=self.root)
                return
            default = next((kind for kind, fields in ALL_FIELDS.items() if tuple(header) == fields), path.stem)
            options = {TITLES[kind]: kind for kind in ALL_FIELDS}
            dialog = ChoiceDialog(self.root, "Importar datos", f"Archivo: {path.name}\nSeleccione el contenido. Los encabezados deben coincidir con el esquema PEA-i.", options, default)
            kind = dialog.result
            if not kind:
                return
            try:
                if isinstance(self.repo, CppRepository):
                    total, expected, errors = self.repo.preview_csv(path, kind)
                else:
                    total, expected, errors = preview_csv_merge(self.repo, path, kind)
            except (DataError, OSError, UnicodeError) as exc:
                self.alert("CSV inválido", str(exc), parent=self.root)
                return
            summary = (f"Archivo {path.name} · {TITLES[kind]}\nFilas: {total}\nAceptables: {expected}\n"
                       f"Rechazadas: {total - expected}")
            if errors:
                summary += "\n\nPrimeros errores:\n" + "\n".join(errors[:20])
            if not self.confirm("Vista previa", summary + "\n\n¿Importar las aceptables?", accept="Importar aceptables",
                                cancel="Cancelar", parent=self.root):
                return
            try:
                if isinstance(self.repo, CppRepository):
                    accepted, errors = self.repo.import_csv(path, kind)
                else:
                    accepted, errors = merge_csv(self.repo, path, kind)
            except (DataError, OSError, UnicodeError) as exc:
                self.alert("No se pudo importar", str(exc), parent=self.root)
                return
            self.refresh()
            self._keep_import(f"Importación: {accepted} filas aceptadas; {total - accepted} rechazadas.")
            self.inform("Resultado de importación", f"Aceptadas: {accepted}\nRechazadas: {total - accepted}\n\n" + "\n".join(errors[:20]), parent=self.root)

        def _accept_external(self, kind: str, row: dict[str, str]) -> None:
            existing = self.repo.get(kind, row["id"]) if row.get("id") else None
            initial = {**existing, **{key: value for key, value in row.items() if value}} if existing else row
            if existing and existing.get("fuente") and row.get("fuente") and row["fuente"] not in existing["fuente"]:
                initial["fuente"] = existing["fuente"] + "; " + row["fuente"]
            if self.edit_record(kind, key=existing["id"] if existing else None, defaults=initial,
                                external=existing is None):
                self._keep_import("Registro importado.")

        def import_url(self, initial_url: str | None = None) -> None:
            """Consulta una URL pública en un diálogo; con initial_url (Consultar fuente) empieza de inmediato."""
            if self._url_busy:
                self.status.set("Ya hay una consulta URL en curso")
                return
            self._url_busy = True
            try:
                dialog = UrlDialog(self.root, self, initial_url or "", start=bool(initial_url))
            finally:
                self._url_busy = False
            if dialog.preview is None:
                return
            self.status.set(f"Página consultada: {dialog.preview.title}")
            self._show_web_preview(dialog.preview)

        def _show_web_preview(self, preview: WebPreview, source: str | None = None) -> Dialog:
            dialog = Dialog(self.root, "Vista previa de URL pública", 920, 700)
            dialog.minsize(780, 480)
            page = ctk.CTkScrollableFrame(dialog.body, fg_color="transparent", corner_radius=0,
                                          scrollbar_button_color=T["line_strong"], scrollbar_button_hover_color=T["muted"])
            page.pack(fill="both", expand=True)
            page._parent_canvas.configure(yscrollcommand=lambda first, last: set_scrollbar(page._scrollbar, first, last))
            wrapped = [ctk.CTkLabel(page, text=preview.title, font=FONTS["h1"], text_color=T["ink"], anchor="w",
                                    justify="left", wraplength=820)]
            wrapped[0].pack(fill="x", padx=(0, 8))
            chips = ctk.CTkFrame(page, fg_color="transparent")
            chips.pack(fill="x", pady=(6, 10))
            for text, tone in ((f"Formato {preview.format}", "muted"), (source or _web_source(preview.url), "primary")):
                Chip(chips, text, tone).pack(side="left", padx=(0, 6))
            url_box = ctk.CTkEntry(page, height=32, corner_radius=8, font=FONTS["small"], fg_color=T["field"],
                                   border_color=T["line_strong"], text_color=T["ink"])
            url_box.insert(0, preview.url)
            url_box.configure(state="readonly")
            url_box.pack(fill="x", padx=(0, 8))
            dialog.diagnostics = preview.diagnostics
            if preview.diagnostics:
                warning_text = (f"Advertencias de extracción: {len(preview.diagnostics)}\n"
                                + "\n".join(preview.diagnostics[:5]))
                if len(preview.diagnostics) > 5:
                    warning_text += "\nConsulte todas en «Ver texto de la página»."
                dialog.diagnostics_label = ctk.CTkLabel(
                    page, text=warning_text, font=FONTS["small"], text_color=T["pending"],
                    anchor="w", justify="left", wraplength=820)
                dialog.diagnostics_label.pack(fill="x", pady=(12, 0))
                wrapped.append(dialog.diagnostics_label)
            detected = bool(preview.related_members or preview.related_products or preview.related_plan)
            dialog.detected = {}
            dialog.category_chips = []
            if detected:
                ctk.CTkLabel(page, text="Datos detectados", font=FONTS["h2"], text_color=T["ink"], anchor="w").pack(
                    fill="x", pady=(16, 8))
                figures = [("Productos", number(len(preview.related_products))),
                           ("Autorías", number(len(preview.related_authorships)))]
                if preview.suggested_kind != "investigadores":
                    figures = [("Integrantes", number(len(preview.related_members))), *figures,
                               ("Plan", "Sí" if preview.related_plan else "No")]
                row = ctk.CTkFrame(page, fg_color="transparent")
                row.pack(fill="x", padx=(0, 8))
                for index, (caption, value) in enumerate(figures):
                    dialog.detected[caption] = tk.StringVar(value=value)
                    Kpi(row, caption, "primary", dialog.detected[caption], bg_color=T["page"]).grid(
                        row=0, column=index, sticky="ew", padx=(0 if index == 0 else 10, 0))
                    row.grid_columnconfigure(index, weight=1, uniform="figures")
                per_class: dict[str, int] = {}
                for product in preview.related_products:
                    category = catalog_by_code.get(product.get("categoria", ""))
                    name = category["clase"] if category else "Sin categoría"
                    per_class[name] = per_class.get(name, 0) + 1
                if per_class:
                    classes = ctk.CTkFrame(page, fg_color="transparent")
                    classes.pack(fill="x", pady=(12, 0))
                    ctk.CTkLabel(classes, text="Categorías estimadas por tipología", font=FONTS["caption"],
                                 text_color=T["muted"]).pack(side="left", padx=(0, 10))
                    for name in (*MODEL_CLASSES, "Sin categoría"):
                        if name in per_class:
                            chip = Chip(classes, f"{name} {number(per_class[name])}",
                                        "unknown" if name == "Sin categoría" else "pending")
                            chip.pack(side="left", padx=(0, 6))
                            dialog.category_chips.append(chip)
            note = ("Revise los datos detectados antes de importarlos: se crean los registros nuevos y se actualizan "
                    "los existentes con lo que diga la fuente." if detected else
                    "Datos detectados para revisar: " + (preview.suggested_kind or "ningún registro de investigación") + ".")
            wrapped.append(ctk.CTkLabel(page, text=note + " Los gráficos usan productos guardados e incluyen los campos "
                                        "desconocidos como «Sin dato».", font=FONTS["small"], text_color=T["muted"],
                                        anchor="w", justify="left", wraplength=820))
            wrapped[-1].pack(fill="x", pady=(14, 0))
            page.bind("<Configure>", lambda event: [label.configure(wraplength=max(300, event.width / scale_of(page) - 24))
                                                    for label in wrapped], add="+")
            details = (["ADVERTENCIAS DE EXTRACCIÓN", *preview.diagnostics, ""] if preview.diagnostics else [])
            details.append("METADATOS DECLARADOS POR LA PÁGINA")
            details.extend(f"{key}: {value}" for key, value in preview.metadata.items())
            if not preview.metadata:
                details.append("Ninguno")
            details.extend(("", "TEXTO VISIBLE (primeras 300 líneas)", *preview.lines))
            if preview.tables:
                details.append("\nTABLAS DETECTADAS (hasta 5 tablas y 100 filas por tabla)")
                for index, table in enumerate(preview.tables, start=1):
                    details.append(f"Tabla {index}")
                    details.extend(" | ".join(row) for row in table)
                    details.append("")
            dialog.page_text = ctk.CTkTextbox(page, height=280, wrap="word", font=FONTS["small"], fg_color=T["surface"],
                                              text_color=T["ink"], border_width=1, border_color=T["line"], corner_radius=8,
                                              scrollbar_button_color=T["line_strong"],
                                              scrollbar_button_hover_color=T["muted"])
            dialog.page_text.insert("1.0", "\n".join(details))
            dialog.page_text.configure(state="disabled")

            def toggle_text() -> None:
                shown = bool(dialog.page_text.winfo_manager())
                if shown:
                    dialog.page_text.pack_forget()
                else:
                    dialog.page_text.pack(fill="x", padx=(0, 8), pady=(8, 0))
                dialog.buttons["text"].configure(text="Ver texto de la página" if shown else "Ocultar texto de la página")
            dialog.buttons["text"] = button(page, "Ver texto de la página", toggle_text, width=210)
            dialog.buttons["text"].pack(anchor="w", pady=(14, 0))

            def create(kind: str) -> None:
                row = (preview.suggested_row or {}).copy() if preview.suggested_kind == kind else {}
                row.setdefault("id", new_id({"grupos": "G", "investigadores": "I", "productos": "P"}[kind]))
                if source is None:
                    row["url"] = _redact_url(preview.url)
                    row.setdefault("fuente", _web_source(preview.url))
                else:
                    row.setdefault("fuente", source)
                dialog.close()
                self._accept_external(kind, row)

            def import_group_profile() -> None:
                dialog.close()
                proposed = (preview.suggested_row or {}).copy()
                if not proposed:
                    return
                kind = preview.suggested_kind
                if kind not in ("grupos", "investigadores"):
                    return
                title = "Ficha GrupLAC" if kind == "grupos" else "Ficha CvLAC"
                reviewed = None
                code = "codigo_gruplac" if kind == "grupos" else "codigo_cvlac"
                try:
                    existing = self.repo.get(kind, proposed["id"]) or self.repo.find_external(kind, proposed[code])
                except BackendError as exc:
                    self.alert(title, str(exc), parent=self.root)
                    return
                if existing is None:
                    review = RecordDialog(self.root, kind, proposed, external=True, repo=self.repo)
                    if not review.result:
                        return
                    reviewed = review.result
                self.status.set(f"Incorporando {title.lower()}...")
                self.root.update_idletasks()
                try:
                    importer = import_gruplac_preview if kind == "grupos" else import_cvlac_preview
                    counts = importer(self.repo, preview, reviewed)
                except (DataError, OSError, ValueError) as exc:
                    self.alert(title, str(exc), parent=self.root)
                    return
                self.refresh()
                self._keep_import(f"Ficha incorporada: {counts['actualizados']} registros actualizados.")
                key = existing["id"] if existing is not None else (reviewed or proposed)["id"]
                self.show_import_result(counts, kind, key, diagnostics=preview.diagnostics)

            for kind, label in (("grupos", "Crear grupo"), ("investigadores", "Crear investigador"),
                                ("productos", "Crear producto")):
                dialog.buttons["create_" + kind] = button(dialog.footer, label, lambda name=kind: create(name), width=120)
                dialog.buttons["create_" + kind].pack(side="left", padx=(0, 6))
            if preview.csv_bytes is not None:
                def import_downloaded_csv() -> None:
                    dialog.close()
                    with tempfile.TemporaryDirectory() as temporary:
                        path = Path(temporary) / "descarga.csv"
                        path.write_bytes(preview.csv_bytes or b"")
                        self.import_csv_path(path)
                dialog.buttons["csv"] = button(dialog.footer, "Importar CSV PEA-i", import_downloaded_csv, width=150)
                dialog.buttons["csv"].pack(side="left", padx=(0, 6))
            if detected:
                dialog.buttons["import"] = button(dialog.footer, "Importar datos detectados", import_group_profile,
                                                  variant="primary", width=200)
                dialog.buttons["import"].pack(side="right")
            dialog.buttons["close"] = button(dialog.footer, "Cerrar", dialog.close, width=100)
            dialog.buttons["close"].pack(side="right", padx=(0, 8) if detected else 0)
            dialog.present(wait=False)
            return dialog

        def show_import_result(self, counts: dict[str, Any], kind: str, key: str,
                               *, diagnostics: tuple[str, ...] = ()) -> Dialog:
            """Resultado de incorporar una ficha; «Ver indicador» abre la ficha del grupo o investigador en su indicador."""
            title = {"grupos": "Ficha GrupLAC", "investigadores": "Ficha CvLAC"}.get(kind, "Resultado de importación")
            dialog = Dialog(self.root, title, 560)
            dialog.diagnostics = diagnostics
            ctk.CTkLabel(dialog.body, text=title + " incorporada", font=FONTS["h2"], text_color=T["ink"],
                         anchor="w").pack(fill="x")
            ctk.CTkLabel(dialog.body, text=self.reference_label(kind, key), font=FONTS["small"], text_color=T["muted"],
                         anchor="w", justify="left", wraplength=500).pack(fill="x", pady=(2, 10))
            dialog.table = DataTable(dialog.body, ("registros", "cantidad"), height_rows=9,
                                     headings={"registros": "Registros", "cantidad": "Cantidad"},
                                     widths={"registros": 320, "cantidad": 120}, anchors={"cantidad": "e"})
            dialog.table.pack(fill="both", expand=True)
            for field, label in (("grupos", "Grupos"), ("investigadores", "Investigadores"), ("membresias", "Membresías"),
                                 ("planes", "Planes"), ("productos", "Productos"),
                                 ("grupos_productos", "Vínculos grupo-producto"), ("autorias", "Autorías"),
                                 ("actualizados", "Actualizados con la fuente"), ("existentes", "Sin cambios")):
                dialog.table.insert("", "end", iid=field, values=(label, number(counts.get(field, 0))))
            if counts.get("errors"):
                ctk.CTkLabel(dialog.body, text=f"Errores: {len(counts['errors'])}\n" + "\n".join(counts["errors"][:5]),
                             font=FONTS["small"], text_color=T["rejected"], anchor="w", justify="left",
                             wraplength=500).pack(fill="x", pady=(10, 0))
            if diagnostics:
                dialog.diagnostics_label = ctk.CTkLabel(
                    dialog.body, text=f"Advertencias de extracción: {len(diagnostics)}\n" + "\n".join(diagnostics[:5]),
                    font=FONTS["small"], text_color=T["pending"], anchor="w", justify="left", wraplength=500)
                dialog.diagnostics_label.pack(fill="x", pady=(10, 0))
                if len(diagnostics) > 5:
                    dialog.buttons["diagnostics"] = button(
                        dialog.body, "Ver todas las advertencias",
                        lambda: self.inform("Advertencias de extracción", "\n".join(diagnostics), parent=dialog),
                        width=220)
                    dialog.buttons["diagnostics"].pack(anchor="w", pady=(6, 0))

            def open_indicator() -> None:
                dialog.close()
                self.open_detail(kind, key, tab="Indicador")
            has_indicator = kind in ("grupos", "investigadores") and self.repo.get(kind, key) is not None
            dialog.indicator_button = button(dialog.footer, "Ver indicador", open_indicator, variant="primary", width=136,
                                             state="normal" if has_indicator else "disabled")
            dialog.indicator_button.pack(side="right")
            button(dialog.footer, "Cerrar", dialog.close, width=112).pack(side="right", padx=(0, 8))
            dialog.present(wait=False, focus=dialog.indicator_button if has_indicator else None)
            return dialog

        def import_pdf(self) -> None:
            path = filedialog.askopenfilename(parent=self.root, title="PDF de GrupLAC o CvLAC", filetypes=[("PDF", "*.pdf")])
            if not path:
                return
            url = self.ask_text("Origen del PDF", "URL original pública de la ficha GrupLAC/CvLAC:", parent=self.root)
            if not url:
                return
            diagnostics: list[str] = []
            try:
                kind, row = parse_scienti_pdf(Path(path), url.strip(), diagnostics=diagnostics)
                row["fuente"] = f"PDF {Path(path).name}, consultado {date.today().isoformat()}"
            except (DataError, OSError, ValueError) as exc:
                self.alert("Importación PDF", str(exc), parent=self.root)
                return
            if diagnostics:
                self.inform("Advertencias de extracción", "\n".join(diagnostics), parent=self.root)
            self._accept_external(kind, row)

        def import_docx(self) -> None:
            path = filedialog.askopenfilename(parent=self.root, title="Documento Word",
                                              filetypes=[("Word", "*.docx")])
            if not path:
                return
            path = Path(path)
            try:
                html = extract_docx_html(path)
            except DataError as exc:
                self.alert("Documento Word", str(exc), parent=self.root)
                return
            preview = parse_web_page(html, path.resolve().as_uri(), "HTML")
            self._show_web_preview(preview, source=f"DOCX {path.name}, consultado {date.today().isoformat()}")

    return App(root, data_file, repository)


def _startup_error(root: Any, title: str, text: str) -> None:
    """Aviso previo a la ventana principal (por ejemplo, un backend C++ indicado que no arranca)."""
    import customtkinter as ctk
    window = ctk.CTkToplevel(root)
    window.title(title)
    ctk.CTkLabel(window, text=text, wraplength=440, justify="left", anchor="w").pack(fill="x", padx=22, pady=(20, 12))
    ctk.CTkButton(window, text="Aceptar", width=112, command=window.destroy).pack(anchor="e", padx=22, pady=(0, 18))
    window.bind("<Escape>", lambda _event: window.destroy())
    window.wait_visibility()
    window.grab_set()
    window.wait_window()


def run_gui(data_file: Path | None = None, *, python_backend: bool = False,
            cpp_binary: Path | None = None) -> None:
    from tkinter import TclError

    if not ensure_customtkinter():
        raise SystemExit(1)
    import customtkinter as ctk

    root = ctk.CTk()
    repository: Repository | CppRepository | None = None
    try:
        repository, fallback_reason = open_gui_repository(
            python_backend=python_backend, cpp_binary=cpp_binary)
        app = build_gui(root, repository, data_file)
        if fallback_reason:
            app.status.set(f"Motor Python activo: {fallback_reason}")
        root.mainloop()
    except DataError as exc:
        _startup_error(root, "Backend C++", str(exc))
        raise
    finally:
        if isinstance(repository, CppRepository):
            repository.close()
        try:
            root.destroy()
        except TclError:
            pass


def main() -> None:
    parser = argparse.ArgumentParser(description="PEA-i UPC con interfaz Tkinter")
    parser.add_argument("--data", type=Path,
                        help="Archivo de datos PEA-i (.csv) que se abre al iniciar; por defecto data/pea_upc.csv. "
                             "También acepta una carpeta anterior para convertirla")
    parser.add_argument("--cpp-binary", type=Path, help="Ejecutable C++ compilado (opcional)")
    parser.add_argument("--python-backend", action="store_true",
                        help="Usar la implementación Python independiente en vez del backend C++")
    args = parser.parse_args()
    try:
        run_gui(args.data, python_backend=args.python_backend, cpp_binary=args.cpp_binary)
    except DataError as exc:
        parser.exit(1, f"Error: {exc}\n")


if __name__ == "__main__":
    main()
