"""Pruebas del núcleo, ejecutables con python -m unittest discover -s tests/python."""

import csv
from collections import Counter
from email.message import Message
import hashlib
import importlib.util
import math
import os
from pathlib import Path
import shutil
import subprocess
import sys
import tempfile
import time
import unittest
from unittest.mock import patch

import fixture_f


SOURCE = Path(__file__).resolve().parents[2] / "src" / "python" / "Taller2_REMR.py"
FIXTURES = Path(__file__).resolve().parents[1] / "fixtures"
SPEC = importlib.util.spec_from_file_location("pea", SOURCE)
pea = importlib.util.module_from_spec(SPEC)
SPEC.loader.exec_module(pea)


def sample() -> pea.Repository:
    repo = pea.Repository()
    repo.create("grupos", {"id": "G1", "nombre": "Grupo uno"})
    repo.create("grupos", {"id": "G2", "nombre": "Grupo dos"})
    repo.create("investigadores", {"id": "I1", "nombre": "Ana"})
    repo.create("investigadores", {"id": "I2", "nombre": "Beto"})
    repo.create("productos", {"id": "P1", "titulo": "Artículo, con coma", "anio": "2025", "tipologia": "Artículo", "url": "https://example.org/p1"})
    repo.create("productos", {"id": "P2", "titulo": "Trabajo\nmultilínea", "anio": "2024", "tipologia": "Libro"})
    repo.create("membresias", {"grupo_id": "G1", "investigador_id": "I1", "rol": "Líder"})
    repo.create("membresias", {"grupo_id": "G1", "investigador_id": "I2", "rol": "Integrante"})
    repo.create("membresias", {"grupo_id": "G2", "investigador_id": "I1", "rol": "Integrante"})
    repo.create("autorias", {"producto_id": "P1", "investigador_id": "I1", "orden": "1"})
    repo.create("autorias", {"producto_id": "P1", "investigador_id": "I2", "orden": "2"})
    repo.create("grupos_productos", {"grupo_id": "G1", "producto_id": "P1"})
    repo.create("grupos_productos", {"grupo_id": "G2", "producto_id": "P1"})
    repo.create("grupos_productos", {"grupo_id": "G2", "producto_id": "P2"})
    repo.create("planes", {"id": "PL1", "grupo_id": "G1", "nombre": "Plan", "objetivo": "Publicar"})
    return repo


def indicator_fixture() -> pea.Repository:
    repo = pea.Repository()
    fixture_f.populate(repo)
    return repo


def write_folder(repo: pea.Repository, path: Path, version: str = "3") -> None:
    """Carpeta anterior al archivo único, con un CSV por tabla como en las versiones 1 a 3."""
    path.mkdir(parents=True, exist_ok=True)
    tables = {"manifest": [{"version": version, "guardado": "2026-09-24T00:00:00"}],
              **{kind: repo.rows(kind) for kind in pea.ALL_FIELDS},
              "cola_validacion": list(repo.queue),
              "historial": [{"orden": str(index), "snapshot": state} for index, state in enumerate(repo.history)]}
    for name, rows in tables.items():
        fields = pea.PRODUCT_FIELDS_V4 if name == "productos" else pea.DATA_SECTIONS[name]
        with (path / f"{name}.csv").open("w", encoding="utf-8", newline="") as stream:
            writer = csv.writer(stream, lineterminator="\n")
            writer.writerow(fields)
            writer.writerows([pea._storage_cell(row.get(field, "")) for field in fields] for row in rows)


def section_markers(path: Path) -> list[str]:
    with path.open(encoding="utf-8", newline="") as stream:
        return [record[0] for record in csv.reader(stream) if len(record) == 1 and record[0].startswith("#")]


class StartupTest(unittest.TestCase):
    def test_windows_bundle_matches_cpp_source(self):
        root = SOURCE.parents[2]
        self.assertEqual(pea.bundled_windows_cpp(root), root / "bin/windows/pea_cpp.exe")
        with tempfile.TemporaryDirectory() as folder:
            checkout = Path(folder)
            for name in ("CMakeLists.txt", "src/cpp/Taller2_REMR.cpp"):
                copied = checkout / name
                copied.parent.mkdir(parents=True, exist_ok=True)
                content = (root / name).read_bytes().replace(b"\r\n", b"\n")
                copied.write_bytes(content.replace(b"\n", b"\r\n"))
            binary = checkout / "bin/windows/pea_cpp.exe"
            binary.parent.mkdir(parents=True)
            binary.write_bytes(b"MZ")
            binary.with_suffix(".source-sha256").write_bytes(
                (root / "bin/windows/pea_cpp.source-sha256").read_bytes())
            (checkout / "bin/windows/pea_cpp.exe.sha256").write_bytes(
                (hashlib.sha256(b"MZ").hexdigest() + "\n").encode("ascii"))
            self.assertEqual(pea.bundled_windows_cpp(checkout), binary)

    def test_python_fallback_and_explicit_cpp_error(self):
        with patch.object(pea, "CppRepository", side_effect=pea.DataError("CMake falló\nNMake no disponible")):
            repository, notice = pea.open_gui_repository()
            self.assertIsInstance(repository, pea.Repository)
            self.assertEqual(notice, "CMake falló")
            with self.assertRaisesRegex(pea.DataError, "CMake falló"):
                pea.open_gui_repository(cpp_binary=Path("backend.exe"))

    @staticmethod
    def pip_result(code: int, stderr: str = ""):
        return lambda command, **_options: subprocess.CompletedProcess(command, code, "", stderr)

    def test_bootstrap_does_nothing_when_customtkinter_is_installed(self):
        with patch("importlib.util.find_spec", return_value=object()):
            self.assertTrue(pea.ensure_customtkinter(ask=lambda _text: self.fail("no debe preguntar"),
                                                     run=lambda *_args, **_kw: self.fail("no debe instalar")))

    def test_bootstrap_installs_when_missing_and_confirmed(self):
        calls, questions = [], []
        def run(command, **_options):
            calls.append(command)
            return subprocess.CompletedProcess(command, 0, "", "")
        with patch("importlib.util.find_spec", side_effect=[None, object()]), \
             patch.object(pea.sys, "prefix", "/usr"), patch.object(pea.sys, "base_prefix", "/usr"), \
             patch.object(pea.site, "addsitedir") as add_site:
            self.assertTrue(pea.ensure_customtkinter(ask=lambda text: questions.append(text) or True, run=run))
        self.assertEqual(calls, [[pea.sys.executable, "-m", "pip", "install", "customtkinter==6.0.0", "--user"]])
        self.assertIn("customtkinter==6.0.0", questions[0])
        add_site.assert_called_once()

    def test_bootstrap_installs_inside_a_virtual_environment_without_user_flag(self):
        calls = []
        def run(command, **_options):
            calls.append(command)
            return subprocess.CompletedProcess(command, 0, "", "")
        with patch("importlib.util.find_spec", side_effect=[None, object()]), \
             patch.object(pea.sys, "prefix", "/proyecto/.venv"), patch.object(pea.sys, "base_prefix", "/usr"):
            self.assertTrue(pea.ensure_customtkinter(ask=lambda _text: True, run=run))
        self.assertNotIn("--user", calls[0])

    def test_bootstrap_declined_does_not_run_pip(self):
        messages = []
        with patch("importlib.util.find_spec", return_value=None), \
             patch.object(pea, "_bootstrap_notice", messages.append):
            self.assertFalse(pea.ensure_customtkinter(ask=lambda _text: False,
                                                      run=lambda *_args, **_kw: self.fail("no debe instalar")))
        self.assertIn("pip install customtkinter==6.0.0", messages[-1])

    def test_bootstrap_explains_externally_managed_python(self):
        messages = []
        with patch("importlib.util.find_spec", return_value=None), \
             patch.object(pea, "_bootstrap_notice", messages.append):
            self.assertFalse(pea.ensure_customtkinter(
                ask=lambda _text: True, run=self.pip_result(1, "error: externally-managed-environment")))
        self.assertIn("python3 -m venv .venv", messages[-1])
        self.assertIn("requirements.txt", messages[-1])

    def test_bootstrap_reports_other_pip_failures_with_the_command(self):
        messages = []
        with patch("importlib.util.find_spec", return_value=None), \
             patch.object(pea, "_bootstrap_notice", messages.append):
            self.assertFalse(pea.ensure_customtkinter(
                ask=lambda _text: True, run=self.pip_result(1, "ERROR: No matching distribution found")))
        self.assertIn("No matching distribution found", messages[-1])
        self.assertIn("-m pip install customtkinter==6.0.0", messages[-1])

    def test_domain_import_does_not_need_customtkinter(self):
        script = ("import runpy, sys; sys.modules['customtkinter'] = None; "
                  "runpy.run_path(sys.argv[1])")
        result = subprocess.run([sys.executable, "-c", script, str(SOURCE)], capture_output=True, text=True, timeout=60)
        self.assertEqual(result.returncode, 0, result.stderr)

    def test_requirements_pin_customtkinter(self):
        requirements = (SOURCE.parents[2] / "requirements.txt").read_text(encoding="utf-8").split()
        self.assertEqual(requirements, [pea.CTK_REQUIREMENT])


class StructuresTest(unittest.TestCase):
    def test_doubly_linked_list_boundaries(self):
        linked = pea.DoublyLinkedList()
        self.assertEqual(list(linked), [])
        for key in "ABC":
            linked.append({"id": key})
        self.assertEqual([v["id"] for v in linked.backwards()], list("CBA"))
        linked.remove("A")
        linked.remove("C")
        linked.remove("B")
        self.assertEqual(linked.length, 0)
        self.assertIsNone(linked.head)
        self.assertIsNone(linked.tail)

    def test_multilist_two_directions_and_unlink(self):
        repo = sample()
        links = repo.relations["membresias"]
        self.assertEqual({r["investigador_id"] for r in links.by_left("G1")}, {"I1", "I2"})
        self.assertEqual({r["grupo_id"] for r in links.by_right("I1")}, {"G1", "G2"})
        repo.delete("membresias", ("G1", "I1"))
        self.assertEqual([r["investigador_id"] for r in links.by_left("G1")], ["I2"])
        self.assertEqual([r["grupo_id"] for r in links.by_right("I1")], ["G2"])

    def test_queue_fifo_stack_lifo(self):
        queue = pea.LinkedQueue()
        queue.enqueue({"id": "A"})
        queue.enqueue({"id": "B"})
        self.assertEqual(queue.dequeue()["id"], "A")
        self.assertEqual(queue.dequeue()["id"], "B")
        self.assertIsNone(queue.dequeue())
        stack = pea.LinkedStack()
        stack.push("A")
        stack.push("B")
        self.assertEqual(stack.pop(), "B")
        self.assertEqual(stack.pop(), "A")
        self.assertIsNone(stack.pop())


class DomainTest(unittest.TestCase):
    def test_model_catalog_matches_anexo(self):
        self.assertEqual(len(pea.MODEL_2024), 169)
        self.assertEqual(pea.MODEL_2024["ART_A2"].peso_global, 550)
        self.assertEqual(pea.MODEL_2024["CAP_LIB_B"].peso_global, 480)
        self.assertEqual(pea.MODEL_2024["PA1"].peso_global, 5000)
        self.assertEqual((pea.MODEL_2024["ART_D"].ventana, pea.MODEL_2024["LIB_B"].ventana,
                          pea.MODEL_2024["PA1"].ventana), (5, 7, 10))
        self.assertEqual(Counter(s.clase for s in pea.MODEL_2024.values()),
                         {"TOP": 18, "A": 26, "B": 48, "ASC": 11, "DPC": 46, "FRH_A": 4, "FRH_B": 16})
        self.assertEqual(pea.MODEL_CLASSES, ("TOP", "A", "B", "ASC", "DPC", "FRH_A", "FRH_B"))
        self.assertEqual(pea.MODEL_COEFFICIENTS,
                         {"TOP": 3.7, "A": 2.3, "B": 0.4, "ASC": 1.5, "DPC": 0.5,
                          "FRH_A": 1.0, "FRH_B": 0.2, "COHE": 0.1, "COOP": 0.3})
        catalog = pea.Repository().catalog()
        self.assertEqual(catalog[0]["codigo"], "ART_OPEN_A1")
        self.assertEqual([row["codigo"] for row in catalog], list(pea.MODEL_2024))
        for row in catalog:
            self.assertEqual(set(row), {"codigo", "clase", "familia", "descripcion", "ventana",
                                        "peso_alta", "peso_relativo", "peso_global"})
            self.assertIsInstance(row["ventana"], int)
            for field in ("peso_alta", "peso_relativo", "peso_global"):
                self.assertIsInstance(row[field], float)
        catalog[0]["descripcion"] = "No debe alterar el catálogo"
        self.assertEqual(pea.Repository().catalog()[0]["descripcion"],
                         "Artículo de investigación A1 acceso abierto")

    def test_product_category_rules(self):
        repo = pea.Repository()
        row = repo.create("productos", {"id": "P", "titulo": "T", "categoria": "ART_A1"})
        self.assertEqual(row["categoria_origen"], "manual")
        with self.assertRaisesRegex(pea.DataError, "Categoría desconocida: XYZ"):
            repo.create("productos", {"id": "Q", "titulo": "T", "categoria": "XYZ"})
        with self.assertRaisesRegex(pea.DataError, "Origen de categoría: estimada o manual"):
            repo.create("productos", {"id": "R", "titulo": "T", "categoria": "ART_D", "categoria_origen": "otra"})
        with self.assertRaisesRegex(pea.DataError, "Origen de categoría: estimada o manual"):
            repo.create("productos", {"id": "R", "titulo": "T", "categoria_origen": "otra"})
        est = repo.create("productos", {"id": "S", "titulo": "T", "categoria": "ART_D", "categoria_origen": "estimada"})
        self.assertEqual(est["categoria_origen"], "estimada")
        self.assertEqual(repo.update("productos", "S", {"observacion": "Conservar estimación"})["categoria_origen"], "estimada")
        self.assertEqual(repo.update("productos", "S", {"categoria": "ART_D"})["categoria_origen"], "estimada")
        self.assertEqual(repo.update("productos", "S", {"categoria": "ART_C"})["categoria_origen"], "manual")
        self.assertEqual(repo.update("productos", "S", {"categoria": ""})["categoria_origen"], "")
        self.assertEqual(repo.statistics()["por_categoria"], {"ART_A1": 1, "Sin dato": 1})
        self.assertEqual(repo.update("productos", "S", {"categoria": "ART_D", "categoria_origen": "estimada"})["categoria_origen"], "estimada")
        self.assertEqual(repo.update("productos", "S", {"categoria": ""})["categoria_origen"], "")
        self.assertEqual(repo.update("productos", "S", {"categoria_origen": "manual"})["categoria_origen"], "")

    def test_undo_keeps_category_and_v4_history(self):
        repo = pea.Repository()
        repo.create("productos", {"id": "P", "titulo": "T", "categoria": "ART_D"})
        repo.update("productos", "P", {"observacion": "nota"})
        self.assertTrue(repo.undo())
        self.assertEqual((repo.get("productos", "P")["categoria"], repo.get("productos", "P")["observacion"]), ("ART_D", ""))
        loaded = pea.load_repository(FIXTURES / "v4_con_historial.csv")
        self.assertEqual(loaded.get("productos", "P1")["categoria"], "")
        self.assertEqual(loaded.get("productos", "P1")["observacion"], "nota")
        self.assertTrue(loaded.undo())
        self.assertEqual(loaded.get("productos", "P1")["observacion"], "")
        with tempfile.TemporaryDirectory() as folder:
            out = Path(folder) / "datos.csv"
            pea.save_repository(loaded, out)
            self.assertEqual(out.read_text(encoding="utf-8").splitlines()[2].split(",")[0], "5")
            self.assertEqual(pea.load_repository(out).snapshot(), loaded.snapshot())
            pea.save_repository(repo, out)
            saved = pea.load_repository(out)
            self.assertEqual(saved.get("productos", "P")["categoria"], "ART_D")
            self.assertTrue(saved.undo())
            self.assertIsNone(saved.get("productos", "P"))

    def test_product_category_delta_history(self):
        repo = pea.Repository()
        repo.create("productos", {"id": "P", "titulo": "T", "categoria": "ART_D", "categoria_origen": "estimada"})
        original = repo.get("productos", "P").copy()
        repo.update("productos", "P", {"categoria": "ART_C"})
        change = {"__op": "replace", "__kind": "productos", **original}
        repo.clear_history()
        repo.history.push(pea.json.dumps({"__undo": [change]}))
        with tempfile.TemporaryDirectory() as folder:
            path = Path(folder) / "datos.csv"
            pea.save_repository(repo, path)
            saved = pea.load_repository(path)
            self.assertTrue(saved.undo())
            product = saved.get("productos", "P")
            self.assertEqual((product["categoria"], product["categoria_origen"], product["observacion"]),
                             ("ART_D", "estimada", ""))

    def test_csv_product_headers_v4_and_v5(self):
        with tempfile.TemporaryDirectory() as folder:
            for fields in (pea.PRODUCT_FIELDS_V4, pea.ENTITY_FIELDS["productos"]):
                with self.subTest(fields=fields):
                    path = Path(folder) / "importacion.csv"
                    row = {"id": "P", "titulo": "T", "anio": "2025"}
                    expected_category = ""
                    if "categoria" in fields:
                        row.update(categoria="ART_D", categoria_origen="estimada")
                        expected_category = "ART_D"
                    with path.open("w", encoding="utf-8", newline="") as stream:
                        writer = csv.DictWriter(stream, fieldnames=fields)
                        writer.writeheader()
                        writer.writerow(row)
                    repo = pea.Repository()
                    self.assertEqual(pea.preview_csv_merge(repo, path, "productos"), (1, 1, []))
                    self.assertEqual(repo.rows("productos"), [])
                    self.assertEqual(pea.merge_csv(repo, path, "productos"), (1, []))
                    self.assertEqual(repo.get("productos", "P")["categoria"], expected_category)
                    self.assertTrue(repo.undo())
                    self.assertEqual(repo.rows("productos"), [])

    def test_counts_dedupe_year_and_views(self):
        repo = sample()
        self.assertEqual(repo.statistics()["total"], 2)
        self.assertEqual(repo.statistics(view="Grupo", selected_id="G2")["total"], 2)
        self.assertEqual(repo.statistics(view="Investigador", selected_id="I1")["total"], 1)
        self.assertEqual(repo.statistics(view="Grupo", selected_id="G1", start=2025, end=2025)["total"], 1)
        self.assertEqual(repo.statistics(view="Grupo", selected_id="")["total"], 0)
        self.assertEqual(repo.statistics(status="validado")["total"], 1)
        self.assertEqual(repo.statistics(status="rechazado")["por_regla"], {"fuente": 1})
        with self.assertRaises(pea.DataError):
            repo.statistics(status="pendiente")

    def test_integrity_validation_and_undo(self):
        repo = sample()
        with self.assertRaises(pea.DataError):
            repo.delete("grupos", "G1")
        with self.assertRaises(pea.DataError):
            repo.create("autorias", {"producto_id": "P1", "investigador_id": "missing"})
        with self.assertRaises(pea.DataError):
            repo.create("productos", {"id": "P9", "titulo": "Con categoría", "categoria": "A1"})
        repo.update("productos", "P1", {"observacion": "Revisado por comité"})
        self.assertTrue(repo.undo())
        self.assertEqual(repo.get("productos", "P1")["observacion"], "")
        repo.toggle("productos", "P1")
        self.assertEqual(repo.statistics()["total"], 1)

    def test_automatic_validation_follows_links(self):
        repo = sample()
        self.assertEqual(repo.get("productos", "P1")["validacion"], "validado")
        # P2 tiene grupo y no autor: el vínculo con un grupo basta.
        self.assertEqual(repo.product_issues("P2"), ["fuente"])
        repo.update("productos", "P2", {"validacion": "rechazado", "url": "https://example.org/p2"})
        self.assertEqual((repo.get("productos", "P2")["validacion"], repo.product_issues("P2")),
                         ("validado", []))
        repo.toggle("grupos_productos", ("G2", "P2"))
        self.assertEqual((repo.get("productos", "P2")["validacion"], repo.product_issues("P2")),
                         ("rechazado", ["vinculo"]))
        repo.create("autorias", {"producto_id": "P2", "investigador_id": "I2"})
        self.assertEqual(repo.get("productos", "P2")["validacion"], "validado")
        repo.toggle("investigadores", "I2")
        self.assertEqual(repo.product_issues("P2"), ["vinculo"])
        self.assertEqual(repo.get("productos", "P1")["validacion"], "validado")
        self.assertTrue(repo.undo())
        self.assertEqual(repo.get("productos", "P2")["validacion"], "validado")
        repo.update("productos", "P1", {"doi": "10.12/corto"})
        self.assertEqual(repo.product_issues("P1"), ["doi"])
        repo.update("productos", "P1", {"doi": "https://doi.org/10.1234/ok", "url": ""})
        self.assertEqual(repo.product_issues("P1"), [])
        repo.delete("autorias", ("P2", "I2"))
        self.assertEqual(repo.product_issues("P2"), ["vinculo"])

    def test_link_rule_accepts_author_or_group(self):
        repo = pea.Repository()
        repo.create("grupos", {"id": "G", "nombre": "Grupo"})
        repo.create("investigadores", {"id": "I", "nombre": "Ana"})
        for key in ("SOLO-GRUPO", "SOLO-AUTOR", "SIN-VINCULO"):
            repo.create("productos", {"id": key, "titulo": key, "anio": "2025", "tipologia": "Evento",
                                      "url": "https://example.org/" + key})
        repo.create("grupos_productos", {"grupo_id": "G", "producto_id": "SOLO-GRUPO"})
        repo.create("autorias", {"producto_id": "SOLO-AUTOR", "investigador_id": "I"})
        self.assertEqual([repo.get("productos", key)["validacion"] for key in ("SOLO-GRUPO", "SOLO-AUTOR")],
                         ["validado", "validado"])
        self.assertEqual(repo.product_issues("SIN-VINCULO"), ["vinculo"])
        self.assertEqual(list(pea.VALIDATION_RULES), ["anio", "tipologia", "vinculo", "doi", "fuente"])
        self.assertEqual(pea.VALIDATION_RULES["vinculo"], "Sin autor ni grupo activos")
        repo.toggle("grupos", "G")
        self.assertEqual(repo.product_issues("SOLO-GRUPO"), ["vinculo"])

    def test_legacy_rule_codes_still_read_from_saved_reasons(self):
        self.assertEqual(pea.auto_reason_codes("reglas:autor,grupo,fuente"), ["autor", "grupo", "fuente"])
        self.assertEqual(pea.RULE_LABELS["autor"], "Sin autor activo registrado")
        self.assertEqual(pea.RULE_LABELS["vinculo"], "Sin autor ni grupo activos")

    def test_review_queue_and_undo(self):
        repo = sample()
        repo.enqueue_review("P1", "Verificar")
        repo.enqueue_review("P2", "Verificar")
        self.assertEqual(repo.queue.peek()["producto_id"], "P1")
        with self.assertRaises(pea.DataError):
            repo.process_review("  ")
        done = repo.process_review("Evidencia suficiente")
        self.assertEqual((done["producto_id"], done["validacion"]), ("P1", "validado"))
        self.assertEqual(repo.queue.peek()["producto_id"], "P2")
        self.assertEqual(repo.get("productos", "P1")["observacion"], "Evidencia suficiente")
        repo.undo()
        self.assertEqual(repo.queue.peek()["producto_id"], "P1")
        self.assertEqual(repo.get("productos", "P1")["observacion"], "")

    def test_rejected_products_enter_queue_once(self):
        repo = sample()
        self.assertEqual(repo.enqueue_rejected(), 1)
        job = repo.queue.peek()
        self.assertEqual((job["producto_id"], job["motivo"]), ("P2", "reglas:fuente"))
        self.assertEqual(pea.auto_reason_codes(job["motivo"]), ["fuente"])
        self.assertIsNone(pea.auto_reason_codes("Revisión manual"))
        self.assertEqual(repo.enqueue_rejected(), 0)
        self.assertTrue(repo.undo())
        self.assertEqual(repo.queue_size(), 0)

    def test_save_load_utf8_and_history(self):
        repo = sample()
        repo.enqueue_review("P1", "Tildes: acción")
        with tempfile.TemporaryDirectory() as temp:
            path = Path(temp) / "datos.csv"
            pea.save_repository(repo, path)
            self.assertEqual(sorted(item.name for item in Path(temp).iterdir()), ["datos.csv"])
            loaded = pea.load_repository(path)
            self.assertEqual(loaded.snapshot(), repo.snapshot())
            self.assertEqual(loaded.history.length, repo.history.length)
            self.assertTrue(loaded.undo())
            self.assertEqual(loaded.queue.length, 0)
            self.assertEqual(loaded.get("productos", "P2")["titulo"], "Trabajo\nmultilínea")

    def test_single_file_sections_in_order(self):
        with tempfile.TemporaryDirectory() as temp:
            path = Path(temp) / "datos.csv"
            pea.save_repository(sample(), path)
            self.assertEqual(section_markers(path), ["#" + name for name in pea.DATA_SECTIONS])
            self.assertEqual(list(pea.DATA_SECTIONS), [
                "manifest", "grupos", "investigadores", "productos", "planes",
                "membresias", "autorias", "grupos_productos", "cola_validacion", "historial"])
            manifest = [row for section, _, row in pea.read_data_sections(path) if section == "manifest"]
            self.assertEqual([row["version"] for row in manifest], ["5"])

    def test_single_file_tolerates_bom_crlf_and_blank_lines(self):
        with tempfile.TemporaryDirectory() as temp:
            path = Path(temp) / "datos.csv"
            repo = sample()
            pea.save_repository(repo, path)
            text = path.read_text(encoding="utf-8")
            path.write_bytes(b"\xef\xbb\xbf" + (text + "\n\n").replace("\n", "\r\n").encode("utf-8"))
            loaded = pea.load_repository(path)
            self.assertEqual(loaded.get("productos", "P2")["titulo"], "Trabajo\r\nmultilínea")
            self.assertEqual(len(loaded.rows("autorias")), 2)

    def test_single_file_rejects_bad_structure(self):
        with tempfile.TemporaryDirectory() as temp:
            good = Path(temp) / "datos.csv"
            pea.save_repository(sample(), good)
            lines = good.read_text(encoding="utf-8").split("\n")
            broken = Path(temp) / "roto.csv"

            def rejects(text: str, message: str) -> None:
                broken.write_text(text, encoding="utf-8")
                with self.assertRaises(pea.DataError) as caught:
                    pea.load_repository(broken)
                self.assertIn(message, str(caught.exception))

            cut = lines.index("#grupos_productos")
            rejects("\n".join(lines[:cut]) + "\n", "falta la sección #grupos_productos")
            first, second = lines.index("#grupos"), lines.index("#investigadores")
            swapped = lines[:first] + lines[second:lines.index("#productos")] + lines[first:second] + lines[lines.index("#productos"):]
            rejects("\n".join(swapped), "línea 4: se esperaba la sección #grupos")
            header = lines.index("#planes") + 1
            rejects("\n".join(lines[:header] + ["id,nombre"] + lines[header + 1:]), "encabezados incorrectos en #planes")
            row = lines.index("#autorias") + 2
            rejects("\n".join(lines[:row] + [lines[row] + ",sobrante"] + lines[row + 1:]),
                    f"línea {row + 1}: fila malformada en #autorias")
            rejects("G1,Grupo\n" + "\n".join(lines), "línea 1: el archivo debe comenzar con #manifest")
            rejects("\n".join(lines).replace("\n5,", "\n3,", 1), "Versión de datos incompatible")
            rejects("\n".join(lines).replace("\n5,", "\n4,", 1), "encabezados incorrectos en #productos")
            rejects("\n".join(lines) + "#extra\n", "sección sobrante #extra")

    def test_failed_save_keeps_previous_file(self):
        with tempfile.TemporaryDirectory() as temp:
            path = Path(temp) / "datos.csv"
            repo = sample()
            pea.save_repository(repo, path)
            before = path.read_bytes()
            repo.update("grupos", "G1", {"nombre": "Cambio que no se guarda"})
            with patch.object(pea.os, "replace", side_effect=OSError("disco lleno")):
                with self.assertRaises(OSError):
                    pea.save_repository(repo, path)
            self.assertEqual(path.read_bytes(), before)
            self.assertEqual([item.name for item in Path(temp).iterdir() if item.name.endswith(".tmp")], [])
            self.assertTrue(repo.dirty)

    def test_previous_save_can_be_recovered(self):
        repo = sample()
        with tempfile.TemporaryDirectory() as temp:
            path = Path(temp) / "datos.csv"
            pea.save_repository(repo, path)
            repo.update("grupos", "G1", {"nombre": "Nuevo nombre"})
            pea.save_repository(repo, path)
            path.write_text("archivo dañado", encoding="utf-8")
            with self.assertRaises(pea.DataError):
                pea.load_repository(path)
            self.assertEqual(pea.backup_of(path), Path(temp) / ".backup" / "datos.csv")
            recovered = pea.load_repository(pea.backup_of(path))
            self.assertEqual(recovered.get("grupos", "G1")["nombre"], "Grupo uno")

    def test_saving_to_a_folder_is_rejected(self):
        with tempfile.TemporaryDirectory() as temp:
            with self.assertRaises(pea.DataError):
                pea.save_repository(sample(), Path(temp))

    def test_legacy_folder_opens_by_folder_or_manifest(self):
        repo = sample()
        repo.enqueue_review("P1", "Revisar")
        with tempfile.TemporaryDirectory() as temp:
            folder = Path(temp) / "anterior"
            write_folder(repo, folder)
            for path in (folder, folder / "manifest.csv"):
                self.assertTrue(pea.is_folder_data(path))
                loaded = pea.load_repository(path)
                self.assertEqual(loaded.snapshot(), repo.snapshot())
                self.assertTrue(loaded.undo())
            self.assertEqual(pea.backup_of(folder), folder / ".backup")
            self.assertFalse(pea.is_folder_data(Path(temp) / "datos.csv"))

    def test_product_year_derived_from_date(self):
        repo = pea.Repository()
        row = repo.create("productos", {"id": "P-DATE", "titulo": "Con fecha", "fecha": "2025-03-12"})
        self.assertEqual(row["anio"], "2025")
        with self.assertRaises(pea.DataError):
            repo.create("productos", {"id": "P-BAD", "titulo": "Fechas distintas", "anio": "2024", "fecha": "2025-03-12"})

    def test_csv_partial_merge_and_bad_headers(self):
        repo = sample()
        with tempfile.TemporaryDirectory() as temp:
            path = Path(temp) / "productos.csv"
            with path.open("w", encoding="utf-8", newline="") as stream:
                writer = csv.DictWriter(stream, fieldnames=pea.ENTITY_FIELDS["productos"])
                writer.writeheader()
                writer.writerow({"id": "P3", "titulo": "Nuevo", "anio": "2026"})
                writer.writerow({"id": "P1", "titulo": "Duplicado", "anio": "2026"})
                writer.writerow({"id": "P4", "titulo": "Año imposible", "anio": "1800"})
            total, anticipated, preview_errors = pea.preview_csv_merge(repo, path, "productos")
            self.assertEqual((total, anticipated, len(preview_errors)), (3, 1, 2))
            self.assertIsNone(repo.get("productos", "P3"))
            accepted, errors = pea.merge_csv(repo, path, "productos")
            self.assertEqual(accepted, 1)
            self.assertEqual(len(errors), 2)
            self.assertIsNotNone(repo.get("productos", "P3"))
            repo.undo()
            self.assertIsNone(repo.get("productos", "P3"))

    def test_scienti_parsing(self):
        group_url = "https://scienti.minciencias.gov.co/gruplac/jsp/visualiza/visualizagr.jsp?nro=0001"
        html = "<html><body><h1>Grupo de prueba</h1><h2>Datos básicos</h2><table><tr><td>Líder</td><td>Ana Pérez</td></tr><tr><td>Clasificación</td><td>A1</td></tr></table></body></html>"
        kind, group = pea.parse_scienti_html(html, group_url)
        self.assertEqual((kind, group["nombre"], group["responsable"], group["categoria"]),
                         ("grupos", "Grupo de prueba", "Ana Pérez", "A1"))
        cv_url = "https://scienti.minciencias.gov.co/cvlac/visualizador/generarCurriculoCv.do?cod_rh=0042"
        kind, person = pea.parse_scienti_html("<h1>Hoja de vida</h1><p>Categoría</p><p>Investigador Junior (IJ) con vigencia</p><p>Nombre</p><p>Juana Álvarez</p>", cv_url)
        self.assertEqual((kind, person["nombre"], person["categoria"]),
                         ("investigadores", "Juana Álvarez", "Investigador Junior (IJ)"))
        with self.assertRaises(pea.DataError):
            pea.parse_scienti_html(html, "https://example.org/gruplac/?nro=0001")

    def test_scienti_download_uses_page_charset(self):
        url = "https://scienti.minciencias.gov.co/cvlac/visualizador/generarCurriculoCv.do?cod_rh=0042"
        html = ('<meta charset="iso-8859-1"><h1>Hoja de vida</h1><p>Categoría</p>'
                '<p>Investigador Junior (IJ) con vigencia</p><p>Nombre</p><p>Juana Álvarez</p>').encode("iso-8859-1")

        class Response:
            def __init__(self):
                self.headers = Message()
                self.headers["Content-Type"] = "text/html"

            def __enter__(self):
                return self

            def __exit__(self, *_args):
                return None

            def geturl(self):
                return url

            def read(self, _limit):
                return html

        class Opener:
            def __init__(self, response):
                self.response = response

            def open(self, _request, timeout):
                return self.response

        public_address = [(2, 1, 6, "", ("93.184.215.14", 443))]
        with patch.object(pea.socket, "getaddrinfo", return_value=public_address):
            with patch.object(pea, "build_opener", return_value=Opener(Response())):
                preview = pea.fetch_web_page(url)
        self.assertEqual(preview.suggested_kind, "investigadores")
        self.assertEqual(preview.suggested_row["nombre"], "Juana Álvarez")

    def test_other_web_pages_are_previewed_without_invented_records(self):
        url = "https://example.org/about"
        html = """<html><head><title>Observatorio público</title>
        <meta name="description" content="Investigación regional"></head>
        <body><h1>Datos disponibles</h1><p>Hay 12 informes.</p>
        <table><tr><th>Año</th><th>Informes</th></tr><tr><td>2024</td><td>12</td></tr></table>
        <script>document.write('999 productos');</script></body></html>"""
        preview = pea.parse_web_page(html, url)
        self.assertEqual(preview.title, "Observatorio público")
        self.assertIn("Hay 12 informes.", preview.lines)
        self.assertNotIn("999 productos", " ".join(preview.lines))
        self.assertEqual(preview.metadata["description"], "Investigación regional")
        self.assertEqual(preview.tables[0][1], ("2024", "12"))
        self.assertIsNone(preview.suggested_kind)
        self.assertIsNone(preview.suggested_row)
        self.assertEqual(pea.chart_values({"Sin dato": 3}), {})

    def test_declared_scholarly_article_only_populates_declared_fields(self):
        url = "https://journal.example.org/article/42"
        html = """<html><head><title>Revista de ciencia</title>
        <meta name="citation_title" content="Artículo de prueba">
        <meta name="citation_publication_date" content="2024-05-14">
        <meta name="citation_doi" content="10.1000/prueba"></head>
        <body><h1>Artículo de prueba</h1></body></html>"""
        preview = pea.parse_web_page(html, url)
        self.assertEqual(preview.suggested_kind, "productos")
        self.assertEqual(preview.suggested_row["titulo"], "Artículo de prueba")
        self.assertEqual(preview.suggested_row["anio"], "2024")
        self.assertEqual(preview.suggested_row["doi"], "10.1000/prueba")
        self.assertNotIn("categoria", preview.suggested_row)
        self.assertNotIn("tipologia", preview.suggested_row)
        self.assertEqual(preview.suggested_row["url"], url)
        schema_only = pea.parse_web_page('''<script type="application/ld+json">
        {"@type":"https://schema.org/ScholarlyArticle","headline":"Segundo artículo",
         "datePublished":"2023-07-01"}</script>''', url)
        self.assertEqual(schema_only.suggested_row["titulo"], "Segundo artículo")
        self.assertEqual(schema_only.suggested_row["anio"], "2023")
        self.assertEqual(schema_only.suggested_row["doi"], "")

    def test_public_url_fetch_formats_limits_and_private_redirects(self):
        class Response:
            def __init__(self, url, body, media_type):
                self.url, self.body = url, body
                self.headers = Message()
                self.headers["Content-Type"] = media_type

            def __enter__(self):
                return self

            def __exit__(self, *_args):
                return None

            def geturl(self):
                return self.url

            def read(self, limit):
                return self.body[:limit]

        class Opener:
            def __init__(self, response):
                self.response = response

            def open(self, _request, timeout):
                self_timeout = timeout
                self_test.assertEqual(self_timeout, 15)
                return self.response

        self_test = self
        public_address = [(2, 1, 6, "", ("93.184.215.14", 443))]
        with patch.object(pea.socket, "getaddrinfo", return_value=public_address):
            with patch.object(pea, "build_opener", return_value=Opener(Response(
                "https://example.org/data.csv", b'id,titulo,anio\nP-1,"Uno, dos",2024\n', "text/csv; charset=utf-8"))):
                preview = pea.fetch_web_page("https://example.org/data.csv")
            self.assertEqual(preview.format, "CSV")
            self.assertEqual(preview.metadata["Columnas CSV"], "id, titulo, anio")
            self.assertIn("Uno, dos", preview.lines[1])
            self.assertIsNotNone(preview.csv_bytes)
            with patch.object(pea, "build_opener", return_value=Opener(Response(
                "http://127.0.0.1/private", b"<html>secret</html>", "text/html"))):
                with self.assertRaises(pea.DataError):
                    pea.fetch_web_page("https://example.org")
            with patch.object(pea, "build_opener", return_value=Opener(Response(
                "https://example.org/huge", b"x" * 8_000_001, "text/plain"))):
                with self.assertRaises(pea.DataError):
                    pea.fetch_web_page("https://example.org/huge")
        with self.assertRaises(pea.DataError):
            pea.fetch_web_page("http://127.0.0.1/private")
        with self.assertRaises(pea.DataError):
            pea.PublicRedirects().redirect_request(None, None, 302, "Found", {}, "http://127.0.0.1/private")

    def test_gruplac_articles_require_declared_title_and_year(self):
        url = "https://scienti.minciencias.gov.co/gruplac/jsp/visualiza/visualizagr.jsp?nro=123"
        html = """<h1>Grupo de prueba</h1><h2>Datos básicos</h2>
        <h2>Artículos publicados</h2>
        <p>1.- Publicado en revista especializada: Estudio verificable</p>
        <p>Colombia, Revista X ISSN: 1234-5678, 2024 vol:2, DOI:10.1000/uno</p>
        <p>Autores: Persona A</p>
        <p>2.- Publicado en revista especializada: Estudio verificable</p>
        <p>Colombia, Revista X ISSN: 1234-5678, 2024 vol:2, DOI:10.1000/uno</p>
        <p>3.- Publicado en revista especializada: Sin fecha</p>
        <p>Revista Y ISSN: 9999-9999, DOI:</p>
        <h2>Libros publicados</h2>
        <p>1.- Libro: No es artículo</p>"""
        rows = pea.parse_gruplac_articles(html, url)
        self.assertEqual(len(rows), 1)
        self.assertEqual((rows[0]["titulo"], rows[0]["anio"], rows[0]["doi"]),
                         ("Estudio verificable", "2024", "10.1000/uno"))
        self.assertEqual((rows[0]["categoria"], rows[0]["categoria_origen"]), ("ART_D", "estimada"))
        self.assertEqual(pea.parse_web_page(html, url).related_products[0]["id"], rows[0]["id"])

    def test_gruplac_repairs_split_doi_and_drops_incomplete_doi(self):
        url = "https://scienti.minciencias.gov.co/gruplac/jsp/visualiza/visualizagr.jsp?nro=123"
        html = """<h1>Grupo de prueba</h1><h2>Datos básicos</h2>
        <h2>Artículos publicados</h2>
        <p>1.- Publicado en revista especializada: DOI partido</p>
        <p>Colombia, GLOBAL JOURNAL ISSN: 1916-9744, 2019 vol:11 fasc: págs: 1 - 13, DOI:10.5539/ gjhs.v11n9p34</p>
        <p>Autores: Persona A</p>
        <p>2.- Publicado en revista especializada: DOI incompleto</p>
        <p>Colombia, REVISTA DE SALUD PUBLICA ISSN: 0124-0064, 2024 vol:26 fasc: págs: 1 - 7, DOI:10.15446/.</p>
        <p>Autores: Persona B</p>
        <h2>Libros publicados</h2>"""
        rows = {row["titulo"]: row for row in pea.parse_gruplac_articles(html, url)}
        self.assertEqual(rows["DOI partido"]["doi"], "10.5539/gjhs.v11n9p34")
        self.assertEqual(rows["DOI incompleto"]["doi"], "")
        self.assertIn("DOI incompleto en la fuente: 10.15446/", rows["DOI incompleto"]["fuente"])
        # El ID sigue derivándose del texto publicado, para no duplicar productos ya importados.
        self.assertEqual(rows["DOI partido"]["id"], pea.product_id(pea.product_identity(
            "DOI partido", "2019", "10.5539/", "Publicado en revista especializada", True)))
        self.assertEqual(rows["DOI incompleto"]["id"], pea.product_id(pea.product_identity(
            "DOI incompleto", "2024", "10.15446/", "Publicado en revista especializada", True)))

    def test_reimport_replaces_an_invalid_stored_doi(self):
        repo = pea.Repository()
        for key, doi in (("P-VACIO", "10.15446/"), ("P-REPARADO", "10.5539/")):
            repo.create("productos", {"id": key, "titulo": key, "anio": "2024", "tipologia": "Artículo", "doi": doi})
        counts = {"existentes": 0, "actualizados": 0, "errors": []}
        for key, doi in (("P-VACIO", ""), ("P-REPARADO", "10.5539/gjhs.v11n9p34")):
            row = pea.clean_row("productos", {"id": key, "titulo": key, "anio": "2024", "tipologia": "Artículo",
                                              "doi": doi, "fuente": "GrupLAC, consultado hoy"})
            self.assertTrue(pea._merge_record(repo, "productos", row, key, counts, {}))
        self.assertEqual((repo.get("productos", "P-VACIO")["doi"], repo.get("productos", "P-REPARADO")["doi"]),
                         ("", "10.5539/gjhs.v11n9p34"))
        self.assertEqual(repo.get("productos", "P-VACIO")["fuente"], "GrupLAC, consultado hoy")
        self.assertEqual(counts["errors"], [])

    def test_reimport_matches_title_that_differs_only_in_final_period(self):
        # Otro grupo lista el mismo artículo sin el punto final; su DOI truncado ya no sirve para emparejarlo.
        repo = pea.Repository()
        repo.create("productos", {"id": "P", "titulo": "Violencia Obstétrica, Valledupar, Colombia.",
                                  "anio": "2024", "tipologia": "Artículo"})
        counts = {"existentes": 0, "actualizados": 0, "errors": []}
        row = pea.clean_row("productos", {"id": "P", "titulo": "Violencia obstétrica, Valledupar, Colombia",
                                          "anio": "2024", "tipologia": "Artículo"})
        self.assertTrue(pea._merge_record(repo, "productos", row, "P", counts, {}))
        self.assertEqual(counts["errors"], [])

    def test_gruplac_roster_plan_products_and_repeat_import(self):
        url = "https://scienti.minciencias.gov.co/gruplac/jsp/visualiza/visualizagr.jsp?nro=123"
        profile = "https://scienti.minciencias.gov.co/cvlac/visualizador/generarCurriculoCv.do?cod_rh=0042"
        html = f"""<h1>Grupo de prueba</h1><h2>Datos básicos</h2>
        <h2>Plan Estratégico</h2><p>Plan de trabajo: publicar libros y software.</p>
        <h2>Estado del arte</h2><p>Investigación previa.</p>
        <h2>Objetivos</h2><p>Compartir conocimiento.</p><h2>Retos</h2>
        <h2>Integrantes del grupo</h2><table>
        <tr><th>Nombre</th><th>Vinculación</th><th>Horas</th><th>Inicio - Fin</th></tr>
        <tr><td><a href="{profile}">1.- Ana Pérez</a></td><td>Investigador</td><td>20</td><td>2020/1 - Actual</td></tr>
        <tr><td>2.- Luis Gómez</td><td>Investigador</td><td>10</td><td>2018/1 - 2021/12</td></tr>
        </table>
        <h2>Artículos publicados</h2>
        <p>1.- Publicado en revista especializada: Estudio verificable</p>
        <p>Colombia, Revista X, 2024 vol:2, DOI:10.1000/uno</p>
        <p>Autores: Ana Perez, Persona Externa</p>
        <h2>Libros publicados</h2>
        <p>1.- Libro resultado de investigación: Libro probado</p>
        <p>Colombia, Editorial X, 2023.</p>
        <p>Autores: Luis Gómez</p>
        <h2>Capítulos de libro publicados</h2>"""
        preview = pea.parse_web_page(html, url)
        self.assertEqual(len(preview.related_members), 2)
        self.assertEqual(preview.related_members[0][0]["url"], profile)
        self.assertEqual([membership["activo"] for _, membership in preview.related_members], ["1", "0"])
        self.assertEqual(len(preview.related_products), 2)
        self.assertEqual({row["tipologia"] for row in preview.related_products},
                         {"Publicado en revista especializada", "Libro resultado de investigación"})
        self.assertIn("publicar libros y software", preview.related_plan["actividad"])
        self.assertEqual(len(preview.related_authorships), 2)
        for factory in (pea.Repository, lambda: pea.CppRepository(SOURCE.parents[2] / "build" / "pea_cpp")):
            if factory is not pea.Repository and not (SOURCE.parents[2] / "build" / "pea_cpp").exists():
                continue
            repo = factory()
            try:
                counts = pea.import_gruplac_preview(repo, preview)
                self.assertEqual([counts[kind] for kind in ("grupos", "investigadores", "membresias",
                                                           "planes", "productos", "grupos_productos", "autorias")],
                                 [1, 2, 2, 1, 2, 2, 2])
                self.assertEqual(counts["errors"], [])
                self.assertEqual(repo.statistics(view="Investigador", selected_id="I-0042")["total"], 1)
                again = pea.import_gruplac_preview(repo, preview)
                self.assertEqual(sum(again[kind] for kind in counts if kind not in ("existentes", "errors")), 0)
                self.assertEqual(again["errors"], [])
            finally:
                if isinstance(repo, pea.CppRepository):
                    repo.close()

    def test_gruplac_reimport_refreshes_existing_records(self):
        url = "https://scienti.minciencias.gov.co/gruplac/jsp/visualiza/visualizagr.jsp?nro=123"
        profile = "https://scienti.minciencias.gov.co/cvlac/visualizador/generarCurriculoCv.do?cod_rh=0042"

        def page(level: str, period: str, typology: str):
            return pea.parse_web_page(f"""<h1>Grupo de prueba</h1><h2>Datos básicos</h2>
            <p>Clasificación: {level}</p>
            <h2>Integrantes del grupo</h2><table>
            <tr><th>Nombre</th><th>Vinculación</th><th>Horas</th><th>Inicio - Fin</th></tr>
            <tr><td><a href="{profile}">1.- Ana Pérez</a></td><td>Investigador</td><td>20</td><td>{period}</td></tr>
            </table>
            <h2>Artículos publicados</h2>
            <p>1.- {typology}: Estudio verificable</p>
            <p>Colombia, Revista X, 2024 vol:2, DOI:10.1000/uno</p>
            <p>Autores: Ana Perez</p>
            <h2>Libros publicados</h2>""", url)

        first = page("B", "2020/1 - Actual", "Publicado en revista especializada")
        product_id = first.related_products[0]["id"]
        binary = SOURCE.parents[2] / "build" / "pea_cpp"
        for factory in (pea.Repository, lambda: pea.CppRepository(binary)):
            if factory is not pea.Repository and not binary.exists():
                continue
            repo = factory()
            try:
                pea.import_gruplac_preview(repo, first)
                repo.update("productos", product_id, {"observacion": "Revisado a mano",
                                                      "url": "https://doi.org/10.1000/uno"})
                repo.update("grupos", "G-123", {"unidad": "Facultad de prueba"})
                repo.update("investigadores", "I-0042", {"fuente": "Fuente anterior"})
                same = pea.import_gruplac_preview(repo, first)
                self.assertEqual((same["actualizados"], same["errors"]), (0, []))
                self.assertEqual(repo.get("investigadores", "I-0042")["fuente"], "Fuente anterior")
                self.assertEqual(repo.get("productos", product_id)["url"], "https://doi.org/10.1000/uno")
                changed = pea.import_gruplac_preview(repo, page("A", "2020/1 - 2025/6", "Artículo de investigación"))
                self.assertEqual((changed["actualizados"], changed["errors"]), (3, []))
                membership = repo.get("membresias", ("G-123", "I-0042"))
                self.assertEqual((membership["activo"], membership["rol"]), ("0", "Investigador · 2020/1 - 2025/6"))
                product = repo.get("productos", product_id)
                self.assertEqual(product["tipologia"], "Artículo de investigación")
                self.assertEqual((product["url"], product["observacion"]),
                                 ("https://doi.org/10.1000/uno", "Revisado a mano"))
                group = repo.get("grupos", "G-123")
                self.assertEqual((group["categoria"], group["unidad"]), ("A", "Facultad de prueba"))
                if isinstance(repo, pea.Repository):
                    before = repo.history_size()
                    quiet = pea.import_gruplac_preview(repo, page("A1", "2020/1 - 2025/6", "Artículo de investigación"),
                                                       remember=False)
                    self.assertEqual(quiet["actualizados"], 1)
                    self.assertEqual(repo.history_size(), before)
            finally:
                if isinstance(repo, pea.CppRepository):
                    repo.close()

    def test_real_dataset_loads_with_expected_relations(self):
        real = pea.load_repository(SOURCE.parents[2] / "data" / "pea_upc.csv")
        self.assertEqual(pea.DEFAULT_DATA_FILE, SOURCE.parents[2] / "data" / "pea_upc.csv")
        self.assertEqual(real.statistics(limit=0)["total"], 21419)
        self.assertEqual(real.statistics(view="Grupo", selected_id="G-00000000002099", limit=0)["total"], 438)
        self.assertEqual(real.statistics(view="Investigador", selected_id="I-0000494917", limit=0)["total"], 86)
        self.assertEqual(len(real.rows("grupos")), 66)
        self.assertEqual(len(real.rows("investigadores")), 2748)
        self.assertEqual(len(real.rows("membresias")), 3303)
        self.assertEqual(len(real.rows("planes")), 66)
        self.assertEqual(len(real.rows("autorias")), 17170)
        self.assertEqual(len(real.rows("grupos_productos")), 22781)
        self.assertEqual(Counter(row["categoria_origen"] for row in real.rows("productos")),
                         {"estimada": 20917, "": 502})
        self.assertEqual(sum(bool(row["categoria"]) for row in real.rows("productos")), 20917)
        self.assertEqual(sum(bool(row["fecha_creacion"]) for row in real.rows("grupos")), 66)
        self.assertEqual(real.statistics(limit=0)["por_validacion"], {"validado": 21419})
        self.assertEqual(real.statistics(limit=0)["por_regla"], {})
        self.assertEqual(sum(bool(row["categoria"]) for row in real.rows("investigadores")), 274)
        self.assertEqual(real.get("investigadores", "I-0000494917")["categoria"], "Investigador Asociado (I)")

    @unittest.skipUnless(shutil.which("pdftotext"), "Poppler no instalado")
    def test_text_pdf_import(self):
        # PDF pequeño generado directamente para probar el adaptador Poppler.
        content = b"BT /F1 12 Tf 50 750 Td (GrupLAC - Plataforma SCienTI - Colombia) Tj 0 -25 Td (Grupo PDF de prueba) Tj 0 -25 Td (Datos basicos) Tj ET"
        objects = [
            b"<< /Type /Catalog /Pages 2 0 R >>",
            b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
            b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Resources << /Font << /F1 4 0 R >> >> /Contents 5 0 R >>",
            b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
            b"<< /Length " + str(len(content)).encode() + b" >>\nstream\n" + content + b"\nendstream",
        ]
        data = bytearray(b"%PDF-1.4\n")
        offsets = [0]
        for index, obj in enumerate(objects, start=1):
            offsets.append(len(data))
            data.extend(f"{index} 0 obj\n".encode() + obj + b"\nendobj\n")
        xref = len(data)
        data.extend(f"xref\n0 {len(offsets)}\n0000000000 65535 f \n".encode())
        for offset in offsets[1:]:
            data.extend(f"{offset:010d} 00000 n \n".encode())
        data.extend(f"trailer\n<< /Size {len(offsets)} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n".encode())
        with tempfile.TemporaryDirectory() as temp:
            path = Path(temp) / "profile.pdf"
            path.write_bytes(data)
            url = "https://scienti.minciencias.gov.co/gruplac/jsp/visualiza/visualizagr.jsp?nro=123"
            kind, group = pea.parse_scienti_pdf(path, url)
        self.assertEqual((kind, group["nombre"]), ("grupos", "Grupo PDF de prueba"))

    def test_demo_contract(self):
        with tempfile.TemporaryDirectory(prefix="pea-fixture-") as folder:
            output = Path(folder) / "demo.csv"
            subprocess.run([sys.executable, str(SOURCE.parents[2] / "scripts/build_demo.py"),
                            "--output", str(output)], check=True, capture_output=True, text=True)
            demo = pea.load_repository(output)
            self.assertEqual(demo.statistics()["total"], 4)
            self.assertEqual(demo.statistics(start=2025, end=2026)["total"], 2)
            self.assertEqual(demo.statistics(start=2022, end=2026)["total"], 3)
            self.assertEqual(demo.statistics(view="Grupo", selected_id="G-DEMO-1")["total"], 3)
            self.assertEqual(demo.queue.peek()["producto_id"], "P-DEMO-2")
            self.assertEqual(demo.statistics(limit=0)["por_validacion"], {"validado": 3, "rechazado": 1})
            self.assertEqual(demo.product_issues("P-DEMO-4"), ["fuente"])


class ScientiImportTest(unittest.TestCase):
    GROUP_URL = "https://scienti.minciencias.gov.co/gruplac/jsp/visualiza/visualizagr.jsp?nro=123"
    CV_URL = "https://scienti.minciencias.gov.co/cvlac/visualizador/generarCurriculoCv.do?cod_rh=0000494917"

    def repositories(self):
        repos = [pea.Repository()]
        binary = SOURCE.parents[2] / "build/pea_cpp"
        if binary.exists():
            cpp = pea.CppRepository(binary)
            self.addCleanup(cpp.close)
            repos.append(cpp)
        return repos

    def group_preview(self):
        return pea.parse_web_page((FIXTURES / "gruplac_nuevas_secciones.html").read_text(encoding="utf-8"), self.GROUP_URL)

    def cv_preview(self):
        return pea.parse_web_page((FIXTURES / "cvlac_productos.html").read_text(encoding="utf-8"), self.CV_URL)

    def test_estimate_category_map(self):
        self.assertEqual(pea.estimate_category("Artículos publicados", "Publicado en revista especializada"), "ART_D")
        self.assertEqual(pea.estimate_category("Trabajos dirigidos/turorías", "Pregrado"), "TP_B")
        self.assertEqual(pea.estimate_category("Otros artículos publicados", "Epílogo"), "")
        # Prensa y divulgación: publicación editorial no especializada, como ya lo clasifica CvLAC.
        for kind in ("Revista de divulgación", "Periódico de noticias", "Columna de opinión"):
            self.assertEqual(pea.estimate_category("Otros artículos publicados", kind), "PEE_C2")
        self.assertEqual(pea.estimate_category("Libros publicados", "Libro pedagógico y/o de divulgación"), "LIB_DIV")
        self.assertTrue(all(code in pea.MODEL_2024 for code in pea.CATEGORY_ESTIMATES.values() if code))

    def test_gruplac_new_sections_and_creation_date(self):
        preview = self.group_preview()
        self.assertEqual(preview.suggested_row["fecha_creacion"], "2001-02-01")
        self.assertEqual({p["categoria"]: p["anio"] for p in preview.related_products},
                         {"TP_B": "2026", "PID_C": "2024", "EC_B": "2025"})
        self.assertTrue(all(p["categoria_origen"] == "estimada" for p in preview.related_products))
        thesis = next(p for p in preview.related_products if p["categoria"] == "TP_B")
        self.assertIn((thesis["id"], "I-0042"),
                      [(r["producto_id"], r["investigador_id"]) for r in preview.related_authorships])
        for repo in self.repositories():
            counts = pea.import_gruplac_preview(repo, preview)
            self.assertEqual((counts["productos"], counts["errors"]), (3, []))

    def test_reimport_keeps_manual_category(self):
        preview = self.group_preview()
        for repo in self.repositories():
            pea.import_gruplac_preview(repo, preview)
            product = next(p for p in preview.related_products if p["categoria"] == "TP_B")
            repo.update("productos", product["id"], {"categoria": "ART_A1"})
            pea.import_gruplac_preview(repo, preview)
            row = repo.get("productos", product["id"])
            self.assertEqual((row["categoria"], row["categoria_origen"]), ("ART_A1", "manual"))
            other = next(p for p in preview.related_products if p["categoria"] == "PID_C")
            repo.update("productos", other["id"], {"categoria": ""})
            refreshed = pea.import_gruplac_preview(repo, preview)
            self.assertEqual((refreshed["actualizados"], refreshed["errors"]), (1, []))
            row = repo.get("productos", other["id"])
            self.assertEqual((row["categoria"], row["categoria_origen"]), ("PID_C", "estimada"))

    def test_gruplac_category_sections_and_boundaries(self):
        for section, stop, kind, code in (
            ("Conceptos técnicos", "Diseños industriales", "Concepto", "CT"),
            ("Generación de Contenido Virtual", "Estrategias de Comunicación del Conocimiento", "Página web", "GC"),
            ("Eventos Artísticos", "Talleres de Creación", "Festival", "ECA_B"),
            ("Programa académico de doctorado", "Programa académico de maestría", "Programa", "AP_A"),
        ):
            with self.subTest(section=section):
                html = f'<h2>{section}</h2><p>1.- {kind}: Producto de prueba</p><p>2025,</p><h2>{stop}</h2>'
                products = pea.parse_gruplac_products(html, self.GROUP_URL)
                self.assertEqual([(p.row["categoria"], p.row["anio"]) for p in products], [(code, "2025")])
        bad = (FIXTURES / "gruplac_nuevas_secciones.html").read_text(encoding="utf-8").replace("2001 - 2", "2001 - 13")
        self.assertEqual(pea.parse_scienti_html(bad, self.GROUP_URL)[1]["fecha_creacion"], "")

    def test_reimport_changed_estimate_keeps_estimated_origin(self):
        original = self.group_preview()
        target = next(p for p in original.related_products if p["categoria"] == "TP_B")
        revised = original._replace(related_products=tuple(
            {**p, "categoria": "TM_B"} if p["id"] == target["id"] else p for p in original.related_products))
        for repo in self.repositories():
            pea.import_gruplac_preview(repo, original)
            for preview, category in ((revised, "TM_B"), (original, "TP_B")):
                counts = pea.import_gruplac_preview(repo, preview)
                self.assertEqual((counts["actualizados"], counts["errors"]), (1, []))
                row = repo.get("productos", target["id"])
                self.assertEqual((row["categoria"], row["categoria_origen"]), (category, "estimada"))

    def test_cvlac_products_parse_and_import(self):
        html = (FIXTURES / "cvlac_productos.html").read_text(encoding="utf-8")
        entries = pea.parse_cvlac_products(html, self.CV_URL)
        self.assertEqual([(p.row["categoria"], p.row["anio"]) for p in entries], [("ART_D", "2023"), ("TP_B", "2005")])
        self.assertEqual(entries[0].row["doi"], "10.15643134/2346030X.3249")
        preview = self.cv_preview()
        self.assertEqual(len(preview.related_authorships), 2)
        self.assertTrue(all(r["investigador_id"] == "I-0000494917" for r in preview.related_authorships))
        for repo in self.repositories():
            counts = pea.import_cvlac_preview(repo, preview)
            self.assertEqual((counts["investigadores"], counts["productos"], counts["autorias"], counts["errors"]), (1, 2, 2, []))
            again = pea.import_cvlac_preview(repo, preview)
            self.assertEqual((again["existentes"], again["actualizados"], again["errors"]), (3, 0, []))
            self.assertEqual(len(repo.rows("productos")), 2)

    def test_cvlac_product_merges_with_gruplac_twin(self):
        preview = self.cv_preview()
        article = preview.related_products[0]
        preview = preview._replace(related_products=(article,), related_authorships=preview.related_authorships[:1])
        html = f'<h1>Grupo de prueba</h1><h2>Datos básicos</h2><h2>Artículos publicados</h2><p>1.- Publicado en revista especializada: {article["titulo"].upper()}</p><p>2023, DOI:{article["doi"]}</p><h2>Libros publicados</h2>'
        group = pea.parse_web_page(html, self.GROUP_URL)
        self.assertEqual(group.related_products[0]["id"], article["id"])
        for repo in self.repositories():
            pea.import_gruplac_preview(repo, group)
            repo.update("productos", article["id"], {"categoria": "ART_A1"})
            counts = pea.import_cvlac_preview(repo, preview)
            self.assertEqual(counts["errors"], [])
            self.assertEqual((len(repo.rows("productos")), len(repo.rows("grupos_productos")), len(repo.rows("autorias"))), (1, 1, 1))
            row = repo.get("productos", article["id"])
            self.assertEqual((row["categoria"], row["categoria_origen"]), ("ART_A1", "manual"))

    def test_cvlac_technical_titles_and_block_boundaries(self):
        html = '''<p>Producción técnica - Softwares - Computacional</p><blockquote>
        <p>ANA PEREZ,</p><p>AG-CASU: Software de prueba,</p><p>Nombre comercial: AG,</p><p>,2025,</p></blockquote>
        <p>Producción técnica - Innovación de proceso o procedimiento</p><blockquote>
        <p>ANA PEREZ,</p><p>Prueba,</p><p>Nombre comercial: ,</p><p>,2024,</p></blockquote>
        <p>Producción técnica - Empresa de base tecnológica - Spin-off</p><blockquote>
        <p>ANA PEREZ,</p><p>Empresa de prueba,,</p><p>Nombre comercial: EMP,</p><p>,2023,</p></blockquote>
        <p>Producción bibliográfica - Artículo - Publicado en revista especializada</p><blockquote>
        <p>ANA PEREZ,</p><p>"Artículo sin fecha"</p></blockquote><h3>Eventos científicos</h3><p>Evento en 2026.</p>'''
        entries = pea.parse_cvlac_products(html, self.CV_URL)
        self.assertEqual([(p.row["titulo"], p.row["categoria"]) for p in entries],
                         [("AG-CASU: Software de prueba", "SF"), ("Prueba", "IPP"), ("Empresa de prueba", "EBT_B")])

    def test_cvlac_technical_twin_has_stable_metadata(self):
        group_html = '<h1>Grupo de prueba</h1><h2>Datos básicos</h2><h2>Softwares</h2><p>1.- Computacional: Software de prueba</p><p>2025,</p><h2>Empresas de base tecnológica</h2>'
        cv_html = '<p>Producción técnica - Softwares - Computacional</p><blockquote><p>ANA PEREZ,</p><p>Software de prueba,</p><p>Nombre comercial: Prueba,</p><p>,2025,</p></blockquote>'
        group = pea.parse_web_page(group_html, self.GROUP_URL)
        entries = pea.parse_cvlac_products(cv_html, self.CV_URL)
        cv = self.cv_preview()._replace(related_products=tuple(p.row for p in entries), related_authorships=())
        self.assertEqual(group.related_products[0]["id"], cv.related_products[0]["id"])
        for repo in self.repositories():
            pea.import_gruplac_preview(repo, group)
            counts = pea.import_cvlac_preview(repo, cv)
            self.assertEqual((counts["productos"], counts["actualizados"], counts["autorias"], counts["errors"]), (0, 0, 1, []))
            counts = pea.import_gruplac_preview(repo, group)
            self.assertEqual((counts["actualizados"], counts["errors"]), (0, []))

    def test_cvlac_review_rejects_different_profile(self):
        preview = self.cv_preview()
        repo = pea.Repository()
        with self.assertRaisesRegex(pea.DataError, "código CvLAC revisado no coincide"):
            pea.import_cvlac_preview(repo, preview, {**preview.suggested_row, "codigo_cvlac": "otro"})
        self.assertEqual(repo.rows("investigadores"), [])

    def test_import_script_accepts_repeated_cvlac(self):
        spec = importlib.util.spec_from_file_location("import_upc_test", SOURCE.parents[2] / "scripts/import_upc_dataset.py")
        script = importlib.util.module_from_spec(spec)
        spec.loader.exec_module(script)
        preview = self.cv_preview()
        with tempfile.TemporaryDirectory() as folder:
            target = Path(folder) / "datos.csv"
            with patch.object(script, "GROUPS", []), patch.object(script, "fetch_web_page", return_value=preview) as fetch:
                with patch.object(sys, "argv", ["import_upc_dataset.py", "--data", str(target), "--cvlac", self.CV_URL,
                                                 "--cvlac", self.CV_URL]), patch("builtins.print"):
                    script.main()
            self.assertEqual(fetch.call_count, 2)
            repo = pea.load_repository(target)
            self.assertEqual((len(repo.rows("productos")), len(repo.rows("autorias")), repo.history_size()), (2, 2, 0))


class IndicatorTest(unittest.TestCase):
    def test_indicator_group_matches_fixture(self):
        repo = indicator_fixture()
        r = repo.indicator("grupos", "G1", end=2026)
        self.assertEqual((r["poblacion"], r["productos"], r["sin_categoria"], r["estimadas"]), (2, 4, 1, 0))
        sub = {s["codigo"]: s for s in r["subtipos"]}
        self.assertEqual(list(sub), ["ART_D", "LIB_B", "WP"])
        self.assertEqual((sub["ART_D"]["cantidad"], sub["ART_D"]["periodo"]), (2, 5))
        self.assertAlmostEqual(sub["ART_D"]["lambda"], 0.336472, 6)
        self.assertAlmostEqual(sub["LIB_B"]["valor"], 320.475342, 6)
        self.assertAlmostEqual(r["cohesion"]["valor"], 1 / 3, 6)
        self.assertAlmostEqual(r["cooperacion"]["valor"], 2 / 3, 6)
        self.assertAlmostEqual(r["clases"][2]["indice"], 0.829843, 6)
        self.assertAlmostEqual(r["indicador"], 3.531937, 6)
        self.assertEqual([c["clase"] for c in r["clases"]], list(pea.MODEL_CLASSES))
        self.assertEqual((r["tipo"], r["id"], r["corte"], r["inicio"], r["estado"]),
                         ("grupos", "G1", 2026, None, ""))
        self.assertTrue(all(p["validacion"] == "rechazado" for p in repo.rows("productos")))
        self.assertEqual(set(r), {"tipo", "id", "corte", "inicio", "estado", "poblacion", "productos",
                                  "sin_categoria", "estimadas", "clases", "cohesion", "cooperacion",
                                  "subtipos", "indicador"})
        for term in (*r["clases"], r["cohesion"], r["cooperacion"]):
            for field in ("coeficiente", "valor", "maximo", "indice", "aporte"):
                self.assertIsInstance(term[field], float)
        self.assertAlmostEqual(r["indicador"], sum(t["aporte"] for t in r["clases"]) +
                               r["cohesion"]["aporte"] + r["cooperacion"]["aporte"], 12)

    def test_indicator_age_window_status_and_researcher(self):
        repo = indicator_fixture()
        g2 = repo.indicator("grupos", "G2", end=2026)
        self.assertEqual([(s["codigo"], s["cantidad"], s["periodo"]) for s in g2["subtipos"]], [("ART_D", 1, 2)])
        self.assertAlmostEqual(g2["indicador"], 0.4, 6)
        narrow = repo.indicator("grupos", "G1", start=2025, end=2026)
        art = next(s for s in narrow["subtipos"] if s["codigo"] == "ART_D")
        self.assertEqual((art["cantidad"], art["periodo"]), (1, 2))
        self.assertAlmostEqual(art["lambda"], 0.405465, 6)
        validated = repo.indicator("grupos", "G1", end=2026, status="validado")
        self.assertEqual((validated["productos"], validated["indicador"]), (0, 0.0))
        self.assertEqual(validated["estado"], "validado")
        self.assertEqual(repo.indicator("grupos", "G1", end=2026, status="rechazado")["indicador"],
                         repo.indicator("grupos", "G1", end=2026)["indicador"])
        i1 = repo.indicator("investigadores", "I1", end=2026)
        self.assertAlmostEqual(i1["cohesion"]["valor"], 0.5, 6)
        self.assertAlmostEqual(i1["cooperacion"]["valor"], 1.0, 6)
        self.assertAlmostEqual(i1["indicador"], 1.3, 6)

    def test_indicator_single_population_and_inactive(self):
        repo = indicator_fixture()
        repo.toggle("grupos", "G1")
        r = repo.indicator("grupos", "G2", end=2026)
        self.assertEqual(r["poblacion"], 1)
        self.assertAlmostEqual(r["indicador"], 0.4, 6)
        for term in (*r["clases"], r["cohesion"], r["cooperacion"]):
            self.assertEqual(term["indice"], 1.0 if term["valor"] else 0.0)
        with self.assertRaisesRegex(pea.DataError, "Registro inactivo"):
            repo.indicator("grupos", "G1", end=2026)

    def test_indicator_group_younger_than_cutoff(self):
        repo = indicator_fixture()
        r = repo.indicator("grupos", "G2", end=2024)
        self.assertEqual((r["productos"], r["subtipos"], r["indicador"]), (0, [], 0.0))
        first_year = repo.indicator("grupos", "G2", end=2025)
        self.assertEqual(first_year["productos"], 0)
        repo.update("grupos", "G1", {"fecha_creacion": "2026-12-31"})
        one_year = repo.indicator("grupos", "G1", end=2026)
        self.assertEqual((one_year["productos"], one_year["indicador"]), (0, 0.0))

    def test_indicator_ignores_category_without_year(self):
        repo = indicator_fixture()
        repo.create("productos", {"id": "P9", "titulo": "Sin año", "tipologia": "Prueba", "categoria": "ART_D"})
        repo.create("grupos_productos", {"grupo_id": "G1", "producto_id": "P9"})
        r = repo.indicator("grupos", "G1", end=2026)
        self.assertEqual((r["productos"], r["sin_categoria"]), (4, 1))

    def test_indicator_errors_and_default_cutoff(self):
        repo = indicator_fixture()
        for kind, key, options, message in (
            ("productos", "P1", {}, "Tipo de indicador: grupos o investigadores"),
            ("grupos", "missing", {}, "Registro no encontrado"),
            ("grupos", "G1", {"start": 2026, "end": 2025}, "El año inicial supera al final"),
            ("grupos", "G1", {"status": "pendiente"}, "Validación: validado o rechazado"),
        ):
            with self.subTest(kind=kind, key=key, options=options):
                with self.assertRaisesRegex(pea.DataError, message):
                    repo.indicator(kind, key, **options)
        repo.toggle("investigadores", "I3")
        with self.assertRaisesRegex(pea.DataError, "Registro inactivo"):
            repo.indicator("investigadores", "I3", end=2026)
        today = pea.date.today().year
        self.assertEqual(repo.indicator("grupos", "G1")["corte"], today)
        self.assertEqual(repo.indicator("grupos", "G1", start=today + 1)["indicador"], 0.0)

    def test_indicator_filters_population_maxima(self):
        repo = indicator_fixture()
        # Con la misma ventana de dos años, G1 y G2 tienen un artículo cada uno.
        narrow = repo.indicator("grupos", "G1", start=2025, end=2026)
        self.assertEqual(narrow["clases"][2]["indice"], 1.0)
        self.assertAlmostEqual(narrow["clases"][2]["maximo"], 45 * math.log(1.5), 12)
        repo.update("productos", "P4", {"url": "https://example.org/p4"})
        validated = repo.indicator("grupos", "G1", end=2026, status="validado")
        self.assertEqual(validated["clases"][2]["valor"], 0.0)
        self.assertAlmostEqual(validated["clases"][2]["maximo"], 45 * math.log(1.5), 12)
        repo.toggle("grupos", "G2")
        alone = repo.indicator("grupos", "G1", end=2026)
        self.assertEqual(alone["poblacion"], 1)
        self.assertEqual(alone["clases"][2]["indice"], 1.0)
        self.assertEqual(alone["cooperacion"]["valor"], 0.0)

    def test_indicator_ignores_inactive_products_and_links(self):
        for kind, key in (("productos", "P1"), ("grupos_productos", ("G1", "P1"))):
            with self.subTest(kind=kind):
                repo = indicator_fixture()
                repo.toggle(kind, key)
                r = repo.indicator("grupos", "G1", end=2026)
                self.assertEqual(r["productos"], 3)
                self.assertEqual(r["cohesion"]["valor"], 0.0)
                self.assertEqual(r["cooperacion"]["valor"], 0.5)
        repo = indicator_fixture()
        repo.toggle("autorias", ("P1", "I1"))
        self.assertEqual(repo.indicator("investigadores", "I1", end=2026)["productos"], 2)
        self.assertEqual(repo.indicator("grupos", "G1", end=2026)["productos"], 4)

    def test_indicator_active_authors_and_memberships(self):
        for kind, key in (("investigadores", "I2"), ("autorias", ("P1", "I2")),
                           ("membresias", ("G1", "I2"))):
            with self.subTest(kind=kind):
                repo = indicator_fixture()
                repo.toggle(kind, key)
                group = repo.indicator("grupos", "G1", end=2026)
                person = repo.indicator("investigadores", "I1", end=2026)
                self.assertEqual(group["productos"], 4)
                self.assertEqual(group["cohesion"]["valor"], 0.0)
                self.assertAlmostEqual(group["cooperacion"]["valor"], 2 / 3, 12)
                self.assertEqual(person["cohesion"]["valor"], 0.0)

    def test_indicator_researcher_cohesion_uses_all_active_groups(self):
        repo = indicator_fixture()
        repo.create("productos", {"id": "P9", "titulo": "Coautor de otro grupo", "anio": "2025", "categoria": "ART_D"})
        for researcher in ("I1", "I3"):
            repo.create("autorias", {"producto_id": "P9", "investigador_id": researcher})
        r = repo.indicator("investigadores", "I1", end=2026)
        self.assertAlmostEqual(r["cohesion"]["valor"], 2 / 3, 12)
        self.assertEqual(r["cooperacion"]["valor"], 1.0)
        repo.toggle("membresias", ("G1", "I1"))
        repo.toggle("membresias", ("G2", "I1"))
        isolated = repo.indicator("investigadores", "I1", end=2026)
        self.assertEqual((isolated["cohesion"]["valor"], isolated["cooperacion"]["valor"]), (0.0, 0.0))

    def test_indicator_uncategorized_and_estimated_counts(self):
        repo = indicator_fixture()
        for product_id, year, active in (("P-OLD", "2016", "1"), ("P-EDGE", "2017", "1"),
                                          ("P-FUTURE", "2027", "1"), ("P-NO-YEAR", "", "1"),
                                          ("P-INACTIVE", "2025", "0")):
            repo.create("productos", {"id": product_id, "titulo": product_id, "anio": year, "activo": active})
            repo.create("grupos_productos", {"grupo_id": "G1", "producto_id": product_id})
        repo.update("productos", "P1", {"categoria_origen": "estimada"})
        repo.update("productos", "P5", {"categoria_origen": "estimada"})
        r = repo.indicator("grupos", "G1", end=2026)
        self.assertEqual((r["sin_categoria"], r["estimadas"]), (2, 1))
        self.assertEqual(repo.indicator("grupos", "G1", start=2025, end=2026)["sin_categoria"], 1)
        self.assertEqual(repo.indicator("grupos", "G1", end=2026, status="validado")["sin_categoria"], 0)
        repo.update("grupos", "G1", {"fecha_creacion": "2025-01-01"})
        self.assertEqual(repo.indicator("grupos", "G1", end=2026)["sin_categoria"], 2)

    def test_indicator_empty_group_and_missing_creation_date(self):
        repo = pea.Repository()
        repo.create("grupos", {"id": "G", "nombre": "Sin fecha de creación"})
        self.assertEqual(repo.indicator("grupos", "G", end=2026)["indicador"], 0.0)
        repo.create("productos", {"id": "P", "titulo": "Sin autores", "anio": "2017", "categoria": "PA1"})
        repo.create("grupos_productos", {"grupo_id": "G", "producto_id": "P"})
        r = repo.indicator("grupos", "G", end=2026)
        self.assertEqual((r["productos"], r["subtipos"][0]["periodo"]), (1, 10))
        self.assertAlmostEqual(r["subtipos"][0]["valor"], 5000 * math.log(1.1), 12)
        self.assertEqual((r["cohesion"]["valor"], r["cooperacion"]["valor"], r["indicador"]), (0.0, 0.0, 3.7))

    def test_indicator_is_read_only_and_recomputes_after_edits(self):
        repo = indicator_fixture()
        snapshot, history, dirty = repo.snapshot(), repo.history_size(), repo.dirty
        original = repo.indicator("grupos", "G1", end=2026)
        self.assertEqual((repo.snapshot(), repo.history_size(), repo.dirty), (snapshot, history, dirty))
        repo.update("productos", "P1", {"categoria": "ART_A1"})
        changed = repo.indicator("grupos", "G1", end=2026)
        self.assertNotEqual(changed["indicador"], original["indicador"])
        self.assertTrue(repo.undo())
        self.assertEqual(repo.indicator("grupos", "G1", end=2026), original)

    def test_indicator_sums_subtypes_and_reaches_ten(self):
        repo = pea.Repository()
        for group in ("G1", "G2"):
            repo.create("grupos", {"id": group, "nombre": group})
        for person in ("I1", "I2"):
            repo.create("investigadores", {"id": person, "nombre": person})
            for group in ("G1", "G2"):
                repo.create("membresias", {"grupo_id": group, "investigador_id": person})
        for category in ("ART_A1", "LIB_B", "ART_D", "SF", "FIS", "WP", "TD_B", "TP_B"):
            repo.create("productos", {"id": category, "titulo": category, "anio": "2026", "categoria": category})
            repo.create("grupos_productos", {"grupo_id": "G1", "producto_id": category})
            for person in ("I1", "I2"):
                repo.create("autorias", {"producto_id": category, "investigador_id": person})
        r = repo.indicator("grupos", "G1", end=2026)
        self.assertEqual(r["productos"], 8)
        self.assertAlmostEqual(r["clases"][2]["valor"], (45 + 280) * math.log(1.2), 12)
        self.assertEqual((r["cohesion"]["valor"], r["cooperacion"]["valor"]), (1.0, 1.0))
        self.assertTrue(all(term["indice"] == 1.0 for term in (*r["clases"], r["cohesion"], r["cooperacion"])))
        self.assertAlmostEqual(r["indicador"], 10.0, 12)

    def test_indicator_observation_window_boundaries(self):
        for category, window in (("ART_D", 5), ("LIB_B", 7), ("PA1", 10)):
            with self.subTest(category=category):
                repo = pea.Repository()
                repo.create("grupos", {"id": "G", "nombre": "Sin límite por edad"})
                for year in (2026 - window, 2026 - window + 1, 2026, 2027):
                    repo.create("productos", {"id": str(year), "titulo": str(year), "anio": str(year), "categoria": category})
                    repo.create("grupos_productos", {"grupo_id": "G", "producto_id": str(year)})
                r = repo.indicator("grupos", "G", end=2026)
                self.assertEqual(r["productos"], 2)
                self.assertEqual(r["subtipos"][0]["periodo"], window)
                self.assertAlmostEqual(r["subtipos"][0]["lambda"], math.log(2 / window + 1), 12)

    def test_real_dataset_indicator_is_fast(self):
        repo = pea.load_repository(SOURCE.parents[2] / "data/pea_upc.csv")
        for kind, key in (("grupos", "G-00000000002099"), ("investigadores", "I-0000494917")):
            with self.subTest(kind=kind):
                started = time.monotonic()
                result = repo.indicator(kind, key, end=2026)
                elapsed = time.monotonic() - started
                self.assertLess(elapsed, 5.0)
                self.assertIsInstance(result["indicador"], float)
                self.assertTrue(0 <= result["indicador"] <= 10)


class SecurityTest(unittest.TestCase):
    def test_public_url_rejects_private_credentials_and_scheme(self):
        for bad in ("http://127.0.0.1/", "http://10.0.0.1/", "http://localhost/x",
                    "http://[::1]/", "ftp://example.org/x",
                    "http://user:pass@example.org/x"):
            with self.assertRaises(pea.DataError):
                pea._validate_public_url(bad)

    def test_public_url_rejects_nonstandard_port(self):
        with self.assertRaises(pea.DataError):
            pea._validate_public_url("https://example.org:8080/x")

    def test_public_url_rejects_private_resolution(self):
        mixed = [(2, 1, 6, "", ("93.184.215.14", 443)),
                 (2, 1, 6, "", ("127.0.0.1", 443))]
        with patch.object(pea.socket, "getaddrinfo", return_value=mixed):
            with self.assertRaises(pea.DataError):
                pea._validate_public_url("https://example.org/x")

    def test_resolve_public_addrs_dedupes(self):
        addresses = [(2, 1, 6, "", ("93.184.215.14", 443)),
                     (2, 1, 6, "", ("93.184.215.14", 443))]
        with patch.object(pea.socket, "getaddrinfo", return_value=addresses):
            self.assertEqual(pea._resolve_public_addrs("example.org", 443), ["93.184.215.14"])

    def test_public_opener_disables_proxy_and_pins(self):
        opener = pea._build_public_opener()
        self.assertTrue(any(isinstance(h, pea._PinnedHTTPHandler) for h in opener.handlers))
        self.assertTrue(any(isinstance(h, pea._PinnedHTTPSHandler) for h in opener.handlers))
        for handler in opener.handlers:
            self.assertFalse(getattr(handler, "proxies", {}))

    def test_redact_url_drops_sensitive_params(self):
        url = ("https://scienti.minciencias.gov.co/gruplac/jsp/visualiza/visualizagr.jsp"
               "?nro=0001&token=SECRET&auth=abc#frag")
        self.assertEqual(pea._redact_url(url),
                         "https://scienti.minciencias.gov.co/gruplac/jsp/visualiza/visualizagr.jsp?nro=0001")
        self.assertEqual(pea._redact_url("https://example.org/a?token=x"), "https://example.org/a")

    def test_csv_neutralizes_formula_cells(self):
        self.assertEqual(pea._neutralize_cell("=SUM(A1)"), "'=SUM(A1)")
        self.assertEqual(pea._neutralize_cell("+cmd"), "'+cmd")
        self.assertEqual(pea._neutralize_cell("-2+3"), "'-2+3")
        self.assertEqual(pea._neutralize_cell("@cmd"), "'@cmd")
        self.assertEqual(pea._neutralize_cell("normal"), "normal")
        self.assertEqual(pea._neutralize_cell(""), "")

    def test_spreadsheet_export_neutralizes(self):
        repo = sample()
        repo.create("grupos", {"id": "G3", "nombre": '=HYPERLINK("http://x")'})
        with tempfile.TemporaryDirectory() as temp:
            path = Path(temp)
            pea.export_repository_spreadsheet(repo, path)
            text = (path / "grupos.csv").read_text(encoding="utf-8")
            self.assertIn("'=HYPERLINK", text)

    def test_saved_csv_neutralizes_and_roundtrips_formula_text(self):
        repo = pea.Repository()
        repo.create("grupos", {"id": "G1", "nombre": "=1+1"}, remember=False)
        repo.create("productos", {"id": "P1", "titulo": "'=texto"}, remember=False)
        with tempfile.TemporaryDirectory() as temp:
            path = Path(temp) / "datos.csv"
            pea.save_repository(repo, path)
            with path.open(encoding="utf-8", newline="") as stream:
                stored = {record[0]: record for record in csv.reader(stream)}
            self.assertEqual(stored["G1"][1], "'=1+1")
            self.assertEqual(stored["P1"][1], "''=texto")
            loaded = pea.load_repository(path)
            self.assertEqual(loaded.get("grupos", "G1")["nombre"], "=1+1")
            self.assertEqual(loaded.get("productos", "P1")["titulo"], "'=texto")

    def test_legacy_v1_variants_load(self):
        with tempfile.TemporaryDirectory() as temp:
            path = Path(temp)
            write_folder(sample(), path, "1")
            groups = pea._csv_read(path / "grupos.csv", pea.ENTITY_FIELDS["grupos"], encoded=True)
            with (path / "grupos.csv").open("w", encoding="utf-8", newline="") as stream:
                fields = ("id", "nombre", "sigla", *pea.ENTITY_FIELDS["grupos"][2:])
                writer = csv.DictWriter(stream, fieldnames=fields)
                writer.writeheader()
                writer.writerows(groups)
            self.assertEqual(pea.load_repository(path).get("grupos", "G1")["nombre"], "Grupo uno")
            products = pea._csv_read(path / "productos.csv", pea.ENTITY_FIELDS["productos"], encoded=True)
            with (path / "productos.csv").open("w", encoding="utf-8", newline="") as stream:
                fields = (*(field for field in pea.PRODUCT_FIELDS_V4 if field != "validacion"), "categoria")
                writer = csv.DictWriter(stream, fieldnames=fields, extrasaction="ignore")
                writer.writeheader()
                writer.writerows(products)
            (path / "cola_validacion.csv").unlink()
            loaded = pea.load_repository(path)
            self.assertEqual(loaded.get("productos", "P1")["validacion"], "validado")
            self.assertEqual(loaded.queue_size(), 0)

    def test_version_2_product_category_moves_to_observation(self):
        with tempfile.TemporaryDirectory() as temp:
            path = Path(temp)
            write_folder(sample(), path, "2")
            products = pea._csv_read(path / "productos.csv", pea.ENTITY_FIELDS["productos"], encoded=True)
            products[0].update(categoria="A1", validacion="pendiente")
            products[1].update(categoria="")
            with (path / "productos.csv").open("w", encoding="utf-8", newline="") as stream:
                writer = csv.DictWriter(stream, fieldnames=(*pea.PRODUCT_FIELDS_V4[:6], "categoria",
                                                            *pea.PRODUCT_FIELDS_V4[6:]))
                writer.writeheader()
                writer.writerows(products)
            loaded = pea.load_repository(path)
            self.assertEqual(loaded.get("productos", "P1")["observacion"],
                             "Categoría registrada antes de la versión 3: A1")
            self.assertEqual(loaded.get("productos", "P1")["validacion"], "validado")
            self.assertEqual(loaded.get("productos", "P2")["observacion"], "")
            converted = path / "datos.csv"
            pea.save_repository(loaded, converted)
            manifest = [row for section, _, row in pea.read_data_sections(converted) if section == "manifest"]
            self.assertEqual(manifest[0]["version"], "5")
            self.assertEqual(pea.load_repository(converted).snapshot(), loaded.snapshot())

    def test_extract_pdf_text_caps_output(self):
        class Stream:
            def __init__(self, data):
                self._data = data

            def read(self, size):
                chunk = self._data[:size]
                self._data = self._data[size:]
                return chunk

            def close(self):
                pass

        class Process:
            def __init__(self, stdout):
                self.stdout = Stream(stdout)
                self.stderr = Stream("")
                self.returncode = 0

            def kill(self):
                pass

            def wait(self, timeout=None):
                return 0

        with tempfile.TemporaryDirectory() as temp:
            path = Path(temp) / "x.pdf"
            path.write_bytes(b"%PDF-1.4")
            with patch.object(pea.shutil, "which", return_value="/usr/bin/pdftotext"):
                with patch.object(pea.subprocess, "Popen",
                                  return_value=Process("x" * (pea.MAX_PDF_TEXT + 1))):
                    with self.assertRaises(pea.DataError):
                        pea.extract_pdf_text(path)

    @unittest.skipIf(os.name == "nt", "El ejecutable falso usa un shebang POSIX")
    def test_pdf_wall_timeout_stops_stalled_converter(self):
        with tempfile.TemporaryDirectory() as temp:
            folder = Path(temp)
            converter = folder / "pdftotext"
            converter.write_text("#!/usr/bin/env python3\nimport time\ntime.sleep(3)\n", encoding="utf-8")
            converter.chmod(0o755)
            pdf = folder / "documento.pdf"
            pdf.write_bytes(b"%PDF-1.4")
            with patch.dict(os.environ, {"PATH": str(folder) + os.pathsep + os.environ["PATH"]}):
                with patch.object(pea, "PDF_TIMEOUT", 1):
                    started = time.monotonic()
                    with self.assertRaisesRegex(pea.DataError, "tiempo"):
                        pea.extract_pdf_text(pdf)
                    self.assertLess(time.monotonic() - started, 2.5)

    def test_field_length_limit(self):
        repo = pea.Repository()
        with self.assertRaises(pea.DataError):
            repo.create("grupos", {"id": "G1", "nombre": "x" * (pea.MAX_FIELD_LEN + 1)})

    def test_csv_row_limit(self):
        with tempfile.TemporaryDirectory() as temp:
            path = Path(temp) / "productos.csv"
            with path.open("w", encoding="utf-8", newline="") as stream:
                writer = csv.DictWriter(stream, fieldnames=pea.ENTITY_FIELDS["productos"])
                writer.writeheader()
                for index in range(10):
                    writer.writerow({"id": f"P{index}", "titulo": "T"})
            with patch.object(pea, "MAX_CSV_ROWS", 5):
                with self.assertRaises(pea.DataError):
                    pea._csv_read(path, pea.ENTITY_FIELDS["productos"])

    def test_deeply_nested_jsonld_does_not_crash(self):
        nested = '{"@type":"ScholarlyArticle","@graph":' + '{"@graph":' * 5000 + '{}' + '}' * 5000 + '}'
        html = f'<script type="application/ld+json">{nested}</script>'
        preview = pea.parse_web_page(html, "https://example.org/x")
        self.assertIsNone(preview.suggested_kind)


def _minimal_docx(text: str) -> bytes:
    content_types = ('<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
                     '<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">'
                     '<Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/>'
                     '<Default Extension="xml" ContentType="application/xml"/>'
                     '<Override PartName="/word/document.xml" '
                     'ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml"/>'
                     '</Types>')
    rels = ('<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
            '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
            '<Relationship Id="rId1" '
            'Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument" '
            'Target="word/document.xml"/></Relationships>')
    document = ('<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
                '<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
                '<w:body><w:p><w:r><w:t>' + text + '</w:t></w:r></w:p></w:body></w:document>')
    import io
    import zipfile
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w") as archive:
        archive.writestr("[Content_Types].xml", content_types)
        archive.writestr("_rels/.rels", rels)
        archive.writestr("word/document.xml", document)
    return buffer.getvalue()


class ImportFormatsTest(unittest.TestCase):
    def test_xlsx_to_csv_roundtrip(self):
        try:
            import openpyxl
        except ImportError:
            self.skipTest("openpyxl no instalado")
        workbook = openpyxl.Workbook()
        workbook.active.append(["id", "nombre"])
        workbook.active.append(["1", "Ana"])
        workbook.active.append(["2", "Beto"])
        with tempfile.TemporaryDirectory() as temp:
            source = Path(temp) / "datos.xlsx"
            target = Path(temp) / "datos.csv"
            workbook.save(source)
            pea._xlsx_to_csv(source, target)
            rows = pea._csv_read(target, ("id", "nombre"))
            self.assertEqual([row["nombre"] for row in rows], ["Ana", "Beto"])

    def test_xlsx_to_csv_requires_headers(self):
        try:
            import openpyxl
        except ImportError:
            self.skipTest("openpyxl no instalado")
        workbook = openpyxl.Workbook()
        with tempfile.TemporaryDirectory() as temp:
            source = Path(temp) / "vacio.xlsx"
            target = Path(temp) / "vacio.csv"
            workbook.save(source)
            with self.assertRaises(pea.DataError):
                pea._xlsx_to_csv(source, target)

    def test_docx_extract_text(self):
        try:
            import mammoth
        except ImportError:
            self.skipTest("mammoth no instalado")
        with tempfile.TemporaryDirectory() as temp:
            path = Path(temp) / "documento.docx"
            path.write_bytes(_minimal_docx("Hola mundo"))
            html = pea.extract_docx_html(path)
            self.assertIn("Hola mundo", html)


if __name__ == "__main__":
    unittest.main()
